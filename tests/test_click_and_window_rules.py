"""JS_Click mimics a person (only what could be seen is clicked); Switch to window fails when the popup it waits for never opens; a test whose
last browser window is gone stops instead of running dozens of steps against nothing."""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from regrunner.engine.runner import RunOptions, execute
from regrunner.events import EventBus
from tests.workbook_factory import Sheet

pytestmark = pytest.mark.browser


def build(path: Path, site: str, page: str, rows: list[tuple[str, str, dict]]) -> Path:
    wb = openpyxl.Workbook()
    ds = wb.active
    ds.title = "DataSheets"
    ds.append(["SheetsToExecute", "blnExecute", "ParameterSheet", "Comments"])
    ds.append(["Flow", "Y", "Params_1", "rules"])
    g = wb.create_sheet("Global")
    g.append(["Parameter", "Value"])
    g.append(["Environment", "UAT"])
    ps = wb.create_sheet("Params_1")
    ps.append(["blnExecute", "DT_URL"])
    ps.append(["Y", site + page])
    s = Sheet(wb.create_sheet("Flow"), "Y")
    s.add("Open", "Open", Page="chrome", Value="DT_URL")
    for method, name, cols in rows:
        s.add(method, name, **cols)
    wb.save(path)
    return path


def at(xpath: str) -> dict:
    return {"FindBy": "xpath", "FindBy_Value": xpath, "Index": 0}


async def run(make_cfg, wb: Path, **overrides):
    cfg = make_cfg(**{"waits.window_grace_s": 0.5, "timeouts.element_s": 3, "timeouts.optional_s": 1.0, **overrides})
    events: list[dict] = []
    bus = EventBus()
    bus.subscribe(events.append)
    result = await execute(RunOptions(workbook=wb, tests=["Flow"], seed=1), cfg, bus)
    return result.tests[0], events


def by_name(test) -> dict:
    return {s.name: s for s in test.steps}


# -- JS_Click --------------------------------------------------------------------------------------------------------------------
async def test_js_click_only_clicks_what_a_person_could_see(site, make_cfg, tmp_path):
    rows = [("JS_Click", "visible", at("//button[@id='shown']")),
            ("JS_Click", "covered by an overlay", at("//button[@id='covered']")),           # a real click cannot reach it, a person still sees it
            ("JS_Click", "waits for a late one", at("//button[@id='late']")),               # appears after 0.7 s
            ("JS_Click", "radio, label shown", at("//input[@id='r1']")),                    # native input hidden, label visible
            ("JS_Click", "display none", at("//button[@id='gone']")),
            ("JS_Click", "visibility hidden", at("//button[@id='invisible']")),
            ("JS_Click", "hidden parent", at("//button[@id='nested']")),
            ("JS_Click", "no size", at("//button[@id='zero']")),
            ("JS_Click", "radio, label hidden too", at("//input[@id='r2']")),
            ("JS_Click", "hidden but optional", {**at("//button[@id='gone']"), "Ignore_not_existing_object": "Y"}),
            ("Output", "what was clicked", {**at("//div[@id='log']"), "Output_Property": "innertext"})]
    test, _ = await run(make_cfg, build(tmp_path / "a.xlsx", site, "click_rules.html", rows), **{"runner.stop_after_failed_steps": 0})   # five failures in a row are the point
    steps = by_name(test)
    for name in ("visible", "covered by an overlay", "waits for a late one", "radio, label shown"):
        assert steps[name].status == "PASSED", (name, steps[name].error)
    for name in ("display none", "visibility hidden", "hidden parent", "no size", "radio, label hidden too"):
        assert steps[name].status == "FAILED" and "not visible, so a person could not click it" in steps[name].error, (name, steps[name].error)
    assert steps["hidden but optional"].status == "PASSED" and steps["hidden but optional"].ignored_error       # optional stays optional
    assert steps["what was clicked"].actual == "shown;covered;late;radio-yes;"                              # the hidden ones were never clicked


# -- Switch to window ---------------------------------------------------------------------------------------------------------------
async def test_switch_to_window_fails_when_the_popup_never_opened(site, make_cfg, tmp_path):
    rows = [("Exist", "page is up, nothing opens a popup", at("//h1")),
            ("SwitchToWindow", "Switch to window", {"Value": "-1"}),
            ("Exist", "still on the first page", at("//h1"))]
    test, _ = await run(make_cfg, build(tmp_path / "a.xlsx", site, "second.html", rows))
    step = by_name(test)["Switch to window"]
    assert step.status == "FAILED" and step.error.startswith("No new window opened") and "only 1 window is open" in step.error


async def test_switch_to_window_still_works_when_the_popup_opens(site, make_cfg, tmp_path):
    rows = [("JS_Click", "Open second window", at("//a[@id='open-win']")),
            ("SwitchToWindow", "Switch to window", {"Value": "-1"}),
            ("Output", "second window heading", {**at("//h1"), "Output_Property": "innertext"}),
            ("SwitchToWindow", "Switch again", {"Value": "-1"}),                              # the popup is still the newest: nothing to fail
            ("Close", "Close popup", {}),
            ("SwitchToMainWindow", "Back to the first window", {})]
    test, _ = await run(make_cfg, build(tmp_path / "a.xlsx", site, "index.html", rows))
    assert test.status == "PASSED", [(s.name, s.error) for s in test.steps if s.status != "PASSED"]
    assert by_name(test)["second window heading"].actual == "Second Window"


# -- no window left ------------------------------------------------------------------------------------------------------------------
async def test_a_test_with_no_window_left_stops_at_the_first_step_that_needs_one(site, make_cfg, tmp_path):
    """The real run: a Close step closed the only window, then 50 steps failed 'No open browser window' and 7 'passed' on a dead browser."""
    rows = [("Close", "Close window", {}),
            ("SwitchToMainWindow", "Back to the main window", {}),
            ("Exist", "heading", at("//h1")),
            ("Wait", "Wait", {"Value": 1}),
            ("Quit", "Quit", {})]
    test, events = await run(make_cfg, build(tmp_path / "a.xlsx", site, "second.html", rows))
    assert [s.name for s in test.steps] == ["Open", "Close window", "Back to the main window"]
    assert test.status == "FAILED" and test.steps[-1].status == "FAILED" and test.skipped == 3
    assert 'step 2 "Close window" closed the last one' in test.error and "steps after step 3 could not run" in test.error


async def test_closing_the_window_at_the_end_of_a_flow_is_fine(site, make_cfg, tmp_path):
    rows = [("Exist", "heading", at("//h1")), ("Close", "Close window", {}), ("Wait", "Wait", {"Value": 1}), ("Quit", "Quit", {})]
    test, _ = await run(make_cfg, build(tmp_path / "a.xlsx", site, "second.html", rows))
    assert test.status == "PASSED" and len(test.steps) == 5 and test.skipped == 0
