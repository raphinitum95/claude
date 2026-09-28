"""What the Workbook Builder (P10) writes into an API sheet, and the runner reads: JSONPath with "the item where", XPath on an XML (SOAP) response,
``{NAME}`` in the URL / headers / a typed ``REQUEST_BODY``, XML request templates, and the builder's extra check kinds."""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

from regrunner.workbook import Workbook
from regrunner.workbook.api import check
from regrunner.workbook.api_compare import compare_response, ignore_list
from regrunner.workbook.api_paths import (MISSING, filter_literal, is_extended_xpath, is_json_path, json_path_all, json_path_first,
                                          json_path_of, xpath_all, xpath_of)
from regrunner.workbook.api_template import json_to_tree, xml_to_tree
from regrunner.workbook.variables import VariablePool
from tests.site import server as mock_server
from tests.test_api_tests import run

RESPONSE = {"quoteId": "Q-1", "premium": {"total": 148.2}, "odd key": {"x": 1},
            "plans": [{"code": "Basic", "eligible": True, "price": 10}, {"code": "Plus", "eligible": True, "price": 20},
                      {"code": "Max", "eligible": False, "price": 30}]}
SOAP = ('<soap:Envelope xmlns:soap="http://schemas.xmlsoap.org/soap/envelope/"><soap:Body><ns:R xmlns:ns="urn:x"><ns:Status code="0">OK</ns:Status>'
        '<ns:Plans><ns:Plan id="A"><ns:Code>Basic</ns:Code></ns:Plan><ns:Plan id="B"><ns:Code>Max</ns:Code></ns:Plan></ns:Plans></ns:R></soap:Body></soap:Envelope>')


# -- JSONPath ------------------------------------------------------------------------------------------------------------------------------------
def test_a_jsonpath_reads_keys_indexes_quoted_keys_and_the_item_where_a_field_has_a_value():
    assert json_path_first(RESPONSE, "$.premium.total") == 148.2 and json_path_first(RESPONSE, "$.plans[2].code") == "Max"
    assert json_path_first(RESPONSE, "$['odd key'].x") == 1 and json_path_first(RESPONSE, "plans[-1].price") == 30
    assert json_path_first(RESPONSE, "$.plans[?(@.code=='Max')].eligible") is False
    assert json_path_first(RESPONSE, "$.plans[?(@.price==20)].code") == "Plus"                  # numbers compare as numbers
    assert json_path_all(RESPONSE, "$.plans[?(@.eligible==true)].code") == ["Basic", "Plus"]
    assert json_path_all(RESPONSE, "$.plans[*].code") == ["Basic", "Plus", "Max"]
    assert json_path_first(RESPONSE, "$.plans[?(@.code=='Nope')].eligible") is MISSING and json_path_first(RESPONSE, "$.a.b") is MISSING
    assert json_path_first(RESPONSE, "$.plans[?(@.code ~ 'x')]") is MISSING                      # an unreadable filter finds nothing


def test_only_paths_in_the_new_form_leave_the_legacy_dotted_reader():
    assert is_json_path("$.a.b") and is_json_path("plans[0].code") and is_json_path("a[?(@.k=='v')].b")
    assert not is_json_path("a.b.[0].c") and not is_json_path("transactionStatus") and not is_json_path("a/b/[0]/c")


def test_the_builder_writes_the_path_of_a_clicked_node_and_quotes_filter_values():
    assert json_path_of(["premium", "total"]) == "$.premium.total" and json_path_of(["plans", 0, "code"]) == "$.plans[0].code"
    assert json_path_of(["odd key", "x"]) == "$['odd key'].x"
    assert [filter_literal(v) for v in ("Basic", "{PLAN}", 5, True)] == ["'Basic'", "'{PLAN}'", "5", "true"]


# -- XPath -----------------------------------------------------------------------------------------------------------------------------------------
def test_an_xpath_starts_at_the_document_or_anywhere_ends_on_an_attribute_and_ignores_namespace_prefixes():
    root = xml_to_tree(SOAP)
    assert xpath_all(root, "/Envelope/Body/R/Status") == ["OK"] and xpath_all(root, "/Envelope/Body/R/Status/@code") == ["0"]
    assert xpath_all(root, "//Plan/Code") == ["Basic", "Max"] and xpath_all(root, "Plans/Plan[2]/Code") == ["Max"]
    assert xpath_all(root, "//Plan[Code='Max']/@id") == ["B"] and xpath_all(root, "/soap:Envelope/soap:Body/ns:R/ns:Status/text()") == ["OK"]
    assert xpath_all(root, "/Body/R") == []                                                      # a document path starts at the document element
    with pytest.raises(ValueError):
        xpath_all(root, "//Plan[position()<2]")
    assert xpath_of([("Envelope", 1, 1), ("Body", 1, 1), ("Plan", 2, 2)], "id") == "/Envelope/Body/Plan[2]/@id"


def test_legacy_output_paths_are_not_mistaken_for_the_new_xpath_forms():
    assert not is_extended_xpath("policyresponses/policydetail/policynumber") and not is_extended_xpath("p[1]/who/last")
    assert is_extended_xpath("/a/b") and is_extended_xpath("//a/@id") and is_extended_xpath("s:Body/x") and is_extended_xpath("a/text()")


def test_an_xpath_reads_a_json_response_the_way_output_paths_see_it():
    assert xpath_all(json_to_tree(RESPONSE), "/root/quoteId") == [] and xpath_all(json_to_tree(RESPONSE), "//premium/total") == ["148.2"]


# -- the builder's check kinds -----------------------------------------------------------------------------------------------------------------------
@pytest.mark.parametrize("kind,expected,actual,ok", [
    ("greater_than", "100", "$148.20", True), ("greater_than", "200", "148.2", False), ("less_than", 5, "4.99", True),
    ("between", "100;200", "148.2", True), ("between", "150 and 100", "148.2", True), ("between", "1;2", "3", False),
    ("between", "100", "148", False), ("greater_than", "1", "n/a", False),
    ("matches", r"^Q-\d+$", "Q-2209", True), ("matches", r"^Q-\d+$", "X-1", False), ("matches", "([", "x", False),
])
def test_greater_than_less_than_between_and_matches_give_a_verdict(kind, expected, actual, ok):
    assert check(kind, expected, actual) is ok


# -- the whole-response check (compare_response) -------------------------------------------------------------------------------------------------------
def test_a_whole_json_response_is_compared_as_data_and_every_difference_is_named_by_its_path():
    expected = json.dumps({"id": 1, "plan": {"code": "B", "price": 12.5}, "items": [{"id": 7, "n": "a"}, {"id": 8, "n": "b"}]})
    same = json.dumps({"plan": {"price": 12.50, "code": "B"}, "items": [{"n": "a", "id": 7}, {"id": 8, "n": "b"}], "id": 1.0})
    assert compare_response(expected, same).same                                              # key order and 12.5 vs 12.50 do not matter
    got = compare_response(expected, json.dumps({"id": 2, "plan": {"code": "B"}, "items": [{"id": 7, "n": "a"}], "extra": True}))
    assert not got.same and got.differences == ["$.id: expected 1, got 2", "$.plan.price: missing (expected 12.5)",
                                                "$.items[1]: missing (the list has 1 items, 2 expected)", "$.extra: not expected (got true)"]
    assert "not valid JSON" in compare_response("{oops", same).problem and not compare_response("{oops", same).same


def test_the_ignore_list_leaves_out_a_field_anywhere_a_jsonpath_with_or_without_the_index_and_an_xpath():
    expected = json.dumps({"quoteId": "Q-1", "createdAt": "t1", "meta": {"requestId": "r1", "createdAt": "t1"}, "items": [{"id": 1, "n": "a"}, {"id": 2, "n": "b"}]})
    actual = json.dumps({"quoteId": "Q-2", "createdAt": "t2", "meta": {"requestId": "r2", "createdAt": "t2"}, "items": [{"id": 3, "n": "a"}, {"id": 4, "n": "b"}]})
    assert compare_response(expected, actual).count == 6
    got = compare_response(expected, actual, ignore="quoteId; createdAt; $.meta.requestId; $.items.id")
    assert got.same and got.ignored == 6
    assert compare_response(expected, actual, ignore=["quoteId", "createdAt", "requestId", "$.items[*].id"]).same
    assert compare_response(expected, actual, ignore="quoteId; createdAt; requestId; $.items[0].id").differences == ["$.items[1].id: expected 2, got 4"]
    xml_expected = '<s:Envelope xmlns:s="urn:s"><s:Body><Quote created="1"><Plan>Basic</Plan><Plan>Max</Plan><Stamp>9</Stamp></Quote></s:Body></s:Envelope>'
    xml_actual = '<Envelope><Body><Quote created="2"><Plan>Basic</Plan><Plan>Plus</Plan><Stamp>10</Stamp></Quote></Body></Envelope>'
    assert compare_response(xml_expected, xml_actual).differences == ['/Envelope/Body/Quote/@created: expected "1", got "2"',
                                                                      '/Envelope/Body/Quote/Plan[2]: expected "Max", got "Plus"',
                                                                      '/Envelope/Body/Quote/Stamp: expected "9", got "10"']
    assert compare_response(xml_expected, xml_actual, ignore="@created; /Envelope/Body/Quote/Plan[2]; //Stamp").same
    assert ignore_list("a; b\nc,, ") == ["a", "b", "c"]


# -- a row the builder wrote --------------------------------------------------------------------------------------------------------------------------
BUILT = ["blnExecute", "TCID", "TC_Name", "WEBSERVICE_METHOD", "WEBSERVICE_URL", "JSON_FORMAT", "REQUEST_BODY", "H_apiKey", "LastName_IN",
         "POLICY_NO", "Holder_EXP", "Holder_OUT", "Count_EXP", "RES_STATUS_CD_EXP"]


def built_workbook(tmp_path: Path, site: str, *, body: str | None = None, key: str = "{SECRET:API_KEY}", url: str = "{DOMAIN}policy/purchase/v2",
                   io_extra: list | None = None) -> Path:
    """An API test the way the Workbook Builder writes one: URL / header / body with {NAME}s, JSONPath outputs, an expected-value check."""
    wb = openpyxl.Workbook()
    ds = wb.active
    ds.title = "DataSheets"
    ds.append(["SheetsToExecute", "blnExecute", "ParameterSheet", "Comments"])
    ds.append(["Purchase", "Y", None, "Purchase API"])
    g = wb.create_sheet("Global")
    g.append(["Parameter", "Value"])
    g.append(["Environment", "UAT"])
    env = wb.create_sheet("_rr_environments")
    for r in (["Variable", "Required", "Secret", "UAT", "PROD"], ["DOMAIN", "Y", "", site, "https://prod.invalid/"]):
        env.append(r)
    io = wb.create_sheet("InputOutput")
    for r in [("Function", "Parameter", "Value"), ("addHeader", "apiKey", "H_apiKey"),
              ("output_json", "$.purchaseResponse.policyResponses[?(@.policyDetail.policyHolder.firstName=='{WHO}')].policyDetail.policyNumber", "POLICY_NO"),
              ("output_json", "$.purchaseResponse.policyResponses[0].policyDetail.policyHolder.lastName", "Holder_OUT"),
              ("compare", "Holder_EXP", "Holder_OUT"), ("compare", "RES_STATUS_CD_EXP", "RES_STATUS_CD_OUT"), *(io_extra or [])]:
        io.append(list(r))
    ws = wb.create_sheet("Purchase")
    ws.append(BUILT)
    values = {"blnExecute": "Y", "TCID": "1", "TC_Name": "Buy", "WEBSERVICE_METHOD": "POST", "WEBSERVICE_URL": url, "JSON_FORMAT": "Y",
              "REQUEST_BODY": body if body is not None else '{"lastName": "{LastName_IN}", "note": "{NOTE}"}', "H_apiKey": key,
              "LastName_IN": 'O\'Neil "Jr"', "Holder_EXP": "{LastName_IN}", "RES_STATUS_CD_EXP": 200}
    ws.append([values.get(c) for c in BUILT])
    wb.save(tmp_path / "built.xlsx")
    return tmp_path / "built.xlsx"


def test_names_in_the_url_header_and_body_come_from_the_row_the_run_and_the_environment_table_and_secrets_are_masked(tmp_path, site, monkeypatch):
    monkeypatch.setenv("RR_SECRET_API_KEY", "KEY-123")
    wb = Workbook(built_workbook(tmp_path, site), seed=1)
    (case,) = wb.discover()
    pool = VariablePool()
    pool.set("NOTE", "from an earlier test")
    rt = wb.api_runtime(case, pool=pool)
    req = rt.request()
    assert (req.method, req.url, req.headers["apiKey"]) == ("POST", site + "policy/purchase/v2", "KEY-123")
    body, missing = rt.request_body(True)
    assert json.loads(body) == {"lastName": 'O\'Neil "Jr"', "note": "from an earlier test"} and missing == []    # a quote cannot break the JSON
    assert rt.mask("sent KEY-123") == "sent ••••••"
    assert rt.pool_reads() == {"NOTE", "WHO"} and rt.pool_sets() == {"POLICY_NO", "HOLDER_OUT"}                  # Needs / Provides
    alone = Workbook(built_workbook(tmp_path, site), seed=1)
    assert alone.api_runtime(alone.discover()[0]).request_body(True)[1] == ["NOTE"]                             # nothing gives {NOTE}


def test_a_url_name_nothing_gives_a_value_for_is_said_instead_of_sending_the_braces(tmp_path, site):
    wb = Workbook(built_workbook(tmp_path, site, url="{NOWHERE}/policy"), seed=1)
    req = wb.api_runtime(wb.discover()[0]).request()
    assert req.url == "" and "{NOWHERE}" in req.url_problem


@pytest.mark.browser
def test_a_builder_written_api_test_sends_its_typed_body_reads_the_item_where_and_hands_the_value_to_the_run(tmp_path, site):
    path = built_workbook(tmp_path, site, body='{"lastName": "{LastName_IN}"}', io_extra=[("between", "Count_EXP", "ELAPSEDTIME_OUT")])
    wb = openpyxl.load_workbook(path)
    wb["Purchase"]["M2"] = "0;60000"                                                             # Count_EXP: the elapsed time is between 0 and 60 s
    wb.save(path)
    proc, results, run_dir = run(tmp_path, path, env={"RR_SECRET_API_KEY": "KEY-123"})
    (test,) = results["tests"]
    assert test["status"] == "PASSED", proc.stdout + proc.stderr
    assert "POLICY_NO" not in {v["name"] for v in test["variables"]}                            # nothing gives {WHO}: that path is not read
    assert mock_server.API_SEEN["body"] == {"lastName": 'O\'Neil "Jr"'} and mock_server.API_SEEN["headers"]["apikey"] == "KEY-123"
    checks = {s["name"]: s for s in test["steps"][1:]}
    assert checks["Check Holder_OUT"]["status"] == "PASSED" and checks["Check Holder_OUT"]["expected"] == 'O\'Neil "Jr"'       # {LastName_IN}
    assert checks["Check RES_STATUS_CD_OUT"]["status"] == "PASSED" and checks["Check ELAPSEDTIME_OUT"]["status"] == "PASSED"
    notes = " ".join(test["steps"][0]["notes"])
    assert "body: the REQUEST_BODY cell" in notes and "nothing gives a value for {WHO}" in notes
    for name in ("results.json", "events.jsonl", "report.html"):
        assert "KEY-123" not in (run_dir / name).read_text(errors="ignore")
    assert "KEY-123" not in (run_dir / "tests" / "Purchase" / "request.txt").read_text()


@pytest.mark.browser
def test_the_item_where_reads_the_policy_of_the_named_holder_and_a_ui_test_of_the_run_can_use_it(tmp_path, site):
    path = built_workbook(tmp_path, site, body='{"lastName": "Smith"}')
    wb = openpyxl.load_workbook(path)
    wb["Purchase"]["K2"] = "Smith"                                                                # Holder_EXP
    wb["_rr_environments"].append(["WHO", "", "", "Duplicate", ""])                               # the item where firstName = {WHO}
    wb.save(path)
    proc, results, run_dir = run(tmp_path, path, env={"RR_SECRET_API_KEY": "KEY-123"})
    (test,) = results["tests"]
    assert test["status"] == "PASSED", proc.stdout + proc.stderr
    assert {v["name"]: v["value"] for v in test["variables"]}["POLICY_NO"] == "P-778"            # the second policy: firstName == Duplicate


# -- XML ------------------------------------------------------------------------------------------------------------------------------------------------
XML_COLUMNS = ["blnExecute", "TC_Name", "WEBSERVICE_URL", "WEBSERVICE_METHOD", "JSON_FORMAT", "XML_LOCATION", "XML_REQUESTFILE", "DT_apiKey",
               "LastName_IN", "PLAN", "Status_OUT", "Code_OUT", "Policy_OUT", "Holder_OUT", "Holder_EXP", "Code_EXP"]
TEMPLATE = ('<soap:Envelope xmlns:soap="http://schemas.xmlsoap.org/soap/envelope/"><soap:Body><Buy><LastName>#LAST#</LastName></Buy>'
            '</soap:Body></soap:Envelope>')


@pytest.mark.browser
def test_an_xml_test_fills_its_template_escaped_and_reads_the_soap_answer_with_xpaths(tmp_path, site):
    wb = openpyxl.Workbook()
    ds = wb.active
    ds.title = "DataSheets"
    ds.append(["SheetsToExecute", "blnExecute", "ParameterSheet", "Comments"])
    ds.append(["Soap", "Y", None, "SOAP"])
    wb.create_sheet("Global").append(["Environment", "UAT"])
    io = wb.create_sheet("InputOutput")
    for r in [("Function", "Parameter", "Value"), ("addHeader", "apiKey", "DT_apiKey"), ("Replace", "#LAST#", "LastName_IN"),
              ("Output", "/Envelope/Body/PolicyResponse/Status", "Status_OUT"), ("Output", "//Status/@code", "Code_OUT"),
              ("Output", "//Policy[Plan='{PLAN}']/@id", "Policy_OUT"), ("Output", "Policies/Policy[1]/LastName", "Holder_OUT"),
              ("compare", "Holder_EXP", "Holder_OUT"), ("compare", "Code_EXP", "Code_OUT")]:
        io.append(list(r))
    ws = wb.create_sheet("Soap")
    ws.append(XML_COLUMNS)
    row = {"blnExecute": "Y", "TC_Name": "soap", "WEBSERVICE_URL": site + "policy/xml/v1", "WEBSERVICE_METHOD": "POST", "JSON_FORMAT": "N",
           "XML_REQUESTFILE": "buy.xml", "DT_apiKey": "KEY-123", "LastName_IN": "Smith & Sons", "PLAN": "Max", "Holder_EXP": "Smith & Sons", "Code_EXP": "0"}
    ws.append([row.get(c) for c in XML_COLUMNS])
    wb.save(tmp_path / "soap.xlsx")
    (tmp_path / "templates").mkdir()
    (tmp_path / "templates" / "buy.xml").write_text(TEMPLATE, encoding="utf-8")
    proc, results, run_dir = run(tmp_path, tmp_path / "soap.xlsx")
    (test,) = results["tests"]
    assert test["status"] == "PASSED", proc.stdout + proc.stderr + json.dumps(test["steps"], indent=1)
    assert "<LastName>Smith &amp; Sons</LastName>" in mock_server.API_SEEN["body"]               # escaped: the XML stays well-formed
    got = {v["name"]: v["value"] for v in test["variables"]}
    assert got["Status_OUT"] == "Success" and got["Policy_OUT"] == "P-902" and got["Holder_OUT"] == "Smith & Sons"
    assert (run_dir / "tests" / "Soap" / "response.xml").is_file()
    assert "Replace: 1 placeholders filled" in " ".join(test["steps"][0]["notes"])
