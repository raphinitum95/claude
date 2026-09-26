"""End-to-end tests: real Chromium, local mock site, synthetic workbook with the real formula tricks."""
from __future__ import annotations

import json
from datetime import datetime, timedelta
from pathlib import Path

import pytest

from regrunner.engine.runner import RunOptions, execute
from regrunner.events import EventBus
from tests.workbook_factory import build_workbook

pytestmark = pytest.mark.browser


def steps_by_row(test) -> dict[int, object]:
    return {s.row: s for s in test.steps}


async def run(cfg, wb_path, **opts):  # noqa: D401
    bus = EventBus()
    events: list[dict] = []
    bus.subscribe(events.append)
    result = await execute(RunOptions(workbook=wb_path, **opts), cfg, bus)
    return result, events


async def test_reference_flow_passes_end_to_end(site, make_cfg, tmp_path):
    wb = tmp_path / "flow.xlsx"
    rows = build_workbook(wb, base_url=site)["Flow"]
    cfg = make_cfg()
    result, events = await run(cfg, wb, tests=["Flow"], seed=3)

    test = result.tests[0]
    failed = [(s.row, s.name, s.error, s.expected, s.actual) for s in test.steps if s.status != "PASSED"]
    assert test.status == "PASSED", failed
    by_row = steps_by_row(test)

    # token substitution + status gating + conditional rows
    assert rows["cost_yes"] in by_row and rows["cost_no"] not in by_row
    assert rows["banner_out"] in by_row                              # gated on the Exist step's PASSED status
    assert by_row[rows["dest_pick"]].locator.startswith("xpath=")   # formula-built XPath resolved (Singapore)

    # cookie reached the server; captured values chain through =S<row> and TEXT()
    assert by_row[rows["bypass"]].actual == "bypass:TEST-UAT-TOKEN"
    assert by_row[rows["premium_capture"]].actual == "$150.00"        # settled value, not the intermediate $123.45
    assert by_row[rows["premium_compare"]].expected == "$150.00"
    tomorrow = (datetime.now() + timedelta(days=1)).strftime("%m/%d/%Y")
    assert by_row[rows["dep_shown"]].expected == tomorrow == by_row[rows["dep_shown"]].actual

    # legacy-compat behaviours
    assert "value is NOT verified" in " ".join(by_row[rows["legacy_noop"]].notes)
    assert by_row[rows["nbsp"]].actual == "Hello World"
    assert by_row[rows["hidden"]].actual == ""
    assert by_row[rows["ignored_missing"]].status == "PASSED" and by_row[rows["ignored_missing"]].ignored_error
    assert any("zoom" in n for n in by_row[rows["zoom"]].notes)
    assert by_row[rows["pay_value"]].actual == "12|2030"
    assert by_row[rows["disabled_state"]].actual == "clicked"        # js_click waited for the disabled control
    assert by_row[rows["second"]].actual == "Second Window"

    # smart waits did not sleep the legacy seconds
    assert by_row[rows["wait_result"]].duration_ms < 4000

    # evidence: a screenshot per step, a named screenshot, results JSON, events JSONL
    run_dir = tmp_path / "runs" / result.run_id
    no_page = {rows["wait_first"], rows["quit"]}                      # nothing to photograph before Open / after Quit
    missing = [s.name for s in test.steps if s.row not in no_page and not (s.screenshot and (run_dir / s.screenshot).is_file())]
    assert not missing, missing
    assert list((run_dir / "tests" / "Flow" / "named").glob("*ConfirmationPage.png"))
    saved = json.loads((run_dir / "results.json").read_text())
    assert saved["summary"]["tests"] == 1 and saved["status"] == "PASSED"
    lines = [json.loads(l) for l in (run_dir / "events.jsonl").read_text().splitlines()]
    kinds = [e["type"] for e in lines]
    assert kinds[0] == "run_started" and kinds[-1] == "run_finished"
    assert {"test_started", "step_passed", "screenshot_saved", "run_progress", "test_finished"} <= set(kinds)

    # review-only capture: never affects status, tagged with test + step
    review = test.review
    assert any(i["type"] == "console_error" and "deliberate console error" in i["message"] for i in review)
    assert any(i["type"] == "console_error" and i["kind"] == "pageerror" for i in review)
    assert any(i["type"] == "network_error" and i["status"] == 404 for i in review)
    assert all(i["test"] == "Flow" and i["step"] >= 1 for i in review)

    # progress reached 100%
    last = [e for e in events if e["type"] == "run_progress"][-1]
    assert last["percent"] == 100.0 and last["tests"]["Flow"]["percent"] == 100.0


async def test_failures_do_not_stop_the_test_and_are_reported(site, make_cfg, tmp_path):
    wb = tmp_path / "fail.xlsx"

    def extra(sheet, rows):
        # insert before Quit by re-adding after; Quit is last so order is fine for this assertion.
        rows["bad_text"] = sheet.add("Output", "Wrong expectation", FindBy="xpath", FindBy_Value="//span[@id='bypass']",
                                     Index=0, Output_Property="innertext", Expected_Value="nope", Exact_Match="Y")
        rows["missing"] = sheet.add("Click", "Required element missing", FindBy="xpath",
                                    FindBy_Value="//button[@id='required-but-absent']", Index=0, Timeout=1)
        rows["after"] = sheet.add("Exist", "Still runs after failures", FindBy="xpath",
                                  FindBy_Value="//span[@id='bypass']", Index=0)

    rows = build_workbook(wb, base_url=site, extra_steps=extra)["Flow"]
    result, _ = await run(make_cfg(), wb, tests=["Flow"])
    test = result.tests[0]
    by_row = steps_by_row(test)
    assert test.status == "FAILED" and test.failed == 2
    assert by_row[rows["bad_text"]].error == "Comparison Failed"
    assert by_row[rows["bad_text"]].expected == "nope" and by_row[rows["bad_text"]].actual.startswith("bypass:")
    assert by_row[rows["missing"]].error == "Object was not found"
    assert by_row[rows["after"]].status == "PASSED"                  # legacy: failures do not stop the test
    assert by_row[rows["bad_text"]].screenshot                       # failed steps keep their screenshot


async def test_quit_break_on_failure_stops_after_first_failure(site, make_cfg, tmp_path):
    import openpyxl
    wb = tmp_path / "qbf.xlsx"

    def extra(sheet, rows):
        rows["missing"] = sheet.add("Click", "Required element missing", FindBy="xpath",
                                    FindBy_Value="//button[@id='required-but-absent']", Index=0, Timeout=1)
        rows["never"] = sheet.add("Exist", "Must not run", FindBy="xpath", FindBy_Value="//span[@id='bypass']", Index=0)

    rows = build_workbook(wb, base_url=site, extra_steps=extra)["Flow"]
    book = openpyxl.load_workbook(wb)
    for r in book["Global"].iter_rows(min_row=2):
        if r[0].value == "QuitBreakOnFailure":
            r[1].value = "Y"
    book.save(wb)
    result, _ = await run(make_cfg(), wb, tests=["Flow"])
    test = result.tests[0]
    assert test.status == "FAILED" and rows["never"] not in steps_by_row(test)
    assert test.skipped >= 1 and "QuitBreakOnFailure" in test.error


async def test_tests_run_concurrently_and_in_isolation(site, make_cfg, tmp_path):
    wb = tmp_path / "multi.xlsx"
    build_workbook(wb, flows=["FlowA", "FlowB", "FlowC"], base_url=site)
    result, events = await run(make_cfg(**{"runner.workers": 3}), wb, seed=1)
    assert [t.status for t in result.tests] == ["PASSED"] * 3
    assert result.workers == 3
    workers_used = {e["worker"] for e in events if e["type"] == "test_started"}
    assert workers_used == {1, 2, 3}
    spans = sorted((datetime.fromisoformat(t.started_at), datetime.fromisoformat(t.ended_at)) for t in result.tests)
    assert spans[1][0] < spans[0][1] and spans[2][0] < spans[0][1]   # all three overlapped in time
    # isolation: each test wrote its own screenshots and got its own random FirstName-independent state
    dirs = {p.name for p in (tmp_path / "runs" / result.run_id / "tests").iterdir()}
    assert dirs == {"FlowA", "FlowB", "FlowC"}


async def test_selection_by_tag_and_by_name(site, make_cfg, tmp_path):
    wb = tmp_path / "sel.xlsx"
    build_workbook(wb, flows=["FlowA", "FlowB"], base_url=site, enabled=["FlowA"])
    cfg = make_cfg()
    result, _ = await run(cfg, wb)                                    # default: workbook Y/N flags
    assert [t.id for t in result.tests] == ["FlowA"]
    result, _ = await run(cfg, wb, tests=["FlowB"])                   # explicit selection overrides the flag
    assert [t.id for t in result.tests] == ["FlowB"]
    result, _ = await run(cfg, wb, tags=["full"])                     # Tags column in DataSheets
    assert [t.id for t in result.tests] == ["FlowB"]


async def test_generic_locators_replace_xpath_with_fallback(site, make_cfg, tmp_path):
    wb = tmp_path / "loc.xlsx"
    rows = build_workbook(wb, base_url=site)["Flow"]
    (tmp_path / "selectors.yaml").write_text(
        "version: 1\nlocators:\n"
        "  \"//input[@name='tripDepartureDate']\":\n    use: ['css=input[name=\"tripDepartureDate\"]']\n"
        "  \"//input[@name='tripReturnDate']\":\n    use: ['css=input[name=\"this-name-was-renamed\"]']\n"   # stale
        "  \"(//input[@name='insuredAge'])[2]\":\n    use: ['css=input[name=\"insuredAge\"]']\n    index: 1\n")
    result, _ = await run(make_cfg(), wb, tests=["Flow"])
    test = result.tests[0]
    by_row = steps_by_row(test)
    assert test.status == "PASSED"
    assert by_row[rows["dep"]].locator == 'css=input[name="tripDepartureDate"]' and not by_row[rows["dep"]].fallback_used
    assert by_row[rows["ret"]].fallback_used and by_row[rows["ret"]].locator.startswith("xpath=")   # stale css -> XPath
    assert by_row[rows["age2"]].locator == 'css=input[name="insuredAge"]'                            # index override
    assert any(i["type"] == "review_item" and i["category"] == "selector_fallback" for i in test.review)


async def test_locator_column_in_the_sheet_is_used_first(site, make_cfg, tmp_path):
    wb = tmp_path / "col.xlsx"

    def extra(sheet, rows):
        sheet.ws.cell(rows["dep"], sheet.columns.index("Locator") + 1,
                      'css=input[name="tripDepartureDate"] || xpath=//nothing')

    rows = build_workbook(wb, base_url=site, extra_columns=["Locator"], extra_steps=extra)["Flow"]
    result, _ = await run(make_cfg(), wb, tests=["Flow"])
    by_row = steps_by_row(result.tests[0])
    assert result.tests[0].status == "PASSED"
    assert by_row[rows["dep"]].locator == 'css=input[name="tripDepartureDate"]'


async def test_missing_captcha_token_is_flagged_not_fatal(site, make_cfg, tmp_path):
    wb = tmp_path / "nocookie.xlsx"
    rows = build_workbook(wb, base_url=site, expect_bypass="none")["Flow"]
    cfg = make_cfg()
    cfg.captcha_bypass.values = {}
    result, _ = await run(cfg, wb, tests=["Flow"])
    test = result.tests[0]
    assert test.status == "PASSED"                                    # 'bypass:none' matches: no cookie was sent
    assert any(i["category"] == "captcha_bypass" for i in test.review)
    assert any("recaptchaBypassToken" in w for w in result.warnings)


async def test_html_and_pdf_report_are_self_contained(site, make_cfg, tmp_path):
    wb = tmp_path / "rep.xlsx"

    def extra(sheet, rows):
        rows["bad"] = sheet.add("Output", "Wrong expectation", FindBy="xpath", FindBy_Value="//p[@id='nb']", Index=0,
                                Output_Property="innertext", Expected_Value="Hello\u00a0World", Exact_Match="Y")

    build_workbook(wb, base_url=site, extra_steps=extra)
    cfg = make_cfg()
    cfg.reports.pdf = True
    result, _ = await run(cfg, wb, tests=["Flow"])
    run_dir = tmp_path / "runs" / result.run_id
    report = (run_dir / "report.html").read_text(encoding="utf-8")
    assert "data:image/jpeg;base64," in report                       # screenshots embedded
    assert "http://" not in report.replace(site, "") or "cdn" not in report   # no external assets
    assert "Things to review" in report and "deliberate console error" in report
    assert "⍽" in report                                             # NBSP made visible in the failed comparison
    assert (run_dir / "report.pdf").read_bytes()[:4] == b"%PDF"
    assert result.artifacts["report_html"] == "report.html" and result.artifacts["report_pdf"] == "report.pdf"


async def test_harvest_learns_stable_locators_from_the_live_dom_and_they_are_then_used(site, make_cfg, tmp_path):
    from regrunner.selectors.harvest import apply_suggestions
    from regrunner.selectors.resolve import SelectorMap
    wb = tmp_path / "harvest.xlsx"
    rows = build_workbook(wb, base_url=site)["Flow"]
    cfg = make_cfg()
    result, _ = await run(cfg, wb, tests=["Flow"], harvest=True)
    run_dir = tmp_path / "runs" / result.run_id
    data = json.loads((run_dir / "selector_suggestions.json").read_text())
    weak = "//button[@type='SUBMIT']"                              # class/type-anchored XPath; the button has an id
    assert data["suggestions"][weak]["use"] == ['css=[id="go"]'] and data["suggestions"][weak]["score"] >= 90
    # already-strong XPaths (name/id anchored) are not "improved"
    assert "//input[@name='tripDepartureDate']" not in data["suggestions"]

    smap = SelectorMap.load(tmp_path / "selectors.yaml")
    assert apply_suggestions(run_dir / "selector_suggestions.json", smap, min_score=85) >= 1
    smap.save()
    second, _ = await run(cfg, wb, tests=["Flow"])
    by_row = steps_by_row(second.tests[0])
    assert second.tests[0].status == "PASSED"
    assert by_row[rows["submit"]].locator == 'css=[id="go"]' and not by_row[rows["submit"]].fallback_used


def test_harvest_refuses_ambiguous_observations():
    from regrunner.selectors.harvest import HarvestStore
    store = HarvestStore()
    good = {"use": 'css=[id="a"]', "kind": "id", "score": 92}
    store.observations = {
        "//x[1]": [{"index": 0, "step": "s", "legacy_score": 40, "candidate": good},
                   {"index": 1, "step": "t", "legacy_score": 40, "candidate": {**good, "use": 'css=[id="b"]'}}],
        "//y": [{"index": 0, "step": "s", "legacy_score": 40, "candidate": good},
                {"index": 0, "step": "t", "legacy_score": 40, "candidate": None}],
        "//z": [{"index": 0, "step": "s", "legacy_score": 40, "candidate": good},
                {"index": 0, "step": "t", "legacy_score": 40, "candidate": good}],
    }
    safe, skipped = store.suggestions()
    assert list(safe) == ["//z"] and safe["//z"]["uses"] == 2
    assert "different elements" in skipped["//x[1]"] and "no unique" in skipped["//y"]


async def test_optional_elements_are_absent_only_once_the_page_is_quiet(site, make_cfg, tmp_path):
    """A consent banner that arrives via a slow request must still be found; a truly absent one must be quick."""
    wb = tmp_path / "opt.xlsx"

    def extra(sheet, rows):
        rows["late"] = sheet.add("Exist", "Late banner (optional)", FindBy="xpath", FindBy_Value="//div[@id='late-banner']",
                                 Index=0, Ignore_not_existing_object="Y")
        rows["absent"] = sheet.add("Exist", "Absent banner (optional)", FindBy="xpath",
                                   FindBy_Value="//div[@id='never-shown']", Index=0, Ignore_not_existing_object="Y")
        rows["absent_again"] = sheet.add("Click", "Absent again", FindBy="xpath", FindBy_Value="//div[@id='never-shown']",
                                         Index=0, Ignore_not_existing_object="Y")

    rows = build_workbook(wb, base_url=site, flows=["Flow"], extra_steps=extra)["Flow"]
    cfg = make_cfg(**{"timeouts.optional_s": 6})              # the legacy-style full wait would cost 6 s each
    result, _ = await run(cfg, wb, tests=["Flow"])
    by_row = steps_by_row(result.tests[0])
    assert result.tests[0].status == "PASSED"
    late = by_row[rows["late"]]
    assert late.locator.startswith("xpath=") and not late.ignored_error       # found, not "ignored"
    assert by_row[rows["absent"]].ignored_error and by_row[rows["absent"]].duration_ms < 2000
    assert by_row[rows["absent_again"]].duration_ms < 1500

    slow = make_cfg(**{"timeouts.optional_s": 2, "timeouts.optional_mode": "full"})
    result, _ = await run(slow, wb, tests=["Flow"])
    assert steps_by_row(result.tests[0])[rows["absent"]].duration_ms >= 1900   # 'full' keeps the exact legacy wait


async def test_missing_frame_is_a_silent_pass_like_the_legacy_runner_and_flagged(site, make_cfg, tmp_path):
    wb = tmp_path / "frame.xlsx"

    def extra(sheet, rows):
        rows["no_frame"] = sheet.add("SWITCHTOFRAME", "Switch to a frame that does not exist", Value="noSuchFrame")
        rows["after"] = sheet.add("Exist", "Still finds top-level elements", FindBy="xpath", FindBy_Value="//span[@id='bypass']", Index=0)

    rows = build_workbook(wb, base_url=site, extra_steps=extra)["Flow"]
    result, _ = await run(make_cfg(), wb, tests=["Flow"])
    by_row = steps_by_row(result.tests[0])
    step = by_row[rows["no_frame"]]
    assert result.tests[0].status == "PASSED" and step.status == "PASSED" and step.duration_ms < 2500
    assert any("continuing in the current document" in n for n in step.notes)
    assert any(i.get("category") == "missing_frame" for i in result.tests[0].review)
    assert by_row[rows["after"]].status == "PASSED"

    strict = make_cfg(**{"behaviour.missing_frame": "fail", "timeouts.element_s": 2})
    result, _ = await run(strict, wb, tests=["Flow"])
    failed = steps_by_row(result.tests[0])[rows["no_frame"]]
    assert failed.status == "FAILED" and "No such frame" in failed.error


async def test_prod_runs_are_refused_unless_explicitly_allowed(site, make_cfg, tmp_path):
    from regrunner.engine.runner import SelectionError
    wb = tmp_path / "prod.xlsx"
    build_workbook(wb, base_url=site)
    with pytest.raises(SelectionError, match="Refusing to run against PROD"):
        await run(make_cfg(), wb, tests=["Flow"], environment="PROD")
    with pytest.raises(SelectionError, match="Refusing to run against PROD"):
        await run(make_cfg(), wb, tests=["Flow"], environment="prod")        # case-insensitive
    assert not list((tmp_path / "runs").glob("*"))                            # refused before anything was created
