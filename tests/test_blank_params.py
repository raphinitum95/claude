"""A step that names a parameter whose Params cell is empty: the legacy runner typed the parameter's *name* into the page (DT_Policy_Out into the
policy number field) and reported a pass.  Now the step fails (or, when a person can answer, asks for the value), and the workbook says so before the run."""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

from regrunner.config import Config
from regrunner.engine.ask import write_answer
from regrunner.engine.runner import RunOptions, execute
from regrunner.events import EventBus
from regrunner.lint import lint
from regrunner.workbook import Workbook
from tests.workbook_factory import Sheet

VALUE = "P-4242"


def build(path: Path, site: str, *, filled: str | None = None, set_earlier: bool = False, ignore: bool = False) -> Path:
    """Sheet Buy writes DT_Policy_Out (a step whose Output_Value names it); sheet View types it into the code field and checks the echo."""
    wb = openpyxl.Workbook()
    ds = wb.active
    ds.title = "DataSheets"
    ds.append(["SheetsToExecute", "blnExecute", "ParameterSheet", "Comments"])
    ds.append(["Buy", "N", "Params_1", "producer"])
    ds.append(["View", "Y", "Params_1", "consumer"])
    g = wb.create_sheet("Global")
    g.append(["Parameter", "Value"])
    g.append(["Environment", "UAT"])
    ps = wb.create_sheet("Params_1")
    ps.append(["blnExecute", "DT_URL", "DT_Policy_Out"])
    ps.append(["Y", f"{site}ask.html", filled])
    buy = Sheet(wb.create_sheet("Buy"), "Y")
    buy.add("Open", "Open", Page="chrome", Value="DT_URL")
    buy.add("Output", "Get Policy Number", FindBy="xpath", FindBy_Value="//h1", Index=0, Output_Property="innertext", Output_Value="DT_Policy_Out")
    view = Sheet(wb.create_sheet("View"), "Y")
    view.add("Open", "Open", Page="chrome", Value="DT_URL")
    if set_earlier:                                       # an earlier step of the same test fills the parameter in
        view.add("Output", "Get Policy Number", FindBy="xpath", FindBy_Value="//h1", Index=0, Output_Property="innertext", Output_Value="DT_Policy_Out")
    view.add("Set", "Enter Policy Number", FindBy="xpath", FindBy_Value="//input[@id='code']", Index=0, Value="DT_Policy_Out",
             **({"Ignore_not_existing_object": "Y"} if ignore else {}))
    view.add("Output", "What the field holds", FindBy="xpath", FindBy_Value="//span[@id='echo']", Index=0, Output_Property="innertext")
    view.add("Output", "Check Policy Number", FindBy="xpath", FindBy_Value="//span[@id='echo']", Index=0, Output_Property="innertext",
             Expected_Value="DT_Policy_Out", Exact_Match="Y")
    view.add("Quit", "Quit")
    wb.save(path)
    return path


async def run(cfg, wb: Path, ask: str = "off", listen=None):
    bus = EventBus()
    events: list[dict] = []
    bus.subscribe(events.append)
    if listen:
        bus.subscribe(listen(cfg))
    result = await execute(RunOptions(workbook=wb, tests=["View"], seed=1, extra={"ask": ask}), cfg, bus)
    return result.tests[0], events


def by_name(test) -> dict:
    return {s.name: s for s in test.steps}


def answers_with(value: str):
    def factory(cfg):
        def listen(event):
            if event["type"] == "user_input_needed" and event.get("kind", "text") == "text":
                write_answer(cfg.path(cfg.runs_dir) / event["run_id"], event["ask"], value)
        return listen
    return factory


# -- before the run ----------------------------------------------------------------------------------------------------------------
def test_the_workbook_says_which_parameter_is_empty_and_who_fills_it(site, tmp_path):
    wb = Workbook(build(tmp_path / "a.xlsx", site), seed=1)
    cases = {c.id: c for c in wb.discover()}
    runtime = wb.runtime(cases["View"])
    runtime.plan()
    (need,) = runtime.needs
    assert need["param"] == "DT_Policy_Out" and need["cell"] == "Params_1!C2" and need["rows"] == [3, 5]   # the Set and the Output check
    assert need["set_by"] == ["Buy"] and need["steps"][0] == "Enter Policy Number"
    said = [f.message for f in lint(wb) if f.test == "View" and "parameter DT_Policy_Out is empty" in f.message]
    assert said and "Params_1!C2" in said[0] and "Buy does, when it runs first" in said[0]
    from regrunner import insight
    (test,) = [t for t in insight.list_tests(wb, Config()) if t["id"] == "View"]
    assert test["needs"][0]["param"] == "DT_Policy_Out"


def test_nothing_is_reported_when_the_cell_has_a_value_or_an_earlier_step_sets_it(site, tmp_path):
    for kwargs in ({"filled": VALUE}, {"set_earlier": True}):
        wb = Workbook(build(tmp_path / "b.xlsx", site, **kwargs), seed=1)
        case = next(c for c in wb.discover() if c.id == "View")
        runtime = wb.runtime(case)
        runtime.plan()
        assert runtime.needs == [], kwargs
        assert not [f for f in lint(wb) if "is empty" in f.message], kwargs


def test_a_value_another_test_wrote_is_seen_by_the_next_test_with_that_scenario_row(site, tmp_path):
    wb = Workbook(build(tmp_path / "a.xlsx", site), seed=1)
    case = next(c for c in wb.discover() if c.id == "View")
    col = wb.data.sheet("Params_1").headers["DT_POLICY_OUT"]
    runtime = wb.runtime(case, shared={("Params_1", 2, col): "POL-9", ("Params_1", 3, col): "another scenario's"})   # only row 2 is this test's
    step = next(s for s in (runtime.prepare_row(r) for r in range(2, runtime.total_rows + 1)) if s is not None and s.name == "Enter Policy Number")
    assert step.text("VALUE") == "POL-9" and step.blank_params == []


# -- during the run ----------------------------------------------------------------------------------------------------------------
@pytest.mark.browser
@pytest.mark.parametrize("ignore", [False, True])
async def test_an_empty_parameter_fails_the_step_instead_of_typing_its_name(site, make_cfg, tmp_path, ignore):
    test, _ = await run(make_cfg(), build(tmp_path / "a.xlsx", site, ignore=ignore))                  # nobody to ask
    steps = by_name(test)
    step = steps["Enter Policy Number"]
    assert step.status == "FAILED" and not step.ignored_error                                        # Ignore_not_existing_object does not cover this
    assert "Parameter DT_Policy_Out is empty" in step.error and "Params_1!C2" in step.error and "Buy fills it in" in step.error
    assert "start the run from the web UI or a terminal" in step.error
    assert steps["What the field holds"].actual == ""                                                # the parameter's name was never typed


@pytest.mark.browser
async def test_a_person_can_supply_the_value_and_the_step_and_the_check_use_it(site, make_cfg, tmp_path):
    test, events = await run(make_cfg(), build(tmp_path / "a.xlsx", site), ask="ui", listen=answers_with(VALUE))
    assert test.status == "PASSED", [(s.name, s.error) for s in test.steps if s.status != "PASSED"]
    steps = by_name(test)
    assert steps["What the field holds"].actual == VALUE                                            # typed into the field
    assert steps["Check Policy Number"].expected == VALUE                                           # the later use of the parameter sees it as well
    (asked,) = [e for e in events if e["type"] == "user_input_needed"]
    assert "DT_Policy_Out is empty for View" in asked["question"] and "Params_1!C2" in asked["question"] and asked["kind"] == "text"
    assert VALUE not in json.dumps([e for e in events if e["type"].startswith("user_input")])           # the answer travels in a file, not in the question's events
    assert [e["by_hand"] for e in events if e["type"] == "variable_set"] == [True]                    # ...and is listed as a value set by hand once it is used


@pytest.mark.browser
async def test_nobody_answering_fails_the_step_with_the_reason(site, make_cfg, tmp_path):
    test, _ = await run(make_cfg(**{"ask.timeout_s": 2}), build(tmp_path / "a.xlsx", site), ask="ui")
    step = by_name(test)["Enter Policy Number"]
    assert step.status == "FAILED" and "Parameter DT_Policy_Out is empty" in step.error and "Nobody answered within 2 s" in step.error


@pytest.mark.browser
async def test_a_parameter_an_earlier_step_of_the_test_fills_in_is_not_a_problem(site, make_cfg, tmp_path):
    test, events = await run(make_cfg(), build(tmp_path / "a.xlsx", site, set_earlier=True))
    assert test.status == "PASSED", [(s.name, s.error) for s in test.steps if s.status != "PASSED"]
    assert by_name(test)["What the field holds"].actual == "Google Authenticator" and not [e for e in events if e["type"] == "user_input_needed"]


def test_steps_that_ask_for_an_empty_value_themselves_are_not_reported_as_empty_parameters(site, tmp_path):
    """GET_GOOGLE_TOKEN with an empty DT_Key asks for the 6-digit code instead: the generic 'parameter is empty' question must not get there first."""
    from tests.test_ask_user import build_ask_workbook
    wb = Workbook(build_ask_workbook(tmp_path / "g.xlsx", site, "google"), seed=1)
    case = next(c for c in wb.discover() if c.id == "Flow")
    runtime = wb.runtime(case)
    runtime.plan()
    assert runtime.needs == []
    assert not [f for f in lint(wb) if "is empty" in f.message]
    step = next(s for s in (runtime.prepare_row(r) for r in range(2, runtime.total_rows + 1)) if s is not None and s.method == "GET_GOOGLE_TOKEN")
    assert step.blank_params                                                       # the cell IS empty: the step deals with it
