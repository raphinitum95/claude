"""The Workbook Builder's model (workbook/builder.py), its ops and the problems engine (lint.builder_problems). No browser."""
from __future__ import annotations

import json
import shutil
import time
from pathlib import Path

import openpyxl
import pytest

from regrunner.workbook.builder import (BLOCK_COLUMN, BuildDocument, BuildError, auto_name, build_model, describe_target, humanize,
                                        make_token, new_workbook, parse_condition)
from regrunner.workbook.model import Workbook
from regrunner.workbook.writer import WorkbookEditor, backups_dir, draft_path
from tests.workbook_factory import build_steps_workbook, build_workbook

ROOT = Path(__file__).resolve().parent.parent
REAL = sorted(p for p in (ROOT / "workbooks").glob("*.xlsx") if not p.name.startswith("~$"))


def model_of(path: Path, **kw) -> dict:
    return build_model(WorkbookEditor.open(path), **kw)


def find_test(model: dict, name: str) -> dict:
    return next(t for t in model["tests"] if t["id"] == name)


def step_at(test: dict, row: int) -> dict:
    return next(s for s in test["steps"] if s["row"] == row)


@pytest.fixture
def mock(tmp_path) -> tuple[Path, dict]:
    path = tmp_path / "mock.xlsx"
    rows = build_workbook(path, flows=["FlowA", "FlowB"])
    return path, rows["FlowA"] if "FlowA" in rows else rows


def small(tmp_path, fill, params=None, name="small.xlsx") -> Path:
    path = tmp_path / name
    build_steps_workbook(path, {"Buy": fill}, params={"Buy": params or {}})
    return path


# ---------------------------------------------------------------------------------------------------
# Reading
# ---------------------------------------------------------------------------------------------------
def test_section_rows_become_blocks_and_a_new_window_is_its_own_block_that_returns(mock):
    path, _ = mock
    t = find_test(model_of(path), "FlowA")
    assert [(b["title"], b["kind"]) for b in t["blocks"]] == [("Open", "page"), ("Quote", "page"), ("Result", "page"),
                                                              ("New window", "window"), ("Result", "page")]
    assert t["blocks"][3]["returnsTo"] == "Result" and all(not b["stored"] for b in t["blocks"])
    assert [s["title"] for s in t["sections"]] == ["Open", "Quote", "Result"]
    assert [s["n"] for s in t["steps"]] == list(range(1, len(t["steps"]) + 1))


def test_steps_carry_frame_and_window_context_legacy_formula_flags_and_data_flags(mock):
    path, rows = mock
    t = find_test(model_of(path), "FlowA")
    assert step_at(t, rows["trip"])["context"] == "frame: aemFormFrame"
    assert step_at(t, rows["second"])["context"] == "window 2"
    cost = step_at(t, rows["cost_yes"])
    assert cost["kind"] == "legacy" and "formula" in cost["legacy"] and cost["enabled"] is None and cost["condition"]["kind"] == "formula"
    plain = step_at(t, rows["dep"])
    assert plain["legacy"] == "" and plain["enabled"] is True and plain["condition"] == {"kind": "flag", "text": "bln1001"}
    assert {"token": "DEPDATE", "column": "VALUE", "form": "cell"} in plain["uses"]


def test_variables_list_who_sets_and_who_uses_them(mock):
    path, rows = mock
    model = model_of(path)
    premium = next(v for v in model["variables"] if v["key"] == "PREMIUM_OUT")
    assert {"test": "FlowA", "row": rows["premium_capture"], "n": step_at(find_test(model, "FlowA"), rows["premium_capture"])["n"]} in premium["setBy"]
    assert sorted(premium["providedBy"]) == ["FlowA", "FlowB"]
    dep = next(v for v in model["variables"] if v["key"] == "DEPDATE")
    assert dep["label"] == "Dep date" and {u["test"] for u in dep["usedBy"]} == {"FlowA", "FlowB"} and dep["kind"] == "data"
    flag = next(v for v in model["variables"] if v["key"] == "BLN1001")
    assert flag["kind"] == "flag"
    assert "PREMIUM_OUT" in find_test(model, "FlowA")["provides"]


def test_automatic_names_read_as_sentences_with_variable_chips():
    variables = {"FIRST_NAME", "PLAN"}
    assert auto_name("SET", target="first name field", value="FIRST_NAME", variables=variables) == "Type {FIRST_NAME} into first name field"
    assert auto_name("CLICK", target='"Choose" button') == 'Click "Choose" button'
    assert auto_name("OUTPUT", target="price", save_as="PRICE") == "Save price text as {PRICE}"
    assert auto_name("OUTPUT", target="plan", expected="PLAN", match="contains", variables=variables) == "Check plan text contains {PLAN}"
    assert auto_name("CALL_TEST", value="policySearch") == "Run policySearch, then carry on"
    assert auto_name("WAIT_UNTIL", target="spinner", output_property="GONE") == "Wait until spinner is gone"
    assert auto_name("SET", target="password field", value="{SECRET:UAT_PASSWORD}") == "Type {UAT_PASSWORD} into password field"


def test_targets_are_described_in_plain_words_from_the_locator():
    assert describe_target("//button[contains(.,'Get Started')]") == '"Get Started" button'
    assert describe_target("//input[@name='primaryTravelerFirstName']") == "primary traveler first name field"
    assert describe_target("#pay-now") == "pay now"
    assert describe_target("PLAN_XPATH", variables={"PLAN_XPATH"}) == "{PLAN_XPATH}"
    assert describe_target("//div", name="checkout.payButton") == "checkout pay button"
    assert humanize("DT_FirstName_IN") == "First name" and make_token("Traveler first name") == "TRAVELER_FIRST_NAME"


def test_if_conditions_follow_the_contract_grammar():
    assert parse_condition("{PROMO_CODE} is filled").op == "filled"
    c = parse_condition("{PLAN} = Max")
    assert (c.left, c.op, c.right) == ("PLAN", "=", "Max")
    assert parse_condition("{COUNT} >= 2").op == ">="
    for bad in ("PLAN = Max", "{PLAN} ~ Max", "{COUNT} > many"):
        with pytest.raises(ValueError):
            parse_condition(bad)


# ---------------------------------------------------------------------------------------------------
# Problems
# ---------------------------------------------------------------------------------------------------
def test_problems_flag_unset_variables_missing_expected_values_and_unmapped_template_variables(tmp_path):
    def fill(s):
        s.add("Set", "Type code", FindBy="xpath", FindBy_Value="//input[@id='code']", Value="{PROMO_CODE}")
        s.add("Set", "Type name", FindBy="xpath", FindBy_Value="//input[@id='n']", Value="FirstName")
        s.add("Output", "Price shows", FindBy="xpath", FindBy_Value="//span[@id='p']", Exact_Match="Y")
        s.add("Set", "Template field", FindBy="xpath", FindBy_Value="//input[@id='d']", Value="{?DESTINATION}")
    model = model_of(small(tmp_path, fill, {"FirstName": None}))
    kinds = {(p["kind"], p["row"], p["severity"]) for p in model["problems"]}
    assert ("unset_variable", 2, "error") in kinds                   # {PROMO_CODE}: nothing has it
    assert ("unset_variable", 3, "error") in kinds                   # FirstName: its Params cell is empty
    assert ("missing_expected", 4, "warning") in kinds
    assert ("unmapped_template_variable", 5, "error") in kinds
    buy = find_test(model, "Buy")
    assert buy["needs"] == ["FIRSTNAME", "PROMO_CODE"] and step_at(buy, 2)["problems"] == ["error"]
    assert buy["blocks"][0]["dot"] == "error" and model["problemCounts"]["error"] == 3


def test_a_variable_another_enabled_test_provides_is_not_a_problem(tmp_path):
    path = tmp_path / "two.xlsx"
    build_steps_workbook(path, {
        "Buy": lambda s: s.add("Output", "Save policy", FindBy="xpath", FindBy_Value="//span[@id='pol']", Output_Value="{POLICY_NO}"),
        "View": lambda s: s.add("Set", "Search policy", FindBy="xpath", FindBy_Value="//input[@id='q']", Value="{POLICY_NO}"),
    })
    model = model_of(path)
    assert find_test(model, "Buy")["provides"] == ["POLICY_NO"] and find_test(model, "View")["needs"] == ["POLICY_NO"]
    assert not [p for p in model["problems"] if p["kind"] == "unset_variable"]
    policy = next(v for v in model["variables"] if v["key"] == "POLICY_NO")
    assert policy["providedBy"] == ["Buy"] and policy["neededBy"] == ["View"] and policy["kind"] == "set"


def test_side_effect_steps_are_blocked_on_production_and_suggested_by_their_name(tmp_path):
    def fill(s):
        s.add("js_click", "Click Purchase", FindBy="xpath", FindBy_Value="//button[contains(.,'Purchase')]")
    path = small(tmp_path, fill)
    doc = BuildDocument(path)
    assert [p["kind"] for p in doc.model()["problems"]] == ["side_effect_suggested"]
    doc.apply([{"op": "update_step", "test": "Buy", "row": 2, "set": {"sideEffects": True}},
               {"op": "set_environments", "names": ["UAT", "PROD"], "production": ["PROD"],
                "rows": [{"variable": "DOMAIN", "required": True, "values": {"UAT": "https://uat.example", "PROD": "https://example"}}]}])
    assert [p["kind"] for p in doc.model("UAT")["problems"]] == []
    prod = doc.model("PROD")["problems"]
    assert [(p["kind"], p["severity"]) for p in prod] == [("side_effect_on_prod", "warning")]


def test_a_required_environment_value_that_is_missing_is_an_error_for_the_chosen_environment(tmp_path):
    path = small(tmp_path, lambda s: s.add("Open", "Open site", Value="{DOMAIN}/home"))
    doc = BuildDocument(path)
    doc.apply([{"op": "set_environments", "names": ["QA", "UAT", "PROD"], "production": ["PROD"],
                "rows": [{"variable": "DOMAIN", "required": True, "values": {"QA": "https://qa.example", "UAT": "", "PROD": "https://example"}}]}])
    model = doc.model()                                              # Global says UAT
    env_problems = [(p["severity"], p["row"]) for p in model["problems"] if p["kind"] == "missing_environment_value"]
    assert ("error", None) in env_problems and ("error", 2) in env_problems
    assert model["environments"] == {"source": "rr", "names": ["QA", "UAT", "PROD"], "production": ["PROD"],
                                     "rows": [{"variable": "DOMAIN", "required": True, "secret": False,
                                               "values": {"QA": "https://qa.example", "UAT": "", "PROD": "https://example"}}]}
    assert not [p for p in doc.model("QA")["problems"] if p["severity"] == "error"]
    assert find_test(model, "Buy")["needs"] == []                      # an environment variable is not a need


def test_structure_problems_unknown_test_fingerprint_bad_condition_and_unclosed_if(tmp_path):
    def fill(s):
        s.add("CALL_TEST", "Call it", Value="Nope")
        s.add("ASSERT_PAGE", "Arrived", Value="Payment page")
        s.add("IF", "If", Value="PROMO is set")
        s.add("Click", "No locator")
    kinds = [p["kind"] for p in model_of(small(tmp_path, fill))["problems"]]
    assert sorted(kinds) == ["bad_condition", "missing_locator", "unbalanced_flow", "unknown_fingerprint", "unknown_test"]


def test_the_last_run_marks_steps_and_a_locator_miss_is_a_passive_warning(tmp_path):
    path = small(tmp_path, lambda s: (s.add("Click", "Next", FindBy="xpath", FindBy_Value="//button[@id='next']"),
                                       s.add("Click", "Ok", FindBy="xpath", FindBy_Value="//button[@id='ok']")))
    runs = tmp_path / "runs"
    for run_id, status in (("20260101-old", "PASSED"), ("20260102-new", "FAILED")):
        (runs / run_id).mkdir(parents=True)
        (runs / run_id / "results.json").write_text(json.dumps({
            "run_id": run_id, "workbook": str(path), "started_at": run_id, "tests": [{"id": "Buy", "sheet": "Buy", "status": status, "steps": [
                {"row": 2, "status": status, "error": "" if status == "PASSED" else "Object was not found"},
                {"row": 3, "status": "PASSED", "error": ""}]}]}))
        time.sleep(0.01)
    model = model_of(path, runs_dir=runs)
    buy = find_test(model, "Buy")
    assert step_at(buy, 2)["lastResult"]["runId"] == "20260102-new" and step_at(buy, 2)["lastResult"]["locatorMiss"] is True
    assert buy["lastRun"]["status"] == "FAILED"
    assert [(p["kind"], p["row"]) for p in model["problems"]] == [("last_run_locator_miss", 2)]


# ---------------------------------------------------------------------------------------------------
# Ops
# ---------------------------------------------------------------------------------------------------
def test_inserting_a_step_names_it_numbers_it_and_places_it_in_the_block(mock):
    path, rows = mock
    doc = BuildDocument(path)
    applied = doc.apply([{"op": "insert_step", "test": "FlowA", "after": rows["dep"],
                          "step": {"method": "Set", "findBy": "xpath", "locator": "//input[@name='promoCode']", "value": "{PROMO_CODE}"}}])
    new_row = applied[0]["row"]
    assert new_row == rows["dep"] + 1
    t = find_test(doc.model(), "FlowA")
    s = step_at(t, new_row)
    assert s["name"] == "Type {PROMO_CODE} into promo code field" and s["nameAuto"] is True and s["enabled"] is True and s["block"] == "Quote"
    assert doc.editor.get("FlowA", new_row, 2) == f"=COUNTA($B$1:B{new_row - 1})"      # Step_Number continues the column's formula
    assert doc.editor.get("FlowA", new_row + 1, 2) == f"=COUNTA($B$1:B{new_row})"      # ...and the rows below moved with their formulas


def test_an_override_name_sticks_until_reset_to_auto_and_legacy_rows_refuse_edits(mock):
    path, rows = mock
    doc = BuildDocument(path)
    doc.apply([{"op": "update_step", "test": "FlowA", "row": rows["dep"], "set": {"name": "Pick the departure"}}])
    doc.apply([{"op": "update_step", "test": "FlowA", "row": rows["dep"], "set": {"value": "RetDate"}}])
    s = step_at(find_test(doc.model(), "FlowA"), rows["dep"])
    assert s["name"] == "Pick the departure" and s["nameAuto"] is False and s["value"] == "RetDate"
    doc.apply([{"op": "update_step", "test": "FlowA", "row": rows["dep"], "set": {"nameAuto": True}}])
    assert step_at(find_test(doc.model(), "FlowA"), rows["dep"])["name"] == "Type {RetDate} into trip departure date field"
    with pytest.raises(BuildError) as err:
        doc.apply([{"op": "update_step", "test": "FlowA", "row": rows["cost_yes"], "set": {"value": "x"}}])
    assert err.value.kind == "legacy"


def test_moving_steps_keeps_their_order_and_can_put_them_in_another_block(mock):
    path, rows = mock
    doc = BuildDocument(path)
    names = lambda: [s["name"] for s in find_test(doc.model(), "FlowA")["steps"]]
    before = names()
    moved = doc.apply([{"op": "move_steps", "test": "FlowA", "rows": [rows["age1"], rows["age2"]], "before": rows["dep"], "block": "Ages"}])
    after = names()
    assert after[after.index("Enter Age"):after.index("Enter Age") + 3] == ["Enter Age", "Enter Age 2", "Departure Date"]
    assert sorted(after) == sorted(before)
    t = find_test(doc.model(), "FlowA")
    assert [step_at(t, r)["block"] for r in moved[0]["rows"]] == ["Ages", "Ages"]
    assert all(b["stored"] for b in t["blocks"]) and "Ages" in [b["title"] for b in t["blocks"]]


def test_blocks_rename_split_and_merge_through_the_block_column(mock):
    path, rows = mock
    doc = BuildDocument(path)
    doc.apply([{"op": "rename_block", "test": "FlowA", "row": rows["dep"], "title": "Trip details"}])
    t = find_test(doc.model(), "FlowA")
    assert [b["title"] for b in t["blocks"]][:3] == ["Open", "Trip details", "Result"]
    assert doc.editor.column("FlowA", BLOCK_COLUMN) is not None
    doc.apply([{"op": "split_block", "test": "FlowA", "row": rows["select"], "title": "Extras"}])
    assert [b["title"] for b in find_test(doc.model(), "FlowA")["blocks"]][:4] == ["Open", "Trip details", "Extras", "Result"]
    doc.apply([{"op": "merge_blocks", "test": "FlowA", "row": rows["select"]}])
    assert [b["title"] for b in find_test(doc.model(), "FlowA")["blocks"]][:3] == ["Open", "Trip details", "Result"]


def test_bulk_edit_changes_several_steps_and_skips_legacy_ones(mock):
    path, rows = mock
    doc = BuildDocument(path)
    applied = doc.apply([{"op": "bulk_edit", "test": "FlowA", "rows": [rows["dep"], rows["ret"], rows["cost_yes"]],
                          "set": {"timeout": 30, "onFail": "continue", "enabled": False}}])
    assert applied[0]["skipped"] == [rows["cost_yes"]]
    t = find_test(doc.model(), "FlowA")
    for r in (rows["dep"], rows["ret"]):
        assert (step_at(t, r)["timeout"], step_at(t, r)["onFail"], step_at(t, r)["enabled"]) == (30, "continue", False)
    with pytest.raises(BuildError):
        doc.apply([{"op": "bulk_edit", "test": "FlowA", "rows": [rows["dep"]], "set": {"value": "x"}}])


def test_renaming_a_variable_updates_params_headers_tokens_and_inline_uses(tmp_path):
    def fill(s):
        s.add("Set", "Type name", FindBy="xpath", FindBy_Value="//input[@id='n']", Value="FirstName")
        s.add("Set", "Greeting", FindBy="xpath", FindBy_Value="//input[@id='g']", Value="Hello {FirstName}!")
    path = small(tmp_path, fill, {"FirstName": "Ann"})
    doc = BuildDocument(path)
    applied = doc.apply([{"op": "rename_variable", "from": "FirstName", "to": "FIRST_NAME"},
                         {"op": "set_variable", "token": "FIRST_NAME", "label": "Traveler first name"}])
    assert applied[0]["changed"] == 3
    model = doc.model()
    buy = find_test(model, "Buy")
    assert step_at(buy, 2)["value"] == "FIRST_NAME" and step_at(buy, 3)["value"] == "Hello {FIRST_NAME}!"
    v = next(v for v in model["variables"] if v["key"] == "FIRST_NAME")
    assert v["label"] == "Traveler first name" and v["labelSet"] and len(v["usedBy"]) == 2 and v["sources"][0]["sheet"] == "Params_1"
    assert not [p for p in model["problems"] if p["kind"] == "unset_variable"]


def test_new_tests_variables_fingerprints_and_datasheets_settings(tmp_path):
    path = small(tmp_path, lambda s: s.add("Open", "Open", Value="https://x"))
    doc = BuildDocument(path)
    doc.apply([{"op": "add_test", "name": "Search", "paramSheet": "Search_Params"},
               {"op": "insert_step", "test": "Search", "step": {"method": "ASSERT_PAGE", "value": "Results page"}},
               {"op": "set_fingerprint", "name": "Results page", "urlContains": "/results", "landmark": "css=h1", "landmarkText": "Results"},
               {"op": "add_variable", "sheet": "Search_Params", "label": "Search text", "value": "shoes"},
               {"op": "set_test", "test": "Search", "tags": ["smoke", "search"], "comment": "made in the builder"}])
    model = doc.model()
    search = find_test(model, "Search")
    assert search["listed"] and search["enabled"] and search["paramSheet"] == "Search_Params" and search["tags"] == ["smoke", "search"]
    assert search["steps"][0]["name"] == "Arrived at: Results page" and search["blocks"][0]["gate"] == "Results page"
    assert model["fingerprints"][0]["usedBy"] == [{"test": "Search", "row": 2, "n": 1}]
    assert next(v for v in model["variables"] if v["key"] == "SEARCH_TEXT")["label"] == "Search text"
    assert not model["problems"]
    roles = {s["name"]: s["role"] for s in model["sheets"]}
    assert roles["Search"] == "test" and roles["Search_Params"] == "params" and roles["_rr_fingerprints"] == "runner"
    doc.save()
    wb = Workbook(path)                                             # the runner still reads it
    assert "Search" in [c.sheet for c in wb.discover()]
    assert openpyxl.load_workbook(path)["_rr_fingerprints"].sheet_state == "hidden"


def test_a_failing_op_applies_nothing_and_ops_see_rows_as_earlier_ops_left_them(mock):
    path, rows = mock
    doc = BuildDocument(path)
    version = doc.version
    with pytest.raises(BuildError) as err:
        doc.apply([{"op": "delete_steps", "test": "FlowA", "rows": [rows["dep"]]}, {"op": "explode"}])
    assert err.value.extra["index"] == 1 and doc.version == version and not doc.modified
    assert step_at(find_test(doc.model(), "FlowA"), rows["dep"])["name"] == "Departure Date"
    with pytest.raises(BuildError) as stale:
        doc.apply([{"op": "delete_steps", "test": "FlowA", "rows": [rows["dep"]]}], version=version - 1)
    assert stale.value.kind == "stale" and stale.value.status == 409


# ---------------------------------------------------------------------------------------------------
# Draft, undo, save, history
# ---------------------------------------------------------------------------------------------------
def test_edits_autosave_a_draft_and_undo_all_the_way_back_leaves_the_file_untouched(mock):
    path, rows = mock
    original = path.read_bytes()
    doc = BuildDocument(path)
    doc.apply([{"op": "update_step", "test": "FlowA", "row": rows["dep"], "set": {"timeout": 9}}])
    assert draft_path(path).is_file() and doc.status()["modified"] and doc.status()["canUndo"]
    again = BuildDocument(path)                                     # a reopened builder finds the draft
    assert step_at(find_test(again.model(), "FlowA"), rows["dep"])["timeout"] == 9 and again.status()["hasDraft"]
    doc.undo()
    assert not draft_path(path).exists() and not doc.modified and doc.status()["canRedo"]
    doc.redo()
    assert step_at(find_test(doc.model(), "FlowA"), rows["dep"])["timeout"] == 9
    doc.undo()
    assert doc.save().backup is None                                # nothing edited: nothing written
    assert path.read_bytes() == original


def test_save_writes_the_workbook_keeps_a_backup_and_history_diff_and_restore_work(mock):
    path, rows = mock
    doc = BuildDocument(path)
    doc.apply([{"op": "delete_steps", "test": "FlowA", "rows": [rows["ret"]]},
               {"op": "update_step", "test": "FlowA", "row": rows["dep"], "set": {"value": "Age"}}])
    changes = doc.diff("disk")
    assert {(c["kind"], (c["before"] or {}).get("name")) for c in changes if c["test"] == "FlowA"} == {("removed", "Return Date"),
                                                                                                    ("changed", "Departure Date")}
    result = doc.save()
    assert result.backup is not None and result.backup.parent == backups_dir(path) and not draft_path(path).exists()
    assert Workbook(path).runtime(Workbook(path).discover()[0]).total_rows > 0
    history = doc.history()
    assert [h["id"] for h in history] == [result.backup.name]
    assert doc.diff(history[0]["id"])                                # the saved file differs from the version before it
    doc.restore(history[0]["id"])
    assert step_at(find_test(doc.model(), "FlowA"), rows["ret"])["name"] == "Return Date" and doc.modified
    with pytest.raises(BuildError):
        doc.restore("../mock.xlsx")


def test_a_file_changed_outside_is_reported_and_save_refuses_unless_forced(mock):
    path, rows = mock
    doc = BuildDocument(path)
    doc.apply([{"op": "update_step", "test": "FlowA", "row": rows["dep"], "set": {"timeout": 7}}])
    other = WorkbookEditor.open(path)
    other.set("Global", 2, 2, "QA")
    other.save(backup=False)
    assert doc.status()["externalChange"]
    with pytest.raises(BuildError) as err:
        doc.save()
    assert err.value.kind == "changed_outside"
    doc.save(force=True)
    assert not doc.status()["externalChange"]


def test_an_unedited_workbook_changed_in_excel_is_reread(mock):
    from regrunner.workbook.builder import BuildStore
    path, _ = mock
    store = BuildStore()
    doc = store.get(path)
    assert doc.model()["environment"] == "UAT"
    other = WorkbookEditor.open(path)
    other.set("Global", 2, 2, "QA")
    other.save(backup=False)
    assert store.get(path).model()["environment"] == "QA"


def test_a_new_workbook_has_global_datasheets_and_the_environment_table(tmp_path):
    path = tmp_path / "fresh.xlsx"
    new_workbook(path, [{"name": "UAT", "domain": "https://uat.example"}, {"name": "PROD", "domain": "https://example", "production": True}])
    model = model_of(path)
    assert model["environment"] == "UAT" and model["tests"] == []
    assert model["environments"]["names"] == ["UAT", "PROD"] and model["environments"]["production"] == ["PROD"]
    assert model["environments"]["rows"][0] == {"variable": "DOMAIN", "required": True, "secret": False,
                                                "values": {"UAT": "https://uat.example", "PROD": "https://example"}}
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Global", "DataSheets", "_rr_environments"] and wb["Global"]["B2"].value == "UAT"
    assert Workbook(path).discover() == []
    with pytest.raises(BuildError):
        new_workbook(path, [])


# ---------------------------------------------------------------------------------------------------
# The real workbooks (read-only: copies are edited)
# ---------------------------------------------------------------------------------------------------
@pytest.mark.realworkbook
@pytest.mark.parametrize("real", REAL, ids=[p.name for p in REAL])
def test_every_real_workbook_loads_and_a_round_trip_through_the_model_changes_nothing(real, tmp_path):
    copy = tmp_path / real.name
    shutil.copy2(real, copy)
    original = copy.read_bytes()
    doc = BuildDocument(copy)
    model = doc.model()
    json.dumps(model)                                                # the server can send it
    web = [t for t in model["tests"] if t["kind"] == "web"]
    assert web and all(t["blocks"] and t["steps"] for t in web)
    assert all(sum(b["count"] for b in t["blocks"]) == len(t["steps"]) for t in web)
    assert doc.save().backup is None and copy.read_bytes() == original
    # one real edit touches that row only: every other row of every sheet keeps its exact XML
    t = next(t for t in web if any(not s["legacy"] and s["method"] == "WAIT" for s in t["steps"]))
    target = next(s for s in t["steps"] if not s["legacy"] and s["method"] == "WAIT")
    before = WorkbookEditor.open(copy)
    doc.apply([{"op": "update_step", "test": t["id"], "row": target["row"], "set": {"timeout": 42}}])
    doc.save()
    after = WorkbookEditor.open(copy)
    for sheet in before.sheet_names():
        for r in range(1, before.max_row(sheet) + 1):
            if (sheet, r) != (t["id"], target["row"]):
                assert before.row_xml(sheet, r) == after.row_xml(sheet, r), (sheet, r)
    assert step_at(find_test(BuildDocument(copy).model(), t["id"]), target["row"])["timeout"] == 42
