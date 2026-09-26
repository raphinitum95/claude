"""The workbook writer: load → edit → save must change only what was edited (P01 of the Workbook Builder).

Real workbooks in ``workbooks/`` are used read-only: every test works on a copy in ``tmp_path`` and a module fixture proves
the originals were not touched.  Formula values are compared with the runner's own evaluator (``BookState``), styles and
sheet structure with openpyxl.
"""
from __future__ import annotations

import io
import os
import re
import shutil
import warnings
import zipfile
from datetime import date, datetime
from pathlib import Path
from xml.etree import ElementTree

import openpyxl
import pytest
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

from regrunner.workbook import Workbook
from regrunner.workbook.sheet import BookState, WorkbookData
from regrunner.workbook.writer import (RowMap, WorkbookChangedOutside, WorkbookEditor, WorkbookLocked, WriterError,
                                       backup_path, draft_path, has_draft, lock_file, map_formula, map_ref_text)

from .workbook_factory import COL, build_workbook

ROOT = Path(__file__).resolve().parents[1]
REAL = sorted(p for p in (ROOT / "workbooks").glob("*.xlsx") if not p.name.startswith("~$"))
NOW = datetime(2026, 1, 15, 9, 30)

real_workbook = pytest.mark.parametrize(
    "real", [pytest.param(p, id=p.stem, marks=pytest.mark.realworkbook) for p in REAL]
    or [pytest.param(None, marks=[pytest.mark.realworkbook, pytest.mark.skip(reason="no workbook in workbooks/")])])


@pytest.fixture(scope="module", autouse=True)
def real_workbooks_are_never_written():
    before = {p: (p.stat().st_mtime_ns, p.read_bytes()) for p in REAL}
    yield
    for p, (mtime, data) in before.items():
        assert p.stat().st_mtime_ns == mtime and p.read_bytes() == data, f"{p.name} in workbooks/ was modified"


@pytest.fixture
def real_copy(real, tmp_path):
    target = tmp_path / real.name
    shutil.copy2(real, target)
    return target


# ---------------------------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------------------------
def zip_parts(data_or_path) -> dict[str, bytes]:
    src = io.BytesIO(data_or_path) if isinstance(data_or_path, bytes) else data_or_path
    with zipfile.ZipFile(src) as z:
        return {n: z.read(n) for n in z.namelist()}


def values(path: Path) -> dict[tuple[str, int, int], str]:
    """Every cell's value as the runner computes it (volatile RAND* skipped: their draw order depends on the layout)."""
    data = WorkbookData.load(path)
    book = BookState(data, now=NOW, seed=1)
    out = {}
    for name in data.order:
        sheet = book.sheet(name)
        for (r, c), raw in data.sheets[name.upper()].cells.items():
            if raw.formula and "RAND" in raw.formula.upper():
                continue
            out[(name.upper(), r, c)] = str(sheet.read(r, c))
    return out


def assert_changed_parts_are_well_formed(original: dict[str, bytes], saved: dict[str, bytes]) -> None:
    """Excel refuses a file with broken XML: every part the writer produced or changed must parse."""
    changed = [n for n, data in saved.items() if original.get(n) != data]
    assert changed
    for name in changed:
        if name.endswith((".xml", ".rels")):
            ElementTree.fromstring(saved[name])


def load_xl(path: Path):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(path)


def test_sheet_with_method(editor: WorkbookEditor) -> str:
    return next(n for n in editor.sheet_names() if editor.column(n, "Method") and editor.max_row(n) > 8)


test_sheet_with_method.__test__ = False   # a helper, not a test


def build_features_workbook(path: Path) -> dict[str, int]:
    """A runnable mock workbook plus everything a round trip can lose: cross-sheet formulas, a defined name, merged
    cells, data validation, conditional formatting, a hyperlink, a comment, column widths, a hidden sheet, legacy rows."""
    rows = build_workbook(path)["Flow"]
    wb = openpyxl.load_workbook(path)
    ws = wb["Flow"]
    last = ws.max_row
    legacy = {}
    for key, method, value, gate in [("goto", "GO_TO_ROW", "5", None), ("snag", "SNAGIT_SCREENSHOT", "after quote", None),
                                     ("gated", "Wait", "1", f'=IF($D${rows["exist"]}="PASSED","bln1001","N")')]:
        last += 1
        ws.cell(last, 1, gate or "bln1001")
        ws.cell(last, COL["Step_Number"], f"=COUNTA($B$1:B{last - 1})")
        ws.cell(last, COL["Step_Name"], f"legacy {method}")
        ws.cell(last, COL["Method"], method)
        ws.cell(last, COL["Value"], value)
        legacy[key] = last
    ws.column_dimensions["C"].width = 41.5
    ws.column_dimensions["L"].width = 22
    ws.cell(1, COL["Notes"]).font = Font(bold=True, color="FF0000FF")
    ws.cell(4, COL["Step_Name"]).fill = PatternFill("solid", fgColor="FFFFFF00")
    ws.row_dimensions[4].height = 30
    ws.merge_cells("AA4:AB4")
    ws["AA4"] = "merged note"
    dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    dv.add(f"A2:A{last}")
    ws.add_data_validation(dv)
    ws.conditional_formatting.add(f"D2:D{last}", CellIsRule(operator="equal", formula=['"FAILED"'],
                                                             fill=PatternFill("solid", fgColor="FFFF0000")))
    ws["AA6"] = "link"
    ws["AA6"].hyperlink = "https://example.invalid/help"
    ws["AA7"] = "has a note"
    ws["AA7"].comment = Comment("explains row 7", "QA")
    summary = wb.create_sheet("Summary")
    summary["A1"], summary["B1"] = "What", "Formula"
    summary["A2"], summary["B2"] = "open step name", f"=Flow!C{rows['open']}"
    summary["A3"], summary["B3"] = "steps with a flag", f"=COUNTA(Flow!A2:A{last})"
    summary["A4"], summary["B4"] = "quoted", f"='Flow'!$L${rows['dest_write']}"
    summary["A5"], summary["B5"] = "after the last step", f"=Flow!C{last + 5}"
    wb.defined_names["FlowStatus"] = DefinedName("FlowStatus", attr_text=f"Flow!$D$2:$D${last}")
    hidden = wb.create_sheet("Hidden")
    hidden["A1"] = "kept hidden"
    hidden.sheet_state = "hidden"
    wb.save(path)
    return {**rows, **legacy, "last": last}


def share_step_numbers(path: Path, first: int, last: int) -> None:
    """Rewrite the Step_Number formulas of Flow rows first..last as one Excel shared-formula group (openpyxl cannot)."""
    editor = WorkbookEditor.open(path)
    part = editor._sheet("Flow").part
    parts = zip_parts(path)
    xml = parts[part].decode()
    col = openpyxl.utils.get_column_letter(COL["Step_Number"])
    for r in range(first, last + 1):
        pattern = re.compile(rf'(<c r="{col}{r}"[^>]*>)<f>([^<]*)</f>')
        repl = (rf'\1<f t="shared" ref="{col}{first}:{col}{last}" si="0">\2</f>' if r == first
                else r'\1<f t="shared" si="0"/>')
        xml, _ = pattern.subn(repl, xml)          # section header rows have no step number: they stay out
    parts[part] = xml.encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in parts.items():
            z.writestr(name, data)


@pytest.fixture
def features(tmp_path):
    path = tmp_path / "features.xlsx"
    rows = build_features_workbook(path)
    return path, rows


# ---------------------------------------------------------------------------------------------
# row maps
# ---------------------------------------------------------------------------------------------
def test_row_maps_grow_ranges_on_insert_shrink_them_on_delete_and_lose_fully_deleted_ones():
    ins, dele = RowMap.insert(5, 2), RowMap.delete(5, 2)
    assert map_ref_text("A4", ins) == "A4" and map_ref_text("$A$5", ins) == "$A$7"
    assert map_ref_text("A1:B10", ins) == "A1:B12" and map_ref_text("A5:A6", ins) == "A7:A8"
    assert map_ref_text("A1:B10", dele) == "A1:B8" and map_ref_text("A5:A6", dele) is None
    assert map_ref_text("A6", dele) is None and map_ref_text("A7", dele) == "A5" and map_ref_text("3:9", dele) == "3:7"
    assert map_ref_text("A:C", dele) == "A:C" and map_ref_text("MyName", dele) == "MyName"
    move = RowMap.move(10, 2, 4)                        # rows 10-11 go above row 4
    assert [move.row(r) for r in (3, 4, 9, 10, 11, 12)] == [3, 6, 11, 4, 5, 12]
    down = RowMap.move(4, 2, 10)                        # rows 4-5 go above row 10
    assert [down.row(r) for r in (3, 4, 5, 6, 9, 10)] == [3, 8, 9, 4, 7, 10]
    assert RowMap.move(4, 2, 5).is_identity()


def test_formula_references_move_only_when_they_point_at_the_edited_sheet():
    rows = RowMap.insert(3, 1)
    assert map_formula('=IF($D$3="PASSED",Other!D3,D2)', "Flow", "Flow", rows) == '=IF($D$4="PASSED",Other!D3,D2)'
    assert map_formula("=SUM('My Flow'!A3:A9)&\"A3\"", "Summary", "My Flow", rows) == "=SUM('My Flow'!A4:A10)&\"A3\""
    assert map_formula("=[1]Flow!A5+Flow!A5", "X", "Flow", rows) == "=[1]Flow!A5+Flow!A6"
    assert map_formula("=Flow!A3", "X", "Flow", RowMap.delete(3, 1)) == "=Flow!#REF!"


# ---------------------------------------------------------------------------------------------
# real workbooks (copies)
# ---------------------------------------------------------------------------------------------
@real_workbook
def test_saving_a_real_workbook_without_changes_keeps_every_part_byte_for_byte(real_copy, real):
    editor = WorkbookEditor.open(real_copy)
    assert not editor.modified
    saved = zip_parts(editor.to_bytes())
    original = zip_parts(real)
    assert list(saved) == list(original)
    assert all(saved[name] == original[name] for name in original)


@real_workbook
def test_inserting_and_deleting_rows_in_a_real_test_sheet_keeps_every_formula_value(real_copy, real, tmp_path):
    editor = WorkbookEditor.open(real_copy)
    sheet = test_sheet_with_method(editor)
    before = values(real)
    xl_before = load_xl(real)
    editor.insert_rows(sheet, 4, 3)
    editor.delete_rows(sheet, 5, 1)                 # one of the new blank rows: nothing real is lost
    out = tmp_path / "edited.xlsx"
    editor.save_as(out)
    assert_changed_parts_are_well_formed(zip_parts(real), zip_parts(out))
    net = RowMap.insert(4, 2)
    after = values(out)
    for (name, r, c), value in before.items():
        target = net.row(r) if name == sheet.upper() else r
        assert after.get((name, target, c)) == value, f"{name}!{openpyxl.utils.get_column_letter(c)}{r}"

    xl_after = load_xl(out)
    assert xl_after.sheetnames == xl_before.sheetnames
    for ws in xl_before.worksheets:
        new = xl_after[ws.title]
        assert new.sheet_state == ws.sheet_state
        assert {k: d.width for k, d in new.column_dimensions.items()} == \
               {k: d.width for k, d in ws.column_dimensions.items()}
        moved = ws.title == sheet
        expected_merges = sorted(map_ref_text(str(m), net) if moved else str(m) for m in ws.merged_cells.ranges)
        assert sorted(str(m) for m in new.merged_cells.ranges) == expected_merges
        assert len(new.data_validations.dataValidation) == len(ws.data_validations.dataValidation)
        for row in ws.iter_rows():
            for cell in row:
                target = new.cell(net.row(cell.row) if moved else cell.row, cell.column)
                assert style_key(target) == style_key(cell), f"{ws.title}!{cell.coordinate}"
        for r, dim in ws.row_dimensions.items():
            assert new.row_dimensions[net.row(r) if moved else r].height == dim.height


@real_workbook
def test_rows_nobody_touched_keep_their_exact_xml_in_a_real_workbook(real_copy):
    editor = WorkbookEditor.open(real_copy)
    sheet = test_sheet_with_method(editor)
    last = editor.max_row(sheet)
    before = {r: editor.row_xml(sheet, r) for r in range(1, last + 1)}
    notes = editor.column(sheet, "Notes") or editor.column(sheet, "Method")
    editor.set(sheet, 3, notes, "edited by the builder")
    editor.save()
    reopened = WorkbookEditor.open(real_copy)
    for r, xml in before.items():
        if r != 3:
            assert reopened.row_xml(sheet, r) == xml, f"row {r} changed"
    assert reopened.get(sheet, 3, notes) == "edited by the builder"


@real_workbook
def test_a_real_workbook_lists_the_same_tests_after_new_columns_rows_and_hidden_sheets(real_copy, real):
    tests_before = [(c.id, c.enabled) for c in Workbook(real, environment="UAT", seed=1).discover()]
    editor = WorkbookEditor.open(real_copy)
    sheet = test_sheet_with_method(editor)
    editor.add_column(sheet, "BLOCK")
    editor.insert_rows(sheet, 3, 1)
    editor.ensure_rr_sheet("_rr_variables")
    editor.write_table("_rr_variables", [["name", "label"], ["DT_URL", "Start page"]])
    editor.save()
    assert_changed_parts_are_well_formed(zip_parts(real), zip_parts(real_copy))
    wb = Workbook(real_copy, environment="UAT", seed=1)
    assert [(c.id, c.enabled) for c in wb.discover()] == tests_before
    assert wb.data.sheet(sheet).headers["BLOCK"] == editor.column(sheet, "BLOCK")


@real_workbook
def test_any_change_drops_the_calc_chain_and_asks_excel_to_recalculate_on_open(real_copy):
    editor = WorkbookEditor.open(real_copy)
    assert "xl/calcChain.xml" in zip_parts(real_copy)
    editor.set(test_sheet_with_method(editor), 2, 1, "N")
    parts = zip_parts(editor.to_bytes())
    assert "xl/calcChain.xml" not in parts
    assert b"calcChain" not in parts["[Content_Types].xml"] and b"calcChain" not in parts["xl/_rels/workbook.xml.rels"]
    assert re.search(rb'<calcPr\b[^>]*fullCalcOnLoad="1"', parts["xl/workbook.xml"])
    untouched = [n for n in zip_parts(real_copy) if n.startswith(("customXml/", "xl/printerSettings/", "xl/theme/",
                                                                  "xl/externalLinks/", "xl/styles.xml",
                                                                  "xl/sharedStrings.xml", "xl/metadata.xml"))]
    original = zip_parts(real_copy)
    assert all(parts[n] == original[n] for n in untouched)


# ---------------------------------------------------------------------------------------------
# synthetic workbooks
# ---------------------------------------------------------------------------------------------
def test_saving_a_synthetic_workbook_without_changes_keeps_every_part_byte_for_byte(features):
    path, _ = features
    assert zip_parts(WorkbookEditor.open(path).to_bytes()) == zip_parts(path)


def test_legacy_rows_survive_byte_for_byte_when_other_rows_are_edited(features):
    path, rows = features
    editor = WorkbookEditor.open(path)
    legacy = {k: editor.row_xml("Flow", rows[k]) for k in ("goto", "snag", "gated")}
    editor.set("Flow", rows["open"], COL["Value"], "DT_URL2")
    editor.set_by_header("Flow", rows["dep"], "Notes", "changed")
    editor.save()
    reopened = WorkbookEditor.open(path)
    assert {k: reopened.row_xml("Flow", rows[k]) for k in legacy} == legacy
    assert reopened.get("Flow", rows["gated"], 1).startswith("=IF(")


def test_inserting_rows_moves_references_in_every_sheet_and_in_defined_names(features):
    path, rows = features
    before = values(path)
    editor = WorkbookEditor.open(path)
    at = rows["exist"]
    editor.insert_rows("Flow", at, 2)
    editor.save()
    wb = load_xl(path)
    assert wb["Summary"]["B2"].value == f"=Flow!C{rows['open']}"                 # above the insert: unchanged
    assert wb["Summary"]["B3"].value == f"=COUNTA(Flow!A2:A{rows['last'] + 2})"   # range grew
    assert wb["Summary"]["B4"].value == f"='Flow'!$L${rows['dest_write'] + 2}"
    assert wb.defined_names["FlowStatus"].attr_text == f"Flow!$D$2:$D${rows['last'] + 2}"
    gate = wb["Flow"].cell(rows["banner_out"] + 2, 1).value
    assert gate == f'=IF($D${rows["exist"] + 2}="PASSED","bln1001","N")'
    after = values(path)
    shift = RowMap.insert(at, 2)
    for (name, r, c), value in before.items():
        assert after[(name, shift.row(r) if name == "FLOW" else r, c)] == value


def test_deleting_rows_turns_references_to_them_into_ref_errors_and_shrinks_ranges(features):
    path, rows = features
    editor = WorkbookEditor.open(path)
    editor.delete_rows("Flow", rows["open"], 1)
    editor.save()
    wb = load_xl(path)
    assert wb["Summary"]["B2"].value == "=Flow!#REF!"
    assert wb["Summary"]["B3"].value == f"=COUNTA(Flow!A2:A{rows['last'] - 1})"
    assert wb.defined_names["FlowStatus"].attr_text == f"Flow!$D$2:$D${rows['last'] - 1}"
    assert wb["Flow"].cell(rows["open"], COL["Method"]).value == "Output"      # the next step moved up


def test_moving_rows_takes_their_references_along(features):
    path, rows = features
    editor = WorkbookEditor.open(path)
    dep, ret = rows["dep"], rows["ret"]
    assert ret == dep + 1
    editor.move_rows("Flow", ret, 1, dep)                  # swap: return date first, then departure date
    original = zip_parts(path)
    editor.save()
    assert_changed_parts_are_well_formed(original, zip_parts(path))
    ws = load_xl(path)["Flow"]
    assert ws.cell(dep, COL["Step_Name"]).value == "Return Date"
    assert ws.cell(ret, COL["Step_Name"]).value == "Departure Date"
    assert ws.cell(dep, COL["Step_Number"]).value == f"=COUNTA($B$1:B{dep - 1})"   # relative refs keep pointing up
    assert load_xl(path)["Summary"]["B4"].value == f"='Flow'!$L${rows['dest_write']}"


def test_shared_formulas_are_expanded_before_rows_shift_so_every_value_stays(features):
    path, rows = features
    first, last = rows["wait_first"], rows["last"]
    share_step_numbers(path, first, last)
    before = values(path)
    editor = WorkbookEditor.open(path)
    assert editor.get("Flow", rows["dep"], COL["Step_Number"]) == f"=COUNTA($B$1:B{rows['dep'] - 1})"
    assert editor.cell_xml("Flow", rows["dep"], COL["Step_Number"]).count('t="shared"') == 1
    editor.insert_rows("Flow", first + 2, 1)
    editor.save()
    shift = RowMap.insert(first + 2, 1)
    after = values(path)
    for (name, r, c), value in before.items():
        assert after[(name, shift.row(r) if name == "FLOW" else r, c)] == value
    assert b't="shared"' not in zip_parts(path)[editor._sheet("Flow").part]


def test_merged_cells_validations_formats_links_and_comments_follow_inserted_rows(features):
    path, rows = features
    editor = WorkbookEditor.open(path)
    editor.insert_rows("Flow", 3, 2)
    editor.save()
    ws = load_xl(path)["Flow"]
    last = rows["last"] + 2
    assert [str(m) for m in ws.merged_cells.ranges] == ["AA6:AB6"]
    assert ws["AA6"].value == "merged note"
    assert [str(dv.sqref) for dv in ws.data_validations.dataValidation] == [f"A2:A{last}"]
    assert [str(cf.sqref) for cf in ws.conditional_formatting] == [f"D2:D{last}"]
    assert ws["AA8"].hyperlink.target == "https://example.invalid/help"
    assert ws["AA9"].comment.text == "explains row 7" and ws["AA7"].comment is None
    vml = next(v for n, v in zip_parts(path).items() if n.endswith(".vml")).decode()
    assert re.search(r"<(\w+:)?Row>8</", vml)             # 0-based: the note of row 9
    assert ws.row_dimensions[6].height == 30 and ws.cell(6, COL["Step_Name"]).fill.fgColor.rgb == "FFFFFF00"


def test_deleting_a_row_removes_its_comment_link_and_merged_cells(features):
    path, _ = features
    editor = WorkbookEditor.open(path)
    editor.delete_rows("Flow", 4, 4)                       # rows 4-7 hold the merge, the link and the comment
    editor.save()
    ws = load_xl(path)["Flow"]
    assert not ws.merged_cells.ranges and not ws._hyperlinks
    assert all(c.comment is None for row in ws.iter_rows() for c in row)
    assert not re.search(r"<(\w+:)?Row>", next(v for n, v in zip_parts(path).items() if n.endswith(".vml")).decode())


def test_inserted_rows_copy_the_formatting_of_the_row_above_but_not_its_values(features):
    path, _ = features
    editor = WorkbookEditor.open(path)
    editor.insert_rows("Flow", 5, 1)
    editor.save()
    ws = load_xl(path)["Flow"]
    assert ws.row_dimensions[5].height == 30
    assert ws.cell(5, COL["Step_Name"]).fill.fgColor.rgb == "FFFFFF00"
    assert all(ws.cell(5, c).value is None for c in range(1, 30))


def test_a_new_column_goes_after_the_last_header_with_the_same_style_and_the_runner_reads_it(features):
    path, _ = features
    editor = WorkbookEditor.open(path)
    col = editor.add_column("Flow", "SIDE_EFFECTS")
    assert col == len(COL) + 1 and editor.column("Flow", "side_effects") == col
    with pytest.raises(WriterError):
        editor.add_column("Flow", "Method")
    editor.set("Flow", 3, col, "Y")
    editor.save()
    ws = load_xl(path)["Flow"]
    assert ws.cell(1, col).value == "SIDE_EFFECTS"
    assert style_key(ws.cell(1, col)) == style_key(ws.cell(1, col - 1))
    assert Workbook(path).data.sheet("Flow").headers["SIDE_EFFECTS"] == col
    assert WorkbookData.load(path).sheet("Flow").cells[(3, col)].value == "Y"


def test_hidden_runner_sheets_can_be_added_written_and_read_back_without_changing_the_tests(features):
    path, _ = features
    tests_before = [c.id for c in Workbook(path).discover()]
    editor = WorkbookEditor.open(path)
    with pytest.raises(WriterError):
        editor.ensure_rr_sheet("fingerprints")
    editor.ensure_rr_sheet("_rr_fingerprints")
    editor.ensure_rr_sheet("_rr_fingerprints")             # second call: no duplicate
    table = [["page", "title", "must_have"], ["Quote", "Get a quote", "#quote-form"], ["Pay", "Payment", 3]]
    editor.write_table("_rr_fingerprints", table)
    editor.save()
    reopened = WorkbookEditor.open(path)
    assert reopened.rr_sheets() == ["_rr_fingerprints"] and reopened.is_hidden("_rr_fingerprints")
    assert reopened.read_table("_rr_fingerprints") == table
    wb = load_xl(path)
    assert wb["_rr_fingerprints"].sheet_state == "hidden" and wb["Hidden"].sheet_state == "hidden"
    assert [c.id for c in Workbook(path).discover()] == tests_before
    reopened.write_table("_rr_fingerprints", [["page"], ["Quote"]])
    assert reopened.read_table("_rr_fingerprints") == [["page"], ["Quote"]]


def test_written_values_come_back_with_their_types_and_keep_the_cells_style(features):
    path, rows = features
    editor = WorkbookEditor.open(path)
    styled = (4, COL["Step_Name"])
    samples = {COL["Notes"]: "  spaced & <tagged> text\nline 2", COL["Timeout"]: 45, COL["Duration"]: 1.25,
               COL["Exact_Match"]: True, COL["Start_Time"]: date(2026, 3, 1), COL["Output_Value"]: "=IFS(A1=1,2)"}
    for col, value in samples.items():
        editor.set("Flow", 3, col, value)
    editor.set("Flow", *styled, None)
    editor.save()
    ws = load_xl(path)["Flow"]
    assert ws.cell(3, COL["Notes"]).value == samples[COL["Notes"]]
    assert ws.cell(3, COL["Timeout"]).value == 45 and ws.cell(3, COL["Duration"]).value == 1.25
    assert ws.cell(3, COL["Exact_Match"]).value is True
    assert ws.cell(3, COL["Output_Value"]).value == "=_xlfn.IFS(A1=1,2)"
    assert ws.cell(*styled).value is None and ws.cell(*styled).fill.fgColor.rgb == "FFFFFF00"
    data = WorkbookData.load(path).sheet("Flow")
    assert data.cells[(3, COL["Notes"])].value == samples[COL["Notes"]]
    assert WorkbookEditor.open(path).get("Flow", 3, COL["Start_Time"]) == 46082   # Excel serial for 2026-03-01


# ---------------------------------------------------------------------------------------------
# saving: backups, locks, external changes, drafts
# ---------------------------------------------------------------------------------------------
def test_save_keeps_a_timestamped_backup_of_the_previous_version(features):
    path, _ = features
    original = path.read_bytes()
    editor = WorkbookEditor.open(path)
    editor.set("Flow", 3, COL["Notes"], "v2")
    first = editor.save(when=datetime(2026, 9, 26, 14, 5))
    assert first.backup == path.parent / "backups" / "features.2026-09-26_1405.xlsx"
    assert first.backup.read_bytes() == original
    editor.set("Flow", 3, COL["Notes"], "v3")
    second = editor.save(when=datetime(2026, 9, 26, 14, 5))
    assert second.backup.name == "features.2026-09-26_1405-2.xlsx"
    assert WorkbookEditor.open(second.backup).get("Flow", 3, COL["Notes"]) == "v2"
    assert backup_path(path, datetime(2026, 9, 26, 14, 5)).name == "features.2026-09-26_1405-3.xlsx"


def test_saving_while_excel_has_the_workbook_open_says_close_it_in_excel(features, monkeypatch):
    path, _ = features
    editor = WorkbookEditor.open(path)
    editor.set("Flow", 3, COL["Notes"], "x")
    lock_file(path).write_bytes(b"owner")
    with pytest.raises(WorkbookLocked, match="Close it in Excel"):
        editor.save()
    lock_file(path).unlink()

    def refuse(src, dst):                                  # what Windows does while Excel holds the file
        raise PermissionError(13, "in use")
    monkeypatch.setattr(os, "replace", refuse)
    with pytest.raises(WorkbookLocked, match="Close it in Excel"):
        editor.save()
    assert not list(path.parent.glob("*.rr-tmp"))


def test_a_workbook_changed_on_disk_is_not_overwritten_unless_forced(features):
    path, _ = features
    editor = WorkbookEditor.open(path)
    editor.set("Flow", 3, COL["Notes"], "mine")
    other = WorkbookEditor.open(path)
    other.set("Flow", 3, COL["Test_Case"], "theirs")
    other.save(backup=False)
    assert editor.external_change()
    with pytest.raises(WorkbookChangedOutside):
        editor.save()
    editor.save(force=True)
    assert not editor.external_change()


def test_touching_the_file_without_changing_it_is_not_an_external_change(features):
    path, _ = features
    editor = WorkbookEditor.open(path)
    os.utime(path, ns=(1_000_000_000, 1_000_000_000))
    assert not editor.external_change()


def test_a_draft_is_autosaved_beside_the_workbook_and_reopens_with_its_edits(features):
    path, _ = features
    editor = WorkbookEditor.open(path)
    editor.set("Flow", 3, COL["Notes"], "draft edit")
    draft = editor.save_draft()
    assert draft == draft_path(path) == path.parent / ".drafts" / path.name and has_draft(path)
    assert [p.name for p in path.parent.iterdir() if p.suffix == ".xlsx"] == [path.name]   # not listed as a workbook
    assert WorkbookEditor.open(path).get("Flow", 3, COL["Notes"]) != "draft edit"        # the original is untouched
    resumed = WorkbookEditor.open_draft(path)
    assert resumed.path == path and resumed.get("Flow", 3, COL["Notes"]) == "draft edit" and not resumed.external_change()
    resumed.save()
    assert not has_draft(path) and WorkbookEditor.open(path).get("Flow", 3, COL["Notes"]) == "draft edit"


def test_a_reopened_draft_notices_the_workbook_changed_since_the_draft_began(features):
    path, _ = features
    editor = WorkbookEditor.open(path)
    editor.set("Flow", 3, COL["Notes"], "draft edit")
    editor.save_draft()
    wb = openpyxl.load_workbook(path)                     # someone edits it in Excel meanwhile
    wb["Flow"].cell(3, COL["Test_Case"], "saved in Excel meanwhile")
    wb.save(path)
    assert WorkbookEditor.open_draft(path).external_change()


def test_save_as_writes_a_new_workbook_and_leaves_the_original_alone(features):
    path, _ = features
    original = path.read_bytes()
    editor = WorkbookEditor.open(path)
    editor.set("Flow", 3, COL["Notes"], "copy")
    target = path.with_name("copy.xlsx")
    editor.save_as(target)
    assert path.read_bytes() == original and editor.path == target
    assert WorkbookEditor.open(target).get("Flow", 3, COL["Notes"]) == "copy"
    with pytest.raises(FileExistsError):
        WorkbookEditor.open(path).save_as(target)
