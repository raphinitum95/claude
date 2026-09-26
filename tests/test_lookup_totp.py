"""What the Qantas StandAlone / Staff workbook needed: VLOOKUP + NUMBERVALUE formulas, TEXT zero-padding, GET_GOOGLE_TOKEN,
and a clear failure (not a quiet ``#NAME?``) for any Excel function regrunner cannot calculate."""
from __future__ import annotations

import asyncio
import base64
from types import SimpleNamespace

import openpyxl
import pytest

from regrunner import totp
from regrunner.config import Config
from regrunner.engine.actions import REGISTRY, UNSUPPORTED, StepContext, run_action
from regrunner.lint import lint
from regrunner.workbook import Workbook
from regrunner.workbook.formula import ExcelError
from regrunner.workbook.model import PreparedStep
from regrunner.workbook.sheet import ErrorText
from tests.test_formula import ev
from tests.workbook_factory import Sheet

TABLE = {(1, 1): "Env", (1, 2): "URL", (1, 3): "API",
         (2, 1): "QA", (2, 2): "https://qa.test", (2, 3): "https://api.qa.test/",
         (3, 1): "UAT", (3, 2): "https://uat.test", (3, 3): None,
         (4, 1): "Prod", (4, 2): "https://prod.test", (4, 3): "https://api.test/"}


# -- VLOOKUP ---------------------------------------------------------------------------------------------------------------------
def test_vlookup_exact_match_returns_the_cell_in_the_wanted_column_case_insensitively():
    assert ev('=VLOOKUP("UAT",A1:C4,2,FALSE)', cells=TABLE) == "https://uat.test"
    assert ev('=VLOOKUP("uat",A1:C4,2,0)', cells=TABLE) == "https://uat.test"
    assert ev('=VLOOKUP("Prod",A1:C4,3,FALSE)', cells=TABLE) == "https://api.test/"


def test_vlookup_a_missing_key_is_na_and_a_bad_column_is_an_error_like_excel():
    with pytest.raises(ExcelError) as err:
        ev('=VLOOKUP("DEV",A1:C4,2,FALSE)', cells=TABLE)
    assert err.value.code == "#N/A"
    with pytest.raises(ExcelError) as err:
        ev('=VLOOKUP("UAT",A1:C4,4,FALSE)', cells=TABLE)
    assert err.value.code == "#REF!"
    with pytest.raises(ExcelError) as err:
        ev('=VLOOKUP("UAT",A1:C4,0,FALSE)', cells=TABLE)
    assert err.value.code == "#VALUE!"


def test_vlookup_an_empty_result_cell_reads_as_zero_and_iferror_can_catch_a_miss():
    assert ev('=VLOOKUP("UAT",A1:C4,3,FALSE)', cells=TABLE) == 0
    assert ev('=IFERROR(VLOOKUP("DEV",A1:C4,2,FALSE),"none")', cells=TABLE) == "none"
    assert ev('=VLOOKUP("UAT",A1:C4,2,FALSE)&"/home"', cells=TABLE) == "https://uat.test/home"


def test_vlookup_wildcards_and_numbers():
    cells = {(1, 1): "alpha", (1, 2): "A", (2, 1): "beta", (2, 2): "B", (3, 1): 10, (3, 2): "ten", (4, 1): 20, (4, 2): "twenty"}
    assert ev('=VLOOKUP("be*",A1:B4,2,FALSE)', cells=cells) == "B"
    assert ev('=VLOOKUP("alph?",A1:B4,2,FALSE)', cells=cells) == "A"
    assert ev("=VLOOKUP(20,A1:B4,2,FALSE)", cells=cells) == "twenty"
    with pytest.raises(ExcelError):
        ev('=VLOOKUP("10",A1:B4,2,FALSE)', cells=cells)                       # text "10" is not the number 10
    assert ev("=VLOOKUP(15,A3:B4,2,TRUE)", cells=cells) == "ten"              # approximate: the last key that is <= 15
    assert ev("=VLOOKUP(25,A3:B4,2)", cells=cells) == "twenty"                # TRUE is the default
    with pytest.raises(ExcelError):
        ev("=VLOOKUP(5,A3:B4,2,TRUE)", cells=cells)


# -- NUMBERVALUE / TEXT ------------------------------------------------------------------------------------------------------------
@pytest.mark.parametrize("formula,expected", [
    ('=NUMBERVALUE("1,234.56",".",",")', 1234.56),
    ('=NUMBERVALUE("1.234,56",",",".")', 1234.56),
    ('=NUMBERVALUE("2 500")', 2500),
    ('=NUMBERVALUE("")', 0),
    ('=NUMBERVALUE("3%")', 0.03),
    ('=NUMBERVALUE("3%%")', 0.0003),
    ('=NUMBERVALUE(0,".",",")', 0),
    ('=NUMBERVALUE("-12.5")', -12.5),
    ('=_xlfn.NUMBERVALUE("99.90",".")', 99.9),                                 # how Excel stores it in the file
])
def test_numbervalue_converts_text_to_a_number_like_excel(formula, expected):
    assert ev(formula) == pytest.approx(expected)


@pytest.mark.parametrize("text", ["abc", "1.2.3", "1.2,3", "12a", "$5", "--1", "."])
def test_numbervalue_rejects_what_excel_rejects(text):
    with pytest.raises(ExcelError) as err:
        ev(f'=NUMBERVALUE("{text}",".",",")')
    assert err.value.code == "#VALUE!"


def test_the_points_to_earn_calculation_chain_of_the_workbook():
    """Wait rows compute with Excel: N171 = NUMBERVALUE(S167), R176 = SUM(R175, N171), R178 = TEXT(R176, "0.00")."""
    assert ev('=TEXT(SUM(NUMBERVALUE("120",".")+NUMBERVALUE("1,234.50",".",",")),"0.00")') == "1354.50"


def test_text_pads_with_zeros_and_a_long_year_code_is_a_four_digit_year():
    assert ev('=TEXT(1,"000")') == "001" and ev('=TEXT(42,"000")') == "042" and ev('=TEXT(1234,"000")') == "1234"
    assert ev('=TEXT(-5,"000")') == "-005" and ev('=TEXT(12.5,"000.00")') == "012.50" and ev('=TEXT(1234.5,"#,##0.00")') == "1,234.50"
    assert ev('="TC"&TEXT(3,"000")&".json"') == "TC003.json"
    assert ev('=TEXT(DATE(2026,5,18),"yyyyy-mmm-dd")') == "2026-May-18"
    assert ev('=TEXT(DATE(2026,5,18),"yy/mm/dd")') == "26/05/18"


# -- the workbook shape: Environment -> VLOOKUP -> URL; a token step feeding a Set --------------------------------------------------
def build(tmp_path, value_for_open="DT_URL", extra=None, key="JBSWY3DPEHPK3PXP", secrets=None):
    wb = openpyxl.Workbook()
    ds = wb.active
    ds.title = "Datasheets"
    ds.append(["SheetsToExecute", "blnExecute", "ParameterSheet", "Comments"])
    ds.append(["Flow", "Y", "Flow_Params", None])
    g = wb.create_sheet("Global")
    g.append(["Parameter", "Value", "Comments", None, "Env", "URL", "API", "Portal"])
    g.append(["Environment", "UAT", None, None, "QA", "https://qa.test", "https://api.qa.test/", "https://portal.qa.test/login"])
    g.append([None, None, None, None, "UAT", "https://uat.test", "https://api.uat.test/", "https://portal.uat.test/login"])
    p = wb.create_sheet("Flow_Params")
    p.append(["blnExecute", "DT_URL", "DT_Portal", "DT_Scenario", "DT_Key"])
    p.append(["Y", "=VLOOKUP(Global!$B$2,Global!$E$1:$F$4,2,FALSE)", "=VLOOKUP(Global!$B$2,Global!$E$1:$H$4,4,FALSE)",
              "Standalone", key])
    sheet = Sheet(wb.create_sheet("Flow"), "Y")
    rows = {"open": sheet.add("Open", "Launch", Page="chrome", Value=value_for_open),
            "portal": sheet.add("Open", "Portal", Page="chrome", Value="DT_Portal"),
            "token": sheet.add("GET_GOOGLE_TOKEN", "Get token", Value="DT_Key")}
    rows["answer"] = sheet.add("Set", "Enter code", FindBy="xpath", FindBy_Value="//input[@name='answer']", Value=f"=S{rows['token']}")
    for name, cols in (extra or {}).items():
        rows[name] = sheet.add(**cols)
    path = tmp_path / "wb.xlsx"
    wb.save(path)
    return Workbook(path, environment="UAT", seed=1, secrets=secrets), rows


def test_urls_come_from_the_environment_table_through_vlookup(tmp_path):
    wb, rows = build(tmp_path)
    rt = wb.runtime(wb.discover()[0])
    assert rt.prepare_row(rows["open"]).text("VALUE") == "https://uat.test"
    assert rt.prepare_row(rows["portal"]).text("VALUE") == "https://portal.uat.test/login"
    assert not [f for f in lint(wb) if f.severity == "error"]                                   # GET_GOOGLE_TOKEN is a supported action now


def test_the_secret_can_live_in_secrets_env_instead_of_the_workbook(tmp_path):
    wb, rows = build(tmp_path, key=None, secrets={"DT_KEY": "MFRGGZDFMZTWQ2LK"})                # Params cell blank, RR_VAR_DT_KEY set
    step = wb.runtime(wb.discover()[0]).prepare_row(rows["token"])
    assert step.text("VALUE") == "MFRGGZDFMZTWQ2LK" and "VALUE" in step.secret_columns             # and reports mask it


def test_the_token_step_output_reaches_the_set_step_that_reads_it(tmp_path):
    wb, rows = build(tmp_path)
    rt = wb.runtime(wb.discover()[0])
    token = rt.prepare_row(rows["token"])
    assert token.text("VALUE") == "JBSWY3DPEHPK3PXP"                                           # DT_Key from the parameter row
    rt.record(token, "PASSED", "", "081804")                                                    # what the action's output becomes (column S)
    assert rt.prepare_row(rows["answer"]).text("VALUE") == "081804"


def test_an_excel_function_that_cannot_be_calculated_is_reported_by_lint_and_fails_the_step_with_the_reason(tmp_path):
    wb, rows = build(tmp_path, extra={"bad": dict(method="Open", name="Open bad", Page="chrome", Value="=WEEKNUM(TODAY())")})
    findings = [f for f in lint(wb) if f.severity == "error"]
    assert len(findings) == 1 and findings[0].row == rows["bad"]
    assert "Value cell evaluates to #NAME?" in findings[0].message and "WEEKNUM" in findings[0].message
    rt = wb.runtime(wb.discover()[0])
    step = rt.prepare_row(rows["bad"])
    assert isinstance(step.values["VALUE"], ErrorText) and "WEEKNUM" in step.values["VALUE"].detail
    ctx = StepContext(step=step, session=SimpleNamespace(is_open=False, typed=set()), cfg=Config(base_dir=tmp_path), selector_map=None, review=None, test_dir=tmp_path)
    asyncio.run(run_action(ctx))
    assert ctx.out.error.startswith("The Value cell evaluates to #NAME?") and "WEEKNUM" in ctx.out.error and "not run" in ctx.out.error


# -- GET_GOOGLE_TOKEN ---------------------------------------------------------------------------------------------------------------
SECRET = base64.b32encode(b"12345678901234567890").decode()          # the RFC 6238 test key


@pytest.mark.parametrize("at,expected", [(59, "287082"), (1111111109, "081804"), (1111111111, "050471"), (1234567890, "005924"),
                                         (2000000000, "279037"), (20000000000, "353130")])
def test_codes_match_the_rfc_6238_test_vectors_including_leading_zeros(at, expected):
    assert totp.code_at(SECRET, at) == expected


def test_the_secret_may_be_written_the_way_authenticator_apps_show_it():
    spaced = " ".join(SECRET[i:i + 4] for i in range(0, len(SECRET), 4)).lower()
    assert totp.code_at(spaced, 59) == "287082"
    assert totp.code_at(spaced.replace(" ", "-"), 59) == "287082"
    assert totp.code_at(SECRET.rstrip("=") + "====", 59) == "287082"


def test_the_window_length_is_reported():
    assert totp.seconds_left(59) == 1 and totp.seconds_left(60) == 30


@pytest.mark.parametrize("secret", ["", "   ", "not a key 1 8 9", "0000", "DT_Key", "DTKEY", "JBSWY3DP!"])
def test_a_bad_secret_is_an_error_not_an_empty_code(secret):
    with pytest.raises(totp.BadSecret):
        totp.code_at(secret, 59)


def token_ctx(tmp_path, value):
    step = PreparedStep(row=5, values={"METHOD": "GET_GOOGLE_TOKEN", "VALUE": value})
    return StepContext(step=step, session=SimpleNamespace(is_open=False, typed=set()), cfg=Config(base_dir=tmp_path), selector_map=None, review=None, test_dir=tmp_path)


def test_the_action_puts_the_code_in_the_step_output(tmp_path, monkeypatch):
    assert "GET_GOOGLE_TOKEN" in REGISTRY and "GET_GOOGLE_TOKEN" not in UNSUPPORTED
    monkeypatch.setattr(totp, "seconds_left", lambda at=None: 20.0)
    monkeypatch.setattr(totp, "code_at", lambda secret, at=None, digits=6: "042042")
    ctx = token_ctx(tmp_path, "ANYSECRET")
    asyncio.run(run_action(ctx))
    assert ctx.out.error == "" and ctx.out.output == "042042"


def test_a_code_about_to_expire_waits_for_the_next_window(tmp_path, monkeypatch):
    monkeypatch.setattr(totp, "seconds_left", lambda at=None: 0.05)
    monkeypatch.setattr(totp, "code_at", lambda secret, at=None, digits=6: "111111")
    ctx = token_ctx(tmp_path, "ANYSECRET")
    asyncio.run(run_action(ctx))
    assert ctx.out.output == "111111" and any("next 30 s code window" in n for n in ctx.out.notes)


def test_the_action_names_the_problem_when_the_value_is_not_a_secret(tmp_path):
    ctx = token_ctx(tmp_path, "DT_Key")                                                         # the token was never replaced by a real key
    asyncio.run(run_action(ctx))
    assert "GET_GOOGLE_TOKEN" in ctx.out.error and "base32" in ctx.out.error and not ctx.out.output
