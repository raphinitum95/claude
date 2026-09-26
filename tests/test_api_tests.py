"""API (web service) tests: a data row of a WEBSERVICE_URL sheet + the InputOutput sheet, run through the same engine, reports and events."""
from __future__ import annotations

import json
import os
import subprocess
import sys
import time
from pathlib import Path

import openpyxl
import pytest

from regrunner import insight
from regrunner.config import Config
from regrunner.lint import lint
from regrunner.workbook import Workbook
from regrunner.workbook.api import check, json_get, json_set, output_json_text, output_text, template_candidates
from tests.site import server as mock_server
from tests.workbook_factory import build_workbook

COLUMNS = ["blnExecute", "TCID", "TC_Name", "WEBSERVICE_URL", "JSON_FORMAT", "WEBSERVICE_METHOD", "XML_LOCATION", "XML_REQUESTFILE",
           "DT_apiKey", "DT_bearerToken", "policyNumber_IN", "lastName_IN", "transactionStatus_EXP", "MedScen01_EXP", "AgentEmail_EXP",
           "PolicyStatus_EXP", "Premium_EXP", "Covered_EXP", "RES_STATUS_CD_EXP", "transactionStatus_OUT", "MedScen01_OUT", "AgentEmail_OUT",
           "PolicyStatus_OUT", "Premium_OUT", "Covered_OUT", "Echo_OUT"]
IO = [("Function", "Parameter", "Value"),
      ("addHeader", "Authorization", "DT_bearerToken"), ("addHeader", "apiKey", "DT_apiKey"),
      ("INPUT_Contract", None, None),
      ("update_json", "searchCriteria.searchParameters.[0].value", "policyNumber_IN"),
      ("update_json", "searchCriteria.searchParameters.[1].value", "lastName_IN"),
      ("OUTPUT_Contract", None, None),
      ("output_json", "transactionStatus", "transactionStatus_OUT"),
      ("output_json", "detailResponse.policyDetail.travelers.[0].customElements.[1].value", "MedScen01_OUT"),
      ("output_json", "detailResponse.policyDetail.accountDetail.agent.email", "AgentEmail_OUT"),
      ("output_json", "detailResponse.policyDetail.displayStatus", "PolicyStatus_OUT"),
      ("output_json", "detailResponse.policyDetail.premium", "Premium_OUT"),
      ("output_json", "detailResponse.policyDetail.covered", "Covered_OUT"),
      ("output_json", "echo.searchCriteria.searchParameters.[0].value", "Echo_OUT"),
      ("COMPARE_Contract", None, None),
      ("compareX", "RES_STATUS_CD_EXP", "RES_STATUS_CD_OUT"),
      ("compare", "transactionStatus_EXP", "transactionStatus_OUT"),
      ("compare", "MedScen01_EXP", "MedScen01_OUT"),
      (None, "AgentEmail_EXP", "AgentEmail_OUT"),                        # no function name: the legacy runner skipped this row
      ("compare_ignore_case", "PolicyStatus_EXP", "PolicyStatus_OUT"),
      ("compare", "Premium_EXP", "Premium_OUT"),
      ("contains", "Covered_EXP", "Covered_OUT"),
      ("DisableCheckpoint", "Generate Token", "PolicyStatus_OUT")]


def build_api_workbook(path: Path, site: str, *, url_tail: str = "policy/v4/P-100", method: str = "GET", exp: dict | None = None,
                       api_key: str = "KEY-123", bearer: str = "Bearer TOKEN-abc", request_file: str = "", location: str = "",
                       policy_in: str = "P-100", url_formula: str | None = None, io_rows: list | None = None,
                       ui_flow: bool = False) -> Path:
    if ui_flow:
        build_workbook(path, flows=["Flow"], base_url=site, enabled=["Flow"])
        wb = openpyxl.load_workbook(path)
        ds = wb["DataSheets"]
        ds.append(["Policy", "Y", None, "Policy API"])
    else:
        wb = openpyxl.Workbook()
        ds = wb.active
        ds.title = "DataSheets"
        ds.append(["SheetsToExecute", "blnExecute", "ParameterSheet", "Comments"])
        ds.append(["Policy", "Y", None, "Policy API"])
        g = wb.create_sheet("Global")
        g.append(["Parameter", "Value", "Comments"])
        g.append(["Environment", "UAT"])
        g.append(["TC_Name_Column", "TC_Name"])
    g = wb["Global"]
    for row, values in enumerate([("Env", "URL", "API"), ("QA", "http://qa.invalid/", "http://qa.invalid/api/"), ("UAT", site, site)], start=1):
        for col, value in enumerate(values, start=5):
            g.cell(row, col, value)
    io = wb.create_sheet("InputOutput")
    for row in io_rows or IO:
        io.append(list(row))
    expected = {"transactionStatus_EXP": "Success", "MedScen01_EXP": "Cover not available", "AgentEmail_EXP": "qa@example.com",
                "PolicyStatus_EXP": "ACTIVE", "Premium_EXP": 123.5, "Covered_EXP": "True", **(exp or {})}
    ws = wb.create_sheet("Policy")
    ws.append(COLUMNS)
    url = url_formula or f'=VLOOKUP(Global!$B$2,Global!$E$1:$G$3,3,FALSE)&"{url_tail}"'
    row = {"blnExecute": "Y", "TCID": "001", "TC_Name": "Policy Search w EMC", "WEBSERVICE_URL": url, "JSON_FORMAT": "Y", "WEBSERVICE_METHOD": method,
           "XML_LOCATION": location or None, "XML_REQUESTFILE": request_file or None, "DT_apiKey": api_key, "DT_bearerToken": bearer,
           "policyNumber_IN": policy_in, "lastName_IN": "Smith", **expected}
    ws.append([row.get(c) for c in COLUMNS])
    wb.save(path)
    return path


def workbook(tmp_path: Path, site: str, **kw) -> Workbook:
    return Workbook(build_api_workbook(tmp_path / "api.xlsx", site, **kw), seed=1)


# -- the pieces ---------------------------------------------------------------------------------------------------------------------
def test_json_paths_read_and_write_dotted_paths_with_list_indexes():
    data = {"a": {"b": [{"c": 1}, {"c": "x"}]}, "n": None}
    assert json_get(data, "a.b.[1].c") == "x" and json_get(data, "a/b/[0]/c") == 1 and json_get(data, "a.b.[2].c") is None and json_get(data, "zzz") is None
    assert json_set(data, "a.b.[0].c", "42") and data["a"]["b"][0]["c"] == 42                 # an existing number stays a number
    assert json_set(data, "a.b.[1].c", "y") and data["a"]["b"][1]["c"] == "y"
    assert not json_set(data, "a.q.[0]", "z") and not json_set(data, "a.b.[9].c", "z")
    assert [output_text(v) for v in (None, True, 12.5, "s", [1, 2])] == ["", "True", "12.5", "s", "[1, 2]"]


def test_an_output_path_that_does_not_resolve_is_the_text_none_like_the_legacy_runner():
    """Legacy ``get_json_value``/``output_json_data``: ``str(None)`` is stored for a dotted path that fails and for a JSON null; a plain absent key stores nothing.
    Sheets rely on it (``MedScen01_EXP = None`` for a policy without medical conditions)."""
    data = {"a": {"b": [{"c": 1}], "n": None, "e": "", "z": 0}, "top": None}
    assert [output_json_text(data, p) for p in ("a.b.[0].c", "a.b.[1].c", "a.q", "a/b/[3]/c", "a.n", "top", "nope", "a.e", "a.z")] == \
        ["1", "None", "None", "None", "None", "None", "", "", "0"]


@pytest.mark.parametrize("kind,expected,actual,ok", [
    ("compare", "Success", " Success ", True), ("compare", "Success", "success", False), ("compare_ignore_case", "Active", "ACTIVE", True),
    ("compare", 123.5, "123.50", True), ("compare", 123.5, "123.6", False), ("compare", 5, "five", False),
    ("contains", "Cover", "Cover not available", True), ("contains", "cover", "Cover not available", False),
    ("contains_ignore_case", "cover", "Cover not available", True), ("compare", "", "", True), ("compare", None, "", True),
])
def test_the_checks_give_the_legacy_verdicts(kind, expected, actual, ok):
    assert check(kind, expected, actual) is ok


def test_templates_are_looked_for_on_the_l_drive_then_in_the_folders_from_config_and_secrets_env_then_beside_the_workbook(tmp_path):
    tried = template_candidates("L:\\Regression\\Templates\\Policy", "search.json", ["", "/data/api", "  "], tmp_path)
    assert tried[0] == Path("L:\\Regression\\Templates\\Policy") / "search.json"
    assert tried[-4:] == [Path("/data/api/search.json"), Path("/data/api/Policy/search.json"), tmp_path / "templates" / "search.json", tmp_path / "search.json"]
    assert len(tried) in (5, 6)                                                            # + the same L: path with / for a Mac / Linux machine
    assert template_candidates("", "search.json", [], tmp_path) == [tmp_path / "templates" / "search.json", tmp_path / "search.json"]


def test_an_api_data_row_is_a_test_and_its_request_and_checks_are_read_from_the_sheets(tmp_path, site):
    wb = workbook(tmp_path, site)
    (case,) = wb.discover()
    assert (case.id, case.kind, case.is_ui, case.runnable, case.enabled, case.param_row, case.scenario) == ("Policy", "api", False, True, True, 2, "Policy Search w EMC")
    assert wb.warnings == []                                                              # an API sheet is not "skipped" any more
    rt = wb.runtime(case)
    req = rt.request()
    assert (req.method, req.url) == ("GET", site + "policy/v4/P-100")                      # VLOOKUP on the Environment table
    assert req.headers == {"Authorization": "Bearer TOKEN-abc", "apiKey": "KEY-123", "Content-Type": "application/json"}
    assert req.secret_headers == {"Authorization", "apiKey"}                               # never written anywhere
    points = [(c.kind, c.expected_column, c.actual_column) for c in rt.checkpoints()]
    assert points == [("compare", "transactionStatus_EXP", "transactionStatus_OUT"), ("compare", "MedScen01_EXP", "MedScen01_OUT"),
                      ("compare", "Premium_EXP", "Premium_OUT"), ("compare_ignore_case", "PolicyStatus_EXP", "PolicyStatus_OUT"),
                      ("contains", "Covered_EXP", "Covered_OUT")]                        # compareX = off; the row without a function name is ignored
    assert [r for r, _ in rt.io.ignored_rows] == [19]
    assert [(m, n) for _, n, m in rt.plan()][:2] == [("GET", f"GET {site}policy/v4/P-100"), ("COMPARE", "Check transactionStatus_OUT")]


def test_disable_checkpoint_switches_off_the_named_checks_for_that_test_name_only(tmp_path, site):
    io = [r for r in IO if r[0] != "DisableCheckpoint"] + [("DisableCheckpoint", "Policy Search w EMC", "PolicyStatus_OUT;Premium_OUT")]
    wb = workbook(tmp_path, site, io_rows=io)
    names = [c.actual_column for c in wb.runtime(wb.discover()[0]).checkpoints()]
    assert "PolicyStatus_OUT" not in names and "Premium_OUT" not in names and "transactionStatus_OUT" in names
    other = workbook(tmp_path, site, io_rows=[r for r in IO if r[0] != "DisableCheckpoint"] + [("DisableCheckpoint", "Some other test", "PolicyStatus_OUT")])
    assert "PolicyStatus_OUT" in [c.actual_column for c in other.runtime(other.discover()[0]).checkpoints()]


def test_outputs_are_written_back_into_the_row_so_formulas_can_read_them(tmp_path, site):
    wb = workbook(tmp_path, site)
    rt = wb.runtime(wb.discover()[0])
    rt.record_output("transactionStatus_OUT", "Success")
    assert rt.actual("transactionStatus_OUT") == "Success" and rt.actual("MedScen01_OUT") == ""
    assert rt.sheet.read(rt.row, rt.columns["TRANSACTIONSTATUS_OUT"]) == "Success"


def test_values_the_ui_tests_produced_reach_the_api_row_through_the_shared_cells(tmp_path, site):
    path = build_api_workbook(tmp_path / "api.xlsx", site, ui_flow=True, url_formula='=Params_1!F2&"policy/v4/"&Params_1!T2')
    wb = Workbook(path, seed=1)
    api = [c for c in wb.discover() if c.kind == "api"][0]
    assert wb.runtime(api).request().url.endswith("policy/v4/")                              # no UI test has run
    shared = {("Params_1", 2, 20): "P-100"}                                                   # what the UI test's Output step wrote into Params_1!T2
    assert wb.api_runtime(api, shared).request().url.endswith("policy/v4/P-100")


def test_lint_reports_what_would_go_wrong_in_an_api_row(tmp_path, site):
    wb = workbook(tmp_path, site, url_formula="=WEEKNUM(TODAY())", method="FETCH")
    said = [(f.severity, f.message) for f in lint(wb)]
    assert any(sev == "error" and "WEBSERVICE_URL cell evaluates to #NAME?" in m and "WEEKNUM" in m for sev, m in said)
    assert any("FETCH" in m for _, m in said)
    io_rows = [f for f in lint(wb) if f.test == "InputOutput"]
    assert len(io_rows) == 1 and "no function name" in io_rows[0].message and io_rows[0].row == 19          # said once, not per test
    ok = lint(workbook(tmp_path, site))
    assert not [f for f in ok if f.severity == "error"]


def test_api_tests_are_listed_planned_and_counted_like_the_others(tmp_path, site):
    build_workbook(tmp_path / "both.xlsx", flows=["Flow"], base_url=site, enabled=["Flow"])
    wb = workbook(tmp_path, site)
    cfg = Config()
    tests = insight.list_tests(wb, cfg)
    assert [(t["id"], t["kind"], t["steps"], t["enabled"], t["runnable"]) for t in tests] == [("Policy", "api", 6, True, True)]
    plan = insight.plan_steps(wb, cfg, "Policy")
    assert plan["count"] == 6 and plan["steps"][0]["action"] == "GET" and plan["steps"][1]["name"] == "Check transactionStatus_OUT"


# -- a real run against the mock API ---------------------------------------------------------------------------------------------------------
def run(tmp_path: Path, wb_path: Path, extra_cfg: str = "", env: dict | None = None, timeout: int = 120, args: list | None = None, stdin: str | None = None):
    cfg = tmp_path / "config.yaml"
    cfg.write_text("runner: {stagger_s: 0, min_page_load_gap_s: 0, block_cooldown_s: 0.5, block_retries: 2}\ntimeouts: {element_s: 6}\n"
                   "captcha_bypass: {values: {UAT: TEST-UAT-TOKEN}}\n" + extra_cfg)
    proc = subprocess.run([sys.executable, "-m", "regrunner", "--config", str(cfg), "run", str(wb_path), "--plain", "--seed", "1", *(args or [])],
                          cwd=tmp_path, capture_output=True, text=True, timeout=timeout, env={**os.environ, **(env or {})}, input=stdin)
    (run_dir,) = sorted((tmp_path / "runs").iterdir())[-1:]
    results = json.loads((run_dir / "results.json").read_text())
    return proc, results, run_dir


@pytest.mark.browser
def test_a_get_test_sends_the_headers_reads_the_response_and_checks_it(tmp_path, site):
    path = build_api_workbook(tmp_path / "api.xlsx", site)
    proc, results, run_dir = run(tmp_path, path)
    assert proc.returncode == 0, proc.stdout + proc.stderr
    (test,) = results["tests"]
    assert test["id"] == "Policy" and test["status"] == "PASSED" and test["passed"] == 6 and test["failed"] == 0
    steps = test["steps"]
    assert steps[0]["action"] == "GET" and steps[0]["name"].startswith("GET http") and "HTTP 200" in " ".join(steps[0]["notes"])
    checks = {s["name"]: s for s in steps[1:]}
    assert checks["Check MedScen01_OUT"]["expected"] == "Cover not available" and checks["Check MedScen01_OUT"]["actual"] == "Cover not available"
    assert checks["Check Premium_OUT"]["actual"] == "123.5" and checks["Check PolicyStatus_OUT"]["actual"] == "Active"     # ignore-case compare
    seen = mock_server.API_SEEN
    assert seen["path"] == "/policy/v4/P-100" and seen["headers"]["apikey"] == "KEY-123" and seen["headers"]["authorization"] == "Bearer TOKEN-abc"
    assert seen["headers"]["content-type"] == "application/json"
    evidence = run_dir / "tests" / "Policy"
    assert json.loads((evidence / "response.json").read_text())["transactionStatus"] == "Success"
    request_text = (evidence / "request.txt").read_text()
    assert "GET " in request_text and "KEY-123" not in request_text and "TOKEN-abc" not in request_text and "••••••" in request_text
    for name in ("results.json", "events.jsonl", "report.html", "run.json"):
        text = (run_dir / name).read_text(errors="ignore")
        assert "KEY-123" not in text and "TOKEN-abc" not in text, name                    # the secrets are in no file the run wrote
    assert "KEY-123" not in proc.stdout + proc.stderr and "TOKEN-abc" not in proc.stdout + proc.stderr
    assert "Policy: step 6/6 passed" in proc.stdout


@pytest.mark.browser
def test_a_path_the_response_does_not_have_is_none_and_a_sheet_that_expects_none_passes(tmp_path, site):
    """PolicySearch "without EMC": the response has fewer customElements than the sheet reads; the sheet expects the text None."""
    rows = [(f, "detailResponse.policyDetail.travelers.[0].customElements.[5].value" if v == "MedScen01_OUT" else p, v) if f == "output_json" else (f, p, v)
            for f, p, v in IO]
    path = build_api_workbook(tmp_path / "api.xlsx", site, exp={"MedScen01_EXP": "None"}, io_rows=rows)
    proc, results, _ = run(tmp_path, path)
    assert proc.returncode == 0, proc.stdout + proc.stderr
    (test,) = results["tests"]
    check_step = next(s for s in test["steps"] if s["name"] == "Check MedScen01_OUT")
    assert test["status"] == "PASSED" and check_step["status"] == "PASSED" and check_step["expected"] == "None" and check_step["actual"] == "None"


@pytest.mark.browser
def test_a_wrong_expected_value_fails_that_check_and_says_what_it_got(tmp_path, site):
    path = build_api_workbook(tmp_path / "api.xlsx", site, exp={"MedScen01_EXP": "Covered"})
    proc, results, _ = run(tmp_path, path)
    assert proc.returncode == 1
    (test,) = results["tests"]
    assert test["status"] == "FAILED" and test["failed"] == 1 and test["passed"] == 5
    (bad,) = [s for s in test["steps"] if s["status"] == "FAILED"]
    assert bad["name"] == "Check MedScen01_OUT" and bad["expected"] == "Covered" and bad["actual"] == "Cover not available" and bad["error"] == "Comparison Failed"


@pytest.mark.browser
def test_missing_credentials_show_up_as_the_apis_own_answer_and_the_env_file_can_supply_them(tmp_path, site):
    path = build_api_workbook(tmp_path / "api.xlsx", site, api_key="", bearer="")
    proc, results, _ = run(tmp_path, path)
    (test,) = results["tests"]
    step = test["steps"][0]
    assert test["status"] == "ERROR" and step["status"] == "FAILED" and len(test["steps"]) == 1              # the API refused: nothing after it can be read
    assert "refused the request: HTTP 401" in step["error"] and "Unauthorized" in step["error"]               # what the API itself said
    assert "Headers sent: Authorization (column DT_bearerToken), apiKey (column DT_apiKey)" in step["error"]   # which credentials, from which columns
    assert "RR_VAR_DT_APIKEY" in step["error"] and "recaptcha" not in step["error"].lower()                    # where to put them; nothing about pages

    proc, results, _ = run(tmp_path, build_api_workbook(tmp_path / "api2.xlsx", site, api_key="", bearer=""),
                           env={"RR_VAR_DT_APIKEY": "KEY-123", "RR_VAR_DT_BEARERTOKEN": "Bearer TOKEN-fromenv"})
    assert results["tests"][0]["status"] == "PASSED"                                                       # RR_VAR_<COLUMN> in secrets.env supplies them
    assert mock_server.API_SEEN["headers"]["authorization"] == "Bearer TOKEN-fromenv"


@pytest.mark.browser
def test_a_cloudfront_403_is_named_and_gets_one_retry_when_nothing_ever_loaded(tmp_path, site):
    path = build_api_workbook(tmp_path / "api.xlsx", site, url_tail="cloudfront")
    before = mock_server.API_SEEN["count"]
    proc, results, run_dir = run(tmp_path, path)
    (test,) = results["tests"]
    assert proc.returncode == 1 and test["status"] == "ERROR" and test["blocked"]
    assert "HTTP 403 from CloudFront" in test["error"] and "standing access rule" in test["error"]
    assert len(test["attempts"]) == 1 and mock_server.API_SEEN["count"] - before == 2                       # one retry, not two
    events = [json.loads(line) for line in (run_dir / "events.jsonl").read_text().splitlines()]
    assert [e["type"] for e in events if e["type"] == "run_paused"] == ["run_paused"]


@pytest.mark.browser
def test_a_post_reads_its_template_from_the_workbooks_location_or_the_fallback_folder_and_fills_it_in(tmp_path, site):
    templates = tmp_path / "templates-elsewhere"
    templates.mkdir()
    (templates / "search.json").write_text(json.dumps({"searchCriteria": {"searchParameters": [{"value": ""}, {"value": ""}]}}), encoding="utf-8-sig")
    path = build_api_workbook(tmp_path / "api.xlsx", site, method="POST", url_tail="policy/v4/search", request_file="search.json",
                              location="L:\\Regression\\Templates\\Policy", exp={"MedScen01_EXP": ""}, io_rows=[
                                  r for r in IO if r[1] not in ("MedScen01_EXP", "AgentEmail_EXP", "PolicyStatus_EXP", "Premium_EXP", "Covered_EXP")
                                  and not (r[0] == "output_json" and "detailResponse" in str(r[1]))] + [("compare", "policyNumber_IN", "Echo_OUT")])
    proc, results, _ = run(tmp_path, path, env={"RR_API_TEMPLATES_DIR": str(templates)})            # the L: path is not there: secrets.env's folder is used
    assert proc.returncode == 0, proc.stdout + proc.stderr
    (test,) = results["tests"]
    assert test["status"] == "PASSED" and "search.json" in " ".join(test["steps"][0]["notes"])
    assert mock_server.API_SEEN["body"] == {"searchCriteria": {"searchParameters": [{"value": "P-100"}, {"value": "Smith"}]}}      # update_json filled it in
    assert mock_server.API_SEEN["headers"]["content-type"] == "application/json"

    proc, results, _ = run(tmp_path, build_api_workbook(tmp_path / "api3.xlsx", site, method="POST", request_file="nowhere.json", location="L:\\x"))
    (test,) = results["tests"]
    assert test["status"] == "ERROR" and "nowhere.json" in test["error"] and "L:\\x" in test["error"] and "RR_API_TEMPLATES_DIR" in test["error"]


@pytest.mark.browser
def test_api_tests_wait_for_the_ui_tests_and_read_what_they_wrote(tmp_path, site):
    """The URL only resolves to a policy that exists when the UI test's output (the page address it captured) is in Params_1!T2 first."""
    path = build_api_workbook(tmp_path / "wb.xlsx", site, ui_flow=True,
                              url_formula=f'="{site}policy/v4/"&IF(ISNUMBER(SEARCH("http",Params_1!T2)),"P-100","MISSING")')
    proc, results, run_dir = run(tmp_path, path, "runner: {workers: 3, stagger_s: 0, min_page_load_gap_s: 0}\n", timeout=180)
    assert proc.returncode == 0, proc.stdout + proc.stderr
    by_id = {t["id"]: t for t in results["tests"]}
    assert by_id["Flow"]["status"] == "PASSED" and by_id["Policy"]["status"] == "PASSED"
    order = [(e["type"], e.get("test")) for e in map(json.loads, (run_dir / "events.jsonl").read_text().splitlines())
             if e["type"] in ("test_started", "test_finished")]
    assert order.index(("test_finished", "Flow")) < order.index(("test_started", "Policy"))      # the API test started after the UI test ended


@pytest.mark.browser
def test_the_apis_own_403_is_a_refusal_with_its_message_and_is_not_waited_out_and_retried(tmp_path, site):
    """The real run: api.travelguard.com answered 403 '{"message": "Invalid key=value pair ... Authorization header"}'.  It was treated like a firewall block:
    a 60 s cool-down, a second attempt, a 'the page did not load' message.  It is the API talking: report it at once, with what it said."""
    path = build_api_workbook(tmp_path / "api.xlsx", site, url_tail="policy/v4/")
    before = mock_server.API_SEEN["count"]
    started = time.monotonic()
    proc, results, run_dir = run(tmp_path, path, extra_cfg="runner: {block_cooldown_s: 30}\n")
    (test,) = results["tests"]
    assert time.monotonic() - started < 25 and mock_server.API_SEEN["count"] - before == 1                    # no 30 s pause, no second request
    assert test["status"] == "ERROR" and not test["blocked"] and test["attempts"] == [] and len(test["steps"]) == 1
    error = test["steps"][0]["error"]
    assert error.startswith("127.0.0.1:") and "refused the request: HTTP 403" in error and "Invalid key=value pair" in error
    assert "Headers sent: Authorization (column DT_bearerToken), apiKey (column DT_apiKey)" in error and "RR_VAR_DT_BEARERTOKEN" in error
    assert "The URL ends in '/'" in error and "policy number" in error                                            # the likely reason, named
    assert "page did not load" not in error and "recaptcha" not in error.lower()
    assert "KEY-123" not in error and "TOKEN-abc" not in error
    events = [json.loads(line) for line in (run_dir / "events.jsonl").read_text().splitlines()]
    assert not [e for e in events if e["type"] == "run_paused"]
    assert "HTTP 403 in" in " ".join(test["steps"][0]["notes"]) and int(test["steps"][0]["duration_ms"]) < 20000


EMPTY_URL = '=VLOOKUP(Global!$B$2,Global!$E$1:$G$3,3,FALSE)&"policy/v4/"&K2'                                       # ...&policyNumber_IN, like the real sheet


@pytest.mark.browser
def test_a_request_built_from_an_empty_cell_is_not_sent_and_says_which_cell(tmp_path, site):
    path = build_api_workbook(tmp_path / "api.xlsx", site, url_formula=EMPTY_URL, policy_in=None)
    before = mock_server.API_SEEN["count"]
    proc, results, _ = run(tmp_path, path)
    (test,) = results["tests"]
    assert mock_server.API_SEEN["count"] == before                                                             # nothing reached the API
    step = test["steps"][0]
    assert test["status"] == "ERROR" and step["status"] == "FAILED" and len(test["steps"]) == 1
    assert "Not sent: policyNumber_IN is empty" in step["error"] and "would go to " in step["error"] and step["error"].count("/policy/v4/") >= 1
    assert "Start the run from the web UI or a terminal" in step["error"]


@pytest.mark.browser
def test_at_a_terminal_the_empty_cell_is_asked_for_and_the_request_goes_out_with_it(tmp_path, site):
    path = build_api_workbook(tmp_path / "api.xlsx", site, url_formula=EMPTY_URL, policy_in=None)
    proc, results, run_dir = run(tmp_path, path, args=["--ask", "terminal"], stdin="P-100\n")
    (test,) = results["tests"]
    assert test["status"] == "PASSED", proc.stdout + proc.stderr
    assert "policyNumber_IN is empty for Policy" in proc.stdout and mock_server.API_SEEN["path"] == "/policy/v4/P-100"
    (var,) = [v for v in test["variables"] if v["by_hand"]]                                                    # (the API's own outputs are recorded as values set too)
    assert var["name"] == "policyNumber_IN" and var["stored"] == "P-100"


@pytest.mark.browser
def test_a_test_that_expects_a_status_may_send_an_empty_input_and_get_a_refusal(tmp_path, site):
    """A negative test (RES_STATUS_CD_EXP with a real compare, not compareX) is about the refusal: nothing is stopped or rewritten."""
    io = [(("compare" if row[1] == "RES_STATUS_CD_EXP" else row[0]),) + tuple(row[1:]) if row[1] == "RES_STATUS_CD_EXP" else row for row in IO]
    path = build_api_workbook(tmp_path / "api.xlsx", site, url_formula=EMPTY_URL, policy_in=None, io_rows=io, exp={"RES_STATUS_CD_EXP": "403"})
    before = mock_server.API_SEEN["count"]
    proc, results, _ = run(tmp_path, path)
    (test,) = results["tests"]
    assert mock_server.API_SEEN["count"] - before == 1                                                         # it was sent, empty, on purpose
    assert test["steps"][0]["status"] == "PASSED" and next(s for s in test["steps"] if s["name"] == "Check RES_STATUS_CD_OUT")["status"] == "PASSED"
