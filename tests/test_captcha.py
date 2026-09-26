"""A captcha challenge is a person's job: the test stops where it meets one (nothing behind it is a result), and when somebody can answer,
the test is run again in a visible browser window that waits for them to solve it."""
from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest
from playwright.async_api import async_playwright

from regrunner.engine import captcha
from regrunner.engine.ask import write_answer
from regrunner.engine.runner import RunOptions, execute
from regrunner.events import EventBus
from tests.workbook_factory import Sheet

pytestmark = pytest.mark.browser


def build_captcha_workbook(path: Path, site: str, query: str = "", wait_s: int = 1) -> Path:
    """Open the login page, type a policy number, click Login (the captcha pops up), wait, expect the welcome text; then two steps that
    would 'pass' whatever happened before them (Exist on the heading, another Wait) - what a test must not report as passed behind a captcha."""
    wb = openpyxl.Workbook()
    ds = wb.active
    ds.title = "DataSheets"
    ds.append(["SheetsToExecute", "blnExecute", "ParameterSheet", "Comments"])
    ds.append(["Flow", "Y", "Params_1", "captcha"])
    g = wb.create_sheet("Global")
    g.append(["Parameter", "Value"])
    g.append(["Environment", "UAT"])
    ps = wb.create_sheet("Params_1")
    ps.append(["blnExecute", "DT_URL"])
    ps.append(["Y", f"{site}captcha.html{query}"])
    s = Sheet(wb.create_sheet("Flow"), "Y")
    s.add("Open", "Open", Page="chrome", Value="DT_URL")
    s.add("Set", "Policy number", FindBy="xpath", FindBy_Value="//input[@id='policy']", Index=0, Value="P-100")
    s.add("Click", "Click Login", FindBy="xpath", FindBy_Value="//button[@id='login']", Index=0)
    s.add("Wait", "Wait", Value=wait_s)
    s.add("Output", "Welcome shown", FindBy="xpath", FindBy_Value="//div[@id='welcome']", Index=0, Output_Property="innertext",
          Expected_Value="Welcome back", Exact_Match="Y")
    s.add("Exist", "Heading is there", FindBy="xpath", FindBy_Value="//h1", Index=0)
    s.add("Wait", "Wait again", Value=1)
    s.add("Quit", "Quit")
    wb.save(path)
    return path


async def run(cfg, wb: Path, ask: str = "off", listen=None):
    bus = EventBus()
    events: list[dict] = []
    bus.subscribe(events.append)
    if listen:
        bus.subscribe(listen(cfg))
    result = await execute(RunOptions(workbook=wb, tests=["Flow"], seed=1, extra={"ask": ask}), cfg, bus)
    return result, events, result.tests[0]


def names(test) -> list[str]:
    return [s.name for s in test.steps]


def gives_up(cfg):
    """A person who looks at the captcha and clicks 'Skip: fail this test'."""
    def listen(event):
        if event["type"] == "user_input_needed" and event.get("kind") == "captcha":
            write_answer(cfg.path(cfg.runs_dir) / event["run_id"], event["ask"], "skip")
    return listen


# -- finding it --------------------------------------------------------------------------------------------------------------------
async def test_only_a_challenge_on_screen_counts_not_the_badge_or_the_hidden_frame(site):
    async with async_playwright() as pw:
        browser = await pw.chromium.launch()
        try:
            page = await browser.new_page()
            await page.goto(site + "captcha.html")
            assert [f.url for f in page.frames if "recaptcha" in f.url]                        # both frames are in the page
            assert await captcha.find([page]) is None                                          # badge visible, challenge hidden: nothing to solve
            await page.click("#login")
            hit = await captcha.find([page])
            assert hit is not None and hit.kind == "reCAPTCHA"
            assert hit.prompt == "Select all images with a fire hydrant"                       # what it asks, without the 'click verify' tail
            assert hit.asks() == 'reCAPTCHA is asking "Select all images with a fire hydrant"'
            await page.frame_locator("iframe[src*='bframe']").locator("#verify").click()       # somebody solves it
            await page.wait_for_timeout(300)                                                   # ...and the page reacts to the message
            assert await captcha.find([page]) is None
            assert await captcha.find([]) is None
        finally:
            await browser.close()


# -- stopping ----------------------------------------------------------------------------------------------------------------------
async def test_a_test_stops_at_the_captcha_and_nothing_after_it_is_reported_as_passed(site, make_cfg, tmp_path):
    cfg = make_cfg()
    result, events, test = await run(cfg, build_captcha_workbook(tmp_path / "a.xlsx", site))
    assert test.status == "ERROR" and result.status == "FAILED"
    assert names(test) == ["Open", "Policy number", "Click Login"]                                # it went no further
    assert [s.status for s in test.steps] == ["PASSED", "PASSED", "FAILED"]                       # the step the captcha came after is the failed one
    assert test.skipped == 5 and test.passed == 2 and test.failed == 1
    assert 'reCAPTCHA is asking "Select all images with a fire hydrant"' in test.error and 'after step 3 "Click Login"' in test.error
    assert test.error == test.captcha == test.steps[-1].error
    assert "the steps after it were not run" in test.error
    assert "bypass cookie was sent, but the site showed the captcha anyway" in test.error         # the token was there: the bypass does not work here
    assert "browser window opens" in test.error                                                  # nobody to ask in this run: says how to solve it by hand
    assert test.attempts == [] and not [e for e in events if e["type"] == "user_input_needed"]
    assert [e["waiting"] for e in events if e["type"] == "captcha_detected"] == [False]


async def test_with_no_bypass_token_the_message_says_a_captcha_is_expected(site, make_cfg, tmp_path):
    cfg = make_cfg(**{"captcha_bypass.values": {}})
    _, _, test = await run(cfg, build_captcha_workbook(tmp_path / "a.xlsx", site))
    assert test.status == "ERROR"
    assert "No recaptchaBypassToken is configured for UAT, so a captcha is expected" in test.error and "RECAPTCHA_BYPASS_TOKEN_UAT" in test.error


async def test_a_captcha_that_pops_up_after_a_spinner_is_caught_at_the_step_it_appeared_during(site, make_cfg, tmp_path):
    cfg = make_cfg(**{"waits.mode": "legacy"})                                                    # a real 2 s sleep: the challenge pops up during it
    _, _, test = await run(cfg, build_captcha_workbook(tmp_path / "a.xlsx", site, "?after=600", wait_s=2))
    assert test.status == "ERROR" and 'after step 4 "Wait"' in test.error
    assert [s.status for s in test.steps] == ["PASSED", "PASSED", "PASSED", "FAILED"] and "Heading is there" not in names(test)


async def test_a_step_that_failed_because_of_the_captcha_names_the_captcha_not_a_missing_element(site, make_cfg, tmp_path):
    """The real run: the step after the click looked for an element the captcha was covering and reported 'Object was not found'."""
    cfg = make_cfg(**{"waits.mode": "legacy"})
    wb = build_captcha_workbook(tmp_path / "a.xlsx", site, "?after=600", wait_s=0)                # no wait: the next step is the Output
    _, _, test = await run(cfg, wb)
    assert test.status == "ERROR" and "captcha" in test.error and "Comparison Failed" not in test.error
    assert test.steps[-1].status == "FAILED" and test.steps[-1].error == test.error


async def test_a_working_bypass_means_no_captcha_and_the_badge_alone_is_not_one(site, make_cfg, tmp_path):
    cfg = make_cfg()
    _, events, test = await run(cfg, build_captcha_workbook(tmp_path / "a.xlsx", site, "?bypass=1"))
    assert test.status == "PASSED" and len(test.steps) == 8 and test.captcha == ""
    assert not [e for e in events if e["type"] == "captcha_detected"]


async def test_detection_can_be_switched_off(site, make_cfg, tmp_path):
    cfg = make_cfg(**{"captcha.detect": False})
    _, _, test = await run(cfg, build_captcha_workbook(tmp_path / "a.xlsx", site))
    assert test.status == "FAILED" and len(test.steps) == 8 and test.captcha == ""                 # the old behaviour: carries on behind it


async def test_solve_fail_ends_the_test_without_opening_a_window(site, make_cfg, tmp_path):
    cfg = make_cfg(**{"captcha.solve": "fail"})
    _, events, test = await run(cfg, build_captcha_workbook(tmp_path / "a.xlsx", site), ask="ui")
    assert test.status == "ERROR" and test.attempts == [] and "browser window opens" not in test.error
    assert not [e for e in events if e["type"] == "user_input_needed"]


# -- a person solves it ------------------------------------------------------------------------------------------------------------
async def test_when_somebody_can_answer_the_test_runs_again_in_a_window_and_carries_on_once_the_captcha_is_solved(site, make_cfg, tmp_path):
    cfg = make_cfg(**{"captcha.headless": True})                                                  # the window is visible for people, headless for the suite
    result, events, test = await run(cfg, build_captcha_workbook(tmp_path / "a.xlsx", site, "?person=2500"), ask="ui")
    assert test.status == "PASSED", [(s.name, s.error) for s in test.steps if s.status != "PASSED"]
    assert len(test.steps) == 8 and test.captcha == ""
    (first,) = test.attempts                                                                      # the headless attempt that met the captcha
    assert first["captcha"] is True and "captcha stopped the test after step 3" in first["error"]
    click = next(s for s in test.steps if s.name == "Click Login")
    assert click.status == "PASSED" and any("solved by hand" in n for n in click.notes)
    detected = [e for e in events if e["type"] == "captcha_detected"]
    assert [e["waiting"] for e in detected] == [False, True]                                       # first stopped, then waited
    (asked,) = [e for e in events if e["type"] == "user_input_needed"]
    assert asked["kind"] == "captcha" and asked["test"] == "Flow" and "Solve it in the browser window that opened" in asked["question"]
    assert 'reCAPTCHA is asking "Select all images with a fire hydrant"' in asked["question"] and asked["timeout_s"] == 300
    assert [e["reason"] for e in events if e["type"] == "user_input_closed"] == ["done"]
    assert result.status == "PASSED"


async def test_skipping_the_captcha_ends_the_test_with_the_reason(site, make_cfg, tmp_path):
    cfg = make_cfg(**{"captcha.headless": True})
    _, events, test = await run(cfg, build_captcha_workbook(tmp_path / "a.xlsx", site), ask="ui", listen=gives_up)
    assert test.status == "ERROR" and "It was skipped, so the test was stopped here" in test.error
    assert names(test) == ["Open", "Policy number", "Click Login"] and test.steps[-1].status == "FAILED"
    assert len(test.attempts) == 1 and len([e for e in events if e["type"] == "user_input_needed"]) == 1        # asked once, not again and again


async def test_nobody_solving_it_within_the_time_ends_the_test(site, make_cfg, tmp_path):
    cfg = make_cfg(**{"captcha.headless": True, "captcha.timeout_s": 2})
    _, events, test = await run(cfg, build_captcha_workbook(tmp_path / "a.xlsx", site), ask="ui")
    assert test.status == "ERROR" and "Nobody solved it within 2 s" in test.error
    assert names(test)[-1] == "Click Login" and "Heading is there" not in names(test)
    assert [e["reason"] for e in events if e["type"] == "user_input_closed"] == ["timeout"]


# -- at a terminal -----------------------------------------------------------------------------------------------------------------
def test_a_terminal_run_prints_the_captcha_notice_and_carries_on_when_it_is_solved(site, tmp_path):
    wb = build_captcha_workbook(tmp_path / "a.xlsx", site, "?person=2500")
    cfg = tmp_path / "config.yaml"
    cfg.write_text("timeouts: {element_s: 6}\ncaptcha_bypass: {values: {UAT: TEST-UAT-TOKEN}}\nreports: {html: false}\ncaptcha: {headless: true}\n")
    proc = subprocess.run([sys.executable, "-m", "regrunner", "--config", str(cfg), "run", str(wb), "--seed", "1", "--ask", "terminal"],
                          cwd=tmp_path, capture_output=True, text=True, input="", timeout=180)
    (run_dir,) = sorted((tmp_path / "runs").iterdir())[-1:]
    test = json.loads((run_dir / "results.json").read_text())["tests"][0]
    assert proc.returncode == 0, proc.stdout + proc.stderr
    assert test["status"] == "PASSED" and test["attempts"][0]["captcha"] is True
    assert "Solve it in the browser window that opened" in proc.stdout


def test_a_run_started_from_a_script_stops_at_the_captcha_without_waiting_for_anyone(site, tmp_path):
    wb = build_captcha_workbook(tmp_path / "a.xlsx", site)
    cfg = tmp_path / "config.yaml"
    cfg.write_text("timeouts: {element_s: 6}\ncaptcha_bypass: {values: {UAT: TEST-UAT-TOKEN}}\nreports: {html: false}\ncaptcha: {headless: true}\n")
    proc = subprocess.run([sys.executable, "-m", "regrunner", "--config", str(cfg), "run", str(wb), "--seed", "1", "--ask", "off", "--plain"],
                          cwd=tmp_path, capture_output=True, text=True, timeout=120)
    (run_dir,) = sorted((tmp_path / "runs").iterdir())[-1:]
    test = json.loads((run_dir / "results.json").read_text())["tests"][0]
    assert proc.returncode == 1 and test["status"] == "ERROR" and test["captcha"] and len(test["steps"]) == 3
