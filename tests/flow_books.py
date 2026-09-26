"""Small workbooks for the flow keywords (SET_VARIABLE, IF, loops, CALL_TEST, JSON_READ) and the environment table: each test is a list of rows."""
from __future__ import annotations

from pathlib import Path

import openpyxl

from tests.workbook_factory import Sheet

Row = tuple  # (method, step name, {column: value}) or (method, step name, {column: value}, blnExecute)


def at(xpath: str) -> dict:
    return {"FindBy": "xpath", "FindBy_Value": xpath, "Index": 0}


def book(path: Path, tests: dict[str, list[Row]], *, params: dict[str, list[dict]] | None = None, enabled: list[str] | None = None,
         environment: str = "UAT", environments: list[list] | None = None, sheets: dict[str, list[list]] | None = None) -> Path:
    """``tests``: {sheet: rows}.  ``params``: {sheet: [{column: value}, ...]} (one Params row each, blnExecute Y unless given).
    ``environments``: the ``_rr_environments`` table (header row first).  ``sheets``: more plain sheets (loop data...), rows as lists."""
    params = params or {}
    enabled = list(tests) if enabled is None else enabled
    wb = openpyxl.Workbook()
    g = wb.active
    g.title = "Global"
    g.append(["Parameter", "Value"])
    g.append(["Environment", environment])
    ds = wb.create_sheet("DataSheets")
    ds.append(["SheetsToExecute", "blnExecute", "ParameterSheet", "Comments"])
    for name, rows in tests.items():
        own = params.get(name)
        ds.append([name, "Y" if name in enabled else "N", f"{name}_Params" if own else "", ""])
        sheet = Sheet(wb.create_sheet(name), "Y")
        for method, step, cols, *flag in rows:
            sheet.add(method, step, gate=flag[0] if flag else None, **cols)
        if own:
            ps = wb.create_sheet(f"{name}_Params")
            headers = ["blnExecute", *dict.fromkeys(k for r in own for k in r if k != "blnExecute")]
            ps.append(headers)
            for r in own:
                ps.append([r.get("blnExecute", "Y"), *[r.get(h) for h in headers[1:]]])
    for name, rows in (sheets or {}).items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    if environments:
        ws = wb.create_sheet("_rr_environments")
        ws.sheet_state = "hidden"
        for r in environments:
            ws.append(r)
    wb.save(path)
    return path
