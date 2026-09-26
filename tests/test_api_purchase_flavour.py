"""The older API layout: ``Environment_Parameter`` (URL from the Environments sheet) + a request template filled by ``Replace`` rows + ``Output`` rows
that read the response, and UI tests that read what the API test filled in (``=PurchaseUS!FB3``).  The real ``PurchaseUS`` sheet of the Claims US /
Expedia workbook is this layout: it used to be skipped as "not a UI keyword sheet", so nothing produced the policy number the others read."""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

from regrunner.engine.order import analyse, plan_order
from regrunner.workbook import Workbook
from regrunner.workbook.api_template import apply_replacements, json_to_tree, output_values, xml_to_tree
from tests.site import server as mock_server
from tests.test_api_tests import run
from tests.workbook_factory import build_workbook

COLUMNS = ["blnExecute", "TCID", "TC_Name", "Environment_Parameter", "XML_LOCATION", "XML_REQUESTFILE", "XML_RESPONSEFILE", "JSON_FORMAT", "DT_apiKey",
           "LastName_IN", "Qty_IN", "Flag_IN", "SiteUrl_OUT", "PolicyNumber_OUT", "LastName_OUT", "Second_OUT", "Status_OUT", "Case_OUT"]
COL = {name: i for i, name in enumerate(COLUMNS, start=1)}
TEMPLATE = {"partner": "partner_value", "lastName": "lastName_value", "trip": {"qty": "qty_value", "primary": "flag_value"}, "untouched": "other_value"}
IO = [("Function", "Parameter", "Value"),
      ("addHeader", "apiKey", "DT_apiKey"),
      ("Replace", "lastName_value", "LastName_IN"), ("Replace", "qty_value", "Qty_IN"), ("Replace", "flag_value", "Flag_IN"),
      ("Replace", "partner_value", "Partner_IN"),                                    # the sheet has no such column: left as it is
      ("x", "qty_value", "Qty_IN"),                                                  # "x" switches a row off
      ("Output", "siteurl", "SiteUrl_OUT"), ("Output", "transactionstatus", "Status_OUT"),
      ("Output", "policyresponses/policydetail/policynumber", "PolicyNumber_OUT"),   # two policies: joined with ;
      ("Output", "policyresponses/policydetail[1]/policyholder/lastname", "LastName_OUT"),
      ("Output", "policyholder/lastname/index:=1", "Second_OUT"),
      ("Output", "policyResponses/policyDetail/policyNumber", "Case_OUT")]           # not lower case: nothing matches when the response is read as HTML


def build(tmp_path: Path, site: str, *, api_key: str = "KEY-123", html: str = "Y", url_column: str | None = None) -> Path:
    path = tmp_path / "purchase.xlsx"
    build_workbook(path, flows=["Flow"], base_url="http://127.0.0.1:1/", enabled=["Flow"])
    wb = openpyxl.load_workbook(path)
    for row in wb["Global"].iter_rows(min_row=1):
        if row[0].value == "isHTMLParsing":
            row[1].value = html
    wb["Params_1"]["F2"] = f"=Purchase!{openpyxl.utils.get_column_letter(COL['SiteUrl_OUT'])}3"       # the UI test opens the address the API answered with
    wb["DataSheets"].append(["Purchase", "Y", None, "Purchase API"])
    env = wb.create_sheet("Environments")
    for row in [("Environment", "Parameter", "Value"), ("QA", "JSON2-Purchase_URL", "http://qa.invalid/x"),
                ("UAT", "JSON2-Purchase_URL", site + "policy/purchase/v2"), ("UAT", "JSON2-Purchase_SoapAction", "urn:purchase")]:
        env.append(list(row))
    io = wb.create_sheet("InputOutput")
    for row in IO:
        io.append(list(row))
    ws = wb.create_sheet("Purchase")
    ws.append(COLUMNS + ([url_column] if url_column else []))
    ws.cell(2, 1, "Purchase 1 traveler")                                             # a label row: no blnExecute, not a test
    data = {"blnExecute": "Y", "TCID": "TC001", "TC_Name": "Purchase_WI", "Environment_Parameter": "JSON2-Purchase_URL", "XML_LOCATION": "L:\\nowhere\\",
            "XML_REQUESTFILE": "purchase.txt", "XML_RESPONSEFILE": "out.txt", "JSON_FORMAT": "Y", "DT_apiKey": api_key,
            "LastName_IN": '="QATESTLN"&RANDBETWEEN(1,9999)', "Qty_IN": 2, "Flag_IN": True}
    for name, value in data.items():
        ws.cell(3, COL[name], value)
    if url_column:
        ws.cell(3, len(COLUMNS) + 1, "http://override.invalid/purchase")
    (tmp_path / "templates").mkdir(exist_ok=True)
    (tmp_path / "templates" / "purchase.txt").write_text(json.dumps(TEMPLATE), encoding="utf-8-sig")
    wb.save(path)
    return path


# -- the pure helpers ---------------------------------------------------------------------------------------------------------------------
def test_replace_swaps_each_placeholder_for_its_column_in_inputoutput_order_and_skips_columns_the_sheet_lacks():
    data, filled = apply_replacements(TEMPLATE, {"lastName_value": "LN", "qty_value": "Q", "flag_value": "F", "partner_value": "P"},
                                      {"LN": "O'Neil \"Jr\"", "Q": 2.0, "F": True})
    assert data == {"partner": "partner_value", "lastName": "O'Neil \"Jr\"", "trip": {"qty": "2", "primary": "True"}, "untouched": "other_value"}
    assert filled == 3                                                                # a quote in a value cannot break the JSON


def test_a_replace_column_that_is_an_excel_error_says_so_instead_of_sending_the_error_text():
    from regrunner.workbook.sheet import ErrorText
    with pytest.raises(ValueError, match="LN cell evaluates to #NAME"):
        apply_replacements(TEMPLATE, {"lastName_value": "LN"}, {"LN": ErrorText("#NAME?", "unsupported function X")})


def test_the_response_becomes_elements_and_output_paths_find_values_like_the_legacy_runner_did():
    response = {"a": {"list": [{"p": {"n": "1", "who": {"last": "Aa"}}}, {"p": {"n": "2", "who": {"last": "Bb"}}}], "empty": "", "nothing": None, "ok": True},
                "Big": {"Key": 5}}
    root = json_to_tree(response, lower=True)
    got = lambda **paths: {c: v for _, c, v in output_values(root, {p: c for c, p in paths.items()})}
    assert got(n="list/p/n") == {"n": "1;2"}                                          # a list is not an element of its own: both p are children of list
    assert got(first="p[1]/who/last", both="p/who/last", second="p/who/last/index:=1") == {"first": "Aa", "both": "Aa;Bb", "second": "Bb"}
    assert got(x="empty", y="nothing") == {}                                          # no text / None: nothing stored
    assert got(ok="ok", big="big/key", missing="zzz", bad="p[0]") == {"ok": "True", "big": "5"}                    # a path that cannot be read is skipped
    assert got(upper="Big/Key") == {}                                                 # paths are case sensitive; the tree is lower case (isHTMLParsing = Y)
    keep = json_to_tree(response, lower=False)
    assert {c: v for _, c, v in output_values(keep, {"Big/Key": "upper", "big/key": "lower"})} == {"upper": "5"}


def test_an_xml_response_can_be_read_too_without_its_namespaces():
    root = xml_to_tree('<s:Envelope xmlns:s="urn:x"><s:Body><Result><Id>7</Id></Result></s:Body></s:Envelope>')
    assert {c: v for _, c, v in output_values(root, {"Result/Id": "id"})} == {"id": "7"}
    assert xml_to_tree("not xml") is None


# -- the workbook -------------------------------------------------------------------------------------------------------------------------
def test_a_sheet_with_environment_parameter_is_an_api_test_and_its_url_comes_from_the_environments_sheet(tmp_path, site):
    wb = Workbook(build(tmp_path, site), seed=1)
    cases = {c.id: c for c in wb.discover()}
    assert cases["Purchase"].kind == "api" and cases["Purchase"].runnable and cases["Purchase"].param_row == 3       # the label row is not a test
    assert cases["Purchase"].scenario == "Purchase_WI" and not any("not a UI keyword sheet" in w for w in wb.warnings)
    req = wb.runtime(cases["Purchase"]).request()
    assert (req.method, req.url, req.url_problem) == ("POST", site + "policy/purchase/v2", "")
    assert req.headers == {"apiKey": "KEY-123", "Content-Type": "application/json", "SOAPAction": "urn:purchase"} and req.secret_headers == {"apiKey"}
    assert req.template_file == "purchase.txt"
    qa = Workbook(build(tmp_path, site), environment="QA", seed=1)
    assert qa.runtime([c for c in qa.discover() if c.kind == "api"][0]).request().url == "http://qa.invalid/x"       # the run's environment picks the row


def test_a_webservice_url_in_the_row_wins_and_a_missing_parameter_is_said(tmp_path, site):
    wb = Workbook(build(tmp_path, site, url_column="WEBSERVICE_URL"), seed=1)
    (api,) = [c for c in wb.discover() if c.kind == "api"]
    assert wb.runtime(api).request().url == "http://override.invalid/purchase"
    other = Workbook(build(tmp_path, site), environment="PROD", seed=1)
    (api,) = [c for c in other.discover() if c.kind == "api"]
    req = other.runtime(api).request()
    assert req.url == "" and "no 'JSON2-Purchase_URL' for environment PROD" in req.url_problem


def test_the_api_test_runs_before_the_ui_test_that_reads_its_output_and_the_ui_test_sees_the_cell_it_filled(tmp_path, site):
    wb = Workbook(build(tmp_path, site), seed=1)
    order = plan_order(analyse(wb), ["Flow", "Purchase"])
    assert order.deps == {"Flow": ["Purchase"], "Purchase": []} and order.after_ui == []               # not "after every UI test"
    assert order.cells["Flow"]["DT_URL"] == ("Purchase", 3, COL["SiteUrl_OUT"])
    (ui,) = [c for c in wb.discover() if c.kind == "ui"]
    assert wb.runtime(ui).params["DT_URL"] in (None, "")                                              # nothing has run yet
    filled = wb.runtime(ui, shared={("Purchase", 3, COL["SiteUrl_OUT"]): "http://from.the.api/"})
    assert filled.params["DT_URL"] == "http://from.the.api/"


# -- a real run against the mock API ---------------------------------------------------------------------------------------------------------
@pytest.mark.browser
def test_the_purchase_runs_first_fills_its_template_and_hands_its_outputs_to_the_ui_test(tmp_path, site):
    proc, results, run_dir = run(tmp_path, build(tmp_path, site), "runner: {workers: 3, stagger_s: 0, min_page_load_gap_s: 0}\n", timeout=180)
    assert proc.returncode == 0, proc.stdout + proc.stderr
    by_id = {t["id"]: t for t in results["tests"]}
    assert by_id["Purchase"]["status"] == "PASSED" and by_id["Flow"]["status"] == "PASSED", by_id["Flow"].get("error")
    order = [(e["type"], e.get("test")) for e in map(json.loads, (run_dir / "events.jsonl").read_text().splitlines()) if e["type"] in ("test_started", "test_finished")]
    assert order.index(("test_finished", "Purchase")) < order.index(("test_started", "Flow"))
    seen = mock_server.API_SEEN
    assert seen["path"] == "/policy/purchase/v2" and seen["headers"]["soapaction"] == "urn:purchase" and seen["headers"]["apikey"] == "KEY-123"
    assert seen["headers"]["content-type"] == "application/json"
    body = seen["body"]
    assert body["lastName"].startswith("QATESTLN") and body["trip"] == {"qty": "2", "primary": "True"}
    assert body["partner"] == "partner_value" and body["untouched"] == "other_value"                  # no column for it: as the template had it
    step = by_id["Purchase"]["steps"][0]
    assert "Replace: 3 placeholders filled" in " ".join(step["notes"]) and "Output: 5 values read" in " ".join(step["notes"])
    filled = {v["name"]: v["value"] for v in by_id["Purchase"]["variables"]}
    assert filled["PolicyNumber_OUT"] == "P-777;P-778" and filled["LastName_OUT"] == body["lastName"] and filled["Second_OUT"] == "Second"
    assert filled["SiteUrl_OUT"] == site and filled["Status_OUT"] == "Success" and "Case_OUT" not in filled
    assert "KEY-123" not in (run_dir / "tests" / "Purchase" / "request.txt").read_text()


@pytest.mark.browser
def test_when_the_purchase_fails_the_ui_test_that_needs_its_output_is_not_run(tmp_path, site):
    proc, results, _ = run(tmp_path, build(tmp_path, site, api_key="WRONG"), "runner: {workers: 3, stagger_s: 0, min_page_load_gap_s: 0}\n", timeout=180)
    by_id = {t["id"]: t for t in results["tests"]}
    assert by_id["Purchase"]["status"] == "ERROR" and "HTTP 401" in by_id["Purchase"]["steps"][0]["error"]
    assert by_id["Flow"]["status"] == "ERROR" and by_id["Flow"]["error"].startswith("Not run: Flow needs DT_URL, which Purchase should have set")
