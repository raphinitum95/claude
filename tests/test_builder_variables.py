"""Variables and environments in the Workbook Builder (builder feedback items 12, 18, 24), no browser:

* the Variables screen's edits: rename (every cell that changes is listed first), delete, a new data variable, a secret's value in secrets.env
  (never in the workbook, never sent back), the values of one variable (``workbook/builder.py``, ``web/variables_api.py``);
* a workbook with its own environment table has no default environment: the runner, the web server and ``regrunner run`` refuse to guess;
* the environment table is a visible sheet right after Global, and Download names the secrets that are not in the file.
"""
from __future__ import annotations

import os
from pathlib import Path

import openpyxl
import pytest
import yaml
from fastapi.testclient import TestClient

from regrunner.cli import main
from regrunner.config import load_config
from regrunner.engine.runner import RunOptions, SelectionError, _load_and_plan
from regrunner.insight import workbook_summary
from regrunner.web.app import create_app
from regrunner.workbook.builder import (BuildDocument, BuildError, choose_environment, new_workbook, rename_preview,
                                        secrets_used, variable_values)
from regrunner.workbook.model import Workbook
from regrunner.workbook.variables import env_choice_problem, read_environment_table
from regrunner.workbook.writer import WorkbookEditor
from tests.flow_books import at, book

HEADERS = {"X-Requested-With": "regrunner"}
ENVS = [["Variable", "Required", "Secret", "QA", "UAT", "PROD"], ["DOMAIN", "Y", "", "https://qa.example", "https://uat.example", "https://example"],
        ["API_HOST", "", "", "qa-api", "uat-api", "api"]]


def sample(path: Path, environments: list[list] | None = ENVS) -> Path:
    """One test that types a data variable (whole cell and inline), a secret, and reads the environment table."""
    return book(path, {"Buy": [("Open", "open", {"Page": "chrome", "Value": "{DOMAIN}/form.html"}),
                               ("Set", "first name", {**at("//input[@id='first']"), "Value": "{FIRST_NAME}"}),
                               ("Set", "again", {**at("//input[@id='again']"), "Value": "FIRST_NAME"}),
                               ("Set", "password", {**at("//input[@id='pw']"), "Value": "{SECRET:PASSWORD}"}),
                               ("Set", "host", {**at("//input[@id='host']"), "Value": "{API_HOST}/x"})]},
                params={"Buy": [{"FIRST_NAME": "Ann", "NOTES": "adult"}, {"FIRST_NAME": "Bob", "NOTES": "child"}]}, environments=environments)


@pytest.fixture
def clean_secrets():
    """secrets.env entries these tests write end up in os.environ (so a replay sees them): take them out again."""
    before = {k for k in os.environ if k.startswith(("RR_SECRET_", "RR_VAR_"))}
    yield
    for k in [k for k in os.environ if k.startswith(("RR_SECRET_", "RR_VAR_")) and k not in before]:
        os.environ.pop(k, None)


# ---------------------------------------------------------------------------------------------------
# Item 24: the environment table is a visible sheet after Global
# ---------------------------------------------------------------------------------------------------
def test_the_environment_table_becomes_a_visible_sheet_right_after_global_the_next_time_the_builder_writes_it(tmp_path):
    path = sample(tmp_path / "a.xlsx")
    wb = openpyxl.load_workbook(path)
    wb["Buy"].print_area = "A1:D5"                          # a defined name scoped to a sheet: counted by position (localSheetId)
    wb.active = wb.sheetnames.index("Buy")
    wb.save(path)
    assert wb.sheetnames[-1] == "_rr_environments" and wb["_rr_environments"].sheet_state == "hidden"
    doc = BuildDocument(path)
    envs = doc.model()["environments"]
    doc.apply([{"op": "set_environments", "names": envs["names"], "production": envs["production"], "rows": envs["rows"]}])
    doc.save()
    after = openpyxl.load_workbook(path)
    assert after.sheetnames == ["Global", "_rr_environments", "DataSheets", "Buy", "Buy_Params"]
    assert after["_rr_environments"].sheet_state == "visible"
    assert after["Buy"].print_area == "'Buy'!$A$1:$D$5" and after.active.title == "Buy"      # renumbered with the move
    assert read_environment_table(Workbook(path).data).names == ["QA", "UAT", "PROD"]      # the engine still finds it


def test_a_new_workbook_has_no_default_environment_and_shows_its_environments(tmp_path):
    path = tmp_path / "new.xlsx"
    new_workbook(path, [{"name": "QA", "domain": "https://qa.example"}, {"name": "Staging", "domain": "https://s.example"}])
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames[:2] == ["Global", "_rr_environments"] and wb["_rr_environments"].sheet_state == "visible"
    assert not any(str(r[0].value).upper() == "ENVIRONMENT" for r in wb["Global"].iter_rows())
    workbook = Workbook(path)
    assert "Pick the environment" in env_choice_problem(workbook.environment_table(), None)
    assert workbook_summary(workbook, load_config(base_dir=tmp_path))["environments"] == {
        "source": "rr", "names": ["QA", "Staging"], "production": [], "required": True}


def test_environment_names_a_run_could_not_name_are_refused(tmp_path):
    doc = BuildDocument(sample(tmp_path / "a.xlsx"))
    with pytest.raises(BuildError, match="cannot be an environment name"):
        doc.apply([{"op": "set_environments", "names": ["QA", "Pre Prod"], "production": [], "rows": []}])


# ---------------------------------------------------------------------------------------------------
# Item 18: no default environment for a workbook with its own table
# ---------------------------------------------------------------------------------------------------
def test_the_build_tab_uses_the_environment_asked_for_else_the_tables_first_but_a_legacy_workbook_keeps_global():
    table = {"source": "rr", "names": ["QA", "Staging"], "production": []}
    assert choose_environment(table, "staging", "UAT") == "Staging"
    assert choose_environment(table, None, "UAT") == "QA"                  # never Global's: there is no default
    assert choose_environment(table, "NOPE", "UAT") == "QA"
    assert choose_environment({"source": "none", "names": ["QA", "UAT", "PROD"]}, None, "uat") == "UAT"
    assert choose_environment({"source": "legacy", "names": ["SIT"]}, "sit", "UAT") == "SIT"


def test_the_runner_refuses_a_workbook_with_its_own_environment_table_until_one_of_its_environments_is_picked(tmp_path):
    path = sample(tmp_path / "a.xlsx")                                     # (Global says UAT: it is not used)
    cfg = load_config(base_dir=tmp_path)
    with pytest.raises(SelectionError, match=r"a\.xlsx: Pick the environment to run on.*QA, UAT, PROD"):
        _load_and_plan(RunOptions(workbook=path, tests=["Buy"]), cfg, True)
    with pytest.raises(SelectionError, match="no environment called SIT"):
        _load_and_plan(RunOptions(workbook=path, tests=["Buy"], environment="SIT"), cfg, True)
    plan = _load_and_plan(RunOptions(workbook=path, tests=["Buy"], environment="qa"), cfg, True)
    assert plan.environment == "QA"
    legacy = sample(tmp_path / "legacy.xlsx", environments=None)
    assert _load_and_plan(RunOptions(workbook=legacy, tests=["Buy"]), cfg, True).environment == "UAT"     # Global, as always


def test_regrunner_run_without_env_names_the_workbooks_environments_and_starts_nothing(tmp_path, capsys):
    path = sample(tmp_path / "a.xlsx")
    (tmp_path / "config.yaml").write_text("reports: {html: false}\n")
    assert main(["--config", str(tmp_path / "config.yaml"), "run", str(path), "--plain", "--no-nice"]) == 2
    err = capsys.readouterr().err
    assert "Pick the environment to run on" in err and "QA, UAT, PROD" in err and "--env" in err
    assert not (tmp_path / "runs").exists() or not any((tmp_path / "runs").iterdir())


# ---------------------------------------------------------------------------------------------------
# Item 12: the Variables screen's edits
# ---------------------------------------------------------------------------------------------------
def test_a_rename_lists_every_cell_it_changes_first_and_then_changes_exactly_those(tmp_path):
    doc = BuildDocument(sample(tmp_path / "a.xlsx"))
    with doc.lock:
        preview = rename_preview(doc.editor, "FIRST_NAME", "GIVEN_NAME")
    assert sorted((c["sheet"], c["header"], c["before"], c["after"]) for c in preview) == [
        ("Buy", "Value", "FIRST_NAME", "GIVEN_NAME"), ("Buy", "Value", "{FIRST_NAME}", "{GIVEN_NAME}"),
        ("Buy_Params", "FIRST_NAME", "FIRST_NAME", "GIVEN_NAME")]
    assert doc.model()["version"] == 1 and not doc.status()["modified"]                  # a preview changes nothing
    applied = doc.apply([{"op": "rename_variable", "from": "FIRST_NAME", "to": "GIVEN_NAME"}])
    assert applied[0]["changed"] == 3
    v = next(v for v in doc.model()["variables"] if v["key"] == "GIVEN_NAME")
    assert len(v["usedBy"]) == 2 and v["sources"] == [{"sheet": "Buy_Params", "column": "GIVEN_NAME"}]
    with pytest.raises(BuildError, match="already a variable called API_HOST"):
        doc.apply([{"op": "rename_variable", "from": "GIVEN_NAME", "to": "API_HOST"}])
    env_rename = doc.apply([{"op": "rename_variable", "from": "API_HOST", "to": "API_SERVER"}])[0]
    assert env_rename["changed"] == 2                                                    # the table's row and the step's {API_HOST}
    assert "API_SERVER" in {r["variable"] for r in doc.model()["environments"]["rows"]}


def test_deleting_a_variable_empties_its_data_column_and_its_environment_row_but_domain_stays(tmp_path):
    doc = BuildDocument(sample(tmp_path / "a.xlsx"))
    doc.apply([{"op": "delete_variable", "token": "FIRST_NAME"}])
    m = doc.model()
    v = next(v for v in m["variables"] if v["key"] == "FIRST_NAME")
    assert v["sources"] == [] and len(v["usedBy"]) == 1          # {FIRST_NAME} stays in its step (the bare word is plain text now)
    assert any(p["kind"] == "unset_variable" and p.get("variable") == "FIRST_NAME" for p in m["problems"])
    with doc.lock:
        assert doc.editor.headers("Buy_Params") == {"blnExecute": 1, "NOTES": 3}         # the column is emptied, not shifted
    doc.apply([{"op": "delete_variable", "token": "API_HOST"}])
    assert [r["variable"] for r in doc.model()["environments"]["rows"]] == ["DOMAIN"]
    with pytest.raises(BuildError, match="DOMAIN cannot be deleted"):
        doc.apply([{"op": "delete_variable", "token": "DOMAIN"}])
    with pytest.raises(BuildError, match="only appears in steps"):
        doc.apply([{"op": "delete_variable", "token": "PASSWORD"}])


def test_a_new_data_variable_for_a_test_without_data_makes_its_params_sheet(tmp_path):
    path = tmp_path / "new.xlsx"
    new_workbook(path, [{"name": "QA", "domain": "https://qa.example"}])
    doc = BuildDocument(path)
    doc.apply([{"op": "add_test", "name": "Login", "kind": "web"}])
    doc.apply([{"op": "add_variable", "test": "Login", "token": "USER_NAME", "value": "ann", "label": "User name"}])
    m = doc.model()
    assert next(t for t in m["tests"] if t["id"] == "Login")["paramSheet"] == "Login_Params"
    v = next(v for v in m["variables"] if v["key"] == "USER_NAME")
    assert v["label"] == "User name" and v["sources"] == [{"sheet": "Login_Params", "column": "USER_NAME"}]
    with doc.lock:
        values = variable_values(doc.editor, "USER_NAME")
    assert [(r["row"], r["value"]) for r in values["sources"][0]["rows"]] == [(2, "ann")]


def test_secrets_are_listed_by_name_and_a_typed_in_secret_value_is_never_shown(tmp_path):
    envs = [*ENVS, ["API_KEY", "", "Y", "hunter2", "{SECRET:API_KEY}", ""]]
    path = sample(tmp_path / "a.xlsx", environments=envs)
    editor = WorkbookEditor.open(path)
    assert secrets_used(editor) == ["API_KEY", "PASSWORD"]
    row = variable_values(editor, "API_KEY")["environment"]
    assert row["values"] == {"QA": "••••••", "UAT": "{SECRET:API_KEY}", "PROD": ""} and row["typedIn"] == ["QA"]
    assert "hunter2" not in repr(variable_values(editor, "API_KEY"))
    as_secret = variable_values(editor, "FIRST_NAME", secret=True)["sources"][0]["rows"]          # a secret typed into a Params column
    assert [(r["value"], r["masked"]) for r in as_secret] == [("••••••", True), ("••••••", True)] and "Ann" not in repr(as_secret)


# ---------------------------------------------------------------------------------------------------
# The routes: /api/runs refuses, the Variables screen's routes, download-info
# ---------------------------------------------------------------------------------------------------
@pytest.fixture
def api(tmp_path):
    (tmp_path / "workbooks").mkdir()
    sample(tmp_path / "workbooks" / "own.xlsx")
    sample(tmp_path / "workbooks" / "legacy.xlsx", environments=None)
    cfg_file = tmp_path / "config.yaml"
    cfg_file.write_text(yaml.safe_dump({"runner": {"workers": 1}}))
    client = TestClient(create_app(load_config(cfg_file, base_dir=tmp_path), str(cfg_file)), base_url="http://127.0.0.1", headers=HEADERS)
    client.root = tmp_path
    return client


def test_the_web_server_refuses_a_run_of_a_workbook_with_its_own_table_without_one_of_its_environments(api):
    info = api.get("/api/workbooks/own.xlsx/tests").json()
    assert info["environment"] == "" and info["environments"] == {"source": "rr", "names": ["QA", "UAT", "PROD"], "production": ["PROD"],
                                                                  "required": True}
    assert api.get("/api/workbooks/legacy.xlsx/tests").json()["environments"]["required"] is False
    none = api.post("/api/runs", json={"workbook": "own.xlsx", "all": True})
    assert none.status_code == 422 and none.json()["kind"] == "env_required" and none.json()["environments"] == ["QA", "UAT", "PROD"]
    assert "own.xlsx: Pick the environment to run on" in none.json()["error"]
    unknown = api.post("/api/runs", json={"workbooks": [{"workbook": "own.xlsx", "all": True}, {"workbook": "legacy.xlsx", "all": True}], "env": "SIT"})
    assert unknown.status_code == 422 and unknown.json()["kind"] == "env_unknown" and "no environment called SIT" in unknown.json()["error"]
    assert api.get("/api/runs").json() == []


def test_the_variables_routes_show_values_rename_with_a_preview_and_keep_secret_values_in_secrets_env(api, clean_secrets):
    values = api.get("/api/build/workbooks/own.xlsx/variables/FIRST_NAME/values").json()
    assert [(r["row"], r["label"], r["value"]) for r in values["sources"][0]["rows"]] == [(2, "adult", "Ann"), (3, "child", "Bob")]
    assert values["secret"] is None
    secret = api.post("/api/build/workbooks/own.xlsx/variables/secret", json={"name": "PASSWORD", "environment": "UAT", "value": "s3cret!"})
    assert secret.status_code == 200 and "s3cret!" not in secret.text
    assert secret.json()["environments"] == {"QA": False, "UAT": True, "PROD": False} and secret.json()["key"] == "RR_SECRET_UAT_PASSWORD"
    assert "RR_SECRET_UAT_PASSWORD=s3cret!" in (api.root / "secrets.env").read_text()
    assert "s3cret!" not in (api.root / "workbooks" / "own.xlsx").read_bytes().decode("latin-1")
    shown = api.get("/api/build/workbooks/own.xlsx/variables/PASSWORD/values").json()
    assert shown["secret"]["environments"]["UAT"] is True and "s3cret!" not in repr(shown)
    preview = api.post("/api/build/workbooks/own.xlsx/variables/rename-preview", json={"from": "PASSWORD", "to": "LOGIN_PASSWORD"}).json()
    assert [(c["before"], c["after"]) for c in preview["changes"]] == [("{SECRET:PASSWORD}", "{SECRET:LOGIN_PASSWORD}")]
    assert preview["secrets"] == ["RR_SECRET_UAT_LOGIN_PASSWORD"]
    version = api.get("/api/build/workbooks/own.xlsx").json()["version"]
    done = api.post("/api/build/workbooks/own.xlsx/variables/rename", json={"from": "PASSWORD", "to": "LOGIN_PASSWORD", "version": version}).json()
    assert done["secretsCopied"] == ["RR_SECRET_UAT_LOGIN_PASSWORD"] and "s3cret!" not in repr(done)
    text = (api.root / "secrets.env").read_text()
    assert "RR_SECRET_UAT_LOGIN_PASSWORD=s3cret!" in text and "RR_SECRET_UAT_PASSWORD=s3cret!" in text     # copied: an undo still works
    assert api.post("/api/build/workbooks/own.xlsx/variables/rename-preview", json={"from": "DOMAIN", "to": "SITE"}).status_code == 400
    again = api.post("/api/build/workbooks/own.xlsx/variables/secret", json={"name": "LOGIN_PASSWORD", "environment": "UAT", "value": "n3w"})
    assert again.status_code == 200 and "RR_SECRET_UAT_LOGIN_PASSWORD=n3w" in (api.root / "secrets.env").read_text()


def test_download_info_names_the_secrets_a_saved_workbook_uses(api):
    assert api.get("/api/workbooks/own.xlsx/download-info").json() == {"name": "own.xlsx", "secrets": ["PASSWORD"]}
    book(api.root / "workbooks" / "plain.xlsx", {"T": [("Open", "open", {"Page": "chrome", "Value": "https://example"})]})
    assert api.get("/api/workbooks/plain.xlsx/download-info").json()["secrets"] == []
