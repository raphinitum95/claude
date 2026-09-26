"""A failed step says why: the browser's whole error, what covers the element (in its frame and on the page around it), where the locator
matches, and the page's HTML - so a run brought back from another computer explains itself."""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

from regrunner.engine.actions import whole_error
from regrunner.engine.failure_capture import last_log_lines, summarize, xpath_prefixes
from regrunner.engine.runner import RunOptions, execute
from regrunner.events import EventBus
from tests.workbook_factory import Sheet

pytestmark = pytest.mark.browser


def build(path: Path, site: str, page: str, rows: list[tuple[str, str, dict]]) -> Path:
    wb = openpyxl.Workbook()
    ds = wb.active
    ds.title = "DataSheets"
    ds.append(["SheetsToExecute", "blnExecute", "ParameterSheet", "Comments"])
    ds.append(["Flow", "Y", "Params_1", "capture"])
    g = wb.create_sheet("Global")
    g.append(["Parameter", "Value"])
    g.append(["Environment", "UAT"])
    ps = wb.create_sheet("Params_1")
    ps.append(["blnExecute", "DT_URL"])
    ps.append(["Y", site + page])
    s = Sheet(wb.create_sheet("Flow"), "Y")
    s.add("Open", "Open", Page="chrome", Value="DT_URL")
    for method, name, cols in rows:
        s.add(method, name, **cols)
    wb.save(path)
    return path


def click(xpath: str, timeout: int = 2) -> dict:
    return {"FindBy": "xpath", "FindBy_Value": xpath, "Index": 0, "Timeout": timeout}


async def run(make_cfg, wb: Path, **overrides):
    cfg = make_cfg(**{"waits.window_grace_s": 0.5, "timeouts.element_s": 2, "timeouts.optional_s": 1.0, **overrides})
    events: list[dict] = []
    bus = EventBus()
    bus.subscribe(events.append)
    result = await execute(RunOptions(workbook=wb, tests=["Flow"], seed=1), cfg, bus)
    return result, events, cfg


def step_named(result, name: str):
    return next(s for s in result.tests[0].steps if s.name == name)


# -- a link in a frame, with a layer of the page over the frame (the agent portal's failure) --------------------------------------
async def test_a_click_that_times_out_says_what_was_on_top_of_the_link(site, make_cfg, tmp_path):
    wb = build(tmp_path / "wb.xlsx", site, "/cover_parent.html", [
        ("SwitchToFrame", "into the frame", {"Value": "aemFormFrame"}),
        ("Click", "click Australia", click("//a[contains(.,'Australia')]")),
        ("Click", "later step", click("//a[@id='never']", timeout=1)),
    ])
    result, events, _ = await run(make_cfg, wb)
    step = step_named(result, "click Australia")
    assert step.status == "FAILED" and step.error.startswith("TimeoutError: Locator.click")

    assert "intercepts pointer events" in step.detail                          # Playwright's own reason, no longer thrown away with all but its first line
    diag = step.diagnosis
    assert diag["element"]["visible"] is True and diag["element"]["disabled"] is False
    assert diag["hit"]["top"]["relation"] == "itself"                          # inside the frame nothing covers the link ...
    assert diag["hit_page"]["top"]["relation"] == "other"                      # ... the page around the frame does
    assert "veil" in diag["hit_page"]["top"]["desc"] and "page-loader" in diag["hit_page"]["top"]["desc"]
    assert diag["hit_page"]["overlays"][0]["z"] == "999"
    assert diag["matches"][0]["where"] == {"main page": 0, "frame 'aemFormFrame'": 1} and diag["matches"][0]["searched"] == "frame 'aemFormFrame'"
    assert diag["frame"]["name"] == "aemFormFrame" and diag["frame"]["ready"] == "complete"
    said = " ".join(diag["summary"])
    assert said.startswith("On the main page, above the frame") and "veil" in said            # the layer that takes the click comes first
    assert "Inside its frame nothing covers the element" in said                                # and it does not pretend the frame is at fault
    assert "positioned layer" not in said                                                       # (the layer was already named: not said twice)
    assert any("intercepts pointer events" in line for line in diag["playwright_last"])

    # it reaches every place a run is read from
    saved = json.loads((tmp_path / "runs" / result.run_id / "results.json").read_text())
    saved_step = next(s for s in saved["tests"][0]["steps"] if s["name"] == "click Australia")
    assert saved_step["diagnosis"]["hit_page"]["top"]["relation"] == "other" and "intercepts" in saved_step["detail"]
    failed = next(e for e in events if e["type"] == "step_failed" and e["name"] == "click Australia")
    assert failed["diagnosis"]["summary"] == diag["summary"] and failed["detail"] == step.detail
    report = (tmp_path / "runs" / result.run_id / "report.html").read_text(encoding="utf-8")
    assert "On the main page, above the frame" in report and "the browser's whole error" in report


async def test_the_html_of_the_page_and_the_frame_is_kept_without_its_scripts(site, make_cfg, tmp_path):
    wb = build(tmp_path / "wb.xlsx", site, "/cover_parent.html", [
        ("SwitchToFrame", "into the frame", {"Value": "aemFormFrame"}),
        ("Click", "click Australia", click("//a[contains(.,'Australia')]")),
    ])
    result, _, _ = await run(make_cfg, wb)
    run_dir = tmp_path / "runs" / result.run_id
    dom = step_named(result, "click Australia").diagnosis["dom"]
    frame_html, page_html = (run_dir / dom["frame"]).read_text(encoding="utf-8"), (run_dir / dom["page"]).read_text(encoding="utf-8")
    assert "Quote &amp; Purchase" in frame_html and 'id="aus"' in frame_html
    assert 'id="veil"' in page_html and "page-loader" in page_html
    assert "not part of a snapshot" not in frame_html and "failure snapshot" in frame_html            # a still picture, not a page that runs again
    assert dom["frame"].startswith("tests/Flow/dom/") and not Path(dom["frame"]).is_absolute()


# -- a link the browser calls "not enabled": a container around it says aria-disabled="true" (the agent portal's second failure) -----------------
async def test_a_link_inside_an_aria_disabled_container_is_named_as_disabled(site, make_cfg, tmp_path):
    wb = build(tmp_path / "wb.xlsx", site, "/disabled_parent.html", [
        ("SwitchToFrame", "into the frame", {"Value": "aemFormFrame"}),
        ("Click", "click Australia", click("//a[contains(.,'Australia')]")),
    ])
    result, _, _ = await run(make_cfg, wb, **{"behaviour.click_ignores_aria_disabled": False})               # (on by default the click goes through: tests/test_aria_disabled_click.py)
    step = step_named(result, "click Australia")
    assert step.status == "FAILED" and step.error.startswith("TimeoutError: Locator.click")
    assert "element is not enabled" in step.detail                                # Playwright's own words
    diag = step.diagnosis
    assert diag["playwright_last"] == ["element is not enabled"]                  # not the vague lines that end each retry
    off = diag["element"]["disabledBy"]
    assert off["kind"] == 'aria-disabled="true"' and off["self"] is False and off["desc"] == "div#guideContainer-rootPanel.guideContainer"
    assert diag["element"]["visible"] is True and diag["hit"]["top"]["relation"] == "itself"      # nothing covers it: only the state is wrong
    said = " ".join(diag["summary"])
    assert "counts as disabled" in said and "div#guideContainer-rootPanel.guideContainer" in said and "Selenium" in said
    assert "Nothing is in the way" not in said                                    # (the conclusion the first version drew from this same page)


async def test_js_click_gets_through_where_the_browsers_enabled_check_does_not(site, make_cfg, tmp_path):
    """The legacy runner (Selenium) never looked at aria-disabled: JS_Click - which mimics a person - clicks the same link."""
    wb = build(tmp_path / "wb.xlsx", site, "/disabled_parent.html", [
        ("SwitchToFrame", "into the frame", {"Value": "aemFormFrame"}),
        ("JS_Click", "click Australia", click("//a[contains(.,'Australia')]")),
    ])
    result, _, _ = await run(make_cfg, wb)
    step = step_named(result, "click Australia")
    assert step.status == "PASSED" and step.diagnosis is None


# -- a link on the page itself, under a banner -----------------------------------------------------------------------------------
async def test_a_banner_over_a_link_is_named(site, make_cfg, tmp_path):
    wb = build(tmp_path / "wb.xlsx", site, "/cover_inside.html", [("Click", "click Continue", click("//a[@id='go']"))])
    result, _, _ = await run(make_cfg, wb)
    diag = step_named(result, "click Continue").diagnosis
    assert diag["hit"]["top"]["relation"] == "other" and diag["hit"]["top"]["desc"] == "div#cookie-banner.banner"
    assert diag["hit"]["top"]["text"] == "We use cookies"
    assert diag["hit"]["overlays"][0]["desc"] == "div#cookie-banner.banner" and "hit_page" not in diag       # not in a frame: no second look
    assert "another element is on top" in " ".join(diag["summary"]) and "We use cookies" in " ".join(diag["summary"])


# -- an element that is not where the step is looking ---------------------------------------------------------------------------
async def test_an_element_in_another_frame_is_reported_as_such(site, make_cfg, tmp_path):
    wb = build(tmp_path / "wb.xlsx", site, "/cover_parent.html", [("Click", "click without switching", click("//a[@id='aus']", timeout=1))])
    result, _, _ = await run(make_cfg, wb)
    step = step_named(result, "click without switching")
    assert step.error == "Object was not found" and "element" not in step.diagnosis
    assert step.diagnosis["matches"][0]["where"] == {"main page": 0, "frame 'aemFormFrame'": 1}
    assert "the step is in the wrong frame" in " ".join(step.diagnosis["summary"]) and "aemFormFrame" in " ".join(step.diagnosis["summary"])
    assert [f["label"] for f in step.diagnosis["frames"]] == ["main page", "frame 'aemFormFrame'"]


async def test_an_element_that_exists_nowhere_says_so(site, make_cfg, tmp_path):
    wb = build(tmp_path / "wb.xlsx", site, "/cover_inside.html", [("Click", "click nothing", click("//a[@id='nope']", timeout=1))])
    result, _, _ = await run(make_cfg, wb)
    assert "matches nothing in any frame" in " ".join(step_named(result, "click nothing").diagnosis["summary"])


# -- a locator that is almost right: which step of it leaves nothing (the agent portal's "Policyholder Name:" vs the page's "Policy Holder Name:") ----------
async def test_the_step_of_the_locator_that_stops_matching_is_named(site, make_cfg, tmp_path):
    wb = build(tmp_path / "wb.xlsx", site, "/cover_inside.html",
               [("Click", "click a label that reads differently", click("//div[@id='cookie-banner']//span[contains(.,'Policyholder')]", timeout=1))])
    result, _, _ = await run(make_cfg, wb)
    step = step_named(result, "click a label that reads differently")
    parts = step.diagnosis["matches"][0]["parts"]
    assert parts == [{"part": "//div[@id='cookie-banner']", "count": 1}, {"part": "//span[contains(.,'Policyholder')]", "count": 0}]
    said = " ".join(step.diagnosis["summary"])
    assert "matches up to //div[@id='cookie-banner'] (1 element)" in said and "nothing once //span[contains(.,'Policyholder')] is added" in said


async def test_a_first_step_that_matches_nothing_is_said_so(site, make_cfg, tmp_path):
    wb = build(tmp_path / "wb.xlsx", site, "/cover_inside.html", [("Click", "wrong id", click("//div[@id='banner']//a", timeout=1))])
    result, _, _ = await run(make_cfg, wb)
    assert "Not even the first part of the locator, //div[@id='banner'], matches anything" in " ".join(step_named(result, "wrong id").diagnosis["summary"])


def test_an_xpath_is_cut_at_its_steps_and_not_inside_brackets_or_quotes():
    assert xpath_prefixes("//div[@id=\"X\"]//span[contains(.,'a/b')]/b") == ['//div[@id="X"]', '//div[@id="X"]//span[contains(.,\'a/b\')]',
                                                                              '//div[@id="X"]//span[contains(.,\'a/b\')]/b']
    assert xpath_prefixes("//a") == [] and xpath_prefixes("(//a)[2]") == [] and xpath_prefixes("id('x')") == []


# -- a comparison that read nothing: hidden, empty, or the text is a value ----------------------------------------------------------------------------------
def read(xpath: str, expected: str) -> dict:
    return {"FindBy": "xpath", "FindBy_Value": xpath, "Index": 0, "Output_Property": "innertext", "Expected_Value": expected, "Contains": "Y", "Timeout": 2}


async def test_a_comparison_that_read_nothing_says_why(site, make_cfg, tmp_path):
    wb = build(tmp_path / "wb.xlsx", site, "/hidden_text.html", [
        ("Output", "hidden copy", read("//p[@id='old']", "the old copy")),
        ("Output", "empty paragraph", read("//p[@id='empty']", "anything")),
        ("Output", "text is a value", read("//input[@id='box']", "E-business Test")),
        ("Output", "visible but different", read("//p[@id='new']", "the old copy")),
    ])
    result, _, _ = await run(make_cfg, wb)
    said = {s.name: " ".join((s.diagnosis or {}).get("summary", [])) for s in result.tests[0].steps if s.name != "Open"}
    assert all(step_named(result, n).error == "Comparison Failed" for n in said)
    assert "element is hidden (display: none" in said["hidden copy"] and "text counts as empty" in said["hidden copy"] and "legacy runner did the same" in said["hidden copy"]
    assert "contains no text and takes no room on the page" in said["empty paragraph"]              # (an empty <p> has no height)
    assert "point a click lands" not in " ".join(said.values())                                        # nothing about clicks in a text comparison
    assert "value is “E-business Test”" in said["text is a value"] and "Output_Property = value" in said["text is a value"]
    assert step_named(result, "visible but different").diagnosis is None                       # it read text and it differs: the two texts say it all


# -- several matches, the wrong one picked (the agent portal's Start Medical Assessment: three buttons, one per traveller) -----------------------------------
async def test_when_the_index_picks_a_hidden_match_the_visible_ones_are_named(site, make_cfg, tmp_path):
    wb = build(tmp_path / "wb.xlsx", site, "/three_buttons.html", [("JS_Click", "start the second", {**click("//button[contains(.,'START ASSESSMENT')]", timeout=1), "Index": 1})])
    result, _, _ = await run(make_cfg, wb)
    step = step_named(result, "start the second")
    diag = step.diagnosis
    assert diag["matches"][0]["visible"] == [False, False, True] and diag["element"]["visible"] is False
    said = " ".join(diag["summary"])
    assert "Index 1 picked one that is hidden; the visible one is Index 2" in said
    assert "point a click lands" not in said                                                  # an element with no size has no click point: nothing "on top of" it


# -- what does not get a diagnosis ------------------------------------------------------------------------------------------------
async def test_a_passing_step_and_a_comparison_failure_carry_no_diagnosis(site, make_cfg, tmp_path):
    wb = build(tmp_path / "wb.xlsx", site, "/cover_inside.html", [
        ("Output", "reads the heading", {"FindBy": "xpath", "FindBy_Value": "//h1", "Index": 0, "Output_Property": "innertext",
                                         "Expected_Value": "Something else", "Exact_Match": "Y"}),
    ])
    result, events, _ = await run(make_cfg, wb)
    opened, read = step_named(result, "Open"), step_named(result, "reads the heading")
    assert opened.status == "PASSED" and opened.diagnosis is None and opened.detail == ""
    assert read.status == "FAILED" and read.error == "Comparison Failed" and read.diagnosis is None
    assert all("diagnosis" not in e for e in events if e["type"] == "step_passed")
    assert not (tmp_path / "runs" / result.run_id / "tests" / "Flow" / "dom").exists()


async def test_capture_can_be_switched_off(site, make_cfg, tmp_path):
    wb = build(tmp_path / "wb.xlsx", site, "/cover_inside.html", [("Click", "click Continue", click("//a[@id='go']"))])
    result, _, _ = await run(make_cfg, wb, **{"failure_capture.enabled": False})
    step = step_named(result, "click Continue")
    assert step.status == "FAILED" and step.diagnosis is None
    assert step.detail                                                          # the browser's whole error costs nothing: it is always kept


async def test_html_is_kept_for_the_first_few_failures_only(site, make_cfg, tmp_path):
    wb = build(tmp_path / "wb.xlsx", site, "/cover_inside.html", [("Click", f"missing {i}", click(f"//a[@id='nope{i}']", timeout=1)) for i in range(4)])
    result, _, _ = await run(make_cfg, wb, **{"failure_capture.dom_max_per_test": 2})
    kept = [bool(step_named(result, f"missing {i}").diagnosis.get("dom")) for i in range(4)]
    assert kept == [True, True, False, False]                                   # the findings stay for all of them, the HTML for the first two


def test_masking_reaches_every_string_in_the_evidence():
    from regrunner.engine.failure_capture import _mask_all
    masked = _mask_all({"a": ["x hunter2 y", {"b": "hunter2"}], "n": 3}, lambda text: text.replace("hunter2", "••••••"))
    assert masked == {"a": ["x •••••• y", {"b": "••••••"}], "n": 3}


# -- pieces --------------------------------------------------------------------------------------------------------------------------
def test_whole_error_keeps_the_call_log_and_leaves_one_liners_to_the_step():
    assert whole_error(Exception("Object was not found")) == ""
    text = "Locator.click: Timeout 30000ms exceeded.\nCall log:\n  - waiting for locator(\"a\")\n  - <div id=veil> intercepts pointer events"
    assert whole_error(Exception(text)) == text
    assert whole_error(Exception("x\n" + "y" * 9000)).endswith("... (cut)")


def test_the_reason_lines_of_the_call_log_are_what_playwright_was_stuck_on():
    log = "\n".join(["Timeout", "Call log:", "  - waiting for locator(a)", "  - element is visible", "  - <div> intercepts pointer events", "  - retrying click action",
                     "  - waiting 20ms", "  - element is visible", "  - <div> intercepts pointer events", "  - retrying click action", "  - waiting 100ms"])
    assert last_log_lines(log) == ["<div> intercepts pointer events"]
    assert last_log_lines("no log at all") == []
    assert last_log_lines("t\nCall log:\n  - waiting for x\n  - retrying click action") == ["waiting for x", "retrying click action"]        # no reason given: its last lines


def test_a_retry_loop_that_ends_on_a_vague_line_still_names_its_reason():
    """The agent portal's log: the reason ("not enabled") is followed by lines that say nothing, 56 times over."""
    log = ("Locator.click: Timeout 30000ms exceeded.\nCall log:\n  - waiting for locator(\"//a\").first\n    - locator resolved to <a>Australia</a>\n"
           "  - attempting click action\n    2 × waiting for element to be visible, enabled and stable\n      - element is not enabled\n    - retrying click action\n"
           "    - waiting 20ms\n    56 × waiting for element to be visible, enabled and stable\n       - element is not enabled\n     - retrying click action\n"
           "       - waiting 500ms\n    - waiting for element to be visible, enabled and stable")
    assert last_log_lines(log) == ["element is not enabled"]


FINE = {"element": {"visible": True, "disabled": False, "inert": False, "pointerEvents": "auto", "moving": False, "inViewport": True, "rect": {"w": 60, "h": 20}},
        "hit": {"top": {"relation": "itself"}}, "frame": {"path": [], "ready": "complete"}}


def test_a_timeout_with_nothing_in_the_way_says_what_that_leaves():
    said = " ".join(summarize(FINE, "TimeoutError: Locator.click: Timeout 30000ms exceeded.", ["retrying click action"]))
    assert "Nothing is in the way now" in said and "only while it was being tried" in said
    waiting = " ".join(summarize(FINE, "TimeoutError: Locator.click: Timeout 30000ms exceeded.", ["waiting for scheduled navigations to finish"]))
    assert "the click was made" in waiting and "navigation it started" in waiting
    assert "Nothing is in the way" not in " ".join(summarize({**FINE, "hit": {"top": {"relation": "other", "desc": "div#x"}}}, "TimeoutError: Locator.click: Timeout"))


def test_summary_of_a_page_that_looks_fine_admits_it():
    assert "nothing obviously wrong" in summarize({}, "TimeoutError: Locator.click")[0]
    assert summarize({}, "") == []
