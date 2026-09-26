"""The Build tab in a real browser: open a workbook, browse its map, edit a test, save to Excel (P04-build-ui.md)."""
from __future__ import annotations

import asyncio
import re

import openpyxl
import pytest
from playwright.async_api import async_playwright

from tests.web_fixtures import web  # noqa: F401  (fixture)

pytestmark = pytest.mark.browser
HEADERS = {"X-Requested-With": "regrunner"}


@pytest.fixture
async def build_copy(web, request):
    """A workbook this file's tests can edit freely, independent of the shared ``mock.xlsx`` and of each other."""
    name = f"build_ui_{request.node.name}"[:80]
    async with web.aclient() as c:
        r = await c.post("/api/build/workbooks", json={"name": name, "from": "mock.xlsx"})
        assert r.status_code == 201, r.text
        return r.json()["name"]


async def open_ui(web, path="/"):
    pw = await async_playwright().start()
    browser = await pw.chromium.launch(headless=True)
    ctx = await browser.new_context(viewport={"width": 1440, "height": 1000})
    await ctx.route(re.compile(r"https://fonts\.(googleapis|gstatic)\.com/.*"), lambda route: route.fulfill(status=200, content_type="text/css", body=""))
    page = await ctx.new_page()
    page.errors = []
    page.on("pageerror", lambda e: page.errors.append(f"pageerror: {e}"))
    noise = re.compile(r"Failed to load resource: the server responded with a status of 4\d\d")
    page.on("console", lambda m: page.errors.append(f"console: {m.text}") if m.type == "error" and not noise.search(m.text) else None)
    await page.goto(web.base + path)
    return pw, browser, page


async def js_until(page, expression: str, timeout: float = 20.0):
    deadline = asyncio.get_running_loop().time() + timeout
    while not await page.evaluate(expression):
        if asyncio.get_running_loop().time() > deadline:
            raise AssertionError(f"page condition not reached: {expression[:120]}")
        await asyncio.sleep(0.1)


async def test_the_build_tab_shows_the_workbook_map_and_a_test_opens_its_editor(web, build_copy):
    pw, browser, page = await open_ui(web, f"/#/build/{build_copy}")
    try:
        await js_until(page, "!!document.querySelector('.tab-btn.on')")
        assert await page.locator('.tab-btn.on').inner_text() == 'Build'
        await js_until(page, "document.querySelectorAll('main a.disp').length >= 2")           # FlowA, FlowB cards
        names = await page.locator('main a.disp').evaluate_all("els => els.map(e => e.textContent)")
        assert set(names) >= {"FlowA", "FlowB"}
        await page.locator('main a.disp', has_text="FlowA").click()
        await js_until(page, "location.hash.includes('/test/FlowA')")
        await js_until(page, "document.querySelectorAll('.scard').length > 0")
        await page.get_by_role("button", name=re.compile("^Grid$")).click()
        await js_until(page, "document.querySelector('table.mono thead') != null")
        assert "Method" in await page.locator("table.mono thead").inner_text()
        assert not page.errors, page.errors
    finally:
        await browser.close()
        await pw.stop()


async def test_editing_a_step_autosaves_a_draft_that_save_to_excel_writes_to_disk(web, build_copy):
    pw, browser, page = await open_ui(web, f"/#/build/{build_copy}/test/FlowA")
    try:
        await js_until(page, "document.querySelectorAll('.scard').length > 0")
        timeout_box = page.locator('input[data-input="build-step-timeout"]')
        await timeout_box.fill("42")
        await timeout_box.blur()
        await js_until(page, "!document.querySelector('[data-act=\"build-undo\"]').disabled", timeout=10)   # the edit landed: undo is now possible
        await js_until(page, "!!document.querySelector('[data-act=\"build-save\"]')")
        await page.get_by_role("button", name=re.compile("Save to Excel")).click()

        async def saved_to_disk():
            r = await page.request.get(f"{web.base}/api/build/workbooks/{build_copy}/status")
            return not (await r.json())["modified"]
        deadline = asyncio.get_running_loop().time() + 15
        while not await saved_to_disk():
            if asyncio.get_running_loop().time() > deadline:
                raise AssertionError("save to Excel did not clear the modified flag in time")
            await asyncio.sleep(0.2)
        assert not page.errors, page.errors
    finally:
        await browser.close()
        await pw.stop()

    wb = openpyxl.load_workbook(web.root / "workbooks" / build_copy)
    ws = wb["FlowA"]
    headers = [c.value for c in ws[1]]
    timeout_col = headers.index("Timeout") + 1
    values = [ws.cell(r, timeout_col).value for r in range(2, ws.max_row + 1)]
    assert 42 in values


async def test_new_workbook_dialog_creates_a_workbook_ready_to_build(web):
    pw, browser, page = await open_ui(web, "/#/build")
    try:
        await js_until(page, "!!document.querySelector('.tab-btn.on')")
        await page.get_by_role("button", name=re.compile("New workbook")).first.click()
        await js_until(page, "!!document.querySelector('[data-input=\"build-nw-name\"]')")
        await page.locator('[data-input="build-nw-name"]').fill("Fresh Build Test")
        await page.get_by_role("button", name="Create workbook").click()
        await js_until(page, "location.hash.includes('/build/Fresh')", timeout=15)
        await js_until(page, "document.querySelector('main .ttl') != null")
        assert not page.errors, page.errors
    finally:
        await browser.close()
        await pw.stop()
