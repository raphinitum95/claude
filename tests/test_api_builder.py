"""The Build tab's API / XML editor (P10, ``build/api_builder.py`` + ``/api/build/api/*``): the view of an API test, the ops that write it
(``InputOutput`` rows, columns, cells), Send now against the mock ``/policy/purchase/v2`` and ``/policy/xml/v1`` APIs, the response tree with
"this item" / "the item where", and the three ways in (cURL, Postman, request templates)."""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest
import yaml
from fastapi.testclient import TestClient

from regrunner.build import api_builder as A
from regrunner.config import load_config
from regrunner.web.app import create_app
from regrunner.workbook import Workbook
from regrunner.workbook.builder import BuildDocument
from tests.site import server as mock_server
from tests.test_api_tests import build_api_workbook, run

HEADERS = {"X-Requested-With": "regrunner"}


def add_environments(path: Path, site: str) -> None:
    wb = openpyxl.load_workbook(path)
    env = wb.create_sheet("_rr_environments")
    for r in (["Variable", "Required", "Secret", "UAT", "PROD"], ["DOMAIN", "Y", "", site, "https://prod.invalid/"]):
        env.append(r)
    wb.save(path)


@pytest.fixture
def api(tmp_path, site):
    (tmp_path / "workbooks").mkdir()
    path = build_api_workbook(tmp_path / "workbooks" / "api.xlsx", site)
    add_environments(path, site)
    cfg_file = tmp_path / "config.yaml"
    cfg_file.write_text(yaml.safe_dump({"runner": {"workers": 1}}))
    app = create_app(load_config(cfg_file, base_dir=tmp_path), str(cfg_file))
    client = TestClient(app, base_url="http://127.0.0.1", headers=HEADERS)
    client.root, client.site = tmp_path, site
    return client


def edit(api, *ops) -> dict:
    r = api.post("/api/build/workbooks/api.xlsx/edit", json={"ops": list(ops)})
    assert r.status_code == 200, r.text
    return r.json()


def new_purchase_test(api, name: str = "Buy") -> dict:
    return edit(api, {"op": "api_add_test", "name": name, "format": "json", "method": "POST", "url": "{DOMAIN}policy/purchase/v2",
                      "headers": [["apiKey", "{SECRET:API_KEY}"], ["Content-Type", "application/json"]],
                      "body": '{"lastName": "{LastName_IN}"}'},
                {"op": "api_set_value", "test": name, "row": 2, "column": "LastName_IN", "value": "Smith"})


# -- the view ------------------------------------------------------------------------------------------------------------------------------------------
def test_a_legacy_api_sheet_shows_its_request_checks_and_saves_and_a_key_typed_in_the_sheet_is_not_sent_to_the_page(api):
    v = api.get("/api/build/api/api.xlsx/tests/Policy").json()
    assert (v["test"], v["row"], v["format"], v["method"]) == ("Policy", 2, "json", "GET") and v["urlFormula"]
    headers = {h["name"]: h for h in v["headers"]}
    assert headers["apiKey"]["value"] == "••••••" and headers["apiKey"]["secret"] and headers["Authorization"]["column"] == "DT_bearerToken"
    checks = {c["actual"]: c for c in v["checks"]}
    assert checks["MedScen01_OUT"]["path"] == "detailResponse.policyDetail.travelers.[0].customElements.[1].value"
    assert checks["MedScen01_OUT"]["expectedValue"] == "Cover not available" and checks["PolicyStatus_OUT"]["kind"] == "compare_ignore_case"
    assert {s["column"] for s in v["saves"]} == {"AgentEmail_OUT", "Echo_OUT"}                  # read from the response, checked by nothing
    assert v["steps"][0]["kind"] == "send" and [s["kind"] for s in v["steps"]].count("check") == len(v["checks"])
    assert "KEY-123" not in json.dumps(v)
    assert api.get("/api/build/api/api.xlsx/tests/DataSheets").status_code == 422             # not an API sheet


# -- the ops ---------------------------------------------------------------------------------------------------------------------------------------------
def test_a_new_api_test_gets_its_sheet_data_row_datasheets_row_and_header_rows_in_one_undo_step(api):
    out = new_purchase_test(api)
    model = out["model"]
    buy = next(t for t in model["tests"] if t["id"] == "Buy")
    assert buy["kind"] == "api" and buy["listed"] and buy["enabled"] and buy["dataRows"][0]["row"] == 2
    v = api.get("/api/build/api/api.xlsx/tests/Buy").json()
    assert v["url"] == "{DOMAIN}policy/purchase/v2" and v["body"] == {"kind": "typed", "text": '{"lastName": "{LastName_IN}"}',
                                                                      "template": v["body"]["template"]}
    assert [(h["name"], h["column"], h["value"]) for h in v["headers"]] == [                  # apiKey reuses the legacy row; Content-Type is a
        ("apiKey", "DT_apiKey", "{SECRET:API_KEY}"), ("Content-Type", "H_Content_Type", "application/json")]   # header row like any other
    assert v["needs"] == [] and "DOMAIN" in v["environmentVariables"]
    api.post("/api/build/workbooks/api.xlsx/undo")
    assert "Buy" not in [t["id"] for t in api.get("/api/build/workbooks/api.xlsx").json()["tests"]]


def test_a_new_header_gets_its_own_inputoutput_row_and_removing_it_takes_the_row_away_when_no_other_sheet_uses_it(api):
    new_purchase_test(api)
    edit(api, {"op": "api_set_header", "test": "Buy", "row": 2, "name": "x-trace", "value": "abc"})
    v = api.get("/api/build/api/api.xlsx/tests/Buy").json()
    assert {h["name"]: h["column"] for h in v["headers"]}["x-trace"] == "H_x_trace"
    edit(api, {"op": "api_remove_header", "test": "Buy", "row": 2, "name": "x-trace"})
    doc = BuildDocument(api.root / "workbooks" / "api.xlsx")
    assert not [r for r in A.read_io(doc.editor) if r.parameter == "x-trace"]
    edit(api, {"op": "api_remove_header", "test": "Buy", "row": 2, "name": "apiKey"})             # Policy sends it too: only this row leaves it out
    doc = BuildDocument(api.root / "workbooks" / "api.xlsx")
    assert [r for r in A.read_io(doc.editor) if r.parameter == "apiKey"]
    assert doc.editor.get("Buy", 2, doc.editor.column("Buy", "DT_apiKey")) == "[BLANK]"


def test_a_check_adds_the_output_row_the_expected_column_and_the_check_and_the_same_path_reuses_its_column(api):
    new_purchase_test(api)
    path = "$.purchaseResponse.policyResponses[0].policyDetail.policyHolder.lastName"
    applied = edit(api, {"op": "api_add_check", "test": "Buy", "row": 2, "path": path, "kind": "compare", "expected": "{LastName_IN}"},
                   {"op": "api_add_check", "test": "Buy", "row": 2, "path": path, "kind": "contains", "expected": "Smi"},
                   {"op": "api_add_check", "test": "Buy", "row": 2, "path": "status", "kind": "compare", "expected": "200"},
                   {"op": "api_add_save", "test": "Buy", "row": 2, "path": "$.purchaseResponse.policyResponses[0].policyDetail.policyNumber",
                    "variable": "Policy no"},
                   {"op": "api_add_save", "test": "Buy", "row": 2, "path": path, "variable": "HOLDER"})["applied"]
    assert (applied[0]["actual"], applied[0]["expected"]) == ("lastName_OUT", "lastName_EXP")
    assert applied[1]["actual"] == "lastName_OUT" and applied[1]["reused"] and applied[1]["expected"] == "lastName_EXP_2"
    assert (applied[2]["actual"], applied[2]["expected"]) == ("RES_STATUS_CD_OUT", "RES_STATUS_CD_EXP")
    assert applied[3]["column"] == "POLICY_NO" and applied[4] == {**applied[4], "column": "lastName_OUT", "reused": True}   # one path, one column
    v = api.get("/api/build/api/api.xlsx/tests/Buy").json()
    assert [(c["label"], c["path"] or c["what"], c["expectedValue"]) for c in v["checks"]] == [
        ("is", path, "{LastName_IN}"), ("contains", path, "Smi"), ("is", "status", "200")]
    assert [s["column"] for s in v["saves"]] == ["POLICY_NO"] and v["provides"] == ["POLICY_NO", "lastName_OUT"]
    doc = BuildDocument(api.root / "workbooks" / "api.xlsx")
    io = [(r.function, r.parameter, r.value) for r in A.read_io(doc.editor)]
    assert io.count(("output_json", path, "lastName_OUT")) == 1 and ("compare", "RES_STATUS_CD_EXP", "RES_STATUS_CD_OUT") in io
    policy = api.get("/api/build/api/api.xlsx/tests/Policy").json()
    assert not any(c["actual"] == "lastName_OUT" for c in policy["checks"])                        # Policy has no such column: not its check


def test_an_io_row_path_stays_editable_and_can_be_deleted(api):
    new_purchase_test(api)
    applied = edit(api, {"op": "api_add_save", "test": "Buy", "row": 2, "path": "$.siteUrl", "variable": "SITE"})["applied"]
    v = api.get("/api/build/api/api.xlsx/tests/Buy").json()
    io_row = next(s["ioRow"] for s in v["saves"] if s["column"] == applied[0]["column"])
    edit(api, {"op": "api_update_io", "ioRow": io_row, "parameter": "$.transactionStatus"})
    assert next(s for s in api.get("/api/build/api/api.xlsx/tests/Buy").json()["saves"] if s["ioRow"] == io_row)["path"] == "$.transactionStatus"
    edit(api, {"op": "api_delete_io", "ioRows": [io_row]})
    assert not api.get("/api/build/api/api.xlsx/tests/Buy").json()["saves"]
    bad = api.post("/api/build/workbooks/api.xlsx/edit", json={"ops": [{"op": "api_set_request", "test": "Buy", "row": 2, "method": "FETCH"}]})
    assert bad.status_code == 422 and "FETCH" in bad.json()["error"]


def io_rows(api) -> list[tuple[str, str, str]]:
    return [(r.function, r.parameter, r.value) for r in A.read_io(BuildDocument(api.root / "workbooks" / "api.xlsx").editor)]


def headers_of(api, test: str) -> list[tuple[str, str]]:
    return [(h["name"], h["value"]) for h in api.get(f"/api/build/api/api.xlsx/tests/{test}").json()["headers"]]


# -- feedback item 13: headers typed into the table, names that can change ---------------------------------------------------------------------------------
def test_a_header_name_can_be_changed_and_its_value_stays_even_a_key_typed_in_the_sheet_that_the_page_never_sees(api):
    edit(api, {"op": "api_set_header", "test": "Policy", "row": 2, "name": "Authorization", "previous": "Authorization"})    # (same name: nothing)
    applied = edit(api, {"op": "api_set_header", "test": "Policy", "row": 2, "name": "X-Authorization", "previous": "Authorization"})["applied"]
    assert applied[0]["renamed"] and applied[0]["column"] == "DT_bearerToken"                  # only Policy sends it: the row is renamed in place
    assert ("addHeader", "X-Authorization", "DT_bearerToken") in io_rows(api) and not [r for r in io_rows(api) if r[1] == "Authorization"]
    assert dict(headers_of(api, "Policy"))["X-Authorization"] == "••••••"
    doc = BuildDocument(api.root / "workbooks" / "api.xlsx")
    assert doc.editor.get("Policy", 2, doc.editor.column("Policy", "DT_bearerToken")) == "Bearer TOKEN-abc"


def test_renaming_a_header_another_sheet_also_sends_moves_only_this_test_to_the_new_name_with_the_same_value(api):
    new_purchase_test(api)
    edit(api, {"op": "api_set_header", "test": "Buy", "row": 2, "name": "x-api-key", "previous": "apiKey"})
    assert dict(headers_of(api, "Buy"))["x-api-key"] == "{SECRET:API_KEY}" and "apiKey" not in dict(headers_of(api, "Buy"))
    assert "apiKey" in dict(headers_of(api, "Policy"))                                          # Policy still sends its apiKey
    doc = BuildDocument(api.root / "workbooks" / "api.xlsx")
    assert doc.editor.get("Buy", 2, doc.editor.column("Buy", "DT_apiKey")) == "[BLANK]"
    taken = api.post("/api/build/workbooks/api.xlsx/edit", json={"ops": [{"op": "api_set_header", "test": "Buy", "row": 2, "name": "Content-Type",
                                                                          "previous": "x-api-key"}]})
    assert taken.status_code == 409 and "already has a Content-Type header" in taken.json()["error"]


# -- feedback item 16: one kind of API test, Content-Type is a header row and JSON_FORMAT follows it ---------------------------------------------------------
def test_a_new_api_test_starts_with_a_json_content_type_row_and_json_format_follows_that_header_or_else_the_body(api):
    model = edit(api, {"op": "api_add_test", "name": "Plain"})["model"]
    assert next(t for t in model["tests"] if t["id"] == "Plain")["kind"] == "api"
    v = api.get("/api/build/api/api.xlsx/tests/Plain").json()
    assert [(h["name"], h["value"], h.get("implied")) for h in v["headers"]] == [("Content-Type", "application/json", None)] and v["format"] == "json"
    json_format = lambda: BuildDocument(api.root / "workbooks" / "api.xlsx").editor.get("Plain", 2, 6)
    edit(api, {"op": "api_set_header", "test": "Plain", "row": 2, "name": "Content-Type", "value": "text/xml; charset=utf-8"})
    assert json_format() == "N" and api.get("/api/build/api/api.xlsx/tests/Plain").json()["format"] == "xml"
    edit(api, {"op": "api_remove_header", "test": "Plain", "row": 2, "name": "Content-Type"},
         {"op": "api_set_request", "test": "Plain", "row": 2, "body": '{"plan": "Basic"}'})
    v = api.get("/api/build/api/api.xlsx/tests/Plain").json()                                   # no header: the body's shape decides
    assert json_format() == "Y" and v["headers"][0]["name"] == "Content-Type" and v["headers"][0]["implied"].startswith("sent for a JSON body")
    edit(api, {"op": "api_set_request", "test": "Plain", "row": 2, "body": "<Quote><Plan>Basic</Plan></Quote>"})
    assert json_format() == "N" and api.get("/api/build/api/api.xlsx/tests/Plain").json()["headers"] == []


def test_an_old_xml_test_opens_as_an_api_test_with_the_content_type_the_runner_sends_for_it(api):
    model = edit(api, {"op": "api_add_test", "name": "Soap", "format": "xml", "url": "{DOMAIN}policy/xml/v1", "body": "<Buy/>"})["model"]
    soap = next(t for t in model["tests"] if t["id"] == "Soap")
    assert soap["kind"] == "api" and api.get("/api/build/api/api.xlsx/tests/Soap").json()["format"] == "xml"      # (was kind "xml")
    assert next(t for t in model["tests"] if t["id"] == "Policy")["kind"] == "api"
    assert headers_of(api, "Policy")[0] == ("Content-Type", "application/json")                 # legacy JSON_FORMAT = Y: shown, as sent


# -- feedback item 14: checks typed by hand, and the whole response ------------------------------------------------------------------------------------------
def test_a_check_typed_by_hand_reads_a_jsonpath_or_an_xpath_from_the_path_itself_without_sending_anything(api):
    edit(api, {"op": "api_add_test", "name": "Quote", "method": "GET", "url": "{DOMAIN}policy/quote/v1"},
         {"op": "api_add_check", "test": "Quote", "row": 2, "path": "$.plan.code", "kind": "compare", "expected": "Basic"},
         {"op": "api_add_check", "test": "Quote", "row": 2, "path": "/Envelope/Body/Status/@code", "kind": "compare", "expected": "0"},
         {"op": "api_add_check", "test": "Quote", "row": 2, "path": "time", "kind": "less_than", "expected": "5000"})
    io = io_rows(api)
    assert ("output_json", "$.plan.code", "code_OUT") in io and ("Output", "/Envelope/Body/Status/@code", "code_OUT_2") in io
    v = api.get("/api/build/api/api.xlsx/tests/Quote").json()
    assert [(c["what"] or c["path"], c["kind"], c["expectedValue"]) for c in v["checks"]] == [
        ("$.plan.code", "compare", "Basic"), ("/Envelope/Body/Status/@code", "compare", "0"), ("time", "less_than", "5000")]
    assert "compare_response" not in [k["kind"] for k in v["checkKinds"]]                    # (the whole response has its own form)


def test_a_whole_response_check_writes_its_expected_body_its_check_row_and_its_ignore_list_and_the_list_can_change(api):
    edit(api, {"op": "api_add_test", "name": "Quote", "method": "GET", "url": "{DOMAIN}policy/quote/v1"},
         {"op": "api_add_check", "test": "Quote", "row": 2, "kind": "compare_response", "expected": '{"plan": {"code": "Basic"}}',
          "ignore": "quoteId; createdAt"},
         {"op": "api_add_test", "name": "Quote2", "method": "GET", "url": "{DOMAIN}policy/quote/v1"},
         {"op": "api_add_response_check", "test": "Quote2", "row": 2, "expected": "{}"})
    io = io_rows(api)
    assert ("compare_response", "RESPONSE_EXP", "RESPONSE_OUT") in io and ("ignore_in_response", "RESPONSE_EXP", "quoteId; createdAt") in io
    assert ("compare_response", "RESPONSE_EXP_2", "RESPONSE_OUT") in io                             # check rows are shared: a column of its own
    v = api.get("/api/build/api/api.xlsx/tests/Quote").json()
    (whole,) = v["checks"]
    assert (whole["what"], whole["ignore"], whole["expectedValue"]) == ("whole", ["quoteId", "createdAt"], '{"plan": {"code": "Basic"}}')
    assert v["steps"][1]["label"] == "the whole response is as expected (ignoring 2)"
    edit(api, {"op": "api_set_response_ignore", "expected": "RESPONSE_EXP", "ignore": ["quoteId", "createdAt", "$.items.id"]})
    assert ("ignore_in_response", "RESPONSE_EXP", "quoteId; createdAt; $.items.id") in io_rows(api)
    edit(api, {"op": "api_set_response_ignore", "expected": "RESPONSE_EXP", "ignore": ""})
    assert not [r for r in io_rows(api) if r[0] == "ignore_in_response"]


# -- feedback item 15: a status expected but never checked ----------------------------------------------------------------------------------------------
def test_an_expected_status_nothing_checks_is_a_problem_and_its_fix_switches_the_check_row_back_on(api):
    model = edit(api, {"op": "api_set_value", "test": "Policy", "row": 2, "column": "RES_STATUS_CD_EXP", "value": "200"})["model"]
    (problem,) = [p for p in model["problems"] if p["kind"] == "status_never_checked"]
    assert problem["test"] == "Policy" and problem["severity"] == "warning" and "compareX" in problem["message"] and "200" in problem["message"]
    state = api.get("/api/build/api/api.xlsx/tests/Policy").json()["statusUnchecked"]
    assert (state["expected"], state["offFunction"], state["sharedWith"]) == ("200", "compareX", [])
    out = edit(api, {"op": "api_add_status_check", "test": "Policy"})
    assert out["applied"][0]["added"] and not [p for p in out["model"]["problems"] if p["kind"] == "status_never_checked"]
    assert ("compare", "RES_STATUS_CD_EXP", "RES_STATUS_CD_OUT") in io_rows(api) and not [r for r in io_rows(api) if r[0] == "compareX"]
    v = api.get("/api/build/api/api.xlsx/tests/Policy").json()
    assert v["statusUnchecked"] is None and any(c["what"] == "status" and c["expectedValue"] == "200" for c in v["checks"])
    assert not edit(api, {"op": "api_add_status_check", "test": "Policy"})["applied"][0]["added"]     # (nothing left to fix)


def test_a_status_check_added_by_hand_on_a_sheet_that_has_the_column_uses_it_instead_of_making_a_second_one(api):
    applied = edit(api, {"op": "api_add_check", "test": "Policy", "row": 2, "path": "status", "kind": "compare", "expected": "200"})["applied"]
    assert (applied[0]["expected"], applied[0]["actual"]) == ("RES_STATUS_CD_EXP", "RES_STATUS_CD_OUT")
    assert ("compare", "RES_STATUS_CD_EXP", "RES_STATUS_CD_OUT") in io_rows(api) and not [r for r in io_rows(api) if r[0] == "compareX"]
    doc = BuildDocument(api.root / "workbooks" / "api.xlsx")
    assert doc.editor.column("Policy", "RES_STATUS_CD_EXP_2") is None


# -- the response tree ----------------------------------------------------------------------------------------------------------------------------------
PURCHASE = {"transactionStatus": "Success", "purchaseResponse": {"policyResponses": [
    {"policyDetail": {"policyNumber": "P-777", "policyHolder": {"firstName": "Claim", "lastName": "Smith"}}},
    {"policyDetail": {"policyNumber": "P-778", "policyHolder": {"firstName": "Duplicate", "lastName": "Second"}}}]}}


def test_a_value_inside_a_list_offers_this_item_and_the_item_where_with_a_column_of_the_row_when_one_holds_that_value():
    nodes, cut = A.json_tree(PURCHASE, {"WHO": "Duplicate", "blnExecute": "Y"})
    node = next(n for n in nodes if n["path"] == "$.purchaseResponse.policyResponses[1].policyDetail.policyNumber")
    assert not cut and node["value"] == '"P-778"' and node["text"] == "P-778"
    arr = node["array"]
    assert (arr["path"], arr["index"], arr["count"], arr["thisItem"]) == ("$.purchaseResponse.policyResponses", 1, 2, node["path"])
    first = arr["where"][0]
    assert first["field"] == "policyDetail.policyHolder.firstName" and first["unique"] and first["variable"] == "WHO"
    assert first["path"] == "$.purchaseResponse.policyResponses[?(@.policyDetail.policyHolder.firstName=='Duplicate')].policyDetail.policyNumber"
    assert first["variablePath"].endswith("[?(@.policyDetail.policyHolder.firstName=='{WHO}')].policyDetail.policyNumber")
    assert all(o["field"] != "policyDetail.policyNumber" for o in arr["where"])                     # the clicked value does not pick its own item
    top = next(n for n in nodes if n["path"] == "$.transactionStatus")
    assert "array" not in top and nodes[0]["type"] == "object"


def test_an_xml_response_tree_has_document_xpaths_attributes_and_the_item_where_a_child_or_attribute_has_a_value():
    xml = ('<soap:Envelope xmlns:soap="urn:s"><soap:Body><Policies><Policy id="P-901"><Plan>Basic</Plan></Policy>'
           '<Policy id="P-902"><Plan>Max</Plan></Policy></Policies></soap:Body></soap:Envelope>')
    nodes, _ = A.xml_tree(xml, {"PLAN": "Max"})
    paths = [n["path"] for n in nodes]
    assert "/Envelope/Body/Policies/Policy[2]/Plan" in paths and "/Envelope/Body/Policies/Policy[1]/@id" in paths
    attr = next(n for n in nodes if n["path"] == "/Envelope/Body/Policies/Policy[2]/@id")
    options = {o["field"]: o for o in attr["array"]["where"]}
    assert "@id" not in options and options["Plan"]["path"] == "/Envelope/Body/Policies/Policy[Plan='Max']/@id"
    assert options["Plan"]["variablePath"] == "/Envelope/Body/Policies/Policy[Plan='{PLAN}']/@id"
    plan = next(n for n in nodes if n["path"] == "/Envelope/Body/Policies/Policy[2]/Plan")
    assert {o["field"]: o["path"] for o in plan["array"]["where"]}["@id"] == "/Envelope/Body/Policies/Policy[@id='P-902']/Plan"
    assert A.xml_tree("not xml") == ([], False)


# -- Send now ---------------------------------------------------------------------------------------------------------------------------------------------
@pytest.mark.browser
def test_send_now_sends_the_draft_with_the_environments_values_and_shows_the_tree_and_what_each_check_says(api, monkeypatch):
    monkeypatch.setenv("RR_SECRET_API_KEY", "KEY-123")
    new_purchase_test(api)
    where = "$.purchaseResponse.policyResponses[?(@.policyDetail.policyHolder.firstName=='Duplicate')].policyDetail.policyNumber"
    edit(api, {"op": "api_add_check", "test": "Buy", "row": 2, "path": "$.purchaseResponse.policyResponses[0].policyDetail.policyHolder.lastName",
               "kind": "compare", "expected": "{LastName_IN}"},
         {"op": "api_add_check", "test": "Buy", "row": 2, "path": "status", "kind": "compare", "expected": "201"},
         {"op": "api_add_save", "test": "Buy", "row": 2, "path": where, "variable": "POLICY_NO"})
    r = api.post("/api/build/api/send", json={"workbook": "api.xlsx", "test": "Buy"})
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["environment"] == "UAT" and out["response"]["status"] == 200 and out["response"]["format"] == "json"
    assert out["request"]["url"] == api.site + "policy/purchase/v2" and json.loads(out["request"]["body"]) == {"lastName": "Smith"}
    assert {h["name"]: h["value"] for h in out["request"]["headers"]}["apiKey"] == "••••••"
    assert mock_server.API_SEEN["headers"]["apikey"] == "KEY-123" and "KEY-123" not in r.text
    assert out["outputs"]["POLICY_NO"] == "P-778"
    assert [(c["actual"], c["passed"]) for c in out["checks"]] == [("lastName_OUT", True), ("RES_STATUS_CD_OUT", False)]
    assert any(n["path"].endswith("policyResponses[1].policyDetail.policyNumber") and n["array"]["where"] for n in out["response"]["tree"])
    assert not api.get("/api/build/workbooks/api.xlsx/status").json()["externalChange"]              # the workbook itself is not written


@pytest.mark.browser
def test_send_now_asks_for_a_value_only_an_earlier_test_would_give_and_refuses_production_without_confirmation(api, monkeypatch):
    monkeypatch.setenv("RR_SECRET_API_KEY", "KEY-123")
    new_purchase_test(api)
    edit(api, {"op": "api_set_request", "test": "Buy", "row": 2, "url": "{DOMAIN}policy/purchase/v2?ref={ORDER_REF}"})
    first = api.post("/api/build/api/send", json={"workbook": "api.xlsx", "test": "Buy"}).json()
    assert first["missing"] == ["ORDER_REF"] and "response" not in first
    second = api.post("/api/build/api/send", json={"workbook": "api.xlsx", "test": "Buy", "values": {"ORDER_REF": "R 1"}}).json()
    assert second["response"]["status"] == 200 and mock_server.API_SEEN["path"] == "/policy/purchase/v2?ref=R%201"
    prod = api.post("/api/build/api/send", json={"workbook": "api.xlsx", "test": "Buy", "env": "PROD"})
    assert prod.status_code == 400 and prod.json()["kind"] == "prod_confirm"


@pytest.mark.browser
def test_send_now_reads_an_xml_answer_into_a_tree_and_a_saved_xpath_takes_its_value(api):
    edit(api, {"op": "api_add_test", "name": "Soap", "format": "xml", "method": "POST", "url": "{DOMAIN}policy/xml/v1",
               "headers": [["apiKey", "KEY-123"]], "body": "<Buy><LastName>{LastName_IN}</LastName></Buy>"},
         {"op": "api_set_value", "test": "Soap", "row": 2, "column": "LastName_IN", "value": "A & B"},
         {"op": "api_set_value", "test": "Soap", "row": 2, "column": "PLAN", "value": "Max"},
         {"op": "api_add_save", "test": "Soap", "row": 2, "path": "//Policy[Plan='{PLAN}']/@id", "variable": "POLICY_ID"},
         {"op": "api_add_check", "test": "Soap", "row": 2, "path": "/Envelope/Body/PolicyResponse/Status/@code", "kind": "compare", "expected": "0"})
    out = api.post("/api/build/api/send", json={"workbook": "api.xlsx", "test": "Soap"}).json()
    assert out["response"]["format"] == "xml" and "<LastName>A &amp; B</LastName>" in mock_server.API_SEEN["body"]
    assert out["outputs"]["POLICY_ID"] == "P-902" and out["checks"] == [{**out["checks"][0], "passed": True}]
    policy = next(n for n in out["response"]["tree"] if n["path"] == "/Envelope/Body/PolicyResponse/Policies/Policy[2]/LastName")
    assert policy["value"] == "Other" and {o["field"] for o in policy["array"]["where"]} >= {"@id", "Plan"}


@pytest.mark.browser
def test_the_whole_response_check_passes_send_now_and_a_run_when_the_fields_that_change_are_ignored_and_names_what_else_differs(api):
    """The mock quote API answers a new quoteId, createdAt and item ids on every call: the check passes only when those are ignored."""
    edit(api, {"op": "api_add_test", "name": "Quote", "method": "GET", "url": "{DOMAIN}policy/quote/v1"})
    first = api.post("/api/build/api/send", json={"workbook": "api.xlsx", "test": "Quote"}).json()
    expected = json.dumps(json.loads(first["response"]["text"]), indent=2)                   # "Use the last response"
    edit(api, {"op": "api_add_check", "test": "Quote", "row": 2, "kind": "compare_response", "expected": expected,
               "ignore": "quoteId; createdAt; $.items.id"},
         {"op": "api_add_test", "name": "Strict", "method": "GET", "url": "{DOMAIN}policy/quote/v1"},
         {"op": "api_add_check", "test": "Strict", "row": 2, "kind": "compare_response", "expected": expected, "ignore": "createdAt"})
    ok = api.post("/api/build/api/send", json={"workbook": "api.xlsx", "test": "Quote"}).json()
    assert [(c["kind"], c["passed"], c["actualValue"]) for c in ok["checks"]] == [("compare_response", True, "the same (4 ignored)")]
    bad = api.post("/api/build/api/send", json={"workbook": "api.xlsx", "test": "Strict"}).json()
    (check,) = bad["checks"]
    assert not check["passed"] and check["actualValue"].startswith("3 differences: $.quoteId: expected")
    assert [d.split(":")[0] for d in check["differences"]] == ["$.quoteId", "$.items[0].id", "$.items[1].id"]
    assert api.post("/api/build/workbooks/api.xlsx/save", json={}).status_code == 200
    proc, results, _ = run(api.root, api.root / "workbooks" / "api.xlsx", args=["--tests", "Quote,Strict"])
    tests = {t["id"]: t for t in results["tests"]}
    assert tests["Quote"]["status"] == "PASSED", proc.stdout + proc.stderr
    step = tests["Quote"]["steps"][1]
    assert (step["name"], step["action"], step["status"]) == ("Check the whole response", "COMPARE_RESPONSE", "PASSED")
    assert step["expected"] == "the whole response in RESPONSE_EXP (ignoring quoteId; createdAt; $.items.id)"
    strict = tests["Strict"]["steps"][1]
    assert tests["Strict"]["status"] == "FAILED" and strict["error"] == "Comparison Failed" and strict["notes"][0].startswith("$.quoteId: expected")


# -- the ways in ------------------------------------------------------------------------------------------------------------------------------------------
def test_a_curl_command_from_bash_or_windows_cmd_is_read_into_method_url_headers_and_body():
    bash = ("curl 'https://uat.example.test/bin/quote' \\\n  -H 'Content-Type: application/json' \\\n  -H 'x-api-key: abc123' \\\n"
            "  --data-raw '{\"plan\":\"Basic\",\"tripCost\":2500}' --compressed")
    got = A.parse_curl(bash)
    assert (got["method"], got["url"], got["format"]) == ("POST", "https://uat.example.test/bin/quote", "json")
    assert got["headers"] == [{"name": "Content-Type", "value": "application/json"}, {"name": "x-api-key", "value": "abc123"}]
    assert json.loads(got["body"]) == {"plan": "Basic", "tripCost": 2500}
    cmd = 'curl "https://h.test/a" ^\n  -X PUT ^\n  -H "Accept: */*" ^\n  --data-raw "^{^\\^"a^\\^":1^}"'
    assert A.parse_curl(cmd)["method"] == "PUT" and A.parse_curl(cmd)["url"] == "https://h.test/a"
    get = A.parse_curl("curl -G https://h.test/s -d q=1 -u me:pw")
    assert get["url"] == "https://h.test/s?q=1" and get["method"] == "GET" and get["headers"][0]["value"].startswith("Basic ")
    assert A.parse_curl("curl --json '{\"a\":1}' h.test/x")["headers"][-1] == {"name": "Content-Type", "value": "application/json"}
    with pytest.raises(ValueError):
        A.parse_curl("wget https://h.test/")


def test_an_imported_request_gets_the_environment_domain_a_secret_for_its_key_and_the_row_columns_for_its_values():
    envs = {"rows": [{"variable": "DOMAIN", "secret": False, "values": {"UAT": "https://uat.example.test/"}}], "production": []}
    request = {"method": "POST", "url": "https://uat.example.test/bin/quote", "format": "json", "body": '{"plan": "Basic", "tripCost": 2500}',
               "headers": [{"name": "x-api-key", "value": "abc123"}, {"name": "Authorization", "value": "Bearer tok-9"}]}
    out = A.suggest(request, environments=envs, environment="UAT", row_values={"PLAN": "Basic", "TRIP_COST": "2500", "blnExecute": "Y"})
    req = out["request"]
    assert req["url"] == "{DOMAIN}bin/quote" and req["body"] == '{"plan": "{PLAN}", "tripCost": {TRIP_COST}}'
    assert req["headers"] == [{"name": "x-api-key", "value": "{SECRET:API_KEY}"}, {"name": "Authorization", "value": "Bearer {SECRET:AUTHORIZATION}"}]
    assert "abc123" not in json.dumps(out) and "tok-9" not in json.dumps(out)
    assert [f["kind"] for f in out["found"]] == ["domain", "secret", "secret", "values"]


def test_a_postman_collection_lists_every_request_in_its_folders_with_its_variables_in_the_workbooks_form():
    collection = {"info": {"name": "Quote API"}, "item": [
        {"name": "Quotes", "item": [
            {"name": "Price", "request": {"method": "POST", "url": {"raw": "{{baseUrl}}/purchase/quote"}, "header": [
                {"key": "x-api-key", "value": "{{apiKey}}"}, {"key": "Old", "value": "x", "disabled": True}],
                "body": {"mode": "raw", "raw": "{\"plan\": \"{{plan}}\"}", "options": {"raw": {"language": "json"}}}}},
            {"name": "Lookup", "request": {"method": "GET", "url": "{{baseUrl}}/purchase/quote/{{quoteId}}"}}]},
        {"name": "Soap", "request": {"method": "POST", "url": "https://h.test/soap", "header": [{"key": "Content-Type", "value": "text/xml"}],
                                     "body": {"mode": "raw", "raw": "<a/>"}}}]}
    got = A.parse_postman(json.dumps(collection))
    assert [(r["name"], r["folder"], r["method"], r["format"]) for r in got] == [
        ("Price", "Quotes", "POST", "json"), ("Lookup", "Quotes", "GET", "json"), ("Soap", "", "POST", "xml")]
    assert got[0]["url"] == "{baseUrl}/purchase/quote" and got[0]["headers"] == [{"name": "x-api-key", "value": "{apiKey}"}]
    assert got[0]["body"] == '{"plan": "{plan}"}' and got[1]["url"].endswith("/{quoteId}")
    with pytest.raises(ValueError):
        A.parse_postman({"not": "a collection"})


def test_the_curl_and_postman_routes_answer_in_the_workbooks_form_without_changing_it(api):
    r = api.post("/api/build/api/curl", json={"workbook": "api.xlsx", "text": f"curl '{api.site}policy/v4/P-100' -H 'apiKey: KEY-123'"}).json()
    assert r["request"]["url"] == "{DOMAIN}policy/v4/P-100" and r["request"]["headers"][0]["value"] == "{SECRET:APIKEY}"
    bad = api.post("/api/build/api/curl", json={"workbook": "api.xlsx", "text": "hello"})
    assert bad.status_code == 422 and bad.json()["kind"] == "curl"
    pm = api.post("/api/build/api/postman", json={"workbook": "api.xlsx", "collection": {"item": [
        {"name": "One", "request": {"method": "GET", "url": api.site + "policy/v4/P-100"}}]}}).json()
    assert pm["requests"][0]["request"]["url"] == "{DOMAIN}policy/v4/P-100"
    assert api.get("/api/build/workbooks/api.xlsx/status").json()["modified"] is False


def test_a_request_template_lists_its_placeholders_with_what_fills_them_and_using_it_writes_the_replace_rows(api):
    folder = api.root / "workbooks" / "templates"
    folder.mkdir()
    (folder / "buy.xml").write_text("<Buy><Last>#LASTNAME#</Last><Plan>#PLAN#</Plan><Id>policyNumber_value</Id></Buy>", encoding="utf-8")
    (folder / "notes.md").write_text("x")
    new_purchase_test(api)
    listed = api.get("/api/build/api/api.xlsx/templates").json()
    assert [(t["file"], t["format"]) for t in listed["templates"] if t["folder"] == str(folder)] == [("buy.xml", "xml")]
    fields = api.get("/api/build/api/api.xlsx/templates/fields", params={"folder": str(folder), "file": "buy.xml", "test": "Buy"}).json()
    by = {f["placeholder"]: f for f in fields["fields"]}
    assert by["#LASTNAME#"]["column"] == "LastName_IN" and by["#PLAN#"]["column"] == "" and by["policyNumber_value"]["column"] == ""
    assert api.get("/api/build/api/api.xlsx/templates/fields", params={"folder": "/etc", "file": "passwd", "test": "Buy"}).status_code == 400
    assert api.get("/api/build/api/api.xlsx/templates/fields", params={"folder": str(folder), "file": "../api.xlsx", "test": "Buy"}).status_code == 400
    edit(api, {"op": "api_use_template", "test": "Buy", "row": 2, "location": str(folder), "file": "buy.xml", "format": "xml",
               "mappings": [{"placeholder": "#LASTNAME#", "column": "LastName_IN"}, {"placeholder": "#PLAN#", "value": "Max"}]})
    v = api.get("/api/build/api/api.xlsx/tests/Buy").json()
    assert v["body"]["kind"] == "template" and v["format"] == "xml" and v["body"]["template"]["file"] == "buy.xml"
    assert {(r["placeholder"], r["column"], r["value"]) for r in v["body"]["template"]["replace"]} == {
        ("#LASTNAME#", "LastName_IN", "Smith"), ("#PLAN#", "PLAN_IN", "Max")}
    edit(api, {"op": "api_add_test", "name": "Buy2", "format": "xml", "url": "{DOMAIN}x"},
         {"op": "api_set_value", "test": "Buy2", "row": 2, "column": "Surname", "value": "Jones"},
         {"op": "api_use_template", "test": "Buy2", "row": 2, "location": str(folder), "file": "buy.xml",
          "mappings": [{"placeholder": "#LASTNAME#", "column": "Surname"}]})
    doc = BuildDocument(api.root / "workbooks" / "api.xlsx")
    assert doc.editor.get("Buy2", 2, doc.editor.column("Buy2", "LastName_IN")).startswith("=")    # the Replace row stays: its column reads Surname
    assert api.post("/api/build/workbooks/api.xlsx/save", json={}).status_code == 200
    wb = Workbook(api.root / "workbooks" / "api.xlsx")
    buy2 = next(c for c in wb.discover() if c.sheet == "Buy2")
    assert wb.api_runtime(buy2).values["LASTNAME_IN"] == "Jones"
