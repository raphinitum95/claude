"""Builds synthetic keyword workbooks that use the same structure and formula tricks as the real one."""
from __future__ import annotations

from pathlib import Path

import openpyxl

COLUMNS = ["blnExecute", "Step_Number", "Step_Name", "Status", "Error_Check", "breakpoint", "Test_Case", "Page",
           "FindBy", "FindBy_Value", "Index", "Method", "Value", "Output_Property", "Ignore_not_existing_object",
           "isDisabled", "Timeout", "Expected_Value", "Output_Value", "Exact_Match", "Contains", "ITERATION_NO",
           "Notes", "Start_Time", "End_Time", "Duration"]
COL = {name: i + 1 for i, name in enumerate(COLUMNS)}
LETTER = {name: openpyxl.utils.get_column_letter(i + 1) for name, i in COL.items()}


class Sheet:
    """Row builder: ``add`` returns the 1-based row so later rows can reference it in formulas."""

    def __init__(self, ws, flag: str, extra_columns: list[str] | None = None):
        self.ws, self.flag = ws, flag
        self.columns = COLUMNS + (extra_columns or [])
        for i, name in enumerate(self.columns, start=1):
            ws.cell(1, i, name)
        self.row = 1

    def col_letter(self, name: str) -> str:
        return openpyxl.utils.get_column_letter(self.columns.index(name) + 1)

    def header(self, text: str) -> int:
        self.row += 1
        self.ws.cell(self.row, 1, text)
        return self.row

    def add(self, method: str, name: str = "", gate: str | None = None, **cols) -> int:
        self.row += 1
        r = self.row
        self.ws.cell(r, 1, gate if gate is not None else self.flag)
        self.ws.cell(r, COL["Step_Number"], f"=COUNTA($B$1:B{r - 1})")
        self.ws.cell(r, COL["Step_Name"], name or method)
        self.ws.cell(r, COL["Method"], method)
        for key, value in cols.items():
            self.ws.cell(r, self.columns.index(key) + 1, value)
        return r


def build_flow(sheet: Sheet, *, expect_bypass: str = "TEST-UAT-TOKEN", before_quit=None) -> dict[str, int]:
    """The reference user flow. Returns named row numbers so tests can assert on specific steps."""
    rows: dict[str, int] = {}
    s = sheet
    s.header("Open")
    rows["wait_first"] = s.add("Wait", "Wait before open", Value=1, Expected_Value="DT_URL", Ignore_not_existing_object="Y")
    rows["open"] = s.add("Open", "Open Browser", Page="DT_BrowserType", Value="DT_URL", Timeout=60)
    rows["bypass"] = s.add("Output", "Bypass cookie reached the server", FindBy="xpath",
                           FindBy_Value="//span[@id='bypass']", Index=0, Output_Property="innertext",
                           Expected_Value=f"bypass:{expect_bypass}", Exact_Match="Y")
    s.header("Quote")
    rows["exist"] = s.add("Exist", "Check banner", FindBy="xpath", FindBy_Value="//button[text()='Accept All Cookies']",
                          Ignore_not_existing_object="Y")
    rows["banner_out"] = s.add("Output", "Output Cookie", gate=f'=IF($D${rows["exist"]}="PASSED","{s.flag}","N")',
                               FindBy="xpath", FindBy_Value="//button[text()='Accept All Cookies']", Index=0,
                               Output_Property="innertext", Expected_Value="Accept All Cookies", Exact_Match="Y",
                               Ignore_not_existing_object="Y")
    rows["banner_click"] = s.add("js_click", "Cookies", FindBy="xpath",
                                 FindBy_Value="//button[@id='onetrust-accept-btn-handler']",
                                 Ignore_not_existing_object="Y")
    rows["banner_again"] = s.add("Exist", "Check banner again (gone)", FindBy="xpath",
                                 FindBy_Value="//button[text()='Accept All Cookies']", Ignore_not_existing_object="Y")
    rows["legacy_noop"] = s.add('getAttribute("value")', "Value present (legacy no-op)", FindBy="xpath",
                                FindBy_Value="//span[@id='bypass']", Index=0)
    rows["frame"] = s.add("SWITCHTOFRAME", "Switch to Frame", Value="aemFormFrame", Notes="tripType")
    rows["trip"] = s.add("js_click", "Select Trip Type", FindBy="xpath",
                         FindBy_Value=f'="//label[contains(.,\'"&{s.col_letter("Notes")}{rows["frame"]}&"\')]"', Index=0)
    rows["dest_write"] = s.add("Write", "Write Destination", FindBy="xpath", FindBy_Value="//input[@id='cmp-Input']",
                               Index=0, Value="Destination")
    rows["dest_pick"] = s.add("js_click", "Pick Destination", FindBy="xpath",
                              FindBy_Value=f'="//a[contains(.,\'"&M{rows["dest_write"]}&"\')]"', Index=0)
    rows["dep"] = s.add("Set", "Departure Date", FindBy="xpath", FindBy_Value="//input[@name='tripDepartureDate']",
                        Index=0, Value="DepDate")
    rows["ret"] = s.add("Set", "Return Date", FindBy="xpath", FindBy_Value="//input[@name='tripReturnDate']",
                        Index=0, Value="RetDate")
    rows["age1"] = s.add("Set", "Enter Age", FindBy="xpath", FindBy_Value="(//input[@name='insuredAge'])[1]",
                         Index=0, Value="Age")
    rows["add"] = s.add("js_click", "Click Add", FindBy="xpath", FindBy_Value="//button[@id='add-container-button__1']",
                        Index=0)
    rows["age2"] = s.add("Set", "Enter Age 2", FindBy="xpath", FindBy_Value="(//input[@name='insuredAge'])[2]",
                         Index=0, Value="Age1")
    rows["cost_yes"] = s.add("Set", "Enter Trip Cost (Comprehensive only)",
                             gate=f'=IF(${s.col_letter("Notes")}${rows["frame"]}="Comprehensive Coverage","Y","N")',
                             FindBy="xpath", FindBy_Value="//input[@name='display_tripCost']", Index=0, Value="Cost")
    rows["cost_no"] = s.add("Set", "Enter Trip Cost (Rental only - must NOT run)",
                            gate=f'=IF(${s.col_letter("Notes")}${rows["frame"]}="Rental Car Coverage","Y","N")',
                            FindBy="xpath", FindBy_Value="//input[@name='display_tripCost']", Index=0, Value="Cost1")
    rows["select"] = s.add("Select", "Select upgrade", FindBy="xpath", FindBy_Value="//select[@name='aoad-options']",
                           Index=0, Value="AOAD3")
    rows["tick"] = s.add("Tick", "I Agree", FindBy="xpath", FindBy_Value="//input[@value='I agree']", Index=0)
    rows["submit"] = s.add("js_click", "Submit", FindBy="xpath", FindBy_Value="//button[@type='SUBMIT']", Index=0)
    rows["wait_result"] = s.add("Wait", "Wait for result", Value=5)
    s.header("Result")
    rows["plan"] = s.add("Output", "Plan displayed", FindBy="xpath",
                         FindBy_Value="//span[@class='handlebar-recursive-replacement productName']", Index=0,
                         Output_Property="innertext", Expected_Value="Travel Med Go", Exact_Match="Y")
    rows["premium_capture"] = s.add("Output", "Capture premium", FindBy="xpath",
                                    FindBy_Value="//span[contains(@class,'totalBaseRate')]", Index=0,
                                    Output_Property="innertext", Output_Value="Premium_OUT")
    rows["premium_compare"] = s.add("Output", "Premium unchanged", FindBy="xpath",
                                    FindBy_Value="//span[contains(@class,'totalBaseRate')]", Index=0,
                                    Output_Property="innertext", Expected_Value=f'=S{rows["premium_capture"]}',
                                    Exact_Match="Y")
    rows["dep_shown"] = s.add("Output", "Departure shown", FindBy="xpath",
                              FindBy_Value="//span[contains(@class,'tripDepartureDate')]", Index=0,
                              Output_Property="innertext",
                              Expected_Value=f'=TEXT(M{rows["dep"]},"mm/dd/yyyy")', Exact_Match="Y")
    rows["default_a"] = s.add("SWITCHTODEFAULT", "Change iFrame")
    rows["pay_frame"] = s.add("switchtoframe", "Change iFrame", Index=1)
    rows["month"] = s.add("Type", "Input Expiry Month", FindBy="xpath", FindBy_Value="//input[@name='tsep-datepicker']",
                          Index=0, Value="ExpiryMonth")
    rows["tab"] = s.add("SendKeys", "Tab", Value="{TAB}")
    rows["year"] = s.add("SENDKEYS", "Type year", Value="ExpiryYear")
    rows["zoom"] = s.add("SENDKEYS", "Zoom out (ignored headless)", Value="^-^-")
    rows["pay_value"] = s.add("Output", "Expiry echoed", FindBy="xpath", FindBy_Value="//span[@id='dp-value']", Index=0,
                              Output_Property="innertext", Expected_Value="12|2030", Exact_Match="Y")
    rows["default_b"] = s.add("switchtodefault", "Change iFrame")
    rows["arm"] = s.add("js_click", "Arm the late button", FindBy="xpath", FindBy_Value="//button[@id='arm']", Index=0)
    rows["disabled_click"] = s.add("js_click", "Click button that enables late", FindBy="xpath",
                                   FindBy_Value="//button[@id='disabled-btn']", Index=0)
    rows["disabled_state"] = s.add("Output", "Late button really clicked", FindBy="xpath",
                                   FindBy_Value="//span[@id='dis-state']", Index=0, Output_Property="innertext",
                                   Expected_Value="clicked", Exact_Match="Y")
    rows["popup"] = s.add("js_click", "Open second window", FindBy="xpath", FindBy_Value="//a[@id='open-win']", Index=0)
    rows["win"] = s.add("SWITCHTOWINDOW", "Switch to window", Value="-1")
    rows["second"] = s.add("Output", "Second window heading", FindBy="xpath", FindBy_Value="//h1[@id='second']", Index=0,
                           Output_Property="innertext", Expected_Value="Second Window", Exact_Match="Y")
    rows["main"] = s.add("SwitchToMainWindow", "Back to main")
    rows["scroll"] = s.add("MOUSE_SCROLL", "Scroll", Value="10, N")
    rows["nbsp"] = s.add("Output", "NBSP is normalised like Selenium", FindBy="xpath", FindBy_Value="//p[@id='nb']",
                         Index=0, Output_Property="innertext", Expected_Value="Hello World", Exact_Match="Y")
    rows["hidden"] = s.add("Output", "Hidden text reads as blank", FindBy="xpath",
                           FindBy_Value="//span[@id='hidden-span']", Index=0, Output_Property="innertext",
                           Exact_Match="Y")
    rows["url"] = s.add("GET_CURRENT_URL", "Verify URL", Output_Property="innertext", Output_Value="URL_OUT")
    rows["shot"] = s.add("SCREENSHOT", "Confirmation screenshot", Value='=Global!B9 & "ConfirmationPage.png"')
    rows["ignored_missing"] = s.add("Click", "Optional element that is absent", FindBy="xpath",
                                    FindBy_Value="//button[@id='does-not-exist']", Index=0,
                                    Ignore_not_existing_object="Y")
    if before_quit:
        before_quit(s, rows)
    rows["quit"] = s.add("Quit", "Quit")
    return rows


def build_workbook(path: Path, *, flows: list[str] | None = None, extra_steps=None, enabled: list[str] | None = None,
                   base_url: str = "http://127.0.0.1:1/", trip_type: str = "Comprehensive Coverage",
                   extra_columns: list[str] | None = None, expect_bypass: str = "TEST-UAT-TOKEN") -> dict[str, dict[str, int]]:
    flows = flows or ["Flow"]
    enabled = enabled if enabled is not None else flows
    wb = openpyxl.Workbook()
    g = wb.active
    g.title = "Global"
    g.append(["Parameter", "Value", "Comments"])
    g.append(["Environment", "UAT"])
    g.append(["ResultFileLocation", '="C:\\Results\\"&B2&"\\"'])
    g.append(["ResultFilename", '=B9&"_x.xlsx"'])
    for name, value in [("isHTMLParsing", "N"), ("removeNamespacesInResponse", "N"), ("sendOutlookMail", "N"),
                        ("CreateHtmlReport", "Y")]:
        g.append([name, value])
    g.append(["ResultTitle", '="Mock Regression_"&B2'])          # row 9 -> referenced as Global!B9
    g.append(["QuitBreakOnFailure", "N"])
    g.append(["HeadlessMode", "N"])
    ds = wb.create_sheet("DataSheets")
    ds.append(["SheetsToExecute", "blnExecute", "ParameterSheet", "Comments", "Tags"])
    named: dict[str, dict[str, int]] = {}
    for i, name in enumerate(flows, start=1):
        ds.append([name, "Y" if name in enabled else "N", f"Params_{i}", f"Mock flow {name}", "smoke" if i == 1 else "full"])
        flag = f"bln{i}00{i}"
        ws = wb.create_sheet(name)
        sheet = Sheet(ws, flag, extra_columns)
        named[name] = build_flow(sheet, expect_bypass=expect_bypass, before_quit=extra_steps)
        ps = wb.create_sheet(f"Params_{i}")
        ps.append(["Scenario", "blnExecute", flag, "DT_BrowserType", "Notes", "DT_URL", "DateToday", "tripType",
                   "Destination", "DepDate", "RetDate", "Age", "Age1", "Cost", "Cost1", "AOAD3", "ExpiryMonth",
                   "ExpiryYear", "Premium_OUT", "URL_OUT", "FirstName"])
        ps.append([flag, "Y", "Y", "Chrome", f"TC_{i}", f'=IF(Global!B2="UAT","{base_url}","http://prod.invalid/")',
                   "=TODAY()", trip_type, "Singapore", '=TEXT(G2+1,"mm/dd/yyyy")', '=TEXT(J2+89,"mm/dd/yyyy")',
                   "30", "25", "1500", "900", "AOAD3", "12", "2030", None, None,
                   '="QAFIRSTNAME"&CHAR(RANDBETWEEN(65,90))'])
    wb.save(path)
    return named
