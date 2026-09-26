"""The legacy row loop, reproduced: substitution, gating, write-back, iterations, secrets."""
from __future__ import annotations

from datetime import datetime, timedelta

import openpyxl
import pytest

from regrunner.workbook import Workbook
from tests.workbook_factory import build_workbook


@pytest.fixture
def flow(tmp_path):
    path = tmp_path / "wb.xlsx"
    rows = build_workbook(path, base_url="http://site.test/")["Flow"]
    wb = Workbook(path, environment="UAT", seed=5)
    return wb, wb.discover()[0], rows, path


def advance(rt, upto: int, status="PASSED", overrides=None):
    """Run the resolution loop up to (excluding) ``upto``, recording each executed step."""
    for row in range(2, upto):
        step = rt.prepare_row(row)
        if step is not None:
            st, out = (overrides or {}).get(row, (status, ""))
            rt.record(step, st, "", out)


def test_discovery_and_plan(flow):
    wb, case, rows, _ = flow
    assert (case.id, case.param_sheet, case.scenario, case.tags, case.enabled) == ("Flow", "Params_1", "TC_1", ["smoke"], True)
    planned = wb.runtime(case).plan()
    assert len(planned) == 48                                      # reference flow with the Rental-only step gated off
    assert rows["cost_no"] not in [p[0] for p in planned] and rows["cost_yes"] in [p[0] for p in planned]


def test_tokens_are_substituted_in_place_and_formulas_see_them(flow):
    wb, case, rows, _ = flow
    rt = wb.runtime(case)
    advance(rt, rows["dest_pick"])
    pick = rt.prepare_row(rows["dest_pick"])
    assert pick.findby_value == "//a[contains(.,'Singapore')]"       # formula built from the substituted Value cell
    trip = None
    rt2 = wb.runtime(case)
    advance(rt2, rows["trip"])
    trip = rt2.prepare_row(rows["trip"])
    assert trip.findby_value == "//label[contains(.,'Comprehensive Coverage')]"   # from the Notes token
    open_step = wb.runtime(case)
    advance(open_step, rows["open"])
    op = open_step.prepare_row(rows["open"])
    assert op.text("VALUE") == "http://site.test/" and op.page == "CHROME"


def test_date_parameters_are_computed_now_not_read_from_the_cache(flow):
    wb, case, rows, _ = flow
    rt = wb.runtime(case)
    advance(rt, rows["dep"])
    dep = rt.prepare_row(rows["dep"]).text("VALUE")
    assert dep == (datetime.now() + timedelta(days=1)).strftime("%m/%d/%Y")


def test_gating_reacts_to_substituted_notes_and_to_step_status(flow):
    wb, case, rows, _ = flow
    rt = wb.runtime(case)
    advance(rt, rows["cost_yes"])
    assert rt.prepare_row(rows["cost_yes"]) is not None            # W<frame> was substituted -> "Comprehensive Coverage"
    assert rt.prepare_row(rows["cost_no"]) is None
    # A failed Exist step switches off the Output gated on `$D$<exist>="PASSED"`.
    failed_rt = wb.runtime(case)
    advance(failed_rt, rows["banner_out"], overrides={rows["exist"]: ("FAILED", "")})
    assert failed_rt.prepare_row(rows["banner_out"]) is None
    ok_rt = wb.runtime(case)
    advance(ok_rt, rows["banner_out"])
    assert ok_rt.prepare_row(rows["banner_out"]) is not None


def test_output_value_flows_into_later_expected_values_and_into_params(flow):
    wb, case, rows, _ = flow
    rt = wb.runtime(case)
    advance(rt, rows["premium_capture"])
    cap = rt.prepare_row(rows["premium_capture"])
    assert cap.raw_output_name == "Premium_OUT"
    rt.record(cap, "PASSED", "", "$150.00")
    cmp_step = rt.prepare_row(rows["premium_compare"])
    assert cmp_step.expected == "$150.00"                          # `=S<row>` sees what the earlier step captured
    assert rt.params["PREMIUM_OUT"] == "$150.00"                   # named output variable written to the param row
    rt.record(cap, "PASSED", "", "$160.00")
    assert rt.params["PREMIUM_OUT"] == "$150.00;$160.00"           # legacy: repeated writes append with ';'


def test_blank_parameters_do_not_overwrite_tokens_and_unknown_tokens_stay_literal(flow):
    wb, case, rows, _ = flow
    rt = wb.runtime(case)
    assert rt.params["URL_OUT"] is None
    advance(rt, rows["url"])
    step = rt.prepare_row(rows["url"])
    assert step.raw_output_name == "URL_OUT"


def test_environment_override_changes_urls_and_gating(tmp_path):
    path = tmp_path / "wb.xlsx"
    build_workbook(path, base_url="http://uat.test/")
    for env, expected in (("UAT", "http://uat.test/"), ("PROD", "http://prod.invalid/")):
        wb = Workbook(path, environment=env)
        rt = wb.runtime(wb.discover()[0])
        assert rt.environment == env and rt.params["DT_URL"] == expected


def test_volatile_values_are_frozen_per_test_and_reproducible_with_a_seed(tmp_path):
    path = tmp_path / "wb.xlsx"
    build_workbook(path, base_url="http://x/")
    a = Workbook(path, seed=11)
    b = Workbook(path, seed=11)
    case = a.discover()[0]
    r1, r2 = a.runtime(case), b.runtime(case)
    assert r1.params["FIRSTNAME"] == r2.params["FIRSTNAME"] and r1.params["FIRSTNAME"].startswith("QAFIRSTNAME")
    other = Workbook(path, seed=12).runtime(case)
    assert other.params["FIRSTNAME"] != r1.params["FIRSTNAME"]     # a different seed gives different data
    first = r1.params["FIRSTNAME"]
    for row in range(2, 12):                                       # every write invalidates Excel-style caches...
        step = r1.prepare_row(row)
        if step is not None:
            r1.record(step, "PASSED", "", "x")
    assert r1.params["FIRSTNAME"] == first                         # ...but the value is frozen (Excel re-rolled it)
    assert r1.params_sheet.read(2, r1.params_sheet.data.headers["FIRSTNAME"]) == first


def test_secret_tokens_are_resolved_but_never_stored_in_the_sheet(tmp_path):
    path = tmp_path / "wb.xlsx"

    def extra(sheet, rows):
        rows["login"] = sheet.add("Set", "Zscaler user", FindBy="xpath", FindBy_Value="//input[@id='u']", Index=0,
                                  Value="DT_ZScalerUser", Ignore_not_existing_object="Y")

    rows = build_workbook(path, base_url="http://x/", extra_steps=extra)["Flow"]
    wb = Workbook(path, secrets={"DT_ZSCALERUSER": "me@example.com"})
    rt = wb.runtime(wb.discover()[0])
    advance(rt, rows["login"])
    step = rt.prepare_row(rows["login"])
    assert step.text("VALUE") == "me@example.com" and "VALUE" in step.secret_columns
    col = rt.columns["VALUE"]
    assert rt.sheet.get(rows["login"], col) == "DT_ZScalerUser"     # the workbook state keeps only the token name
    rt.record(step, "PASSED", "", "")
    assert "me@example.com" not in repr(rt.sheet.overlay) and "me@example.com" not in repr(rt.params)


def test_multiple_enabled_parameter_rows_become_iterations(tmp_path):
    path = tmp_path / "wb.xlsx"
    build_workbook(path, base_url="http://x/")
    book = openpyxl.load_workbook(path)
    ps = book["Params_1"]
    second = [c.value for c in ps[2]]
    ps.append(second)
    ps.cell(3, 5, "TC_1b")
    book.save(path)
    cases = Workbook(path).discover()
    assert [c.id for c in cases] == ["Flow#1", "Flow#2"] and cases[1].param_row == 3 and cases[1].iterations == 2
    assert cases[1].scenario == "TC_1b"


def test_disabled_parameter_row_is_skipped(tmp_path):
    path = tmp_path / "wb.xlsx"
    build_workbook(path, base_url="http://x/")
    book = openpyxl.load_workbook(path)
    book["Params_1"].cell(2, 2, "N")
    book.save(path)
    wb = Workbook(path)
    (case,) = wb.discover()
    assert case.enabled is False and case.param_enabled is False   # legacy ran nothing for a disabled parameter row
    assert wb.runtime(case).plan() == [] and any("no row with blnExecute=Y" in w for w in wb.warnings)
