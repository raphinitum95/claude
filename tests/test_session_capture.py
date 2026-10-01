"""What a test records of the browser session (capture.*): the console at each step, what its server calls sent and answered, the form fields
and storage after each step, and a warning for a switched-off field that holds nothing.  Plain logic first (fast), then one run on the mock site."""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

from regrunner.capture.state import StateTracker, locked_without_value, replay
from regrunner.engine.runner import RunOptions, execute
from regrunner.engine.session import mask_payload
from regrunner.events import EventBus
from tests.workbook_factory import Sheet


def field(label="Date of Birth", **over):
    return {"label": label, "type": "text", "value": "", "empty": True, "checked": None, "disabled": False, "readonly": False,
            "ariaDisabled": False, "placeholder": "", "visible": True, **over}


def snap(fields, url="https://site/a?x=1", storage=None, messages=None):
    return {"url": url, "title": "t", "fields": fields, "storage": storage or {}, "cookies": [], "messages": messages or []}


# -- the warning: a switched-off field with nothing in it ------------------------------------------------------------------------------------
def test_a_disabled_field_that_shows_only_its_placeholder_is_reported():
    found = locked_without_value({"dob2": field(disabled=True, placeholder="DD/MM/YYYY")})
    assert [w["key"] for w in found] == ["dob2"]
    assert found[0]["text"] == '"Date of Birth" is disabled and has no value (it shows the placeholder "DD/MM/YYYY")'


def test_a_disabled_field_that_holds_the_placeholder_text_as_its_value_is_reported():
    found = locked_without_value({"d": field(disabled=True, value="DD/MM/YYYY", empty=False, placeholder="DD/MM/YYYY")})
    assert "shows its placeholder" in found[0]["text"]


def test_read_only_and_aria_disabled_fields_count_as_switched_off():
    found = locked_without_value({"a": field("A", readonly=True), "b": field("B", ariaDisabled=True)})
    assert [w["how"] for w in found] == ["read-only", "not editable"]


@pytest.mark.parametrize("over", [
    {"disabled": True, "value": "28/01/2006", "empty": False},          # filled in: fine
    {"disabled": False},                                                 # a person can type in it: not the site's job to fill it
    {"disabled": True, "visible": False},                                # not on screen
    {"disabled": True, "type": "checkbox"},                              # an unticked box is an answer
])
def test_fields_that_are_fine_are_not_reported(over):
    assert locked_without_value({"f": field(**over)}) == []


# -- the line written after each step -------------------------------------------------------------------------------------------------
def test_the_first_look_at_a_page_is_in_full_and_later_ones_hold_only_what_changed():
    tracker = StateTracker("changes")
    a = field("First name", value="Ann", empty=False)
    first, _ = tracker.observe(snap({"fn": a, "ln": field("Last")}), seq=1, row=3, name="open")
    assert first["full"] and set(first["fields"]) == {"fn", "ln"}
    second, _ = tracker.observe(snap({"fn": a, "ln": field("Last", value="Lee", empty=False)}), seq=2, row=4, name="type")
    assert not second["full"] and set(second["fields"]) == {"ln"}
    third, _ = tracker.observe(snap({"fn": a, "ln": field("Last", value="Lee", empty=False)}), seq=3, row=5, name="wait")
    assert "fields" not in third and "storage" not in third


def test_a_different_page_is_looked_at_in_full_again_but_a_new_query_string_is_the_same_page():
    tracker = StateTracker("changes")
    tracker.observe(snap({"a": field()}, url="https://site/a"), seq=1, row=1, name="x")
    same, _ = tracker.observe(snap({"a": field()}, url="https://site/a?step=2"), seq=2, row=2, name="y")
    other, _ = tracker.observe(snap({"b": field()}, url="https://site/b"), seq=3, row=3, name="z")
    assert not same["full"] and other["full"]


def test_every_step_mode_writes_the_whole_page_each_time():
    tracker = StateTracker("every_step")
    tracker.observe(snap({"a": field()}), seq=1, row=1, name="x")
    again, _ = tracker.observe(snap({"a": field()}), seq=2, row=2, name="y")
    assert again["full"] and "a" in again["fields"]


def test_an_empty_switched_off_field_is_warned_once_then_again_on_a_failed_step_and_after_it_recovers():
    tracker = StateTracker("changes")
    empty = {"dob2": field(disabled=True, placeholder="DD/MM/YYYY")}
    filled = {"dob2": field(disabled=True, value="28/01/2006", empty=False)}
    _, w1 = tracker.observe(snap(empty), seq=1, row=1, name="a")
    _, w2 = tracker.observe(snap(empty), seq=2, row=2, name="b")
    _, w3 = tracker.observe(snap(empty), seq=3, row=3, name="submit", failed=True)
    _, w4 = tracker.observe(snap(filled), seq=4, row=4, name="c")
    _, w5 = tracker.observe(snap(empty), seq=5, row=5, name="d")
    assert [len(w) for w in (w1, w2, w3, w4, w5)] == [1, 0, 1, 0, 1]


def test_replaying_the_lines_gives_the_page_as_it_was_after_a_step():
    tracker = StateTracker("changes")
    lines = []
    for seq, fields in enumerate([{"a": field("A")}, {"a": field("A", value="1", empty=False)}, {"a": field("A", value="2", empty=False), "b": field("B")}], start=1):
        line, _ = tracker.observe(snap(fields, storage={"session:q": str(seq)}), seq=seq, row=seq, name=f"s{seq}")
        lines.append(line)
    assert replay(lines, 2)["fields"]["a"]["value"] == "1" and "b" not in replay(lines, 2)["fields"]
    assert replay(lines, 3)["storage"] == {"session:q": "3"} and "b" in replay(lines, 3)["fields"]


# -- bodies -------------------------------------------------------------------------------------------------------------------------------
def test_secret_looking_fields_of_a_body_are_hidden_but_dates_and_postcodes_stay():
    assert mask_payload('{"dob":"28/01/2006","postcode":"123456","password":"abc","cvv":123}') == '{"dob":"28/01/2006","postcode":"123456","password":"***","cvv":"***"}'
    assert mask_payload("dob=28%2F01%2F2006&password=zz&cardNo=4111") == "dob=28%2F01%2F2006&password=***&cardNo=***"


def test_a_long_body_is_cut():
    assert mask_payload("x" * 50, limit=10) == "x" * 10 + "... (cut)"


# -- one run on the mock site -------------------------------------------------------------------------------------------------------------
def build(path: Path, site: str) -> Path:
    wb = openpyxl.Workbook()
    ds = wb.active
    ds.title = "DataSheets"
    ds.append(["SheetsToExecute", "blnExecute", "ParameterSheet", "Comments"])
    ds.append(["Flow", "Y", "Params_1", "capture"])
    g = wb.create_sheet("Global")
    g.append(["Parameter", "Value"])
    g.append(["Environment", "UAT"])
    ps = wb.create_sheet("Params_1")
    ps.append(["blnExecute", "DT_URL"])
    ps.append(["Y", site + "/session_capture.html"])
    s = Sheet(wb.create_sheet("Flow"), "Y")
    s.add("Open", "Open", Page="chrome", Value="DT_URL")
    s.add("Set", "type the first name", FindBy="xpath", FindBy_Value="//input[@id='fname']", Index=0, Value="Ann")
    s.add("Set", "type the password", FindBy="xpath", FindBy_Value="//input[@id='pw']", Index=0, Value="Hunter2-Secret")
    s.add("Click", "save", FindBy="xpath", FindBy_Value="//button[@id='save']", Index=0)
    s.add("Wait", "let the answer show", Value=1)
    wb.save(path)
    return path


def lines(path: Path) -> list[dict]:
    return [json.loads(x) for x in path.read_text(encoding="utf-8").splitlines() if x.strip()]


@pytest.mark.browser
async def test_a_run_keeps_the_console_the_calls_with_their_bodies_and_the_fields_after_each_step(site, make_cfg, tmp_path):
    cfg = make_cfg(**{"waits.window_grace_s": 0.5, "timeouts.element_s": 3, "timeouts.optional_s": 1.0})
    events: list[dict] = []
    bus = EventBus()
    bus.subscribe(events.append)
    result = await execute(RunOptions(workbook=build(tmp_path / "wb.xlsx", site), tests=["Flow"], seed=1), cfg, bus)
    folder = tmp_path / "runs" / result.run_id / "tests" / "Flow"

    console = lines(folder / "console.jsonl")
    assert any(c["level"] == "log" and c["text"] == "page loaded" for c in console)                      # not only errors
    assert any(c["level"] == "warning" and "not ready" in c["text"] for c in console)
    error = next(c for c in console if c["level"] == "pageerror")
    assert "datepickerInit" in error["text"] and "session_capture.html" in error["stack"] and error["step"] >= 1

    call = next(c for c in lines(folder / "network.jsonl") if c["url"].startswith("/bin/capture/submit"))
    assert call["status"] == 200 and call["step"] >= 1
    assert json.loads(call["request_body"]) == {"dob": "28/01/2006", "password": "***"}                    # what was sent: the password never
    assert "system error" in call["response_body"] and "server-side-secret" not in call["response_body"]

    state = lines(folder / "state.jsonl")
    assert state[0]["full"] and state[0]["fields"]["dob2"]["disabled"] and state[0]["fields"]["dob2"]["empty"]
    after_typing = replay(state, 3)
    assert after_typing["fields"]["fname"]["value"] == "Ann"
    assert after_typing["fields"]["pw"]["value"] == "(hidden)"                                             # a password field is never recorded
    assert after_typing["storage"]["session:authToken"] == "(hidden)" and "dob" in after_typing["storage"]["session:quote"]
    assert any("system error" in m for m in replay(state, 5)["messages"])                                  # the banner the page showed
    assert "Hunter2-Secret" not in (folder / "state.jsonl").read_text(encoding="utf-8")
    assert "Hunter2-Secret" not in (folder / "network.jsonl").read_text(encoding="utf-8")

    first = result.tests[0].steps[0]
    assert any("Date of Birth" in n and n.startswith("Warning:") for n in first.notes)                      # on the step, in the results
    assert any(e["type"] == "review_item" and e.get("category") == "disabled_field_empty" for e in events)
    assert not any(n.startswith("Warning:") for n in result.tests[0].steps[1].notes)                        # said once, not on every step


@pytest.mark.browser
async def test_capture_can_be_switched_off(site, make_cfg, tmp_path):
    cfg = make_cfg(**{"waits.window_grace_s": 0.5, "timeouts.element_s": 3, "capture.console": "off", "capture.bodies": "off", "capture.state": "off"})
    bus = EventBus()
    result = await execute(RunOptions(workbook=build(tmp_path / "wb.xlsx", site), tests=["Flow"], seed=1), cfg, bus)
    folder = tmp_path / "runs" / result.run_id / "tests" / "Flow"
    assert not (folder / "console.jsonl").exists() and not (folder / "state.jsonl").exists()
    assert "request_body" not in (folder / "network.jsonl").read_text(encoding="utf-8")


@pytest.mark.browser
async def test_a_trace_is_kept_when_asked_for_always(site, make_cfg, tmp_path):
    cfg = make_cfg(**{"waits.window_grace_s": 0.5, "timeouts.element_s": 3, "capture.trace": "always"})
    result = await execute(RunOptions(workbook=build(tmp_path / "wb.xlsx", site), tests=["Flow"], seed=1), cfg, EventBus())
    assert (tmp_path / "runs" / result.run_id / "tests" / "Flow" / "trace.zip").stat().st_size > 1000


@pytest.mark.browser
async def test_a_trace_of_a_passing_test_is_thrown_away_when_asked_for_on_failure(site, make_cfg, tmp_path):
    cfg = make_cfg(**{"waits.window_grace_s": 0.5, "timeouts.element_s": 3, "capture.trace": "on_failure"})
    result = await execute(RunOptions(workbook=build(tmp_path / "wb.xlsx", site), tests=["Flow"], seed=1), cfg, EventBus())
    assert result.tests[0].status == "PASSED"
    assert not (tmp_path / "runs" / result.run_id / "tests" / "Flow" / "trace.zip").exists()


# -- config -------------------------------------------------------------------------------------------------------------------------------
def test_a_bare_off_in_config_yaml_means_the_word_off_and_a_wrong_choice_is_refused(tmp_path):
    from regrunner.config import load_config
    (tmp_path / "config.yaml").write_text("capture:\n  trace: off\n  console: off\n  state: every_step\n")       # (YAML reads a bare `off` as false)
    cfg = load_config(tmp_path / "config.yaml", base_dir=tmp_path)
    assert (cfg.capture.trace, cfg.capture.console, cfg.capture.state, cfg.capture.bodies) == ("off", "off", "every_step", "servlets")
    (tmp_path / "config.yaml").write_text("capture:\n  state: sometimes\n")
    with pytest.raises(ValueError, match="capture.state"):
        load_config(tmp_path / "config.yaml", base_dir=tmp_path)
