"""The legacy runner starts every step with ``OV = the row's Output_Value cell`` (a formula Excel had calculated) and only an action that captures something replaces
it: a ``Wait`` row can therefore carry a comparison - the agent portal's "Verify points to earn" is a Wait row whose Output_Value is a formula over the previous row's
capture and whose Expected_Value is what was calculated before."""
from __future__ import annotations

import openpyxl
import pytest

from regrunner.engine.runner import RunOptions, execute
from regrunner.events import EventBus
from tests.workbook_factory import Sheet

pytestmark = pytest.mark.browser


def build(path, site: str, expected: str, tail=()):
    wb = openpyxl.Workbook()
    ds = wb.active
    ds.title = "DataSheets"
    ds.append(["SheetsToExecute", "blnExecute", "ParameterSheet", "Comments"])
    ds.append(["Flow", "Y", "Params_1", "actual"])
    g = wb.create_sheet("Global")
    g.append(["Parameter", "Value"])
    g.append(["Environment", "UAT"])
    ps = wb.create_sheet("Params_1")
    ps.append(["blnExecute", "DT_URL", "DT_Points"])
    ps.append(["Y", site + "/points.html", None])
    s = Sheet(wb.create_sheet("Flow"), "Y")
    s.add("Open", "Open", Page="chrome", Value="DT_URL")
    got = s.add("Output", "Get points to earn", FindBy="xpath", FindBy_Value="//span[@id='pts']", Index=0, Output_Property="innertext")
    s.add("Wait", "Verify points to earn", Value=1, Expected_Value=expected, Exact_Match="Y",
          Output_Value=f'=TEXT(_xlfn.NUMBERVALUE(S{got},"."),"0.00")')
    for method, name, cols in tail:
        s.add(method, name, **cols)
    wb.save(path)


async def run(make_cfg, wb):
    events = []
    bus = EventBus()
    bus.subscribe(events.append)
    result = await execute(RunOptions(workbook=wb, tests=["Flow"], seed=1), make_cfg(**{"waits.window_grace_s": 0.5, "timeouts.element_s": 3}), bus)
    return {s.name: s for s in result.tests[0].steps}


async def test_a_wait_row_is_compared_with_its_own_output_value_cell(site, make_cfg, tmp_path):
    build(tmp_path / "wb.xlsx", site, "129.00")
    steps = await run(make_cfg, tmp_path / "wb.xlsx")
    verify = steps["Verify points to earn"]
    assert verify.status == "PASSED" and verify.actual == "129.00" and verify.expected == "129.00"


async def test_and_a_real_mismatch_still_fails_and_shows_what_the_cell_held(site, make_cfg, tmp_path):
    build(tmp_path / "wb.xlsx", site, "130.00")
    verify = (await run(make_cfg, tmp_path / "wb.xlsx"))["Verify points to earn"]
    assert verify.status == "FAILED" and verify.error == "Comparison Failed" and verify.actual == "129.00" and verify.expected == "130.00"


async def test_a_cell_that_only_names_a_parameter_is_not_a_value(site, make_cfg, tmp_path):
    """Output_Value = DT_Points means "store the output there"; a step that captured nothing must not compare (or report) the name as if it were text."""
    build(tmp_path / "wb.xlsx", site, "129.00", tail=[("Wait", "names a parameter", {"Value": 1, "Output_Value": "DT_Points", "Expected_Value": "x", "Exact_Match": "Y"})])
    steps = await run(make_cfg, tmp_path / "wb.xlsx")
    assert steps["names a parameter"].actual == "" and steps["names a parameter"].status == "FAILED"          # blank vs "x": same as before, the name is not the actual
