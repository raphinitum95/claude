"""What a run set: parameters a step fills in (Output_Value names a Params column - Quote_OUT, DT_Policy_Out) are listed in results.json, the HTML report,
the shared summary and the web UI, so the value a test produced does not have to be dug out of a step."""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

from regrunner.engine.runner import RunOptions, execute
from regrunner.events import EventBus
from regrunner.publish import summary_text
from regrunner.reporting.from_events import results_from_events
from regrunner.reporting.results import variables_of
from tests.test_blank_params import VALUE, answers_with, build

pytestmark = pytest.mark.browser

SECRET = "s3cr3t-Quote-77"


async def run(cfg, wb: Path, tests: list[str], ask: str = "off", listen=None):
    bus = EventBus()
    events: list[dict] = []
    bus.subscribe(events.append)
    if listen:
        bus.subscribe(listen(cfg))
    result = await execute(RunOptions(workbook=wb, tests=tests, seed=1, extra={"ask": ask}), cfg, bus)
    return result, events


async def test_the_parameters_a_step_set_are_listed_with_their_value_and_where_they_came_from(site, make_cfg, tmp_path):
    cfg = make_cfg()
    result, events = await run(cfg, build(tmp_path / "a.xlsx", site), ["Buy"])
    (test,) = result.tests
    assert test.status == "PASSED"
    (var,) = test.variables
    assert var["name"] == "DT_Policy_Out" and var["value"] == "Google Authenticator" and var["stored"] == "Google Authenticator"
    assert var["cell"] == "Params_1!C2" and var["step"] == "Get Policy Number" and var["row"] == 3 and var["by_hand"] is False
    step = next(s for s in test.steps if s.name == "Get Policy Number")
    assert step.sets == [{"name": "DT_Policy_Out", "value": "Google Authenticator", "stored": "Google Authenticator", "cell": "Params_1!C2"}]
    assert next(e for e in events if e["type"] == "step_passed" and e["name"] == "Get Policy Number")["sets"] == step.sets      # the live view reads this

    run_dir = cfg.path(cfg.runs_dir) / result.run_id
    saved = json.loads((run_dir / "results.json").read_text())["tests"][0]
    assert variables_of(saved)[0]["name"] == "DT_Policy_Out"
    html = (run_dir / "report.html").read_text()
    assert "Values set by the tests" in html and "DT_Policy_Out" in html and "Google Authenticator" in html and "sets DT_Policy_Out" in html
    text = summary_text(json.loads((run_dir / "results.json").read_text()))
    assert "Values set by the tests:" in text and "Buy: DT_Policy_Out = Google Authenticator" in text and "Get Policy Number" in text

    rebuilt = results_from_events(run_dir)                                              # a run that stopped early keeps them as well
    assert [v["name"] for v in rebuilt.tests[0].variables] == ["DT_Policy_Out"]


async def test_a_test_that_sets_nothing_has_no_section(site, make_cfg, tmp_path):
    cfg = make_cfg()
    wb = build(tmp_path / "a.xlsx", site, filled=VALUE)
    result, _ = await run(cfg, wb, ["View"])                                            # only reads DT_Policy_Out
    assert result.tests[0].variables == []
    html = (cfg.path(cfg.runs_dir) / result.run_id / "report.html").read_text()
    assert "Values set by the tests" not in html


async def test_a_secret_a_person_typed_is_never_shown_as_a_value(site, make_cfg, tmp_path):
    """ASK_USER (secret) whose Output_Value names a parameter: later steps get the real text, every report shows the mask."""
    wb = openpyxl.load_workbook(build(tmp_path / "a.xlsx", site))
    ws = wb["Buy"]
    hdr = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        if row[hdr.index("Method")].value == "Output":                                   # turn the capture into a secret question
            row[hdr.index("Method")].value = "ASK_USER"
            row[hdr.index("Value")].value = "Quote number?"
            row[hdr.index("Output_Property")].value = "SECRET"
    wb.save(tmp_path / "s.xlsx")
    cfg = make_cfg()
    result, events = await run(cfg, tmp_path / "s.xlsx", ["Buy"], ask="ui", listen=answers_with(SECRET))
    (var,) = result.tests[0].variables
    assert var["name"] == "DT_Policy_Out" and var["stored"] == "••••••" and var["value"] == "••••••"
    run_dir = cfg.path(cfg.runs_dir) / result.run_id
    everything = "".join(p.read_text(errors="ignore") for p in run_dir.rglob("*") if p.is_file() and p.suffix in (".json", ".jsonl", ".html", ".log", ".txt"))
    assert SECRET not in everything


async def test_a_value_a_person_supplied_is_listed_as_entered_by_hand(site, make_cfg, tmp_path):
    cfg = make_cfg()
    result, events = await run(cfg, build(tmp_path / "a.xlsx", site), ["View"], ask="ui", listen=answers_with(VALUE))
    (var,) = result.tests[0].variables
    assert var["name"] == "DT_Policy_Out" and var["stored"] == VALUE and var["by_hand"] is True and var["cell"] == "Params_1!C2"
    assert [e["by_hand"] for e in events if e["type"] == "variable_set"] == [True]
    assert "entered by hand" in (cfg.path(cfg.runs_dir) / result.run_id / "report.html").read_text()
