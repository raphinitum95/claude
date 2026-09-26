"""ASK_USER: the run stops, asks a person (run screen or terminal), and the answer lands where the workbook says - typed into a field, or in
Output_Value for a later step.  What a person typed as a secret is never stored anywhere."""
from __future__ import annotations

import asyncio
import json
import subprocess
import sys
import time
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest

from regrunner.config import Config
from regrunner.engine.actions import ASK_METHODS, StepContext, ask_person, run_action
from regrunner.engine.ask import AskCancelled, AskTimeout, AskUnavailable, Asker, write_answer
from regrunner.lint import lint
from regrunner.workbook import Workbook
from regrunner.workbook.model import PreparedStep
from tests.workbook_factory import Sheet

SECRET = "s3cr3t-Answer-77"


class Bus:
    def __init__(self):
        self.events: list[dict] = []

    def emit(self, type_: str, **fields):
        self.events.append({"type": type_, **fields})


def asker(tmp_path, mode="ui", timeout=5.0, cancel=None):
    bus = Bus()
    return Asker(tmp_path, bus, cancel or asyncio.Event(), mode=mode, default_timeout_s=timeout), bus


# -- the broker ------------------------------------------------------------------------------------------------------------------
async def test_a_run_with_nobody_to_ask_says_so_at_once(tmp_path):
    a, bus = asker(tmp_path, mode="off")
    with pytest.raises(AskUnavailable, match="nobody to ask"):
        await a.ask("T", 3, "Code?")
    assert bus.events == []                                                            # nothing is announced that nobody can answer


async def test_a_question_waits_for_the_answer_file_and_the_answer_never_travels_in_an_event(tmp_path):
    a, bus = asker(tmp_path)
    task = asyncio.create_task(a.ask("Portal", 8, "Code from your phone?", secret=True))
    for _ in range(50):
        await asyncio.sleep(0.05)
        if bus.events:
            break
    needed = bus.events[0]
    assert needed["type"] == "user_input_needed" and needed["test"] == "Portal" and needed["step"] == 8 and needed["secret"] is True
    assert needed["question"] == "Code from your phone?" and needed["mode"] == "ui" and needed["timeout_s"] == 5
    assert write_answer(tmp_path, needed["ask"], SECRET) == ""
    assert await asyncio.wait_for(task, 3) == SECRET
    assert [e["type"] for e in bus.events] == ["user_input_needed", "user_input_received"]
    assert SECRET not in json.dumps(bus.events)
    assert not list((tmp_path / "answers").glob("*.json"))                             # the file is gone as soon as it is read
    a.cleanup()
    assert not (tmp_path / "answers").exists()


async def test_an_unanswered_question_times_out_and_a_late_answer_is_refused_not_left_on_disk(tmp_path):
    a, bus = asker(tmp_path, timeout=0.6)
    with pytest.raises(AskTimeout, match="within 0.6 s"):
        await a.ask("T", 1, "Name?")
    closed = bus.events[-1]
    assert closed["type"] == "user_input_closed" and closed["reason"] == "timeout"
    assert write_answer(tmp_path, closed["ask"], "too late") == "closed"
    assert not list((tmp_path / "answers").glob("*.json"))
    assert write_answer(tmp_path / "elsewhere", "a1-abcdef", "x") == "unknown"


async def test_a_cancelled_run_stops_waiting(tmp_path):
    cancel = asyncio.Event()
    a, bus = asker(tmp_path, timeout=30, cancel=cancel)
    task = asyncio.create_task(a.ask("T", 1, "Name?"))
    await asyncio.sleep(0.3)
    cancel.set()
    with pytest.raises(AskCancelled):
        await asyncio.wait_for(task, 3)
    assert bus.events[-1]["reason"] == "cancelled"


# -- the step ----------------------------------------------------------------------------------------------------------------------
class FakeAsker:
    available = True

    def __init__(self, answer="", error=None):
        self.answer, self.error, self.asked = answer, error, []

    async def ask(self, test, step, question, *, secret=False, timeout_s=None):
        self.asked.append((test, step, question, secret, timeout_s))
        if self.error:
            raise self.error
        return self.answer


def step_ctx(tmp_path, values, asker=None):
    step = PreparedStep(row=8, values={"METHOD": "ASK_USER", **values})
    return StepContext(step=step, session=SimpleNamespace(is_open=False, typed=set()), cfg=Config(), selector_map=None, review=None, test_dir=tmp_path,
                       test_id="Portal", seq=4, asker=asker)


def test_ask_user_puts_the_question_and_the_answer_becomes_the_output(tmp_path):
    fake = FakeAsker("Ada")
    ctx = step_ctx(tmp_path, {"VALUE": "What is your first name?", "TIMEOUT": 45}, fake)
    asyncio.run(run_action(ctx))
    assert fake.asked == [("Portal", 4, "What is your first name?", False, 45)]              # the Timeout column is the wait
    assert ctx.out.error == "" and ctx.out.output == "Ada" and not ctx.out.secret and ctx.secret_values == set()


def test_a_secret_answer_is_flagged_and_remembered_so_it_can_be_masked_everywhere(tmp_path):
    fake = FakeAsker(SECRET)
    ctx = step_ctx(tmp_path, {"VALUE": "Code?", "OUTPUT_PROPERTY": "secret"}, fake)
    asyncio.run(run_action(ctx))
    assert fake.asked[0][3] is True and ctx.out.output == SECRET and ctx.out.secret and ctx.secret_values == {SECRET}


def test_ask_user_without_a_question_uses_the_step_name(tmp_path):
    fake = FakeAsker("x")
    ctx = step_ctx(tmp_path, {"STEP_NAME": "Type the voucher"}, fake)
    asyncio.run(run_action(ctx))
    assert fake.asked[0][2] == "Type the voucher"


@pytest.mark.parametrize("error,fragment", [(AskTimeout("Nobody answered within 5 s."), "No answer to"),
                                            (AskUnavailable("needs a person"), "needs a person"),
                                            (AskCancelled("The run was cancelled while waiting for an answer."), "cancelled")])
def test_a_question_that_gets_no_answer_fails_the_step_with_the_reason(tmp_path, error, fragment):
    ctx = step_ctx(tmp_path, {"VALUE": "Code?"}, FakeAsker(error=error))
    asyncio.run(run_action(ctx))
    assert fragment in ctx.out.error and ctx.out.output is None


def test_the_google_token_step_asks_for_the_code_when_it_has_no_usable_key_and_someone_can_answer(tmp_path):
    fake = FakeAsker(" 123456 ")
    ctx = step_ctx(tmp_path, {"METHOD": "GET_GOOGLE_TOKEN", "VALUE": "DT_Key"}, fake)          # the token was never replaced by a real key
    asyncio.run(run_action(ctx))
    assert ctx.out.error == "" and ctx.out.output == "123456" and ctx.out.secret and "asked for the code instead" in " ".join(ctx.out.notes)
    assert fake.asked[0][2].startswith("Enter the 6-digit code from Google Authenticator.") and fake.asked[0][3] is True
    assert "Each code works once" in fake.asked[0][2]                     # a person reading it off the phone cannot be coordinated: told instead
    ctx = step_ctx(tmp_path, {"METHOD": "GET_GOOGLE_TOKEN", "VALUE": "DT_Key"}, None)            # nobody to ask: the old, clear error
    asyncio.run(run_action(ctx))
    assert "base32" in ctx.out.error and "ASK_USER" in ctx.out.error
    key = FakeAsker("999999")
    ctx = step_ctx(tmp_path, {"METHOD": "GET_GOOGLE_TOKEN", "VALUE": "JBSWY3DPEHPK3PXP"}, key)   # a real key: computed, nobody bothered
    asyncio.run(run_action(ctx))
    assert key.asked == [] and len(ctx.out.output) == 6 and ctx.out.output != "999999"


def test_steps_that_may_wait_for_a_person_are_not_cut_off_by_the_step_time_limit():
    assert {"ASK_USER", "PROMPT", "ASK", "GET_GOOGLE_TOKEN"} <= ASK_METHODS


# -- workbooks: what the Excel looks like -------------------------------------------------------------------------------------------------
def build_ask_workbook(path: Path, site: str, style: str) -> Path:
    wb = openpyxl.Workbook()
    ds = wb.active
    ds.title = "DataSheets"
    ds.append(["SheetsToExecute", "blnExecute", "ParameterSheet", "Comments"])
    ds.append(["Flow", "Y", "Params_1", "ask"])
    g = wb.create_sheet("Global")
    g.append(["Parameter", "Value"])
    g.append(["Environment", "UAT"])
    ps = wb.create_sheet("Params_1")
    ps.append(["blnExecute", "DT_URL", "DT_Key"])
    ps.append(["Y", f"{site}ask.html", None])
    s = Sheet(wb.create_sheet("Flow"), "Y")
    s.add("Open", "Open", Page="chrome", Value="DT_URL")
    if style == "into_field":            # ask and type it here: FindBy says which field
        ask = s.add("ASK_USER", "Ask for the code", FindBy="xpath", FindBy_Value="//input[@id='code']", Index=0,
                    Value="Enter the code from your phone", Output_Property="SECRET")
        s.add("Output", "Field holds the answer", FindBy="xpath", FindBy_Value="//span[@id='echo']", Index=0, Output_Property="innertext",
              Expected_Value=f"=S{ask}", Exact_Match="Y")           # the page echoes the field as it is typed into
    elif style == "then_set":            # ask, keep the answer in Output_Value, a later Set types it
        ask = s.add("ASK_USER", "Ask for the code", Value="Enter the code from your phone", Output_Property="SECRET")
        s.add("Set", "Type it", FindBy="xpath", FindBy_Value="//input[@id='code']", Index=0, Value=f"=S{ask}")
        s.add("Output", "Echoed", FindBy="xpath", FindBy_Value="//span[@id='echo']", Index=0, Output_Property="innertext", Expected_Value=f"=S{ask}",
              Exact_Match="Y")
    elif style == "google":              # the authenticator step with no key in DT_Key
        tok = s.add("get_google_token", "Get Okta PW", Value="DT_Key")
        s.add("Set", "Google PW", FindBy="xpath", FindBy_Value="//input[@id='code']", Index=0, Value=f"=S{tok}")
        s.add("Output", "Echoed", FindBy="xpath", FindBy_Value="//span[@id='echo']", Index=0, Output_Property="innertext", Expected_Value=f"=S{tok}",
              Exact_Match="Y")
    s.add("Quit", "Quit")
    wb.save(path)
    return path


def test_lint_says_that_the_step_needs_a_person_and_flags_a_meaningless_output_property(tmp_path, site):
    wb = Workbook(build_ask_workbook(tmp_path / "a.xlsx", site, "then_set"), seed=1)
    said = [(f.severity, f.message) for f in lint(wb)]
    assert any(sev == "info" and "asks a person while the run is going" in m for sev, m in said)
    assert not [1 for sev, _ in said if sev == "error"]
    sheet = openpyxl.load_workbook(tmp_path / "a.xlsx")
    ws = sheet["Flow"]
    for row in ws.iter_rows(min_row=2):
        if row[11].value == "ASK_USER":
            row[13].value = "innertext"
    sheet.save(tmp_path / "b.xlsx")
    assert any(sev == "warning" and "means nothing for ASK_USER" in m for sev, m in [(f.severity, f.message) for f in lint(Workbook(tmp_path / "b.xlsx"))])


def test_the_test_list_counts_the_questions_so_the_ui_can_warn_before_the_run_starts(tmp_path, site):
    from regrunner import insight
    (test,) = insight.list_tests(Workbook(build_ask_workbook(tmp_path / "a.xlsx", site, "then_set"), seed=1), Config())
    assert test["asks"] == 1
    (test,) = insight.list_tests(Workbook(build_ask_workbook(tmp_path / "g.xlsx", site, "google"), seed=1), Config())
    assert test["asks"] == 0


# -- a real run at the terminal ------------------------------------------------------------------------------------------------------
def run_terminal(tmp_path, wb_path, stdin: str, ask: str = "terminal", extra_cfg: str = "", timeout: int = 120):
    cfg = tmp_path / "config.yaml"
    cfg.write_text("timeouts: {element_s: 6}\ncaptcha_bypass: {values: {UAT: TEST-UAT-TOKEN}}\nreports: {html: true}\n" + extra_cfg)
    proc = subprocess.run([sys.executable, "-m", "regrunner", "--config", str(cfg), "run", str(wb_path), "--seed", "1", "--ask", ask],
                          cwd=tmp_path, capture_output=True, text=True, input=stdin, timeout=timeout)
    (run_dir,) = sorted((tmp_path / "runs").iterdir())[-1:]
    return proc, json.loads((run_dir / "results.json").read_text()), run_dir


def everything_the_run_wrote(run_dir: Path, proc) -> str:
    text = proc.stdout + proc.stderr
    for path in run_dir.rglob("*"):
        if path.is_file() and path.suffix in (".json", ".jsonl", ".html", ".log", ".txt"):
            text += path.read_text(errors="ignore")
    return text


@pytest.mark.browser
@pytest.mark.parametrize("style", ["into_field", "then_set"])
def test_the_answer_typed_at_the_terminal_lands_in_the_field_and_is_stored_nowhere(tmp_path, site, style):
    proc, results, run_dir = run_terminal(tmp_path, build_ask_workbook(tmp_path / "a.xlsx", site, style), SECRET + "\n")
    assert proc.returncode == 0, proc.stdout + proc.stderr
    (test,) = results["tests"]
    assert test["status"] == "PASSED" and test["failed"] == 0                             # the page echoed what landed in the field
    ask_step = next(s for s in test["steps"] if s["action"] == "ASK_USER")
    assert ask_step["value"] == "Enter the code from your phone" and ask_step["actual"] == "••••••"          # the question is shown, the answer never
    assert "Enter the code from your phone" in proc.stdout                                 # the person was asked at the terminal
    assert SECRET not in everything_the_run_wrote(run_dir, proc).replace(f"{SECRET}\n", "", 0)
    assert not (run_dir / "answers").exists()
    if style == "then_set":
        typed = next(s for s in test["steps"] if s["name"] == "Type it")
        assert typed["value"] == "••••••"                                                  # the Set that used it does not show it either
        expected = next(s for s in test["steps"] if s["name"] == "Echoed")["expected"]
        assert expected == "••••••"


@pytest.mark.browser
def test_the_authenticator_step_asks_at_the_terminal_when_the_key_cell_is_empty(tmp_path, site):
    proc, results, run_dir = run_terminal(tmp_path, build_ask_workbook(tmp_path / "a.xlsx", site, "google"), "654321\n")
    assert proc.returncode == 0, proc.stdout + proc.stderr
    assert results["tests"][0]["status"] == "PASSED" and "Enter the 6-digit code from Google Authenticator" in proc.stdout
    assert "654321" not in everything_the_run_wrote(run_dir, proc)


@pytest.mark.browser
def test_a_run_started_from_a_script_fails_the_step_at_once_instead_of_hanging(tmp_path, site):
    started = time.monotonic()
    proc, results, _ = run_terminal(tmp_path, build_ask_workbook(tmp_path / "a.xlsx", site, "then_set"), "", ask="off")
    assert time.monotonic() - started < 60
    (test,) = results["tests"]
    (bad,) = [s for s in test["steps"] if s["action"] == "ASK_USER"]
    assert bad["status"] == "FAILED" and "nobody to ask" in bad["error"]


@pytest.mark.browser
def test_nobody_answering_fails_the_step_after_the_timeout(tmp_path, site):
    proc, results, _ = run_terminal(tmp_path, build_ask_workbook(tmp_path / "a.xlsx", site, "then_set"), "", extra_cfg="ask: {timeout_s: 2}\n")
    (bad,) = [s for s in results["tests"][0]["steps"] if s["action"] == "ASK_USER"]
    assert bad["status"] == "FAILED" and "Nobody answered within 2 s" in bad["error"]
