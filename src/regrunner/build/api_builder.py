"""The Build tab's API / XML test editor (P10; design Q18/Q19): what an API test is in the workbook, the edits the form makes, "Send now", the
clickable response tree, and the three ways in (paste cURL, a Postman collection, an existing request template).

An API test stays what the runner already reads (``workbook/api.py``): one data row of an API sheet plus the workbook's shared ``InputOutput``
sheet.  The builder writes only that format, plus a few additions the new runner understands (CONTRACT.md: P10 section):

* ``REQUEST_BODY`` column: the body typed in the form, with ``{NAME}``s (a template file still works; a typed body wins);
* ``{NAME}`` / ``{SECRET:NAME}`` in ``WEBSERVICE_URL``, header cells and response paths; a secret header's cell holds ``{SECRET:NAME}``, never
  the value (secrets.env gives it: ``RR_SECRET_<NAME>``);
* a clicked value becomes an ``output_json`` (JSONPath, ``$.a.b``) or ``Output`` (XPath, ``/a/b/@c``) row into a column; a check adds the
  expected column (``<name>_EXP``, the value in the row) and a ``compare`` / ``contains`` / ``greater_than`` / ``less_than`` / ``between`` /
  ``matches`` row; a save is an output row into a column named for the variable, which the run's shared variables get.

``InputOutput`` rows are shared by every API sheet of the workbook (a row applies to a sheet that has its columns), and the reader keys them by
path / header / placeholder.  So a path, header or placeholder that already has a row reuses that row's column instead of adding a second row that
would take it away from the other sheets.

Every edit is an op (``api_*``) applied through ``BuildDocument.apply`` like the P02 ops: one call = one undo step, autosaved as the draft.
"Send now" reads the draft, builds the request with the runner's own code (``engine/api_runner.prepare_body`` / ``fetch`` / ``read_response``),
sends it with Playwright's request context (the proxy / HTTPS settings of a run) and never writes the workbook.
"""
from __future__ import annotations

import dataclasses
import json
import os
import re
import shlex
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

from ..config import Config, workbook_secrets
from ..workbook import builder as B
from ..workbook.api import CHECK_KINDS, METHODS, check
from ..workbook.api_paths import filter_literal, json_path_of, xpath_of
from ..workbook.api_template import xml_to_tree
from ..workbook.builder import BuildDocument, BuildError, _Grid, _Ops, _sheet_key, is_formula, is_true, text_of
from ..workbook.writer import WorkbookEditor, WriterError

MASK = "••••••"
IO_SHEET = "InputOutput"
IO_HEADERS = ("Function", "Parameter", "Value")
API_COLUMNS = ("blnExecute", "TCID", "TC_Name", "WEBSERVICE_METHOD", "WEBSERVICE_URL", "JSON_FORMAT", "REQUEST_BODY")
OUTPUT_FUNCTIONS = {"output_json": "output_json", "output": "Output"}
CHECK_LABELS = {"compare": "is", "compare_ignore_case": "is (any case)", "contains": "contains", "contains_ignore_case": "contains (any case)",
                "greater_than": "is greater than", "less_than": "is less than", "between": "is between", "matches": "matches"}
STATUS_COLUMNS = {"status": ("RES_STATUS_CD_EXP", "RES_STATUS_CD_OUT"), "time": ("ELAPSEDTIME_EXP", "ELAPSEDTIME_OUT")}
SECRET_HEADER = re.compile(r"authori[sz]ation|api[-_ ]?key|token|secret|password|passwd|cookie|session|signature", re.I)
TREE_LIMIT = 4000                    # nodes of a response tree sent to the page (a huge response is cut, and says so)
TEXT_LIMIT = 400_000                 # characters of a raw response sent to the page

_FUNCTION_KIND = {"addheader": "add_header", "add_header": "add_header", "update_json": "update_json", "output_json": "output_json",
                  "output": "output", "replace": "replace", "disablecheckpoint": "disable", "disable_checkpoint": "disable",
                  **{k: k for k in CHECK_KINDS}}


# ---------------------------------------------------------------------------------------------------------------------------------------------
# Reading
# ---------------------------------------------------------------------------------------------------------------------------------------------
@dataclasses.dataclass
class IoRow:
    row: int
    kind: str              # add_header | update_json | output_json | output | replace | disable | a CHECK_KINDS name | "" (ignored)
    function: str          # as written
    parameter: str
    value: str

    def to_json(self) -> dict:
        return {"ioRow": self.row, "kind": self.kind, "function": self.function, "parameter": self.parameter, "value": self.value}


def io_sheet(editor: WorkbookEditor) -> str | None:
    return _sheet_key(editor, IO_SHEET)


def read_io(editor: WorkbookEditor) -> list[IoRow]:
    """Every row of ``InputOutput`` that the runner acts on (a row with no function or no parameter is skipped by the runner too)."""
    name = io_sheet(editor)
    if name is None:
        return []
    out = []
    for r, values in enumerate(editor.rows(name)[1:], start=2):
        cells = [text_of(values[i]).strip() if i < len(values) else "" for i in range(3)]
        function, parameter, value = cells
        if not function or not parameter:
            continue
        out.append(IoRow(r, _FUNCTION_KIND.get(function.lower(), ""), function, parameter, value))
    return out


def _cols(editor: WorkbookEditor, sheet: str) -> dict[str, int]:
    return {h.upper(): c for h, c in editor.headers(sheet).items()}


def api_sheet(editor: WorkbookEditor, name: Any) -> str:
    key = _sheet_key(editor, str(name or ""))
    if key is None:
        raise BuildError(f"There is no sheet called {name!r}.", "not_found", 404)
    if not B._is_api_sheet(_cols(editor, key)):
        raise BuildError(f"{key} is not an API test sheet (it needs blnExecute and WEBSERVICE_URL or ENVIRONMENT_PARAMETER).", "kind")
    return key


def _data_row(editor: WorkbookEditor, sheet: str, row: Any) -> int:
    rows = B._data_rows(_Grid(editor, sheet))
    if row not in (None, ""):
        r = B._int(row, "row")
        if r < 2:
            raise BuildError("Row 1 holds the headers; a data row starts at row 2.")
        return r
    first = next((d["row"] for d in rows if d["enabled"] is not False), None) or next((d["row"] for d in rows), None)
    return first or 2


def _is_secret_header(name: str, value: str = "") -> bool:
    return bool(SECRET_HEADER.search(name or "")) or "{SECRET:" in (value or "").upper()


def _shown(name: str, value: str) -> str:
    """A header value as the page shows it: a secret typed straight into the sheet (the legacy way) is not sent to the page."""
    if not value or "{" in value or not _is_secret_header(name):
        return value
    return MASK


def api_view(editor: WorkbookEditor, sheet_name: str, row: Any = None, *, environments: dict | None = None,
             environment: str = "") -> dict:
    """The API editor's view of one data row: request (method, URL, headers, body), the steps (send, checks, saves), and what the variable picker
    offers.  ``environments``: ``builder.read_environments`` (read once by the caller); ``environment``: the one "Send now" uses."""
    sheet = api_sheet(editor, sheet_name)
    grid = _Grid(editor, sheet)
    row = _data_row(editor, sheet, row)
    cols = grid.cols

    def raw(header: str) -> str:
        return grid.raw(row, header.upper())

    json_format = is_true(grid.get(row, "JSON_FORMAT")) if "JSON_FORMAT" in cols else False
    io = read_io(editor)
    has = lambda column: column.strip().upper() in cols
    outputs_by_col: dict[str, IoRow] = {}
    for r in io:
        if r.kind in ("output_json", "output") and has(r.value):
            outputs_by_col.setdefault(r.value.upper(), r)

    headers = [{"name": r.parameter, "column": r.value, "value": _shown(r.parameter, raw(r.value)), "secret": _is_secret_header(r.parameter, raw(r.value)),
                "ioRow": r.row} for r in io if r.kind == "add_header" and has(r.value)]
    template = {"location": raw("XML_LOCATION") or raw("TEMPLATES_PATH"), "file": raw("XML_REQUESTFILE") or raw("REQUESTFILE") or raw("TEMPLATE_FILE"),
                "replace": [{"placeholder": r.parameter, "column": r.value, "value": raw(r.value), "ioRow": r.row}
                            for r in io if r.kind == "replace" and has(r.value)],
                "updates": [{"path": r.parameter, "column": r.value, "value": raw(r.value), "ioRow": r.row}
                            for r in io if r.kind == "update_json" and has(r.value)]}
    typed = raw("REQUEST_BODY") if "REQUEST_BODY" in cols else ""
    body_kind = "typed" if typed else "template" if template["file"] else "none"

    checks, used_outputs = [], set()
    skip = _disabled_columns(io, raw(_global_text(editor, "TC_Name_Column") or "TC_Name"))
    for r in io:
        if r.kind not in CHECK_KINDS or not has(r.parameter):
            continue
        actual = r.value.strip().upper()
        source = outputs_by_col.get(actual)
        if source is not None:
            used_outputs.add(source.row)
        what = "status" if actual == "RES_STATUS_CD_OUT" else "time" if actual == "ELAPSEDTIME_OUT" else ""
        checks.append({"ioRow": r.row, "kind": r.kind, "label": CHECK_LABELS.get(r.kind, r.kind), "expected": r.parameter,
                       "expectedValue": raw(r.parameter), "actual": r.value, "path": source.parameter if source else "", "what": what,
                       "outputRow": source.row if source else None, "function": source.kind if source else "", "disabled": actual in skip,
                       "lastValue": raw(r.value) if has(r.value) else ""})
    saves = [{"ioRow": r.row, "function": r.kind, "path": r.parameter, "column": r.value, "lastValue": raw(r.value)}
             for r in outputs_by_col.values() if r.row not in used_outputs]
    steps = [{"n": 1, "kind": "send", "label": f"{(raw('WEBSERVICE_METHOD') or 'POST').upper()} request"}]
    steps += [{"n": len(steps) + i + 1, "kind": "check", "ioRow": c["ioRow"],
               "label": f"{c['what'] or c['path'] or c['actual']} {c['label']} {c['expectedValue'] or '(empty)'}"} for i, c in enumerate(checks)]
    steps += [{"n": len(steps) + i + 1, "kind": "save", "ioRow": s["ioRow"], "label": f"{s['path']} as {{{s['column']}}}"} for i, s in enumerate(saves)]

    env = environments or B.read_environments(editor)
    env_vars = {r["variable"].upper() for r in env.get("rows", [])}
    texts = [raw("WEBSERVICE_URL"), typed, *(h["value"] for h in headers), *(c["path"] for c in checks), *(s["path"] for s in saves),
             *(c["expectedValue"] for c in checks)]
    names = sorted({m.upper() for t in texts for m in B.INLINE_RE.findall(t or "")})
    row_values = {name: grid.raw(row, name.upper()) for name in grid.header_names if name}
    needs = [n for n in names if not (n in cols and grid.raw(row, n)) and n not in env_vars]
    provides = sorted({s["column"] for s in saves} | {c["actual"] for c in checks if c["path"]})
    return {
        "test": sheet, "sheet": sheet, "row": row, "dataRows": B._data_rows(grid), "format": "json" if json_format else "xml",
        "method": (raw("WEBSERVICE_METHOD") or "POST").upper(), "methods": list(METHODS),
        "url": raw("WEBSERVICE_URL"), "urlFormula": is_formula(grid.get(row, "WEBSERVICE_URL")),
        "environmentParameter": raw("ENVIRONMENT_PARAMETER"), "hasUrlColumn": "WEBSERVICE_URL" in cols,
        "headers": headers, "body": {"kind": body_kind, "text": typed, "template": template},
        "checks": checks, "saves": saves, "steps": steps, "checkKinds": [{"kind": k, "label": CHECK_LABELS[k]} for k in CHECK_KINDS],
        "columns": [h for h in grid.header_names if h],
        "rowValues": {k: _shown(k, v) for k, v in row_values.items()},
        "needs": needs, "provides": provides, "environment": environment,
        "environmentVariables": sorted(r["variable"] for r in env.get("rows", [])),
        "production": [n.upper() for n in env.get("production", [])],
    }


def _global_text(editor: WorkbookEditor, name: str) -> str:
    g = _sheet_key(editor, "Global")
    if g is None:
        return ""
    for values in editor.rows(g)[1:]:
        if values and text_of(values[0]).strip().upper() == name.upper():
            return text_of(values[1]).strip() if len(values) > 1 else ""
    return ""


def _disabled_columns(io: list[IoRow], test_name: str) -> set[str]:
    out: set[str] = set()
    for r in io:
        if r.kind == "disable" and test_name and r.parameter.strip().upper() == test_name.strip().upper():
            out.update(c.strip().upper() for c in r.value.split(";") if c.strip())
    return out


# ---------------------------------------------------------------------------------------------------------------------------------------------
# Edits (ops, applied through BuildDocument.apply; registered into builder.EXTRA_OPS below)
# ---------------------------------------------------------------------------------------------------------------------------------------------
def _ensure_io(e: WorkbookEditor) -> str:
    name = io_sheet(e)
    if name is None:
        e.add_sheet(IO_SHEET)
        for c, h in enumerate(IO_HEADERS, start=1):
            e.set(IO_SHEET, 1, c, h)
        name = IO_SHEET
    return name


def _append_io(e: WorkbookEditor, function: str, parameter: str, value: str) -> int:
    name = _ensure_io(e)
    rows = e.rows(name)
    row = len(rows) + 1
    for c, v in enumerate((function, parameter, value), start=1):
        e.set(name, row, c, v)
    return row


def _column(e: WorkbookEditor, sheet: str, header: str) -> int:
    col = e.column(sheet, header)
    return col if col is not None else e.add_column(sheet, header)


def _put(e: WorkbookEditor, sheet: str, row: int, header: str, value: Any) -> None:
    if value in (None, "") and e.column(sheet, header) is None:
        return
    e.set(sheet, row, _column(e, sheet, header), None if value == "" else value)


def _column_name(text: str) -> str:
    """A column name from a header / path / label: letters, digits and ``_`` (``x-api-key`` -> ``x_api_key``)."""
    t = re.sub(r"[^A-Za-z0-9_]+", "_", str(text)).strip("_")
    if not t:
        raise BuildError("A column name needs letters or digits.")
    return t if not t[0].isdigit() else "C_" + t


def _base_of_path(path: str) -> str:
    """The last name in a JSONPath / XPath (``$.premium.total`` -> ``total``, ``/a/b[2]/@id`` -> ``id``)."""
    names = re.findall(r"[A-Za-z_][A-Za-z0-9_-]*", re.sub(r"\[[^\]]*\]", "", path or "").replace("text()", ""))
    return names[-1] if names else "value"


def _other_api_sheets_with(e: WorkbookEditor, column: str, but: str) -> list[str]:
    out = []
    for name in e.sheet_names():
        if name.upper() == but.upper() or name.upper().startswith(B.RR_PREFIX.upper()):
            continue
        cols = _cols(e, name)
        if B._is_api_sheet(cols) and column.upper() in cols:
            out.append(name)
    return out


def _unique_column(e: WorkbookEditor, sheet: str, wanted: str, taken: set[str]) -> str:
    base, n = wanted, 2
    cols = _cols(e, sheet)
    while wanted.upper() in cols or wanted.upper() in taken:
        wanted = f"{base}_{n}"
        n += 1
    return wanted


def _row_of(e: WorkbookEditor, sheet: str, op: dict) -> int:
    return _data_row(e, sheet, op.get("row"))


def op_set_request(e: WorkbookEditor, op: dict) -> dict:
    """``{test, row, method?, url?, format?: json|xml, body?}``: the request cells of the data row (a missing column is added)."""
    sheet = api_sheet(e, op.get("test"))
    row = _row_of(e, sheet, op)
    if "method" in op:
        method = str(op["method"] or "").strip().upper()
        if method not in METHODS:
            raise BuildError(f"{method or 'An empty method'} is not one of {', '.join(METHODS)}.")
        _put(e, sheet, row, "WEBSERVICE_METHOD", method)
    if "url" in op:
        url = str(op["url"] or "").strip()
        if url.startswith("="):
            raise BuildError("Type the address itself (a formula is edited in the Excel grid).")
        _put(e, sheet, row, "WEBSERVICE_URL", url)
    if "format" in op:
        _put(e, sheet, row, "JSON_FORMAT", "Y" if str(op["format"]).lower() == "json" else "N")
    if "body" in op:
        body = op["body"]
        _put(e, sheet, row, "REQUEST_BODY", str(body) if body not in (None, "") else "")
    return {"op": "api_set_request", "test": sheet, "row": row}


def _header_row(io: list[IoRow], name: str) -> IoRow | None:
    return next((r for r in io if r.kind == "add_header" and r.parameter.strip().lower() == name.strip().lower()), None)


def op_set_header(e: WorkbookEditor, op: dict) -> dict:
    """``{test, row, name, value, previous?}``: send header ``name`` with ``value`` (``{SECRET:X}`` for a key).  A header another sheet already
    sends keeps its ``InputOutput`` row and column; a new one gets ``addHeader | name | H_<name>``."""
    sheet = api_sheet(e, op.get("test"))
    row = _row_of(e, sheet, op)
    name = str(op.get("name") or "").strip()
    if not name or re.search(r"[\s:]", name):
        raise BuildError("A header name has no spaces or colons (Content-Type, x-api-key).")
    previous = str(op.get("previous") or "").strip()
    if previous and previous.lower() != name.lower():
        op_remove_header(e, {"test": sheet, "row": row, "name": previous})
    value = op.get("value")
    value = "" if value is None else str(value)
    if value == MASK:
        raise BuildError("Type the new value (or {SECRET:NAME} to take it from secrets.env).")
    io = read_io(e)
    found = _header_row(io, name)
    if found is None:
        taken = {r.value.upper() for r in io}
        column = _unique_column(e, sheet, "H_" + _column_name(name), taken)
        _append_io(e, "addHeader", name, column)
    else:
        column = found.value
    e.set(sheet, row, _column(e, sheet, column), value if value != "" else "None")      # (the runner sends "None" as an empty value)
    return {"op": "api_set_header", "test": sheet, "row": row, "column": column}


def op_remove_header(e: WorkbookEditor, op: dict) -> dict:
    """``{test, row, name}``: stop sending it.  The ``InputOutput`` row goes when no other API sheet has its column; otherwise this row's cell
    says ``[BLANK]`` (the runner's "leave this header out")."""
    sheet = api_sheet(e, op.get("test"))
    row = _row_of(e, sheet, op)
    found = _header_row(read_io(e), str(op.get("name") or ""))
    if found is None or e.column(sheet, found.value) is None:
        return {"op": "api_remove_header", "test": sheet, "row": row, "removed": False}
    if _other_api_sheets_with(e, found.value, sheet):
        e.set(sheet, row, e.column(sheet, found.value), "[BLANK]")
    else:
        e.set(sheet, row, e.column(sheet, found.value), None)
        e.delete_rows(io_sheet(e), found.row, 1)
    return {"op": "api_remove_header", "test": sheet, "row": row, "removed": True}


def _function_of(op: dict, e: WorkbookEditor, sheet: str, row: int) -> str:
    """``output_json`` for a JSON response, ``Output`` for XML: given, else from the row's format."""
    f = str(op.get("function") or "").strip().lower()
    if f in ("output_json", "json"):
        return "output_json"
    if f in ("output", "xml", "xpath"):
        return "Output"
    fmt = e.get(sheet, row, e.column(sheet, "JSON_FORMAT")) if e.column(sheet, "JSON_FORMAT") else None
    return "output_json" if is_true(fmt) else "Output"


def _output_column(e: WorkbookEditor, sheet: str, function: str, path: str, wanted: str) -> tuple[str, bool]:
    """The column ``path`` is read into: the one an existing row already reads it into (reused: the runner keys rows by path), else ``wanted``
    made unique.  ``(column, reused)``; the output row exists afterwards and the column is in the sheet."""
    kind = _FUNCTION_KIND[function.lower()]
    io = read_io(e)
    found = next((r for r in io if r.kind == kind and r.parameter == path), None)
    if found is not None:
        _column(e, sheet, found.value)
        return found.value, True
    taken = {r.value.upper() for r in io if r.kind in ("output_json", "output")} | {r.parameter.upper() for r in io if r.kind in CHECK_KINDS}
    column = wanted if wanted.upper() not in taken and e.column(sheet, wanted) is None else _unique_column(e, sheet, wanted, taken)
    _append_io(e, function, path, column)
    _column(e, sheet, column)
    return column, False


def op_add_check(e: WorkbookEditor, op: dict) -> dict:
    """``{test, row, path, function?, kind, expected, name?}``: check a value of the response.  ``path`` ``"status"`` / ``"time"`` checks the HTTP
    status / the time taken (``RES_STATUS_CD_OUT`` / ``ELAPSEDTIME_OUT``).  Adds the output row (or reuses it), the expected column
    (``<name>_EXP``) holding ``expected`` in this row, and the check row."""
    sheet = api_sheet(e, op.get("test"))
    row = _row_of(e, sheet, op)
    kind = str(op.get("kind") or "compare").strip().lower()
    if kind not in CHECK_KINDS:
        raise BuildError(f"{kind!r} is not a check ({', '.join(CHECK_KINDS)}).")
    path = str(op.get("path") or "").strip()
    if not path:
        raise BuildError("Which value? Click one in the response, or type its path.")
    io = read_io(e)
    reused = False
    if path.lower() in STATUS_COLUMNS:
        exp_wanted, actual = STATUS_COLUMNS[path.lower()]
    else:
        base = _column_name(op.get("name") or _base_of_path(path))
        actual, reused = _output_column(e, sheet, _function_of(op, e, sheet, row), path, base + "_OUT")
        exp_wanted = re.sub(r"_OUT$", "", actual, flags=re.I) + "_EXP"
    taken = {r.parameter.upper() for r in read_io(e) if r.kind in CHECK_KINDS} | {r.value.upper() for r in io}
    expected_col = exp_wanted if exp_wanted.upper() not in taken and e.column(sheet, exp_wanted) is None else _unique_column(e, sheet, exp_wanted, taken)
    _put(e, sheet, row, expected_col, str(op.get("expected") if op.get("expected") is not None else ""))
    _column(e, sheet, expected_col)
    io_row = _append_io(e, kind, expected_col, actual)
    return {"op": "api_add_check", "test": sheet, "row": row, "ioRow": io_row, "expected": expected_col, "actual": actual, "reused": reused}


def op_add_save(e: WorkbookEditor, op: dict) -> dict:
    """``{test, row, path, function?, variable}``: keep a value of the response under a name the rest of the run can use as ``{NAME}``.  When the
    path is already read into a column, that column is the name (said in ``reused``)."""
    sheet = api_sheet(e, op.get("test"))
    row = _row_of(e, sheet, op)
    path = str(op.get("path") or "").strip()
    if not path:
        raise BuildError("Which value? Click one in the response, or type its path.")
    variable = B.make_token(str(op.get("variable") or _base_of_path(path)))
    column, reused = _output_column(e, sheet, _function_of(op, e, sheet, row), path, variable)
    return {"op": "api_add_save", "test": sheet, "row": row, "column": column, "reused": reused and column.upper() != variable.upper()}


def op_update_io(e: WorkbookEditor, op: dict) -> dict:
    """``{ioRow, function?, parameter?, value?}``: edit one ``InputOutput`` row (the path written for the person stays editable)."""
    name = io_sheet(e)
    if name is None:
        raise BuildError("This workbook has no InputOutput sheet.", "not_found", 404)
    r = B._int(op.get("ioRow"), "ioRow")
    if r < 2 or r > e.max_row(name):
        raise BuildError(f"InputOutput has no row {r}.", "not_found", 404)
    for c, key in enumerate(("function", "parameter", "value"), start=1):
        if key in op:
            text = str(op[key] or "").strip()
            if not text and key != "value":
                raise BuildError(f"The {key} of an InputOutput row cannot be empty.")
            e.set(name, r, c, text or None)
    return {"op": "api_update_io", "ioRow": r}


def op_delete_io(e: WorkbookEditor, op: dict) -> dict:
    """``{ioRows: [..]}``: delete ``InputOutput`` rows (a check's output row goes with it only when asked: it may feed other checks)."""
    name = io_sheet(e)
    if name is None:
        raise BuildError("This workbook has no InputOutput sheet.", "not_found", 404)
    rows = sorted({B._int(r, "ioRow") for r in (op.get("ioRows") or [])}, reverse=True)
    for r in rows:
        if r < 2 or r > e.max_row(name):
            raise BuildError(f"InputOutput has no row {r}.", "not_found", 404)
        e.delete_rows(name, r, 1)
    return {"op": "api_delete_io", "ioRows": rows}


def op_set_value(e: WorkbookEditor, op: dict) -> dict:
    """``{test, row, column, value}``: one cell of the data row (an expected value, a column the body names); the column is added if missing."""
    sheet = api_sheet(e, op.get("test"))
    row = _row_of(e, sheet, op)
    column = str(op.get("column") or "").strip()
    if not column:
        raise BuildError("Which column?")
    value = op.get("value")
    if value == MASK:
        raise BuildError("Type the new value.")
    e.set(sheet, row, _column(e, sheet, _column_name(column) if e.column(sheet, column) is None else column),
          None if value in (None, "") else value)
    return {"op": "api_set_value", "test": sheet, "row": row, "column": column}


def op_add_test(e: WorkbookEditor, op: dict) -> dict:
    """``{name, format?: json|xml, method?, url?, headers?: [[name, value]], body?}``: a new API test (sheet, first data row, DataSheets row)."""
    name = str(op.get("name") or "").strip()
    if not name:
        raise BuildError("Give the test a name.")
    if _sheet_key(e, name) is not None:
        raise BuildError(f"A sheet called {name!r} already exists.", "exists", 409)
    try:
        e.add_sheet(name)
    except WriterError as err:
        raise BuildError(str(err)) from err
    for c, h in enumerate(API_COLUMNS, start=1):
        e.set(name, 1, c, h)
    for h, v in (("blnExecute", "Y"), ("TCID", "1"), ("TC_Name", name)):
        e.set(name, 2, e.column(name, h), v)
    fmt = str(op.get("format") or "json").lower()
    op_set_request(e, {"test": name, "row": 2, "method": op.get("method") or ("GET" if not op.get("body") else "POST"),
                       "url": op.get("url") or "", "format": fmt, "body": op.get("body") or ""})
    for header in op.get("headers") or []:
        hname, hvalue = (header.get("name"), header.get("value")) if isinstance(header, dict) else header
        if str(hname).lower() == "content-type" and fmt == "json" and "json" in str(hvalue).lower():
            continue                                                                 # (JSON_FORMAT = Y sends it)
        op_set_header(e, {"test": name, "row": 2, "name": hname, "value": hvalue})
    _Ops(e).set_test({"test": name, "enabled": True, **({"comment": op["comment"]} if op.get("comment") else {})})
    return {"op": "api_add_test", "test": name}


def op_use_template(e: WorkbookEditor, op: dict) -> dict:
    """``{test, row, location, file, format?, mappings: [{placeholder, column?, value?}]}``: build the body from a request template file.  Each
    placeholder is filled from ``column`` (a column of the sheet: a variable) or ``value`` (typed).  A placeholder another sheet already fills
    keeps its ``Replace`` row; this row then gets that row's column, set to ``=<column>`` of this row or the typed value."""
    sheet = api_sheet(e, op.get("test"))
    row = _row_of(e, sheet, op)
    file = str(op.get("file") or "").strip()
    if not file:
        raise BuildError("Which template file?")
    _put(e, sheet, row, "XML_LOCATION", str(op.get("location") or "").strip())
    _put(e, sheet, row, "XML_REQUESTFILE", file)
    _put(e, sheet, row, "REQUEST_BODY", "")
    if op.get("format"):
        _put(e, sheet, row, "JSON_FORMAT", "Y" if str(op["format"]).lower() == "json" else "N")
    from openpyxl.utils import get_column_letter
    filled = []
    for m in op.get("mappings") or []:
        placeholder = str(m.get("placeholder") or "")
        if not placeholder:
            continue
        source = str(m.get("column") or "").strip()
        value = m.get("value")
        if not source and value in (None, ""):
            continue                                                                  # left for later: stays as written in the template
        io = read_io(e)
        found = next((r for r in io if r.kind == "replace" and r.parameter == placeholder), None)
        if found is None:
            column = source or _unique_column(e, sheet, _column_name(placeholder.strip("#{}$ ")) + "_IN", {r.value.upper() for r in io})
            _append_io(e, "Replace", placeholder, column)
            if not source or e.column(sheet, column) is None:
                _put(e, sheet, row, column, "" if source else value)
            _column(e, sheet, column)
        else:
            column = found.value
            if source and source.upper() != column.upper():
                src_col = _column(e, sheet, source)
                e.set(sheet, row, _column(e, sheet, column), f"={get_column_letter(src_col)}{row}")
            elif not source:
                e.set(sheet, row, _column(e, sheet, column), value)
        filled.append({"placeholder": placeholder, "column": column})
    return {"op": "api_use_template", "test": sheet, "row": row, "filled": filled}


OPS = {"api_set_request": op_set_request, "api_set_header": op_set_header, "api_remove_header": op_remove_header, "api_add_check": op_add_check,
       "api_add_save": op_add_save, "api_update_io": op_update_io, "api_delete_io": op_delete_io, "api_set_value": op_set_value,
       "api_add_test": op_add_test, "api_use_template": op_use_template}
B.EXTRA_OPS.update(OPS)


# ---------------------------------------------------------------------------------------------------------------------------------------------
# The response as a clickable tree (Q19)
# ---------------------------------------------------------------------------------------------------------------------------------------------
def _display(value: Any) -> str:
    if isinstance(value, str):
        text = json.dumps(value, ensure_ascii=False)
    elif value is None:
        text = "null"
    elif isinstance(value, bool):
        text = "true" if value else "false"
    else:
        text = str(value)
    return text if len(text) <= 200 else text[:197] + "…\""


def _type(value: Any) -> str:
    return ("object" if isinstance(value, dict) else "array" if isinstance(value, list) else "string" if isinstance(value, str)
            else "bool" if isinstance(value, bool) else "number" if isinstance(value, (int, float)) else "null")


def _leaves(item: Any, prefix: tuple = (), depth: int = 0) -> list[tuple[tuple, Any]]:
    """Scalar fields of an array item, up to three levels down: what "the item where ..." can name."""
    out: list[tuple[tuple, Any]] = []
    if isinstance(item, dict) and depth < 3:
        for k, v in item.items():
            if isinstance(v, (dict, list)):
                out += _leaves(v, prefix + (k,), depth + 1) if isinstance(v, dict) else []
            elif re.match(r"^[A-Za-z_][A-Za-z0-9_-]*$", str(k)):
                out.append((prefix + (k,), v))
    return out


def json_tree(data: Any, row_values: dict[str, str] | None = None) -> tuple[list[dict], bool]:
    """``(nodes, cut)``: every node of a JSON response, depth first.  Each has ``path`` (JSONPath).  A node inside a list also has ``array``: the
    list's path, the item's index, the rest of the path after the item, and ``where``: fields of the item that tell it from the others (unique
    values first), each with the path "the item where field = value" and, when a column of the data row holds that value, the same path with
    ``{COLUMN}`` - what the pop-up offers for "this item" / "the item where ..."."""
    nodes: list[dict] = []
    values = {v.strip(): k for k, v in (row_values or {}).items() if v and v.strip() and len(v.strip()) > 1 and MASK not in v}
    cut = False

    def walk(value: Any, key: Any, steps: list, depth: int, arrays: list) -> None:
        nonlocal cut
        if len(nodes) >= TREE_LIMIT:
            cut = True
            return
        node = {"id": len(nodes), "depth": depth, "key": key, "type": _type(value), "path": json_path_of(steps)}
        if isinstance(value, (dict, list)):
            node["count"] = len(value)
        else:
            node["value"] = _display(value)
            node["text"] = "" if value is None else ("true" if value is True else "false" if value is False else str(value))
        if arrays and not isinstance(value, (dict, list)):
            node["array"] = arrays[-1]["describe"](steps)
        nodes.append(node)
        if isinstance(value, dict):
            for k, v in value.items():
                walk(v, k, steps + [k], depth + 1, arrays)
        elif isinstance(value, list):
            items = value
            for i, v in enumerate(items):
                walk(v, i, steps + [i], depth + 1, arrays + [_array_frame(items, steps, i, values)])

    walk(data, "$", [], 0, [])
    return nodes, cut


def _array_frame(items: list, steps: list, index: int, values: dict[str, str]) -> dict:
    """What a node inside item ``index`` of the list at ``steps`` needs to offer "this item" and "the item where"."""
    at = len(steps)

    def describe(node_steps: list) -> dict:
        rest = node_steps[at + 1:]
        array_path = json_path_of(steps)
        rest_path = json_path_of(rest)[1:]
        item = items[index]
        where = []
        for field_steps, value in _leaves(item):
            if tuple(field_steps) == tuple(rest):
                continue                                               # (the value clicked on does not pick its own item)
            same = sum(1 for other in items if _at(other, field_steps) == value)
            field = ".".join(str(s) for s in field_steps)
            literal = filter_literal(value)
            option = {"field": field, "value": "" if value is None else str(value), "unique": same == 1,
                      "path": f"{array_path}[?(@.{field}=={literal})]{rest_path}"}
            column = values.get(str(value).strip()) if value is not None else None
            if column:
                option["variable"] = column
                option["variablePath"] = f"{array_path}[?(@.{field}=='{{{column}}}')]{rest_path}"
            where.append(option)
        where.sort(key=lambda o: (not o.get("variable"), not o["unique"]))
        return {"path": array_path, "index": index, "count": len(items), "rest": rest_path, "thisItem": json_path_of(node_steps),
                "where": where[:8]}

    return {"describe": describe}


def _at(item: Any, steps: tuple) -> Any:
    node = item
    for s in steps:
        if not isinstance(node, dict) or s not in node:
            return object()
        node = node[s]
    return node


def xml_tree(text: str, row_values: dict[str, str] | None = None) -> tuple[list[dict], bool]:
    """``(nodes, cut)`` of an XML response (namespaces dropped, like the runner reads it): elements and their attributes, each with its XPath
    (``/Envelope/Body/Plan[2]/Code``, ``.../@id``).  An element repeated among its siblings gives the nodes inside it ``array``: "this item"
    (the positional path) and "the item where" a child's text or an attribute has a value (``Plan[Code='Max']``, ``Plan[@id='B']``)."""
    root = xml_to_tree(text)
    if root is None or not len(root):
        return [], False
    values = {v.strip(): k for k, v in (row_values or {}).items() if v and v.strip() and len(v.strip()) > 1 and MASK not in v}
    nodes: list[dict] = []
    cut = False

    def walk(el: ET.Element, steps: list, depth: int, frame: dict | None) -> None:
        nonlocal cut
        if len(nodes) >= TREE_LIMIT:
            cut = True
            return
        children = list(el)
        text_value = (el.text or "").strip()
        node = {"id": len(nodes), "depth": depth, "key": el.tag, "type": "element" if children else "text", "path": xpath_of(steps)}
        if children:
            node["count"] = len(children)
        else:
            node["value"], node["text"] = text_value, text_value
        if frame and not children:
            node["array"] = frame["describe"](steps, "")
        nodes.append(node)
        for name, value in el.attrib.items():
            attr = {"id": len(nodes), "depth": depth + 1, "key": "@" + name, "type": "attribute", "path": xpath_of(steps, name),
                    "value": value, "text": value}
            if frame:
                attr["array"] = frame["describe"](steps, name)
            nodes.append(attr)
        tags = [c.tag for c in children]
        for c in children:
            pos = sum(1 for t in tags[:children.index(c)] if t == c.tag) + 1
            count = tags.count(c.tag)
            sub = frame
            if count > 1:
                sub = _xml_frame(children, c, steps + [(c.tag, pos, count)], values)
            walk(c, steps + [(c.tag, pos, count)], depth + 1, sub)

    walk(root[0], [(root[0].tag, 1, 1)], 0, None)
    return nodes, cut


def _xml_frame(siblings: list, item: ET.Element, item_steps: list, values: dict[str, str]) -> dict:
    same = [s for s in siblings if s.tag == item.tag]
    at = len(item_steps)

    def describe(node_steps: list, attr: str) -> dict:
        base = xpath_of(item_steps[:-1]) + "/" + item.tag
        rest = xpath_of(node_steps[at:], attr)
        where = []
        fields = [("@" + k, v) for k, v in item.attrib.items()] + [(c.tag, (c.text or "").strip()) for c in item if not len(c) and (c.text or "").strip()]
        clicked = ("@" + attr) if attr and len(node_steps) == at else node_steps[at][0] if len(node_steps) == at + 1 and not attr else None
        for field, value in fields:
            if field == clicked:
                continue                                               # (the value clicked on does not pick its own item)
            count = sum(1 for s in same if (s.attrib.get(field[1:]) if field.startswith("@") else (s.findtext(field) or "").strip()) == value)
            option = {"field": field, "value": value, "unique": count == 1, "path": f"{base}[{field}='{value}']{rest}"}
            column = values.get(value)
            if column:
                option["variable"] = column
                option["variablePath"] = f"{base}[{field}='{{{column}}}']{rest}"
            where.append(option)
        where.sort(key=lambda o: (not o.get("variable"), not o["unique"]))
        return {"path": base, "index": item_steps[-1][1] - 1, "count": len(same), "rest": rest, "thisItem": xpath_of(node_steps, attr), "where": where[:8]}

    return {"describe": describe}


# ---------------------------------------------------------------------------------------------------------------------------------------------
# Send now
# ---------------------------------------------------------------------------------------------------------------------------------------------
def _draft_workbook(doc: BuildDocument, folder: Path, environment: str):
    from ..workbook import Workbook
    folder.mkdir(parents=True, exist_ok=True)
    with doc.lock:
        data = doc.editor.to_bytes()
    target = folder / ("api-" + re.sub(r"[^A-Za-z0-9._-]+", "_", doc.path.name))
    tmp = target.with_suffix(".tmp")
    tmp.write_bytes(data)
    tmp.replace(target)
    return Workbook(target, environment=environment or None, secrets=workbook_secrets())


def _case(wb, sheet: str, row: int):
    from ..workbook.model import TestCase
    cases = [c for c in wb.discover() if c.sheet.upper() == sheet.upper() and c.kind == "api"]
    found = next((c for c in cases if c.param_row == row), None)
    if found is not None:
        return found
    if cases:
        return dataclasses.replace(cases[0], param_row=row)
    return TestCase(id=sheet, sheet=sheet, param_sheet=None, param_row=row, title=sheet, is_ui=False, kind="api")


async def send_now(doc: BuildDocument, cfg: Config, folder: Path, *, test: str, row: int | None, environment: str,
                   values: dict[str, str] | None = None) -> dict:
    """Send the request of one data row as the draft has it, with the environment's values; nothing is written to the workbook.  ``values``:
    what the person typed for ``{NAME}``s only an earlier test of a run would give.  Returns the request as sent (secrets masked), the response
    (status, time, headers, text, tree) and what each check of the test would say."""
    from playwright.async_api import async_playwright

    from ..engine.api_runner import fetch, is_an_error_page, prepare_body, read_response
    from ..workbook.variables import VariablePool
    with doc.lock:
        sheet = api_sheet(doc.editor, test)
        row = _data_row(doc.editor, sheet, row)
    wb = _draft_workbook(doc, folder, environment)
    pool = VariablePool()
    for k, v in (values or {}).items():
        if str(v).strip():
            pool.set(str(k).upper(), str(v))
    runtime = wb.api_runtime(_case(wb, sheet, row), pool=pool)
    request = runtime.request()
    out: dict[str, Any] = {"test": sheet, "row": row, "environment": runtime.environment, "missing": [], "notes": []}
    problem = runtime.error_in("WEBSERVICE_URL") or runtime.error_in("WEBSERVICE_METHOD")
    if problem:
        raise BuildError(f"Cannot build the request: {problem}.", "request")
    missing = sorted(runtime.pool_reads() - {k.upper() for k, v in (values or {}).items() if str(v).strip()})
    if not request.url:
        if missing:
            return {**out, "missing": missing, "error": "Give a value for " + ", ".join("{%s}" % m for m in missing) + " to send it."}
        raise BuildError(f"There is no URL: {request.url_problem}." if request.url_problem else "Type the address first.", "request")
    if urlparse(request.url).scheme not in ("http", "https"):
        raise BuildError(f"{runtime.mask(request.url)!r} is not an http(s) address.", "request")
    body, notes, problem = prepare_body(runtime, request, cfg, wb.path.parent)
    if problem:
        body_missing = re.findall(r"\{([A-Za-z_][A-Za-z0-9_]*)\}", problem) if problem.startswith("Nothing gives") else []
        if body_missing:
            return {**out, "missing": sorted(set(missing) | {m.upper() for m in body_missing}), "error": problem}
        raise BuildError(problem, "request")
    shown_headers = [{"name": k, "value": MASK if k in request.secret_headers and _is_secret_header(k, v) else runtime.mask(v)}
                     for k, v in request.headers.items()]
    out["request"] = {"method": request.method, "url": runtime.mask(request.url), "headers": shown_headers,
                      "body": runtime.mask(body.decode("utf-8", "replace")) if body else ""}
    secrets = {v for k, v in request.headers.items() if k in request.secret_headers and _is_secret_header(k, v) and len(v) >= 3}
    runtime.secret_values.update(secrets)
    pw = await async_playwright().start()
    try:
        response = await fetch(pw, cfg, request.method, request.url, request.headers, body)
    except Exception as err:
        from ..engine.api_runner import tls_hint
        reason = runtime.mask((str(err).strip().splitlines() or [type(err).__name__])[0][:300])
        return {**out, "error": f"The request could not be completed: {reason}.{tls_hint(reason)}"}
    finally:
        await pw.stop()
    notes += read_response(runtime, response)
    text = runtime.mask(response.text)
    row_values = {k: text_of(v) for k, v in runtime.values.items() if not isinstance(v, (dict, list))}
    if response.is_json:
        tree, cut = json_tree(json.loads(runtime.mask(json.dumps(response.data, ensure_ascii=False))), row_values)
        fmt = "json"
    else:
        tree, cut = xml_tree(text, row_values)
        fmt = "xml" if tree else "text"
    results = []
    for point in runtime.checkpoints():
        expected_raw = runtime.values.get(point.expected_column.upper())
        if isinstance(expected_raw, str) and "{" in expected_raw:
            expected_raw = runtime.fill(expected_raw).text
        actual = runtime.actual(point.actual_column)
        results.append({"expected": point.expected_column, "actual": point.actual_column, "kind": point.kind,
                        "expectedValue": runtime.mask(text_of(expected_raw)), "actualValue": runtime.mask(actual),
                        "passed": check(point.kind, expected_raw, actual)})
    content_type = next((v for k, v in response.headers.items() if k.lower() == "content-type"), "")
    out.update({
        "response": {"status": response.status, "ms": response.elapsed_ms, "contentType": content_type, "format": fmt,
                     "headers": [{"name": k, "value": runtime.mask(v)} for k, v in response.headers.items() if k.lower() != "set-cookie"],
                     "text": text[:TEXT_LIMIT], "textCut": len(text) > TEXT_LIMIT, "tree": tree, "treeCut": cut,
                     "blocked": response.status in cfg.runner.block_statuses and is_an_error_page(response.status, response.headers, response.text)},
        "outputs": {k: runtime.mask(v) for k, v in runtime.outputs.items()}, "checks": results, "notes": [runtime.mask(n) for n in notes],
        "missing": missing})
    return out


# ---------------------------------------------------------------------------------------------------------------------------------------------
# The ways in: cURL, Postman, request templates
# ---------------------------------------------------------------------------------------------------------------------------------------------
def parse_curl(text: str) -> dict:
    """A ``curl`` command (copied from a browser's "Copy as cURL" - bash or Windows cmd - or typed) as ``{method, url, headers: [{name, value}],
    body, format}``.  ``ValueError`` when it is not one."""
    t = (text or "").strip()
    t = re.sub(r"\^\r?\n", " ", t)                              # cmd continuation
    t = re.sub(r"\\\r?\n", " ", t)                              # bash continuation
    if re.search(r"\^\"", t):                                  # cmd quoting: ^" ... ^" and ^^ / ^&
        t = t.replace('^"', '"').replace("^^", "^").replace("^&", "&").replace("^%", "%").replace("^\\", "\\")
    try:
        words = shlex.split(t, posix=True)
    except ValueError as err:
        raise ValueError(f"the command could not be read ({err})") from None
    if not words or Path(words[0]).name.lower() not in ("curl", "curl.exe"):
        raise ValueError("it does not start with curl")
    method, url, headers, data, is_json, get = "", "", [], [], False, False
    takes_value = {"-o", "--output", "-e", "--referer", "--connect-timeout", "-m", "--max-time", "--proxy", "-x", "--cacert", "--cert", "--key",
                   "-w", "--write-out", "-A", "--user-agent", "-b", "--cookie", "-F", "--form", "--retry", "-c", "--cookie-jar", "--resolve"}
    i = 1
    while i < len(words):
        w = words[i]
        nxt = words[i + 1] if i + 1 < len(words) else ""
        if w in ("-X", "--request"):
            method, i = nxt.upper(), i + 2
            continue
        if w.startswith("-X") and len(w) > 2:
            method, i = w[2:].upper(), i + 1
            continue
        if w in ("-H", "--header"):
            name, _, value = nxt.partition(":")
            if name.strip():
                headers.append({"name": name.strip(), "value": value.strip()})
            i += 2
            continue
        if w in ("-d", "--data", "--data-raw", "--data-binary", "--data-ascii", "--data-urlencode"):
            data.append(nxt[1:] if nxt.startswith("$") and w == "--data-raw" else nxt)
            i += 2
            continue
        if w == "--json":
            data.append(nxt)
            is_json = True
            i += 2
            continue
        if w in ("-u", "--user"):
            import base64
            headers.append({"name": "Authorization", "value": "Basic " + base64.b64encode(nxt.encode()).decode()})
            i += 2
            continue
        if w == "--url":
            url, i = nxt, i + 2
            continue
        if w in ("-G", "--get"):
            get, i = True, i + 1
            continue
        if w == "-A" or w == "--user-agent":
            headers.append({"name": "User-Agent", "value": nxt})
            i += 2
            continue
        if w in ("-b", "--cookie"):
            headers.append({"name": "Cookie", "value": nxt})
            i += 2
            continue
        if w in takes_value:
            i += 2
            continue
        if w.startswith("-"):
            i += 1
            continue
        if not url:
            url = w
        i += 1
    if not url:
        raise ValueError("there is no address in it")
    if not re.match(r"^https?://", url, re.I):
        url = "https://" + url
    body = "&".join(data) if data else ""
    if get and body:
        url += ("&" if "?" in url else "?") + body
        body = ""
    method = method or ("POST" if body else "GET")
    ctype = next((h["value"] for h in headers if h["name"].lower() == "content-type"), "")
    stripped = body.lstrip()
    fmt = "json" if is_json or "json" in ctype.lower() or stripped.startswith(("{", "[")) else "xml" if "xml" in ctype.lower() or stripped.startswith("<") else "json"
    if is_json and not ctype:
        headers.append({"name": "Content-Type", "value": "application/json"})
    return {"method": method, "url": url, "headers": headers, "body": body, "format": fmt}


def _postman_text(value: Any) -> str:
    """``{{var}}`` (Postman) -> ``{var}`` (the workbook's form)."""
    return re.sub(r"\{\{\s*([A-Za-z_][A-Za-z0-9_]*)\s*\}\}", r"{\1}", str(value or ""))


def parse_postman(data: Any) -> list[dict]:
    """The requests of a Postman collection (v2.0 / v2.1 JSON), folders flattened: ``[{name, folder, method, url, headers, body, format}]``.
    Postman's ``{{variables}}`` become ``{variables}``."""
    if isinstance(data, (str, bytes)):
        data = json.loads(data)
    if not isinstance(data, dict) or "item" not in data:
        raise ValueError("this is not a Postman collection (no 'item' list)")
    out: list[dict] = []

    def walk(items: list, folder: str) -> None:
        for item in items or []:
            if not isinstance(item, dict):
                continue
            if "item" in item:
                walk(item["item"], (folder + " / " if folder else "") + str(item.get("name") or ""))
                continue
            req = item.get("request")
            if isinstance(req, str):
                req = {"url": req, "method": "GET"}
            if not isinstance(req, dict):
                continue
            url = req.get("url")
            if isinstance(url, dict):
                url = url.get("raw") or ((url.get("protocol") + "://" if url.get("protocol") else "") + ".".join(url.get("host") or [])
                                         + ("/" + "/".join(url.get("path") or []) if url.get("path") else ""))
            headers = [{"name": _postman_text(h.get("key")), "value": _postman_text(h.get("value"))}
                       for h in req.get("header") or [] if isinstance(h, dict) and h.get("key") and not h.get("disabled")]
            body, fmt = "", "json"
            b = req.get("body") or {}
            if isinstance(b, dict):
                mode = b.get("mode")
                if mode == "raw":
                    body = _postman_text(b.get("raw"))
                    lang = str(((b.get("options") or {}).get("raw") or {}).get("language") or "")
                    fmt = "xml" if lang == "xml" or body.lstrip().startswith("<") else "json"
                elif mode == "urlencoded":
                    body = "&".join(f"{p.get('key')}={_postman_text(p.get('value'))}" for p in b.get("urlencoded") or [] if not p.get("disabled"))
            ctype = next((h["value"] for h in headers if h["name"].lower() == "content-type"), "")
            if "xml" in ctype.lower():
                fmt = "xml"
            out.append({"name": str(item.get("name") or "Request"), "folder": folder, "method": str(req.get("method") or "GET").upper(),
                        "url": _postman_text(url), "headers": headers, "body": body, "format": fmt})

    walk(data["item"], "")
    return out


def suggest(request: dict, *, environments: dict, environment: str, row_values: dict[str, str] | None = None,
            variables: dict[str, str] | None = None) -> dict:
    """Turn an imported request into the workbook's form, and say what was found: the start of the URL that an environment variable holds
    becomes ``{VAR}``; a credential header (Authorization, x-api-key, a token) becomes ``{SECRET:NAME}`` and its value is dropped (it goes in
    secrets.env, never the sheet); body values equal to a column of the data row become ``{COLUMN}``.  ``variables``: Postman variables ->
    the workbook names chosen for them."""
    req = {**request, "headers": [dict(h) for h in request.get("headers") or []]}
    found: list[dict] = []
    env = environment.upper()
    best = ("", "")
    for r in environments.get("rows", []):
        value = str((r.get("values") or {}).get(env) or "").strip()
        for candidate in (value, value.rstrip("/")):                   # (the value's own trailing slash, if the address has it too)
            if candidate and len(candidate) > len(best[1]) and req["url"].lower().startswith(candidate.lower()) and not r.get("secret"):
                best = (r["variable"], candidate)
                break
    if best[0]:
        req["url"] = "{" + best[0] + "}" + req["url"][len(best[1]):]
        found.append({"kind": "domain", "variable": best[0], "text": f"The start of the address is {{{best[0]}}} ({environment})"})
    for h in req["headers"]:
        if _is_secret_header(h["name"]) and h["value"] and "{" not in h["value"]:
            name = B.make_token(re.sub(r"^x[-_]", "", h["name"], flags=re.I))
            scheme = re.match(r"^(Bearer|Basic|Token)\s+", h["value"], re.I)
            h["value"] = (scheme.group(0) if scheme else "") + "{SECRET:" + name + "}"
            found.append({"kind": "secret", "variable": name, "header": h["name"],
                          "text": f"{h['name']} -> {{SECRET:{name}}}: not stored in the sheet. Put it in secrets.env as RR_SECRET_{name}=..."})
    for old, new in (variables or {}).items():
        if old and new and old != new:
            req["url"] = req["url"].replace("{" + old + "}", "{" + new + "}")
            req["body"] = (req.get("body") or "").replace("{" + old + "}", "{" + new + "}")
            for h in req["headers"]:
                h["value"] = h["value"].replace("{" + old + "}", "{" + new + "}")
    matched = []
    body = req.get("body") or ""
    for column, value in sorted((row_values or {}).items(), key=lambda kv: -len(str(kv[1] or ""))):
        v = str(value or "").strip()
        if len(v) < 2 or MASK in v or v.upper() in ("Y", "N", "TRUE", "FALSE", "NONE") or column.upper() in ("BLNEXECUTE", "TCID", "TC_NAME"):
            continue
        quoted = json.dumps(v)
        if quoted in body:
            body = body.replace(quoted, '"{' + column + '}"')
            matched.append(column)
        elif re.search(r"(?<![\w.])" + re.escape(v) + r"(?![\w.])", body) and (re.fullmatch(r"-?\d+(\.\d+)?", v) or req.get("format") == "xml"):
            body = re.sub(r"(?<![\w.])" + re.escape(v) + r"(?![\w.])", "{" + column + "}", body)
            matched.append(column)
    if matched:
        req["body"] = body
        found.append({"kind": "values", "columns": matched, "text": f"{len(matched)} body values match this data row -> "
                      + ", ".join("{%s}" % c for c in matched)})
    return {"request": req, "found": found}


_PLACEHOLDER = re.compile(r"#[A-Za-z_][A-Za-z0-9_]*#|\{\{\s*[A-Za-z_][A-Za-z0-9_]*\s*\}\}|\$\{[A-Za-z_][A-Za-z0-9_]*\}|\b[A-Za-z][A-Za-z0-9]*_value\b")


def template_folders(cfg: Config, workbook_dir: Path, location: str = "") -> list[Path]:
    """Where request templates may be (the runner's own search order, ``workbook/api.template_candidates``), folders that exist only."""
    out: list[Path] = []
    candidates = [location, location.replace("\\", "/") if os.sep == "/" else "", cfg.api.templates_dir, os.environ.get("RR_API_TEMPLATES_DIR", ""),
                  str(workbook_dir / "templates"), str(workbook_dir)]
    for c in candidates:
        if c and c.strip():
            p = Path(c.strip()).expanduser()
            if p.is_dir() and p not in out:
                out.append(p)
    return out


def list_templates(cfg: Config, workbook_dir: Path, location: str = "") -> list[dict]:
    """Request template files (``.json`` / ``.xml`` / ``.txt``) in those folders: ``[{file, folder, format}]`` (no workbook files)."""
    out = []
    for folder in template_folders(cfg, workbook_dir, location):
        try:
            entries = sorted(folder.iterdir())
        except OSError:
            continue
        for p in entries:
            if p.is_file() and p.suffix.lower() in (".json", ".xml", ".txt") and p.stat().st_size < 2_000_000:
                out.append({"file": p.name, "folder": str(folder), "format": _template_format(p)})
    return out[:500]


def _template_format(path: Path) -> str:
    try:
        head = path.read_text(encoding="utf-8-sig", errors="replace")[:200].lstrip()
    except OSError:
        return "json"
    return "xml" if head.startswith("<") else "json"


def template_fields(path: Path, editor: WorkbookEditor, sheet: str, row: int) -> dict:
    """The placeholders of a template file and what fills each one: the ``Replace`` row that already names it (and its column), else a column
    of the sheet whose name matches, else nothing yet (the person chooses)."""
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    io = [r for r in read_io(editor) if r.kind == "replace"]
    cols = {h.upper(): h for h in editor.headers(sheet)}
    grid = _Grid(editor, sheet)
    fields: list[dict] = []
    seen: set[str] = set()
    for r in io:
        if r.parameter in text and r.parameter not in seen:
            seen.add(r.parameter)
            fields.append({"placeholder": r.parameter, "count": text.count(r.parameter), "column": cols.get(r.value.upper(), r.value),
                           "mapped": True, "value": grid.raw(row, r.value.upper())})
    for m in _PLACEHOLDER.finditer(text):
        p = m.group(0)
        if p in seen:
            continue
        seen.add(p)
        bare = re.sub(r"^[#{$]+\s*|\s*[#}]+$|_value$", "", p).upper()
        match = next((h for u, h in cols.items() if re.sub(r"^(DT_|TC_)|(_IN|_OUT)$", "", u) == bare or u == bare), "")
        fields.append({"placeholder": p, "count": text.count(p), "column": match, "mapped": False,
                       "value": grid.raw(row, match.upper()) if match else ""})
    return {"file": path.name, "folder": str(path.parent), "format": "xml" if text.lstrip().startswith("<") else "json",
            "preview": text[:4000], "fields": fields}
