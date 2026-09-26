"""Workbook snapshot (immutable) and per-test cell state (mutable overlay).

``WorkbookData`` is read once from disk with openpyxl and shared by every test.
``BookState`` is created per test: it owns the *overlay* of values written at runtime (token
substitution, Status/Error_Check/Output_Value write-back, environment overrides) and evaluates
formulas lazily against snapshot + overlay.  Because every test owns its own ``BookState`` there
is no shared mutable state, which is what makes concurrent execution safe.
"""
from __future__ import annotations

import random
import warnings
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

import openpyxl

from .formula import (Evaluator, ExcelDate, ExcelError, RangeValue, collapse, number_to_text)


class ErrorText(str):
    """A cell whose formula evaluated to an Excel error (``#N/A`` ...). Re-raised when referenced."""

    detail: str = ""                     # why (e.g. "unsupported function X"), for messages; never part of the text

    def __new__(cls, code: str, detail: str = ""):
        obj = super().__new__(cls, code)
        obj.detail = detail
        return obj

    @property
    def code(self) -> str:
        return str(self)


@dataclass(frozen=True)
class RawCell:
    value: Any = None
    formula: str | None = None
    cached: Any = None  # value Excel last saved (diagnostics / test oracle only)


class SheetData:
    def __init__(self, name: str):
        self.name = name
        self.cells: dict[tuple[int, int], RawCell] = {}
        self.max_row = 0
        self.max_col = 0
        self.headers: dict[str, int] = {}          # UPPER header -> column (last wins, like the legacy runner)
        self.header_names: dict[int, str] = {}     # column -> header as written

    def finalize(self) -> None:
        self.headers = {}
        self.header_names = {}
        for col in range(1, self.max_col + 1):
            cell = self.cells.get((1, col))
            if cell is not None and cell.value not in (None, "") and cell.formula is None:
                text = str(cell.value).strip()
                self.headers[text.upper()] = col
                self.header_names[col] = text


class WorkbookData:
    """Read-only snapshot of an .xlsx file (formulas preserved)."""

    def __init__(self, path: Path):
        self.path = path
        self.sheets: dict[str, SheetData] = {}   # UPPER name -> data
        self.order: list[str] = []               # names as written

    @classmethod
    def load(cls, path: str | Path, *, with_cached: bool = False) -> "WorkbookData":
        path = Path(path)
        if not path.is_file():
            raise FileNotFoundError(f"Workbook not found: {path}")
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = openpyxl.load_workbook(path, data_only=False, keep_links=False)
            wb_values = openpyxl.load_workbook(path, data_only=True, keep_links=False) if with_cached else None
        data = cls(path)
        for ws in wb.worksheets:
            sheet = SheetData(ws.title)
            cached_ws = wb_values[ws.title] if wb_values else None
            for row in ws.iter_rows():
                for c in row:
                    v = c.value
                    if v is None:
                        continue
                    formula = None
                    if hasattr(v, "text"):                       # ArrayFormula
                        formula, v = v.text, None
                    elif isinstance(v, str) and v.startswith("=") and len(v) > 1:
                        formula, v = v, None
                    cached = cached_ws[c.coordinate].value if cached_ws is not None else None
                    sheet.cells[(c.row, c.column)] = RawCell(v, formula, cached)
                    sheet.max_row = max(sheet.max_row, c.row)
                    sheet.max_col = max(sheet.max_col, c.column)
            sheet.finalize()
            data.sheets[ws.title.upper()] = sheet
            data.order.append(ws.title)
        return data

    def sheet(self, name: str) -> SheetData | None:
        return self.sheets.get(str(name).upper().strip())


# ---------------------------------------------------------------------------------------------
# Runtime state
# ---------------------------------------------------------------------------------------------
def cell_text(value: Any) -> str:
    """Text as the legacy runner saw it (``str()`` of the COM cell value, blanks as empty)."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, ExcelDate):
        from .textfmt import to_datetime
        return str(to_datetime(value))
    if isinstance(value, float):
        return number_to_text(value)
    if isinstance(value, (datetime, date, time)):
        return str(value)
    return str(value)


class _SheetCtx:
    """EvalContext bound to the sheet a formula lives on (unqualified refs resolve here)."""

    def __init__(self, book: "BookState", sheet: "SheetState"):
        self.book, self.sheet = book, sheet

    def get_cell(self, sheet, row, col):
        target = self.book.sheet(sheet) if sheet else self.sheet
        value = target.get(row, col)
        if isinstance(value, ErrorText):
            raise ExcelError(value.code, value.detail)
        return value

    def get_range(self, sheet, r1, c1, r2, c2):
        target = self.book.sheet(sheet) if sheet else self.sheet
        return RangeValue([[target.get_raw_for_range(r, c) for c in range(c1, c2 + 1)] for r in range(r1, r2 + 1)])

    def now(self):
        return self.book.now()

    def randint(self, low, high):
        return self.book.rng.randint(low, high)


class SheetState:
    def __init__(self, book: "BookState", data: SheetData):
        self.book = book
        self.data = data
        self.name = data.name
        self.overlay: dict[tuple[int, int], Any] = {}
        self._ctx = _SheetCtx(book, self)

    # -- reading -----------------------------------------------------------------------------
    def get(self, row: int, col: int) -> Any:
        key = (row, col)
        if key in self.overlay:
            return self.overlay[key]
        raw = self.data.cells.get(key)
        if raw is None:
            return None
        if raw.formula is None:
            return raw.value
        return self.book.eval_cell(self, key, raw.formula)

    def get_raw_for_range(self, row: int, col: int) -> Any:
        """Range members: errors do not propagate (COUNTA etc. tolerate them)."""
        try:
            return self.get(row, col)
        except ExcelError:
            return None

    def col(self, header: str) -> int | None:
        return self.data.headers.get(header.upper().strip())

    def set(self, row: int, col: int, value: Any) -> None:
        self.overlay[(row, col)] = value
        self.book.invalidate()

    def is_formula(self, row: int, col: int) -> bool:
        return (row, col) not in self.overlay and (raw := self.data.cells.get((row, col))) is not None \
            and raw.formula is not None

    def read(self, row: int, col: int) -> Any:
        """Like ``get`` but an Excel error becomes an ``ErrorText`` value instead of raising."""
        try:
            return self.get(row, col)
        except ExcelError as err:
            return ErrorText(err.code, err.detail)

    def freeze(self) -> None:
        """Evaluate every formula once and store the result as a constant.

        Used for ``Global`` and ``Params_*`` sheets so volatile functions (TODAY, RANDBETWEEN) yield
        one consistent value for the whole test - Excel would silently re-roll them on every recalc.
        All cells are evaluated *before* the cache is invalidated so each formula runs exactly once.
        """
        frozen: dict[tuple[int, int], Any] = {}
        for (r, c), raw in list(self.data.cells.items()):
            if raw.formula is not None and (r, c) not in self.overlay:
                frozen[(r, c)] = self.read(r, c)
        self.overlay.update(frozen)
        self.book.invalidate()


class BookState:
    """Mutable, per-test view of a workbook."""

    def __init__(self, data: WorkbookData, *, now: datetime | None = None, seed: int | None = None):
        self.data = data
        self._now = now
        self.rng = random.Random(seed)
        self._sheets: dict[str, SheetState] = {}
        self._cache: dict[tuple[str, int, int], Any] = {}
        self._active: set[tuple[str, int, int]] = set()

    def now(self) -> datetime:
        return self._now or datetime.now()

    def sheet(self, name: str) -> SheetState:
        key = str(name).upper().strip()
        if key not in self._sheets:
            data = self.data.sheets.get(key)
            if data is None:
                raise ExcelError("#REF!", f"unknown sheet {name!r}")
            self._sheets[key] = SheetState(self, data)
        return self._sheets[key]

    def has_sheet(self, name: str) -> bool:
        return str(name).upper().strip() in self.data.sheets

    def invalidate(self) -> None:
        self._cache.clear()

    def eval_cell(self, sheet: SheetState, key: tuple[int, int], formula: str) -> Any:
        ckey = (sheet.name.upper(), key[0], key[1])
        if ckey in self._cache:
            cached = self._cache[ckey]
            if isinstance(cached, ExcelError):
                raise cached
            return cached
        if ckey in self._active:
            raise ExcelError("#REF!", "circular reference")
        self._active.add(ckey)
        try:
            try:
                value = collapse(Evaluator(sheet._ctx).evaluate(formula))
            except ExcelError as err:
                self._cache[ckey] = err
                raise
            if isinstance(value, ExcelError):
                raise value
            self._cache[ckey] = value
            return value
        finally:
            self._active.discard(ckey)
