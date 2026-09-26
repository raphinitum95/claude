"""Edit a workbook in place without losing anything Excel put in it.

Why not openpyxl for writing: measured on the four real workbooks, an openpyxl load/save drops the printer settings, the
``customXml`` parts, ``metadata.xml`` (dynamic-array formulas turn into ``@`` implicit-intersection ones in Excel), the calc
chain, every cached formula value (Excel users would see blanks until a recalculation) and renames the comment parts.
``WorkbookEditor`` instead keeps the ``.xlsx`` package as it is and patches only the XML of the parts an edit touches:

* an unchanged part is written back byte-for-byte; an unchanged row inside a changed sheet keeps its exact XML text (legacy
  rows such as ``GO_TO_ROW``, formula-driven ``blnExecute`` or ``SNAGIT_SCREENSHOT`` survive untouched, Q28);
* inserting, deleting or moving rows renumbers the rows below and rewrites every reference to them: formulas in every sheet,
  defined names, merged cells, data validations, conditional formats, hyperlinks, tables, comments and their note shapes;
* new text goes into the cell as an inline string, so ``sharedStrings.xml`` never changes;
* after any change the calc chain is dropped and ``fullCalcOnLoad`` is set, so Excel recomputes every cached value when the
  file is opened (the runner evaluates formulas itself and never needs the cached values). Viewers that do not calculate
  (Quick Look, some previewers) may show stale values for formulas that depend on edited cells: a documented gap.

Saving guards the user's file: Excel's lock file ``~$name.xlsx`` (or Windows refusing the write) raises ``WorkbookLocked``
("close it in Excel"); a file changed on disk since it was opened (mtime + size + hash) raises ``WorkbookChangedOutside``
unless the caller forces it; the previous version is copied to ``backups/<name>.<yyyy-mm-dd_hhmm>.xlsx`` first. Drafts
(autosave) live in ``.drafts/`` next to the workbook, like ``.trash/`` and ``.chains/``, so they never show up in the list
of workbooks.
"""
from __future__ import annotations

import hashlib
import html
import json
import io
import os
import posixpath
import re
import shutil
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path, PurePosixPath
from typing import Any, Iterable, Iterator

from openpyxl.formula.tokenizer import Token, Tokenizer
from openpyxl.formula.translate import Translator
from openpyxl.utils import column_index_from_string, get_column_letter

MAX_ROW = 1048576
MAX_COL = 16384
RR_PREFIX = "_rr_"                       # hidden sheets the runner owns (CONTRACT.md section 1)

NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL_WORKSHEET = NS_REL + "/worksheet"
CT_WORKSHEET = "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"

# Functions newer than Excel 2007 are stored with a ``_xlfn.`` prefix; without it Excel shows #NAME?.  Only functions that
# return a single value are listed: dynamic-array functions also need cell metadata this writer does not create.
FUTURE_FUNCTIONS = frozenset({"IFS", "SWITCH", "TEXTJOIN", "CONCAT", "MAXIFS", "MINIFS", "IFNA", "NUMBERVALUE", "DAYS",
                              "ISOWEEKNUM", "XLOOKUP", "XMATCH", "CEILING.MATH", "FLOOR.MATH", "STDEV.S", "STDEV.P"})


class WriterError(ValueError):
    """An edit that cannot be done (unknown sheet, bad name, a header that already exists...)."""


class WorkbookLocked(RuntimeError):
    """The workbook is open in Excel (lock file present, or the OS refused the write)."""


class WorkbookChangedOutside(RuntimeError):
    """The file on disk changed after it was opened for editing: saving would throw those changes away."""


# ---------------------------------------------------------------------------------------------
# Files: stamps, lock file, backups, drafts
# ---------------------------------------------------------------------------------------------
@dataclass(frozen=True)
class FileStamp:
    """What the file looked like when it was read. mtime alone is not trusted (sync clients touch files)."""

    mtime_ns: int
    size: int
    sha256: str

    @classmethod
    def of(cls, path: Path) -> "FileStamp":
        data = path.read_bytes()
        stat = path.stat()
        return cls(stat.st_mtime_ns, stat.st_size, hashlib.sha256(data).hexdigest())

    def matches(self, path: Path) -> bool:
        """True when ``path`` still holds the same bytes (cheap mtime/size check first, hash only when they differ)."""
        try:
            stat = path.stat()
        except FileNotFoundError:
            return False
        if stat.st_mtime_ns == self.mtime_ns and stat.st_size == self.size:
            return True
        return stat.st_size == self.size and hashlib.sha256(path.read_bytes()).hexdigest() == self.sha256

    def to_json(self) -> dict:
        return {"mtime_ns": self.mtime_ns, "size": self.size, "sha256": self.sha256}

    @classmethod
    def from_json(cls, data: dict) -> "FileStamp":
        return cls(int(data["mtime_ns"]), int(data["size"]), str(data["sha256"]))


def lock_file(path: Path) -> Path:
    """Excel's owner file, created while the workbook is open in Excel."""
    path = Path(path)
    return path.with_name("~$" + path.name)


def is_locked(path: Path) -> bool:
    return lock_file(path).exists()


def backups_dir(path: Path) -> Path:
    return Path(path).parent / "backups"


def backup_path(path: Path, when: datetime | None = None) -> Path:
    """``backups/<name>.<yyyy-mm-dd_hhmm>.xlsx``; a second save in the same minute gets ``-2``, ``-3``..."""
    path = Path(path)
    stamp = (when or datetime.now()).strftime("%Y-%m-%d_%H%M")
    folder = backups_dir(path)
    target = folder / f"{path.stem}.{stamp}{path.suffix}"
    n = 2
    while target.exists():
        target = folder / f"{path.stem}.{stamp}-{n}{path.suffix}"
        n += 1
    return target


def draft_path(path: Path) -> Path:
    path = Path(path)
    return path.parent / ".drafts" / path.name


def _draft_meta_path(path: Path) -> Path:
    draft = draft_path(path)
    return draft.with_name(draft.name + ".json")


def has_draft(path: Path) -> bool:
    return draft_path(path).is_file()


def discard_draft(path: Path) -> None:
    draft_path(path).unlink(missing_ok=True)
    _draft_meta_path(path).unlink(missing_ok=True)


def _locked_message(path: Path) -> str:
    return f"{path.name} is open in Excel. Close it in Excel, then save again."


def _atomic_write(target: Path, data: bytes) -> None:
    """Write next to the target, then swap: a crash never leaves half a workbook. Windows refuses the swap while Excel
    has the file open, which is reported the same way as the lock file."""
    target.parent.mkdir(parents=True, exist_ok=True)
    tmp = target.with_name(target.name + ".rr-tmp")
    tmp.write_bytes(data)
    try:
        os.replace(tmp, target)
    except PermissionError as err:
        tmp.unlink(missing_ok=True)
        raise WorkbookLocked(_locked_message(target)) from err


# ---------------------------------------------------------------------------------------------
# Small XML helpers (regex based on purpose: the untouched bytes must stay exactly as Excel wrote them)
# ---------------------------------------------------------------------------------------------
_ATTR = re.compile(r'([\w:.-]+)\s*=\s*"([^"]*)"')
_CELL_REF = re.compile(r"^\$?([A-Za-z]{1,3})\$?(\d+)$")
_XML_ESCAPE = re.compile(r"_x([0-9A-Fa-f]{4})_")
_BAD_XML_CHARS = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _attrs(tag: str) -> dict[str, str]:
    return {k: html.unescape(v) for k, v in _ATTR.findall(tag)}


def _set_attr(tag: str, name: str, value: str | None) -> str:
    """Change, add (``value``) or remove (``None``) one attribute of an opening tag, leaving the rest untouched."""
    pattern = re.compile(r'(\s)' + re.escape(name) + r'="[^"]*"')
    if value is None:
        return pattern.sub("", tag, count=1)
    escaped = html.escape(value, quote=True)
    if pattern.search(tag):
        return pattern.sub(lambda m: f'{m.group(1)}{name}="{escaped}"', tag, count=1)
    end = -2 if tag.endswith("/>") else -1
    return f'{tag[:end]} {name}="{escaped}"{tag[end:]}'


def _unescape_text(text: str) -> str:
    return _XML_ESCAPE.sub(lambda m: chr(int(m.group(1), 16)), html.unescape(text))


def _escape_text(text: str) -> str:
    text = _BAD_XML_CHARS.sub(lambda m: f"_x{ord(m.group(0)):04X}_", text)
    return html.escape(text, quote=False)


def _split_ref(ref: str) -> tuple[int, int]:
    m = _CELL_REF.match(ref)
    if not m:
        raise WriterError(f"not a cell reference: {ref!r}")
    return int(m.group(2)), column_index_from_string(m.group(1).upper())


def _ref(row: int, col: int) -> str:
    return f"{get_column_letter(col)}{row}"


# ---------------------------------------------------------------------------------------------
# Row maps: what an insert / delete / move does to row numbers
# ---------------------------------------------------------------------------------------------
class RowMap:
    """Piecewise map of old row numbers to new ones.  ``segments`` = (first, last, offset or None=deleted); rows outside
    every segment keep their number.  Ranges map to the smallest block holding what is left of them, which is what Excel
    does for inserts and deletes (a range grows when rows are inserted inside it and shrinks when rows are deleted)."""

    def __init__(self, segments: list[tuple[int, int, int | None]], moved: tuple[int, int, int] | None = None):
        self.segments = sorted(segments)
        self.moved = moved                      # (first, last, displacement) of the rows a move carried

    @classmethod
    def insert(cls, at: int, count: int) -> "RowMap":
        return cls([(at, MAX_ROW, count)])

    @classmethod
    def delete(cls, at: int, count: int) -> "RowMap":
        return cls([(at, at + count - 1, None), (at + count, MAX_ROW, -count)])

    @classmethod
    def move(cls, first: int, count: int, before: int) -> "RowMap":
        """Rows ``first..first+count-1`` move to sit just above old row ``before`` (Excel's cut + insert cut cells)."""
        last = first + count - 1
        if first <= before <= last + 1:
            return cls([])
        if before < first:
            return cls([(before, first - 1, count), (first, last, before - first)], (first, last, before - first))
        return cls([(last + 1, before - 1, -count), (first, last, before - count - first)],
                   (first, last, before - count - first))

    def is_identity(self) -> bool:
        return not any(off != 0 for _, _, off in self.segments)

    def row(self, r: int) -> int | None:
        for lo, hi, off in self.segments:
            if lo <= r <= hi:
                return None if off is None else r + off
        return r

    def span(self, r1: int, r2: int) -> tuple[int, int] | None:
        if r1 > r2:
            r1, r2 = r2, r1
        lows, highs = [], []
        cursor = r1
        for lo, hi, off in self.segments + [(MAX_ROW + 1, MAX_ROW + 1, 0)]:
            if hi < cursor:
                continue
            if cursor < lo:                                   # identity gap before this segment
                gap_hi = min(lo - 1, r2)
                if cursor <= gap_hi:
                    lows.append(cursor)
                    highs.append(gap_hi)
                cursor = lo
            if cursor > r2:
                break
            if lo > MAX_ROW:
                break
            seg_lo, seg_hi = max(lo, cursor), min(hi, r2)
            if seg_lo <= seg_hi and off is not None:
                lows.append(seg_lo + off)
                highs.append(seg_hi + off)
            cursor = hi + 1
            if cursor > r2:
                break
        if not lows:
            return None
        return min(lows), max(highs)


_REF_CELL = r"\$?[A-Za-z]{1,3}\$?\d+"
_REF_PATTERN = re.compile(rf"^(?P<a>{_REF_CELL})(?::(?P<b>{_REF_CELL}))?$|^(?P<ra>\$?\d+):(?P<rb>\$?\d+)$")
_ROW_PART = re.compile(r"^(\$?[A-Za-z]{1,3})(\$?)(\d+)$")


def _map_cell_text(text: str, row: int) -> str:
    m = _ROW_PART.match(text)
    return f"{m.group(1)}{m.group(2)}{row}"


def map_ref_text(ref: str, rows: RowMap) -> str | None:
    """Map one reference (``A1``, ``$A$1:B7``, ``3:5``) through ``rows``. None = everything it pointed at was deleted.
    Anything that is not a plain reference (whole columns, names...) comes back unchanged."""
    m = _REF_PATTERN.match(ref)
    if not m:
        return ref
    if m.group("ra"):
        ra, rb = m.group("ra"), m.group("rb")
        span = rows.span(int(ra.lstrip("$")), int(rb.lstrip("$")))
        if span is None:
            return None
        return f"{'$' if ra.startswith('$') else ''}{span[0]}:{'$' if rb.startswith('$') else ''}{span[1]}"
    a, b = m.group("a"), m.group("b")
    ra = int(_ROW_PART.match(a).group(3))
    if b is None:
        if ra > MAX_ROW:
            return ref
        target = rows.row(ra)
        return None if target is None else _map_cell_text(a, target)
    rb = int(_ROW_PART.match(b).group(3))
    if ra <= rb:
        span = rows.span(ra, rb)
        return None if span is None else f"{_map_cell_text(a, span[0])}:{_map_cell_text(b, span[1])}"
    span = rows.span(rb, ra)
    return None if span is None else f"{_map_cell_text(a, span[1])}:{_map_cell_text(b, span[0])}"


def _map_ref_points(ref: str, point) -> str | None:
    """Like ``map_ref_text`` but each end of the reference is mapped on its own by ``point(row, absolute)``."""
    m = _REF_PATTERN.match(ref)
    if not m:
        return ref
    if m.group("ra"):
        ends = [m.group("ra"), m.group("rb")]
        mapped = [point(int(e.lstrip("$")), e.startswith("$")) for e in ends]
        if None in mapped:
            return None
        return ":".join(("$" if e.startswith("$") else "") + str(r) for e, r in zip(ends, mapped))
    out = []
    for end in (m.group("a"), m.group("b")):
        if end is None:
            continue
        pm = _ROW_PART.match(end)
        target = point(int(pm.group(3)), pm.group(2) == "$")
        if target is None:
            return None
        out.append(_map_cell_text(end, target))
    return ":".join(out)


def map_sqref(sqref: str, rows: RowMap) -> str:
    """Space separated list of ranges (data validation, conditional format); deleted pieces disappear."""
    out = [mapped for part in sqref.split() if (mapped := map_ref_text(part, rows)) is not None]
    return " ".join(out)


def _sheet_prefix(operand: str) -> tuple[str | None, str, str]:
    """Split ``'My sheet'!A1`` into (sheet name, prefix as written, reference). External refs (``[1]Sheet!A1``) and 3-D
    refs get sheet ``"\0"`` so they never match a sheet of this workbook."""
    if "!" not in operand:
        return None, "", operand
    prefix, _, ref = operand.rpartition("!")
    name = prefix
    if name.startswith("'") and name.endswith("'"):
        name = name[1:-1].replace("''", "'")
    if name.startswith("[") or ":" in name:
        return "\0", prefix + "!", ref
    return name, prefix + "!", ref


def map_formula(formula: str, own_sheet: str | None, target_sheet: str, rows: RowMap, point=None) -> str:
    """Rewrite the references in ``formula`` (with or without the leading ``=``) that point into ``target_sheet``.
    ``own_sheet`` is where the formula lives (its unqualified references point there; None for defined names).
    ``point(row, absolute)``, when given, maps each end of a reference instead of ``rows`` (used for moved rows)."""
    has_eq = formula.startswith("=")
    try:
        tok = Tokenizer(formula if has_eq else "=" + formula)
    except Exception:                                          # not parseable: leave it exactly as it was
        return formula
    target = target_sheet.upper()
    own = own_sheet.upper() if own_sheet else None
    changed = False
    for item in tok.items:
        if item.type != Token.OPERAND or item.subtype != Token.RANGE:
            continue
        sheet, prefix, ref = _sheet_prefix(item.value)
        points_to = sheet.upper() if sheet is not None else own
        if points_to != target:
            continue
        mapped = _map_ref_points(ref, point) if point is not None else map_ref_text(ref, rows)
        new = prefix + ("#REF!" if mapped is None else mapped)
        if new != item.value:
            item.value = new
            changed = True
    if not changed:
        return formula
    text = tok.render()
    return text if has_eq else text[1:]


def add_future_prefixes(formula: str) -> str:
    """``=IFS(...)`` → ``=_xlfn.IFS(...)`` so Excel recognises functions added after 2007 (see ``FUTURE_FUNCTIONS``)."""
    try:
        tok = Tokenizer(formula)
    except Exception:
        return formula
    changed = False
    for item in tok.items:
        if item.type == Token.FUNC and item.subtype == Token.OPEN:
            name = item.value[:-1]
            if name.upper() in FUTURE_FUNCTIONS:
                item.value = "_xlfn." + name.upper() + "("
                changed = True
    return tok.render() if changed else formula


# ---------------------------------------------------------------------------------------------
# Cells and rows (kept as XML text; parsed only when an edit needs it)
# ---------------------------------------------------------------------------------------------
_ROW_RE = re.compile(r"<row\b[^>]*?/>|<row\b[^>]*>.*?</row>", re.S)
_CELL_RE = re.compile(r"<c\b[^>]*?/>|<c\b[^>]*>.*?</c>", re.S)
_F_RE = re.compile(r"<f\b([^>]*?)(?:/>|>(.*?)</f>)", re.S)
_V_RE = re.compile(r"<v>(.*?)</v>", re.S)
_T_RE = re.compile(r"<t\b[^>]*>(.*?)</t>", re.S)
_RPH_RE = re.compile(r"<rPh\b.*?</rPh>", re.S)


def _open_tag(xml: str) -> str:
    return xml[:xml.index(">") + 1]


class _Row:
    __slots__ = ("r", "xml", "tag", "cells", "dirty")

    def __init__(self, r: int, xml: str):
        self.r = r
        self.xml = xml                                  # original text; returned unchanged while not dirty
        self.tag = _open_tag(xml)
        self.cells: dict[int, str] | None = None        # parsed lazily
        self.dirty = False

    def parsed(self) -> dict[int, str]:
        if self.cells is None:
            self.cells = {}
            if not self.tag.endswith("/>"):
                inner = self.xml[len(self.tag):-len("</row>")]
                for m in _CELL_RE.finditer(inner):
                    cell = m.group(0)
                    ref = _attrs(_open_tag(cell)).get("r")
                    col = _split_ref(ref)[1] if ref else (max(self.cells) + 1 if self.cells else 1)
                    self.cells[col] = cell
        return self.cells

    def render(self) -> str:
        if not self.dirty:
            return self.xml
        cells = self.parsed()
        tag = _set_attr(self.tag, "r", str(self.r))
        if cells:
            spans = _attrs(tag).get("spans")
            if spans and ":" in spans:
                lo, hi = (int(x) for x in spans.split(":")[:2])
                lo, hi = min(lo, min(cells)), max(hi, max(cells))
                tag = _set_attr(tag, "spans", f"{lo}:{hi}")
        if tag.endswith("/>"):
            tag = tag[:-2].rstrip() + ">"
        if not cells:
            return tag[:-1] + "/>"
        body = "".join(cells[c] for c in sorted(cells))
        return f"{tag}{body}</row>"


@dataclass
class CellInfo:
    """One cell as stored: ``formula`` without ``=`` (None for constants), the constant or cached value, style index."""

    value: Any = None
    formula: str | None = None
    style: str | None = None
    shared_master: str | None = None    # shared-formula group id when this cell holds the group's text
    shared_follower: str | None = None  # shared-formula group id when this cell only points at the group


class _Sheet:
    def __init__(self, book: "WorkbookEditor", name: str, part: str, state: str | None):
        self.book = book
        self.name = name
        self.part = part
        self.state = state
        self._xml: str | None = None
        self.head = self.tail = ""
        self.rows: dict[int, _Row] | None = None
        self.masters: dict[str, tuple[str, str]] | None = None   # shared formulas, cached for reading
        self.dirty = False

    # -- parsing -------------------------------------------------------------------------------
    @property
    def xml(self) -> str:
        if self._xml is None:
            self._xml = self.book._read_text(self.part)
        return self._xml

    def load(self) -> dict[int, _Row]:
        if self.rows is None:
            xml = self.xml
            m = re.search(r"<sheetData\s*/>", xml)
            if m:
                self.head, body, self.tail = xml[:m.start()] + "<sheetData>", "", "</sheetData>" + xml[m.end():]
            else:
                start = xml.index("<sheetData")
                start = xml.index(">", start) + 1
                end = xml.rindex("</sheetData>")
                self.head, body, self.tail = xml[:start], xml[start:end], xml[end:]
            self.rows = {}
            last = 0
            for m in _ROW_RE.finditer(body):
                tag = _open_tag(m.group(0))
                r = int(_attrs(tag).get("r", last + 1))
                self.rows[r] = _Row(r, m.group(0))
                last = r
        return self.rows

    def render(self) -> str:
        rows = self.load()
        body = "".join(rows[r].render() for r in sorted(rows))
        head = self.head
        dim = self._dimension()
        if dim:
            head = re.sub(r'(<dimension\b[^>]*\bref=")[^"]*(")', lambda m: m.group(1) + dim + m.group(2), head, count=1)
        return head + body + self.tail

    def _dimension(self) -> str:
        cells = [(r, c) for r, row in self.load().items() for c in row.parsed()]
        if not cells:
            return "A1"
        top = _ref(min(r for r, _ in cells), min(c for _, c in cells))
        bottom = _ref(max(r for r, _ in cells), max(c for _, c in cells))
        return top if top == bottom else f"{top}:{bottom}"

    # -- reading -------------------------------------------------------------------------------
    def cell_xml(self, row: int, col: int) -> str | None:
        r = self.load().get(row)
        return None if r is None else r.parsed().get(col)

    def info(self, row: int, col: int) -> CellInfo:
        xml = self.cell_xml(row, col)
        return CellInfo() if xml is None else self.book._decode(xml)

    def max_row(self) -> int:
        """Last row holding a value or formula (rows of empty formatted cells do not count)."""
        rows = self.load()
        used = [r for r, row in rows.items() if any(_has_content(self.book._decode(x)) for x in row.parsed().values())]
        return max(used) if used else 0

    def max_col(self, row: int | None = None) -> int:
        rows = self.load()
        pick = [rows[row]] if row is not None and row in rows else ([] if row is not None else rows.values())
        cols = [c for rw in pick for c, xml in rw.parsed().items() if _has_content(self.book._decode(xml))]
        return max(cols) if cols else 0

    # -- writing -------------------------------------------------------------------------------
    def put(self, row: int, col: int, xml: str | None) -> None:
        rows = self.load()
        r = rows.get(row)
        if r is None:
            if xml is None:
                return
            r = rows[row] = _Row(row, f'<row r="{row}"/>')
        cells = r.parsed()
        if xml is None:
            cells.pop(col, None)
        else:
            cells[col] = xml
        r.dirty = True
        self.dirty = True
        self.masters = None


def _has_content(info: CellInfo) -> bool:
    return info.formula is not None or info.shared_follower is not None or info.value not in (None, "")


# ---------------------------------------------------------------------------------------------
# The editor
# ---------------------------------------------------------------------------------------------
@dataclass
class SaveResult:
    path: Path
    backup: Path | None = None


class WorkbookEditor:
    """An ``.xlsx`` opened for editing.  Rows and columns are 1-based, sheet names are matched case-insensitively."""

    def __init__(self, path: Path, data: bytes, stamp: FileStamp | None):
        self.path = Path(path)
        self.stamp = stamp
        self._zip_infos: list[zipfile.ZipInfo] = []
        self._parts: dict[str, bytes] = {}
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            for info in z.infolist():
                self._zip_infos.append(info)
                self._parts[info.filename] = z.read(info.filename)
        self._changed_parts: dict[str, bytes] = {}
        self._removed: set[str] = set()
        self._shared: list[str] | None = None
        self._sheets: dict[str, _Sheet] = {}
        self._order: list[str] = []
        self._modified = False
        self._load_sheet_list()

    # -- opening -------------------------------------------------------------------------------
    @classmethod
    def open(cls, path: str | Path) -> "WorkbookEditor":
        path = Path(path)
        if not path.is_file():
            raise FileNotFoundError(f"Workbook not found: {path}")
        data = path.read_bytes()
        stat = path.stat()
        return cls(path, data, FileStamp(stat.st_mtime_ns, stat.st_size, hashlib.sha256(data).hexdigest()))

    @classmethod
    def open_draft(cls, path: str | Path) -> "WorkbookEditor":
        """Reopen the autosaved draft of ``path``.  The stamp is the original's at the time the draft was started, so
        ``external_change()`` still tells whether someone changed the workbook since."""
        path = Path(path)
        draft = draft_path(path)
        if not draft.is_file():
            raise FileNotFoundError(f"No draft for {path.name}")
        stamp = None
        meta = _draft_meta_path(path)
        if meta.is_file():
            stamp = FileStamp.from_json(json.loads(meta.read_text(encoding="utf-8"))["original"])
        editor = cls(path, draft.read_bytes(), stamp)
        editor._modified = True
        return editor

    # -- package parts -------------------------------------------------------------------------
    def _read_bytes(self, part: str) -> bytes:
        if part in self._changed_parts:
            return self._changed_parts[part]
        return self._parts[part]

    def _read_text(self, part: str) -> str:
        return self._read_bytes(part).decode("utf-8")

    def _write_text(self, part: str, text: str) -> None:
        self._changed_parts[part] = text.encode("utf-8")
        self._removed.discard(part)
        self._modified = True

    def _has_part(self, part: str) -> bool:
        return (part in self._parts or part in self._changed_parts) and part not in self._removed

    @staticmethod
    def _rels_of(part: str) -> str:
        p = PurePosixPath(part)
        return str(p.parent / "_rels" / (p.name + ".rels"))

    @staticmethod
    def _resolve(base_part: str, target: str) -> str:
        if target.startswith("/"):
            return target.lstrip("/")
        parts: list[str] = list(PurePosixPath(base_part).parent.parts)
        for piece in target.split("/"):
            if piece == "..":
                if parts:
                    parts.pop()
            elif piece not in ("", "."):
                parts.append(piece)
        return "/".join(parts)

    def _relationships(self, part: str) -> list[dict[str, str]]:
        rels = self._rels_of(part)
        if not self._has_part(rels):
            return []
        return [_attrs(m.group(0)) for m in re.finditer(r"<Relationship\b[^>]*>", self._read_text(rels))]

    def _workbook_part(self) -> str:
        for rel in self._relationships(""):
            if rel.get("Type", "").endswith("/officeDocument"):
                return self._resolve("", rel["Target"])
        return "xl/workbook.xml"

    def _load_sheet_list(self) -> None:
        self._wb_part = self._workbook_part()
        targets = {rel["Id"]: rel for rel in self._relationships(self._wb_part)}
        for m in re.finditer(r"<sheet\b[^>]*>", self._read_text(self._wb_part)):
            a = _attrs(m.group(0))
            rel = targets.get(a.get("r:id", ""))
            if rel is None or not rel.get("Type", "").endswith("/worksheet"):
                continue                                             # chart sheets and the like are left alone
            sheet = _Sheet(self, a["name"], self._resolve(self._wb_part, rel["Target"]), a.get("state"))
            self._sheets[a["name"].upper()] = sheet
            self._order.append(a["name"])

    # -- shared strings ------------------------------------------------------------------------
    def _shared_strings(self) -> list[str]:
        if self._shared is None:
            self._shared = []
            for rel in self._relationships(self._wb_part):
                if rel.get("Type", "").endswith("/sharedStrings"):
                    xml = self._read_text(self._resolve(self._wb_part, rel["Target"]))
                    for m in re.finditer(r"<si\b[^>]*?(?:/>|>(.*?)</si>)", xml, re.S):
                        inner = _RPH_RE.sub("", m.group(1) or "")
                        self._shared.append("".join(_unescape_text(t) for t in _T_RE.findall(inner)))
        return self._shared

    def _decode(self, xml: str) -> CellInfo:
        tag = _open_tag(xml)
        a = _attrs(tag)
        info = CellInfo(style=a.get("s"))
        if tag.endswith("/>"):
            return info
        fm = _F_RE.search(xml)
        if fm:
            fa = _attrs(fm.group(1))
            if fa.get("t") == "shared":
                if fm.group(2):
                    info.shared_master = fa.get("si")
                else:
                    info.shared_follower = fa.get("si")
            if fm.group(2) is not None:
                info.formula = html.unescape(fm.group(2))
        t = a.get("t", "n")
        if t == "inlineStr":
            m = re.search(r"<is>(.*?)</is>", xml, re.S)
            info.value = "".join(_unescape_text(x) for x in _T_RE.findall(_RPH_RE.sub("", m.group(1)))) if m else ""
            return info
        vm = _V_RE.search(xml)
        if vm is None:
            return info
        raw = html.unescape(vm.group(1))
        if t == "s":
            info.value = self._shared_strings()[int(raw)]
        elif t == "b":
            info.value = raw == "1"
        elif t in ("str", "e"):
            info.value = _unescape_text(raw)
        else:
            number = float(raw)
            info.value = int(number) if number.is_integer() and "." not in raw and "E" not in raw.upper() else number
        return info

    # -- sheets --------------------------------------------------------------------------------
    def sheet_names(self) -> list[str]:
        return list(self._order)

    def has_sheet(self, name: str) -> bool:
        return str(name).upper().strip() in self._sheets

    def is_hidden(self, name: str) -> bool:
        return self._sheet(name).state in ("hidden", "veryHidden")

    def _sheet(self, name: str) -> _Sheet:
        sheet = self._sheets.get(str(name).upper().strip())
        if sheet is None:
            raise WriterError(f"no sheet named {name!r}")
        return sheet

    # -- reading cells -------------------------------------------------------------------------
    def get(self, sheet: str, row: int, col: int) -> Any:
        """What the cell holds as written: ``"=formula"`` for formulas (shared ones expanded), else the constant."""
        s = self._sheet(sheet)
        info = s.info(row, col)
        if info.shared_follower is not None:
            return "=" + self._follower_formula(s, row, col, info.shared_follower)
        return "=" + info.formula if info.formula is not None else info.value

    def cached(self, sheet: str, row: int, col: int) -> Any:
        """The value Excel last calculated (None for formulas this editor wrote)."""
        info = self._sheet(sheet).info(row, col)
        return info.value

    def style_of(self, sheet: str, row: int, col: int) -> str | None:
        return self._sheet(sheet).info(row, col).style

    def cell_xml(self, sheet: str, row: int, col: int) -> str | None:
        """The raw XML of one cell (tests use it to prove untouched cells are byte-identical)."""
        return self._sheet(sheet).cell_xml(row, col)

    def row_xml(self, sheet: str, row: int) -> str | None:
        r = self._sheet(sheet).load().get(row)
        return None if r is None else r.render()

    def max_row(self, sheet: str) -> int:
        return self._sheet(sheet).max_row()

    def headers(self, sheet: str) -> dict[str, int]:
        """Row-1 headers as written → column (the reader matches them case-insensitively, last one wins)."""
        s = self._sheet(sheet)
        out: dict[str, int] = {}
        row = s.load().get(1)
        for col in sorted(row.parsed() if row else {}):
            info = s.info(1, col)
            if info.formula is None and info.value not in (None, ""):
                out[str(info.value).strip()] = col
        return out

    def column(self, sheet: str, header: str) -> int | None:
        wanted = header.upper().strip()
        found = None
        for name, col in self.headers(sheet).items():
            if name.upper() == wanted:
                found = col
        return found

    def rows(self, sheet: str) -> list[list[Any]]:
        """Every row from 1 to the last used one as lists of ``get`` values (blank rows are empty lists)."""
        s = self._sheet(sheet)
        out = []
        for r in range(1, s.max_row() + 1):
            row = s.load().get(r)
            cols = row.parsed() if row else {}
            width = max(cols) if cols else 0
            out.append([self.get(sheet, r, c) for c in range(1, width + 1)])
        while out and all(v in (None, "") for v in out[-1]):
            out.pop()
        return out

    # -- writing cells -------------------------------------------------------------------------
    def set(self, sheet: str, row: int, col: int, value: Any, *, style: str | None = None) -> None:
        """Write a constant, or a formula when ``value`` is text starting with ``=``.  The cell keeps its style (or takes
        ``style``, a style index as found in another cell).  ``None`` clears the value and keeps the formatting."""
        if not (1 <= row <= MAX_ROW and 1 <= col <= MAX_COL):
            raise WriterError(f"cell out of range: row {row}, column {col}")
        s = self._sheet(sheet)
        old = s.cell_xml(row, col)
        info = self._decode(old) if old else CellInfo()
        if info.shared_master is not None:
            self._expand_shared(s, only={info.shared_master})
            old = s.cell_xml(row, col)
        keep_style = style if style is not None else info.style
        extra = {}
        if old:
            a = _attrs(_open_tag(old))
            extra = {k: v for k, v in a.items() if k not in ("r", "s", "t", "cm", "vm")}
        s.put(row, col, _cell_xml(_ref(row, col), keep_style, value, extra))
        self._touch()

    def set_by_header(self, sheet: str, row: int, header: str, value: Any) -> None:
        col = self.column(sheet, header)
        if col is None:
            raise WriterError(f"sheet {sheet!r} has no column {header!r}")
        self.set(sheet, row, col, value)

    def add_column(self, sheet: str, header: str) -> int:
        """Add ``header`` after the last used header of row 1; it takes the previous header's style.  Returns its column."""
        if self.column(sheet, header) is not None:
            raise WriterError(f"sheet {sheet!r} already has a column {header!r}")
        s = self._sheet(sheet)
        last = s.max_col(1)
        col = last + 1
        style = s.info(1, last).style if last else None
        self.set(sheet, 1, col, header, style=style)
        return col

    # -- rows ----------------------------------------------------------------------------------
    def insert_rows(self, sheet: str, at: int, count: int = 1, *, style_from: int | None = None) -> None:
        """Insert ``count`` blank rows above row ``at``.  They copy the formatting (row height, cell styles; not values
        and not hidden) of ``style_from``, by default the row above, or the row below when the row above is the header."""
        if at < 1 or count < 1:
            raise WriterError("insert_rows needs at >= 1 and count >= 1")
        s = self._sheet(sheet)
        if style_from is None:
            style_from = at - 1 if at > 2 else at
        source = s.load().get(style_from)
        template = None
        if source is not None:
            template = (source.tag, {c: _attrs(_open_tag(x)).get("s") for c, x in source.parsed().items()})
        rows = RowMap.insert(at, count)
        self._apply_rows(s, rows)
        if template is not None:
            tag, styles = template
            for r in range(at, at + count):
                row_tag = tag
                for name in ("hidden", "collapsed", "outlineLevel", "thickBot", "thickTop"):
                    row_tag = _set_attr(row_tag, name, None)
                row_tag = _set_attr(row_tag, "r", str(r))
                row_tag = row_tag if row_tag.endswith("/>") else row_tag[:-1] + "/>"
                new = s.load()[r] = _Row(r, row_tag)
                cells = new.parsed()
                for c, st in styles.items():
                    if st is not None:
                        cells[c] = f'<c r="{_ref(r, c)}" s="{st}"/>'
                new.dirty = True
        s.dirty = True
        self._touch()

    def delete_rows(self, sheet: str, at: int, count: int = 1) -> None:
        if at < 1 or count < 1:
            raise WriterError("delete_rows needs at >= 1 and count >= 1")
        self._apply_rows(self._sheet(sheet), RowMap.delete(at, count))
        self._touch()

    def move_rows(self, sheet: str, first: int, count: int, before: int) -> None:
        """Move rows ``first..first+count-1`` so they sit just above old row ``before`` (references follow the cells)."""
        if first < 1 or count < 1 or before < 1:
            raise WriterError("move_rows needs first, count and before >= 1")
        rows = RowMap.move(first, count, before)
        if rows.is_identity():
            return
        self._apply_rows(self._sheet(sheet), rows)
        self._touch()

    # -- new sheets ----------------------------------------------------------------------------
    def add_sheet(self, name: str, *, hidden: bool = False) -> None:
        """Append an empty sheet at the end of the tab order (so ``localSheetId`` of defined names stays valid)."""
        name = str(name)
        if not name or len(name) > 31 or re.search(r"[\[\]:*?/\\]", name) or name.startswith("'") or name.endswith("'"):
            raise WriterError(f"not a valid sheet name: {name!r}")
        if self.has_sheet(name):
            raise WriterError(f"a sheet named {name!r} already exists")
        existing = {s.part for s in self._sheets.values()} | set(self._parts) | set(self._changed_parts)
        n = 1
        while f"xl/worksheets/sheet{n}.xml" in existing:
            n += 1
        part = f"xl/worksheets/sheet{n}.xml"
        self._write_text(part, '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                               f'<worksheet xmlns="{NS_MAIN}" xmlns:r="{NS_REL}"><dimension ref="A1"/>'
                               '<sheetViews><sheetView workbookViewId="0"/></sheetViews>'
                               '<sheetFormatPr defaultRowHeight="15"/><sheetData/>'
                               '<pageMargins left="0.7" right="0.7" top="0.75" bottom="0.75" header="0.3" footer="0.3"/>'
                               '</worksheet>')
        rels_part = self._rels_of(self._wb_part)
        rels = self._read_text(rels_part)
        ids = [int(x) for x in re.findall(r'Id="rId(\d+)"', rels)]
        rid = f"rId{max(ids, default=0) + 1}"
        target = posixpath.relpath(part, posixpath.dirname(self._wb_part))
        rels = rels.replace("</Relationships>",
                            f'<Relationship Id="{rid}" Type="{REL_WORKSHEET}" Target="{target}"/></Relationships>')
        self._write_text(rels_part, rels)
        wb = self._read_text(self._wb_part)
        ids = [int(x) for x in re.findall(r'<sheet\b[^>]*\bsheetId="(\d+)"', wb)]
        state = ' state="hidden"' if hidden else ""
        wb = wb.replace("</sheets>", f'<sheet name="{html.escape(name, quote=True)}" sheetId="{max(ids, default=0) + 1}"'
                                     f'{state} r:id="{rid}"/></sheets>', 1)
        self._write_text(self._wb_part, wb)
        ct = self._read_text("[Content_Types].xml")
        ct = ct.replace("</Types>", f'<Override PartName="/{part}" ContentType="{CT_WORKSHEET}"/></Types>')
        self._write_text("[Content_Types].xml", ct)
        sheet = _Sheet(self, name, part, "hidden" if hidden else None)
        self._sheets[name.upper()] = sheet
        self._order.append(name)
        self._touch()

    def ensure_rr_sheet(self, name: str) -> None:
        """A hidden runner sheet (``_rr_fingerprints``, ``_rr_variables``...), created on first use."""
        if not name.startswith(RR_PREFIX):
            raise WriterError(f"runner sheets must start with {RR_PREFIX!r}: {name!r}")
        if not self.has_sheet(name):
            self.add_sheet(name, hidden=True)

    def rr_sheets(self) -> list[str]:
        return [n for n in self._order if n.upper().startswith(RR_PREFIX.upper())]

    def read_table(self, sheet: str) -> list[list[Any]]:
        return self.rows(sheet)

    def write_table(self, sheet: str, rows: Iterable[Iterable[Any]]) -> None:
        """Replace everything in ``sheet`` with ``rows`` (used for the hidden ``_rr_*`` tables)."""
        s = self._sheet(sheet)
        s.load()
        s.rows = {}
        s.masters = None
        for r, values in enumerate(rows, start=1):
            for c, value in enumerate(values, start=1):
                if value not in (None, ""):
                    s.put(r, c, _cell_xml(_ref(r, c), None, value, {}))
        s.dirty = True
        self._touch()

    # -- structural edits ----------------------------------------------------------------------
    def _touch(self) -> None:
        self._modified = True

    @property
    def modified(self) -> bool:
        return self._modified

    def _apply_rows(self, s: _Sheet, rows: RowMap) -> None:
        if rows.is_identity():
            return
        self._expand_shared(s)                         # shared formulas cannot follow a split group: make them plain
        for other in self._sheets.values():
            if other is not s and self._mentions(other, s.name):
                self._expand_shared(other, referencing=s.name)
        # 1. formulas inside moved rows, then renumber the rows of the sheet itself
        carried = self._map_moved_formulas(s, rows) if rows.moved else set()
        old = s.load()
        new: dict[int, _Row] = {}
        for r, row in old.items():
            target = rows.row(r)
            if target is None:
                continue
            if target != r:
                cells = row.parsed()
                row.cells = {c: _set_attr(_open_tag(x), "r", _ref(target, c)) + x[len(_open_tag(x)):]
                             for c, x in cells.items()}
                row.r = target
                row.dirty = True
            new[target] = row
        s.rows = new
        s.masters = None
        s.dirty = True
        # 2. formulas everywhere that point into this sheet
        for other in self._sheets.values():
            if other is s or self._mentions(other, s.name):
                self._map_sheet_formulas(other, s.name, rows, skip=carried if other is s else set())
        # 3. the sheet's own ranges (merges, validations, formats, links, tables, comments)
        self._map_sheet_ranges(s, rows)
        # 4. defined names
        wb = self._read_text(self._wb_part)
        new_wb = re.sub(r"(<definedName\b[^>]*>)(.*?)(</definedName>)",
                        lambda m: m.group(1) + html.escape(map_formula(html.unescape(m.group(2)), None, s.name, rows),
                                                           quote=False) + m.group(3), wb, flags=re.S)
        if new_wb != wb:
            self._write_text(self._wb_part, new_wb)

    def _map_moved_formulas(self, s: _Sheet, rows: RowMap) -> set[int]:
        """Formulas in rows a move carries: a relative row reference keeps its distance (a running count such as
        ``=COUNTA($B$1:B15)`` still ends just above its own row, as it would after copying), an absolute one (``$D$13``)
        keeps pointing at the same cell wherever that cell went.  Excel's cut/insert would make the running count include
        its own cell (a circular reference).  Returns the new row numbers so the general pass leaves them alone."""
        first, last, disp = rows.moved

        def point(r: int, absolute: bool) -> int | None:
            target = rows.row(r) if absolute else r + disp
            return target if target is not None and 1 <= target <= MAX_ROW else None

        done = set()
        for r, rw in s.load().items():
            if not first <= r <= last:
                continue
            done.add(r + disp)
            cells = rw.parsed()
            for c, xml in list(cells.items()):
                fm = _F_RE.search(xml)
                if not fm or fm.group(2) is None:
                    continue
                text = html.unescape(fm.group(2))
                mapped = map_formula(text, s.name, s.name, rows, point=point)
                ftag = "<f" + fm.group(1)
                if "ref" in _attrs(fm.group(1)):
                    ftag = _set_attr(ftag + ">", "ref", _map_ref_points(_attrs(fm.group(1))["ref"], point)
                                     or _ref(r + disp, c))[:-1]
                if mapped != text or ftag != "<f" + fm.group(1):
                    cells[c] = xml[:fm.start()] + f"{ftag}>{html.escape(mapped, quote=False)}</f>" + xml[fm.end():]
                    rw.dirty = True
        return done

    def _mentions(self, sheet: _Sheet, name: str) -> bool:
        """Cheap pre-check before parsing a sheet: could any of its formulas name sheet ``name``?"""
        if sheet.dirty:
            return True
        low = sheet.xml.lower()
        candidates = {name, name.replace("'", "''")}
        return any(html.escape(c, quote=False).lower() + "!" in low or html.escape(c, quote=False).lower() + "'!" in low
                   for c in candidates)

    def _shared_masters(self, s: _Sheet) -> dict[str, tuple[str, str]]:
        """Shared-formula group id → (formula text, cell holding it)."""
        masters: dict[str, tuple[str, str]] = {}
        for r, rw in s.load().items():
            if 't="shared"' not in rw.xml and not rw.dirty:
                continue
            for c, xml in rw.parsed().items():
                if 't="shared"' in xml:
                    info = self._decode(xml)
                    if info.shared_master is not None and info.formula is not None:
                        masters[info.shared_master] = (info.formula, _ref(r, c))
        return masters

    def _follower_formula(self, s: _Sheet, row: int, col: int, si: str) -> str:
        if s.masters is None:
            s.masters = self._shared_masters(s)
        if si not in s.masters:
            raise WriterError(f"shared formula {si} has no master cell in {s.name!r}")
        text, origin = s.masters[si]
        return Translator("=" + text, origin=origin).translate_formula(_ref(row, col))[1:]

    def _expand_shared(self, s: _Sheet, *, only: set[str] | None = None, referencing: str | None = None) -> None:
        """Turn shared formulas into one plain formula per cell (same meaning; Excel may share them again on save)."""
        rows = s.load()
        masters = self._shared_masters(s)
        if only is not None:
            masters = {k: v for k, v in masters.items() if k in only}
        if referencing is not None:
            masters = {k: v for k, v in masters.items() if _formula_mentions(v[0], referencing)}
        if not masters:
            return
        for r, rw in rows.items():
            cells = rw.parsed()
            for c, xml in list(cells.items()):
                if 't="shared"' not in xml:
                    continue
                fm = _F_RE.search(xml)
                fa = _attrs(fm.group(1))
                si = fa.get("si")
                if si not in masters:
                    continue
                text, origin = masters[si]
                formula = text if _ref(r, c) == origin else \
                    Translator("=" + text, origin=origin).translate_formula(_ref(r, c))[1:]
                ftag = "<f"
                for name, value in fa.items():
                    if name not in ("t", "ref", "si"):
                        ftag += f' {name}="{html.escape(value, quote=True)}"'
                cells[c] = xml[:fm.start()] + f"{ftag}>{html.escape(formula, quote=False)}</f>" + xml[fm.end():]
                rw.dirty = True
                s.dirty = True
        s.masters = None

    def _map_sheet_formulas(self, s: _Sheet, target: str, rows: RowMap, skip: set[int] = frozenset()) -> None:
        for rw in s.load().values():
            if rw.r in skip or ("<f" not in rw.xml and not rw.dirty):
                continue
            cells = rw.parsed()
            for c, xml in list(cells.items()):
                fm = _F_RE.search(xml)
                if not fm or fm.group(2) is None:
                    continue
                text = html.unescape(fm.group(2))
                mapped = map_formula(text, s.name, target, rows)
                ftag = "<f" + fm.group(1)
                fa = _attrs(fm.group(1))
                if "ref" in fa and s.name.upper() == target.upper():
                    new_ref = map_ref_text(fa["ref"], rows)
                    ftag = _set_attr(ftag + ">", "ref", new_ref or _ref(rw.r, c))[:-1]
                if mapped != text or ftag != "<f" + fm.group(1):
                    cells[c] = xml[:fm.start()] + f"{ftag}>{html.escape(mapped, quote=False)}</f>" + xml[fm.end():]
                    rw.dirty = True
                    s.dirty = True
        # formulas in data validations and conditional formats of this sheet
        for attr in ("head", "tail"):
            text = getattr(s, attr)
            new = re.sub(r"(<(formula|formula1|formula2|xm:f)>)(.*?)(</\2>)",
                         lambda m: m.group(1) + html.escape(map_formula(html.unescape(m.group(3)), s.name, target, rows),
                                                            quote=False) + m.group(4), text, flags=re.S)
            if new != text:
                setattr(s, attr, new)
                s.dirty = True

    def _map_sheet_ranges(self, s: _Sheet, rows: RowMap) -> None:
        tail = s.tail

        def drop_or_map(pattern: str, attr: str, text: str) -> str:
            def fix(m: re.Match) -> str:
                mapped = map_sqref(_attrs(m.group(0)).get(attr, ""), rows)
                return _set_attr(m.group(0), attr, mapped) if mapped else ""
            return re.sub(pattern, fix, text, flags=re.S)

        tail = drop_or_map(r"<mergeCell\b[^>]*/>", "ref", tail)
        tail = drop_or_map(r"<hyperlink\b[^>]*/>", "ref", tail)
        tail = re.sub(r"(<dataValidation\b[^>]*>)(.*?</dataValidation>)",
                      lambda m: _map_block(m, "sqref", rows), tail, flags=re.S)
        tail = re.sub(r"(<conditionalFormatting\b[^>]*>)(.*?</conditionalFormatting>)",
                      lambda m: _map_block(m, "sqref", rows), tail, flags=re.S)
        tail = re.sub(r"(<x14:dataValidation\b[^>]*>)(.*?</x14:dataValidation>)",
                      lambda m: _map_xm_block(m, rows), tail, flags=re.S)
        tail = re.sub(r"(<x14:conditionalFormatting\b[^>]*>)(.*?</x14:conditionalFormatting>)",
                      lambda m: _map_xm_block(m, rows), tail, flags=re.S)
        tail = re.sub(r'(<autoFilter\b[^>]*\bref=")([^"]*)(")',
                      lambda m: m.group(1) + (map_sqref(m.group(2), rows) or m.group(2)) + m.group(3), tail)
        tail = _fix_counts(tail, "mergeCells", "mergeCell")
        tail = _fix_counts(tail, "dataValidations", "dataValidation")
        tail = _fix_counts(tail, "x14:dataValidations", "x14:dataValidation")
        tail = re.sub(r"<(mergeCells|hyperlinks|dataValidations)\b[^>]*>\s*</\1>", "", tail)
        tail = re.sub(r"<x14:dataValidations\b[^>]*>\s*</x14:dataValidations>", "", tail)
        head = re.sub(r"(<selection\b[^>]*>)", lambda m: _map_selection(m.group(1), rows), s.head)
        if tail != s.tail or head != s.head:
            s.tail, s.head = tail, head
            s.dirty = True
        for rel in self._relationships(s.part):
            kind = rel.get("Type", "")
            part = self._resolve(s.part, rel.get("Target", ""))
            if rel.get("TargetMode") == "External" or not self._has_part(part):
                continue
            if kind.endswith("/table"):
                xml = self._read_text(part)
                new = re.sub(r'(<(?:table|autoFilter|sortState)\b[^>]*\bref=")([^"]*)(")',
                             lambda m: m.group(1) + (map_sqref(m.group(2), rows) or m.group(2)) + m.group(3), xml)
                if new != xml:
                    self._write_text(part, new)
            elif kind.endswith("/comments"):
                self._map_comments(s, part, rows)

    def _map_comments(self, s: _Sheet, part: str, rows: RowMap) -> None:
        xml = self._read_text(part)
        gone: set[tuple[int, int]] = set()
        moved: dict[tuple[int, int], int] = {}

        def fix(m: re.Match) -> str:
            tag = _open_tag(m.group(0))
            r, c = _split_ref(_attrs(tag)["ref"])
            target = rows.row(r)
            if target is None:
                gone.add((r - 1, c - 1))
                return ""
            if target != r:
                moved[(r - 1, c - 1)] = target - 1
                return _set_attr(tag, "ref", _ref(target, c)) + m.group(0)[len(tag):]
            return m.group(0)

        new = re.sub(r"<comment\b.*?</comment>", fix, xml, flags=re.S)
        if new != xml:
            self._write_text(part, new)
        if not gone and not moved:
            return
        for rel in self._relationships(s.part):
            if rel.get("Type", "").endswith("/vmlDrawing"):
                vml_part = self._resolve(s.part, rel["Target"])
                vml = self._read_bytes(vml_part).decode("utf-8", errors="surrogateescape")
                new_vml = re.sub(r"<((?:\w+:)?)shape\b.*?</\1shape>",
                                 lambda m: _map_note_shape(m.group(0), gone, moved), vml, flags=re.S)
                if new_vml != vml:
                    self._changed_parts[vml_part] = new_vml.encode("utf-8", errors="surrogateescape")

    # -- saving --------------------------------------------------------------------------------
    def _finalize_parts(self) -> None:
        """Serialise dirty sheets; after any change drop the calc chain and ask Excel to recalculate on open."""
        for s in self._sheets.values():
            if s.dirty:
                self._write_text(s.part, s.render())
                s._xml = self._read_text(s.part)
                s.rows = None
                s.masters = None
                s.dirty = False
        if not self._changed_parts:
            return
        for rel in self._relationships(self._wb_part):
            if rel.get("Type", "").endswith("/calcChain"):
                part = self._resolve(self._wb_part, rel["Target"])
                self._removed.add(part)
                rels_part = self._rels_of(self._wb_part)
                self._write_text(rels_part, re.sub(r"<Relationship\b[^>]*/calcChain\"[^>]*/>", "",
                                                   self._read_text(rels_part)))
                self._write_text("[Content_Types].xml", re.sub(
                    r'<Override\b[^>]*PartName="/' + re.escape(part) + r'"[^>]*/>', "",
                    self._read_text("[Content_Types].xml")))
        wb = self._read_text(self._wb_part)
        if "<calcPr" in wb:
            new = re.sub(r"<calcPr\b[^>]*?/?>", lambda m: _set_attr(m.group(0), "fullCalcOnLoad", "1"), wb, count=1)
        else:
            anchor = next(a for a in ("</definedNames>", "</externalReferences>", "</sheets>") if a in wb)
            new = wb.replace(anchor, anchor + '<calcPr fullCalcOnLoad="1"/>', 1)
        if new != wb:
            self._write_text(self._wb_part, new)

    def to_bytes(self) -> bytes:
        self._finalize_parts()
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w") as z:
            seen = set()
            for info in self._zip_infos:
                name = info.filename
                seen.add(name)
                if name in self._removed:
                    continue
                data = self._changed_parts.get(name, self._parts[name])
                zi = zipfile.ZipInfo(name, date_time=info.date_time)
                zi.compress_type = info.compress_type
                zi.external_attr = info.external_attr
                z.writestr(zi, data)
            for name, data in self._changed_parts.items():
                if name not in seen and name not in self._removed:
                    zi = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
                    zi.compress_type = zipfile.ZIP_DEFLATED
                    z.writestr(zi, data)
        return out.getvalue()

    def external_change(self) -> bool:
        """True when the workbook on disk is no longer the one this editor was opened from."""
        return self.stamp is not None and not self.stamp.matches(self.path)

    def save(self, *, backup: bool = True, force: bool = False, when: datetime | None = None) -> SaveResult:
        """Save over the original: refuse while Excel has it open, refuse if it changed on disk (unless ``force``), keep
        the previous version in ``backups/``, then drop the draft."""
        if is_locked(self.path):
            raise WorkbookLocked(_locked_message(self.path))
        if not force and self.external_change():
            raise WorkbookChangedOutside(f"{self.path.name} was changed outside the builder since it was opened. "
                                         "Reload it, or save anyway to overwrite those changes.")
        data = self.to_bytes()
        saved_backup = None
        if backup and self.path.exists():
            saved_backup = backup_path(self.path, when)
            saved_backup.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(self.path, saved_backup)
        _atomic_write(self.path, data)
        self.stamp = FileStamp.of(self.path)
        discard_draft(self.path)
        return SaveResult(self.path, saved_backup)

    def save_as(self, target: str | Path, *, overwrite: bool = False) -> SaveResult:
        """Write to another path (a new workbook); the editor then edits that file."""
        target = Path(target)
        if target.exists() and not overwrite:
            raise FileExistsError(f"{target.name} already exists")
        if is_locked(target):
            raise WorkbookLocked(_locked_message(target))
        _atomic_write(target, self.to_bytes())
        self.path = target
        self.stamp = FileStamp.of(target)
        return SaveResult(target)

    def save_draft(self) -> Path:
        """Autosave: ``.drafts/<name>`` next to the workbook, plus the original's stamp so a reopened draft still
        notices changes made to the workbook meanwhile."""
        target = draft_path(self.path)
        _atomic_write(target, self.to_bytes())
        meta = {"original": self.stamp.to_json() if self.stamp else None, "saved": datetime.now().isoformat()}
        _draft_meta_path(self.path).write_text(json.dumps(meta), encoding="utf-8")
        return target


# ---------------------------------------------------------------------------------------------
# Helpers for the pieces above
# ---------------------------------------------------------------------------------------------
def _formula_mentions(formula: str, sheet: str) -> bool:
    low = formula.lower()
    return any(n.lower() + "!" in low or n.lower() + "'!" in low for n in (sheet, sheet.replace("'", "''")))


def _number_text(value: float | int) -> str:
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, int):
        return str(value)
    text = repr(float(value))
    return text[:-2] if text.endswith(".0") else text


def _excel_serial(value: date | datetime) -> float:
    base = datetime(1899, 12, 30)
    if not isinstance(value, datetime):
        value = datetime(value.year, value.month, value.day)
    delta = value - base
    return delta.days + delta.seconds / 86400 + delta.microseconds / 86400e6


def _cell_xml(ref: str, style: str | None, value: Any, extra: dict[str, str]) -> str:
    attrs = f'r="{ref}"' + (f' s="{style}"' if style not in (None, "0", "") else "")
    for name, v in extra.items():
        attrs += f' {name}="{html.escape(v, quote=True)}"'
    if value is None or value == "":
        return f"<c {attrs}/>"
    if isinstance(value, str) and value.startswith("=") and len(value) > 1:
        formula = add_future_prefixes(value)[1:]
        return f"<c {attrs}><f>{html.escape(formula, quote=False)}</f></c>"
    if isinstance(value, bool):
        return f'<c {attrs} t="b"><v>{"1" if value else "0"}</v></c>'
    if isinstance(value, (int, float)):
        return f"<c {attrs}><v>{_number_text(value)}</v></c>"
    if isinstance(value, (date, datetime)):
        return f"<c {attrs}><v>{_number_text(_excel_serial(value))}</v></c>"
    text = str(value)
    space = ' xml:space="preserve"' if text != text.strip() or "\n" in text else ""
    return f'<c {attrs} t="inlineStr"><is><t{space}>{_escape_text(text)}</t></is></c>'


def _map_block(m: re.Match, attr: str, rows: RowMap) -> str:
    tag = m.group(1)
    value = _attrs(tag).get(attr, "")
    mapped = map_sqref(value, rows)
    if not mapped:
        return ""
    return _set_attr(tag, attr, mapped) + m.group(2)


def _map_xm_block(m: re.Match, rows: RowMap) -> str:
    block = m.group(0)
    sq = re.search(r"<xm:sqref>(.*?)</xm:sqref>", block)
    if not sq:
        return block
    mapped = map_sqref(sq.group(1), rows)
    if not mapped:
        return ""
    return block[:sq.start(1)] + mapped + block[sq.end(1):]


def _fix_counts(text: str, outer: str, inner: str) -> str:
    def fix(m: re.Match) -> str:
        if "count" not in _attrs(m.group(1)):
            return m.group(0)
        n = len(re.findall(rf"<{inner}\b", m.group(2)))
        return _set_attr(m.group(1), "count", str(n)) + m.group(2)
    return re.sub(rf"(<{outer}\b[^>]*>)(.*?</{outer}>)", fix, text, flags=re.S)


def _map_selection(tag: str, rows: RowMap) -> str:
    a = _attrs(tag)
    if "sqref" in a:
        mapped = map_sqref(a["sqref"], rows)
        tag = _set_attr(tag, "sqref", mapped or "A1")
    if "activeCell" in a:
        mapped = map_ref_text(a["activeCell"], rows)
        tag = _set_attr(tag, "activeCell", mapped or "A1")
    return tag


def _map_note_shape(shape: str, gone: set[tuple[int, int]], moved: dict[tuple[int, int], int]) -> str:
    """A comment's note box in the VML drawing: ``x:Row``/``x:Column`` are 0-based; the anchor holds rows too.
    Excel writes the ``x:`` prefix, openpyxl ``ns2:``: any prefix is accepted."""
    rm = re.search(r"<((?:\w+:)?)Row>(\d+)</\1Row>", shape)
    cm = re.search(r"<((?:\w+:)?)Column>(\d+)</\1Column>", shape)
    if not rm or not cm:
        return shape
    key = (int(rm.group(2)), int(cm.group(2)))
    if key in gone:
        return ""
    if key not in moved:
        return shape
    delta = moved[key] - key[0]
    shape = shape[:rm.start(2)] + str(moved[key]) + shape[rm.end(2):]

    def anchor(m: re.Match) -> str:
        nums = [x.strip() for x in m.group(2).split(",")]
        if len(nums) == 8:
            nums[2] = str(int(nums[2]) + delta)
            nums[6] = str(int(nums[6]) + delta)
            lead = re.match(r"\s*", m.group(2)).group(0)
            return m.group(1) + lead + ", ".join(nums) + m.group(3)
        return m.group(0)

    return re.sub(r"(<(?:\w+:)?Anchor>)(.*?)(</(?:\w+:)?Anchor>)", anchor, shape, flags=re.S)


def iter_sheet_cells(editor: WorkbookEditor, sheet: str) -> Iterator[tuple[int, int, Any]]:
    """Every stored cell of ``sheet`` as (row, column, ``get`` value), in row order."""
    s = editor._sheet(sheet)
    for r in sorted(s.load()):
        for c in sorted(s.load()[r].parsed()):
            yield r, c, editor.get(sheet, r, c)
