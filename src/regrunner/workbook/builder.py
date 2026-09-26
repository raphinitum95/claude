"""The Workbook Builder's model of a workbook: what the Build tab shows and edits (dev/plan/CONTRACT.md sections 1-3).

Why a separate model: ``workbook/model.py`` evaluates a workbook for *running* it (formulas, one data row, token substitution); the builder
needs the workbook *as written* (formulas as text, every data row, legacy rows kept as they are) and must be rebuilt after every edit in well
under a second. So this module reads cells straight from ``WorkbookEditor`` (P01) and never evaluates a formula.

Pieces:

* ``build_model(editor)``: workbook → tests → blocks → steps, variables (set-by / used-by, needs / provides), environments, fingerprints.
  The problems (``lint.builder_problems``) run on the JSON it returns, so the same rules serve the UI and any later CLI.
* ``apply_ops(editor, ops)``: the small edits the UI sends (insert / move / delete / update steps, blocks, variables, environments...).
* ``BuildDocument``: one open workbook on the server: undo/redo (one snapshot per batch of ops), autosaved draft, save with backup, history,
  diff against a backup, restore. ``BuildStore`` keeps one per workbook.

Nothing here writes a cell unless an op asks for it: opening a workbook, building the model and saving without edits leaves the file untouched.
"""
from __future__ import annotations

import difflib
import json
import re
import threading
import zipfile
import io
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter

from .writer import (RR_PREFIX, RowMap, SaveResult, WorkbookEditor, WriterError, backups_dir, discard_draft, draft_path, has_draft,
                     is_locked)

# ---------------------------------------------------------------------------------------------
# File format (CONTRACT.md section 1)
# ---------------------------------------------------------------------------------------------
BLOCK_COLUMN = "BLOCK"
SIDE_EFFECTS_COLUMN = "SIDE_EFFECTS"
BACKUP_LOCATORS_COLUMN = "BACKUP_LOCATORS"
STEP_NAME_AUTO_COLUMN = "STEP_NAME_AUTO"
NEW_STEP_COLUMNS = (BLOCK_COLUMN, SIDE_EFFECTS_COLUMN, BACKUP_LOCATORS_COLUMN, STEP_NAME_AUTO_COLUMN)

RR_VARIABLES = "_rr_variables"
RR_ENVIRONMENTS = "_rr_environments"
RR_FINGERPRINTS = "_rr_fingerprints"
RR_SCENARIOS = "_rr_scenarios"
VARIABLE_HEADERS = ("Token", "Label", "Secret", "EnvSpecific", "Notes")
FINGERPRINT_HEADERS = ("Name", "UrlContains", "Landmark", "LandmarkText", "Notes")
ENVIRONMENT_FIXED_HEADERS = ("Variable", "Required", "Secret")
PRODUCTION_ROW = "#PRODUCTION"
DEFAULT_ENVIRONMENTS = ("QA", "UAT", "PROD")

# the 26 columns every keyword sheet of the real workbooks has (a new test gets them)
STANDARD_COLUMNS = ("blnExecute", "Step_Number", "Step_Name", "Status", "Error_Check", "breakpoint", "Test_Case", "Page", "FindBy",
                    "FindBy_Value", "Index", "Method", "Value", "Output_Property", "Ignore_not_existing_object", "isDisabled", "Timeout",
                    "Expected_Value", "Output_Value", "Exact_Match", "Contains", "ITERATION_NO", "Notes", "Start_Time", "End_Time",
                    "Duration")
UI_COLUMNS = ("METHOD", "PAGE", "FINDBY", "FINDBY_VALUE")
RESULT_COLUMNS = {"STATUS", "ERROR_CHECK", "START_TIME", "END_TIME", "DURATION"}
# columns whose cell may name a variable that the step *reads* (legacy substitutes any column; these are the ones that matter to a person)
USE_COLUMNS = ("VALUE", "FINDBY_VALUE", "EXPECTED_VALUE", "PAGE", "LOCATOR", "INDEX", "TIMEOUT", "OUTPUT_PROPERTY", BACKUP_LOCATORS_COLUMN)

TOKEN_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")
INLINE_RE = re.compile(r"\{([A-Za-z_][A-Za-z0-9_]*)\}")
SECRET_RE = re.compile(r"\{SECRET:([A-Za-z_][A-Za-z0-9_]*)\}", re.I)
UNMAPPED_RE = re.compile(r"\{\?([A-Za-z_][A-Za-z0-9_]*)\}")

FLOW_KEYWORDS = ("IF", "ELSE", "END_IF", "ITERATION_START", "ITERATION_END")
CHECK_KEYWORDS = ("CHECK_VALUE", "CHECK_REGEX", "CHECK_COMPARE", "CHECK_COUNT", "CHECK_ENABLED", "CHECK_CHECKED", "CHECK_SELECTED",
                  "CHECK_DATE_FORMAT")
NEW_KEYWORDS = ("SET_VARIABLE", *FLOW_KEYWORDS, "CALL_TEST", "ASSERT_PAGE", "WAIT_UNTIL", "DISMISS_IF_SHOWN", "PICK_DATE",
                "CHOOSE_SUGGESTION", *CHECK_KEYWORDS)
NEW_ELEMENT_KEYWORDS = {"WAIT_UNTIL", "DISMISS_IF_SHOWN", "PICK_DATE", "CHOOSE_SUGGESTION", *CHECK_KEYWORDS}
# rows the builder shows as locked cards (Q28): the new runner cannot do them, or they only made sense on the old Windows runner
LEGACY_METHODS = {
    "GO_TO_ROW": "GO_TO_ROW jumped to another row of the sheet on the old runner. The builder keeps it as it is; use IF on a variable instead.",
    "SNAGIT_SCREENSHOT": "SNAGIT_SCREENSHOT used the SnagIt desktop app on the old Windows runner. The new runner skips it; use a normal screenshot.",
    "DB_CONNECT": "Database steps ran on the old runner only.", "DB_DISCONNECT": "Database steps ran on the old runner only.",
    "DB_QUERY": "Database steps ran on the old runner only.", "DB_SAVE_JSON_RESULT": "Database steps ran on the old runner only.",
    "SEND_EMAIL": "SEND_EMAIL sent Outlook mail from the old runner; the new runner does not send mail.",
    "FILE_COMPARE": "FILE_COMPARE compared files on the old Windows runner only.",
    "RUN_SCRIPT": "RUN_SCRIPT ran a local script on the old Windows runner only.",
    "UPDATE_SYSTEM_TIMEZONE": "UPDATE_SYSTEM_TIMEZONE changed the Windows clock on the old runner only.",
    "BROKEN_LINK_CHECK": "BROKEN_LINK_CHECK ran on the old runner only.", "BROKEN_LINK_CHECK_CRAWL": "BROKEN_LINK_CHECK_CRAWL ran on the old runner only.",
    "BS_CAPABILITIES": "BS_CAPABILITIES configured BrowserStack on the old runner only.",
    "FOCUSWINDOW": "FOCUSWINDOW used the Windows desktop on the old runner only.",
    "GET_MOUSE_POS": "GET_MOUSE_POS used the Windows mouse on the old runner only.",
    "DRAGANDDROP": "DRAGANDDROP is not supported by the new runner yet.",
    "JSON_READ": "JSON_READ is not supported by the new runner yet; API tests read responses through their InputOutput sheet.",
}
FORMULA_FLAG_REASON = "blnExecute is a formula here: {text}. The builder keeps it byte-for-byte; edit it in the Excel grid."
NOT_PAGE_NAMES = {"", "CHROME", "FIREFOX", "EDGE", "MSEDGE", "IE", "SAFARI", "WEBKIT", "CHROMIUM", "APPLICATION", "BROWSER", "WEB", "NONE"}
_TRUE = {"1", "Y", "YES", "TRUE"}
_FALSE = {"0", "N", "NO", "FALSE"}
SIDE_EFFECT_WORDS = re.compile(r"\b(purchase|pay( now)?|place order|submit order|buy( now)?|confirm (and|&) pay|complete (the )?(purchase|order|payment)|"
                               r"make payment)\b", re.I)

_NAV = {"OPEN", "NAVIGATE", "DRIVER_GET", "BACK", "REFRESH", "MAXIMIZE", "CLOSE", "QUIT", "SWITCHTOFRAME", "SWITCHTODEFAULT", "SWITCHTOWINDOW",
        "SWITCHTOMAINWINDOW", "SET_PAGE_LOAD_TIMEOUT"}
_ACT = {"CLICK", "JS_CLICK", "DOUBLECLICK", "CONTEXTCLICK", "MOVETOELEMENT", "HIGHLIGHT", "MOUSE_SCROLL", "MOUSE_MOVE", "MOUSE_CLICK",
        "MOUSE_LEFT_CLICK", "MOUSE_DOUBLE_CLICK", "MOUSE_DOUBLE_LEFT_CLICK", "MOUSE_RIGHT_CLICK", "ALERT_OK", "ALERT_CANCEL", "EXECUTESCRIPT",
        "DISMISS_IF_SHOWN", "CLEAR_CLIPBOARD"}
_INPUT = {"SET", "INPUT", "WRITE", "TYPE", "CLEAR", "TICK", "UNTICK", "SELECT", "SENDKEYS", "SENDKEY", "SEND_KEY", "SEND_KEYS", "SPECIALKEY",
          "CLICK_SENDKEYS", "CUSTOMINPUTDATA", "PICK_DATE", "CHOOSE_SUGGESTION", "ASK_USER", "PROMPT", "ASK", "GET_GOOGLE_TOKEN"}
_CHECK = {"EXIST", "NOT_EXIST", 'GETATTRIBUTE("VALUE")', "ASSERT_PAGE", *CHECK_KEYWORDS}
_SAVE = {"SCREENSHOT", "PNG_SCREENSHOT", "BASE64_SCREENSHOT", "SET_VARIABLE", "ALERT_TEXT_OUT", "GET_CURRENT_URL"}
_WAIT = {"WAIT", "WAIT_UNTIL"}
_FLOW = {*FLOW_KEYWORDS, "BREAK", "STOP"}

FIELD_COLUMNS = {
    "name": "Step_Name", "method": "Method", "page": "Page", "findBy": "FindBy", "locator": "FindBy_Value", "index": "Index",
    "locatorName": "Locator", "value": "Value", "expected": "Expected_Value", "saveAs": "Output_Value", "output": "Output_Value",
    "outputProperty": "Output_Property", "onFail": "Ignore_not_existing_object", "timeout": "Timeout", "enabled": "blnExecute",
    "sideEffects": SIDE_EFFECTS_COLUMN, "backups": BACKUP_LOCATORS_COLUMN, "block": BLOCK_COLUMN, "notes": "Notes",
    "nameAuto": STEP_NAME_AUTO_COLUMN, "match": None,
}
BULK_FIELDS = {"timeout", "enabled", "onFail", "page", "block", "sideEffects"}


class BuildError(ValueError):
    """An op or request the builder refuses; ``kind`` and ``status`` go to the HTTP answer."""

    def __init__(self, message: str, kind: str = "op", status: int = 422, **extra: Any):
        super().__init__(message)
        self.kind, self.status, self.extra = kind, status, extra


# ---------------------------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------------------------
def text_of(value: Any) -> str:
    """A cell as the person reads it: numbers without a trailing .0, booleans as TRUE/FALSE."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=") and len(value) > 1


def is_true(value: Any) -> bool:
    return text_of(value).strip().upper() in _TRUE


def humanize(token: str) -> str:
    """``DT_FirstName_IN`` → ``First name``: what a person calls a variable nobody gave a label (Q37)."""
    t = re.sub(r"^(DT|TC|RR)_", "", str(token).strip(), flags=re.I)
    t = re.sub(r"_(IN|OUT|IH|OH)$", "", t, flags=re.I)
    t = re.sub(r"([a-z0-9])([A-Z])", r"\1 \2", t)
    t = re.sub(r"[_\-.]+", " ", t)
    t = re.sub(r"\s+", " ", t).strip().lower()
    return (t[:1].upper() + t[1:]) if t else str(token)


def make_token(label: str) -> str:
    """A friendly label → an UPPER_SNAKE token (``Traveler first name`` → ``TRAVELER_FIRST_NAME``)."""
    t = re.sub(r"[^A-Za-z0-9]+", "_", str(label)).strip("_").upper()
    if not t:
        raise BuildError("A variable needs a name made of letters or digits.")
    return t if not t[0].isdigit() else "V_" + t


@dataclass
class Condition:
    left: str
    op: str             # filled | empty | = | != | contains | > | >= | < | <=
    right: str = ""


_COND_RE = re.compile(r"^\s*(\{[A-Za-z_][A-Za-z0-9_]*\})\s*(?:(is\s+filled|is\s+empty)|(!=|>=|<=|=|>|<|contains)\s*(.*?))\s*$", re.I)


def parse_condition(text: str) -> Condition:
    """The IF grammar of CONTRACT.md 1.3: ``{A} is filled``, ``{A} = x``, ``{A} contains x``, ``{A} > 3``...  Raises ``ValueError``."""
    m = _COND_RE.match(str(text or ""))
    if not m:
        raise ValueError(f"not a condition the builder understands: {text!r} (write e.g. {{PROMO_CODE}} is filled, or {{PLAN}} = Max)")
    left = m.group(1)[1:-1]
    if m.group(2):
        return Condition(left, "filled" if "filled" in m.group(2).lower() else "empty")
    op, right = m.group(3).lower(), m.group(4).strip()
    if op in (">", ">=", "<", "<=") and not INLINE_RE.fullmatch(right):
        try:
            float(right)
        except ValueError:
            raise ValueError(f"{op} needs a number on the right: {text!r}") from None
    return Condition(left, op, right)


def legacy_reason(method: str, flag: Any) -> str:
    if is_formula(flag):
        return FORMULA_FLAG_REASON.format(text=flag)
    return LEGACY_METHODS.get(method, "")


def _element_methods() -> set[str]:
    try:
        from ..engine.actions import REGISTRY
        return {name for name, spec in REGISTRY.items() if spec.element} | NEW_ELEMENT_KEYWORDS
    except Exception:                                          # the model must load even if the engine cannot be imported
        return set(NEW_ELEMENT_KEYWORDS)


def _known_methods() -> set[str]:
    try:
        from ..engine.actions import REGISTRY
        return set(REGISTRY) | set(NEW_KEYWORDS) | {"BREAK", "STOP"} | set(LEGACY_METHODS)
    except Exception:
        return set(NEW_KEYWORDS) | set(LEGACY_METHODS)


ELEMENT_METHODS = _element_methods()
KEY_METHODS = {"SENDKEYS", "SENDKEY", "SEND_KEY", "SEND_KEYS", "SPECIALKEY", "CLICK_SENDKEYS", "CUSTOMINPUTDATA"}
try:
    from ..engine.keys import SPECIAL as _SPECIAL_KEYS
    KEY_NAMES = set(_SPECIAL_KEYS)
except Exception:
    KEY_NAMES = {"TAB", "ENTER", "ESC", "DOWN", "UP", "LEFT", "RIGHT", "BACKSPACE", "DELETE", "HOME", "END", "SPACE"}
KNOWN_METHODS = _known_methods()


# ---------------------------------------------------------------------------------------------
# Plain-words targets and automatic step names (Q38)
# ---------------------------------------------------------------------------------------------
_TEXT_PATTERNS = (
    re.compile(r"contains\(\s*(?:\.|text\(\)|normalize-space\([^)]*\))\s*,\s*['\"]([^'\"]{1,40})['\"]"),
    re.compile(r"(?:text\(\)|normalize-space\([^)]*\)|\.)\s*=\s*['\"]([^'\"]{1,40})['\"]"),
    re.compile(r"@(?:placeholder|aria-label|title|alt)\s*=\s*['\"]([^'\"]{1,40})['\"]"),
    re.compile(r"^text\s*=\s*['\"]?([^'\"]{1,40})['\"]?$"),
    re.compile(r":has-text\(['\"]([^'\"]{1,40})['\"]\)"),
)
_NAME_PATTERNS = (
    re.compile(r"@(?:name|id|data-testid|data-test|for|formcontrolname)\s*=\s*['\"]([^'\"]{1,60})['\"]"),
    re.compile(r"contains\(\s*@(?:name|id|class)\s*,\s*['\"]([^'\"]{1,60})['\"]"),
    re.compile(r"\[(?:name|id|data-testid)\s*=\s*['\"]?([^'\"\]]{1,60})['\"]?\]"),
    re.compile(r"#([A-Za-z][\w-]{0,60})"),
)
_TAG_NOUNS = {"button": "button", "a": "link", "input": "field", "select": "list", "textarea": "box", "label": "label", "img": "image",
              "iframe": "frame", "h1": "heading", "h2": "heading", "h3": "heading", "li": "item", "option": "option", "table": "table"}


def describe_target(locator: str, name: str = "", variables: set[str] | None = None) -> str:
    """Plain words for what a step acts on: ``"Get Started" button``, ``last name field``, ``{DT_PLAN}``."""
    if name:
        return humanize(name).lower()
    loc = str(locator or "").strip()
    if not loc:
        return ""
    if variables and loc.upper() in variables:
        return "{" + loc + "}"
    tags = re.findall(r"//([a-zA-Z][a-zA-Z0-9]*)", loc) or re.findall(r"^([a-zA-Z][a-zA-Z0-9]*)(?:[.#\[:]|$)", loc)
    noun = _TAG_NOUNS.get(tags[-1].lower(), "") if tags else ""
    for pattern in _TEXT_PATTERNS:
        m = pattern.search(loc)
        if m:
            return f'"{m.group(1).strip()}" {noun}'.strip()
    for pattern in _NAME_PATTERNS:
        m = pattern.search(loc)
        if m:
            words = humanize(m.group(1)).lower()
            return f"{words} {noun}".strip() if noun and not words.endswith(noun) else words
    return noun or "the element"


def _show(value: str, variables: set[str]) -> str:
    v = str(value or "").strip()
    if not v:
        return ""
    if v.upper() in variables:
        return "{" + v + "}"
    m = SECRET_RE.fullmatch(v)
    if m:
        return "{" + m.group(1) + "}"
    if INLINE_RE.fullmatch(v) or is_formula(v):
        return v
    return f'"{v if len(v) <= 40 else v[:37] + "..."}"'


def auto_name(method: str, *, target: str = "", value: str = "", expected: str = "", save_as: str = "", match: str = "",
              output_property: str = "", variables: set[str] | None = None) -> str:
    """The automatic sentence for a step (Q38): "Type {FIRST_NAME} into first name field"."""
    variables = variables or set()
    m = method.upper().strip()
    t = target or "the element"
    v = _show(value, variables)
    e = _show(expected, variables)
    s = "{" + save_as + "}" if save_as else ""
    prop = {"": "text", "INNERTEXT": "text", "TEXT": "text", "VALUE": "value"}.get(output_property.upper(), output_property.lower())
    how = "contains" if match == "contains" else "is"
    names = {
        "OPEN": f"Open {v}".strip(), "NAVIGATE": f"Go to {v}".strip(), "DRIVER_GET": f"Go to {v}".strip(), "BACK": "Go back",
        "REFRESH": "Reload the page", "MAXIMIZE": "Maximise the window", "CLOSE": "Close the window", "QUIT": "Close the browser",
        "SWITCHTOFRAME": f"Switch into the frame {v or target}".strip(), "SWITCHTODEFAULT": "Leave the frame",
        "SWITCHTOWINDOW": "Switch to the new window", "SWITCHTOMAINWINDOW": "Back to the main window",
        "CLICK": f"Click {t}", "JS_CLICK": f"Click {t}", "DOUBLECLICK": f"Double-click {t}", "CONTEXTCLICK": f"Right-click {t}",
        "MOVETOELEMENT": f"Scroll to {t}", "HIGHLIGHT": f"Highlight {t}", "SET": f"Type {v} into {t}", "INPUT": f"Type {v} into {t}",
        "WRITE": f"Add {v} to {t}", "TYPE": f"Add {v} to {t}", "CLEAR": f"Clear {t}", "TICK": f"Tick {t}", "UNTICK": f"Untick {t}",
        "SELECT": f"Choose {v} in {t}", "SENDKEYS": f"Press {v} in {t}", "SEND_KEYS": f"Press {v} in {t}", "SPECIALKEY": f"Press {v} in {t}",
        "EXIST": f"Check {t} shows", 'GETATTRIBUTE("VALUE")': f"Check {t} shows", "NOT_EXIST": f"Check {t} is gone",
        "WAIT": f"Wait {text_of(value)} s", "SCREENSHOT": "Take a screenshot", "PNG_SCREENSHOT": "Take a screenshot",
        "SET_VARIABLE": f"Set {s or '{?}'} to {v}", "CALL_TEST": f"Run {text_of(value) or '?'}, then carry on",
        "IF": f"If {text_of(value)}", "ELSE": "Otherwise", "END_IF": "End of if", "ITERATION_START": f"For each row of {text_of(value) or 'the data'}",
        "ITERATION_END": "End of the loop", "ASSERT_PAGE": f"Arrived at: {text_of(value)}", "DISMISS_IF_SHOWN": f"If {t} shows, close it",
        "PICK_DATE": f"Pick date {v} in {t}", "CHOOSE_SUGGESTION": f"Choose {e or v} from the suggestions of {t}",
        "CHECK_VALUE": f"Check {t} value {how} {e}", "CHECK_REGEX": f"Check {t} matches {e}", "CHECK_COMPARE": f"Check {t} is {output_property.lower() or 'eq'} {e}",
        "CHECK_COUNT": f"Check there are {e} {t}", "CHECK_ENABLED": f"Check {t} is {'enabled' if is_true(expected) or not expected else 'disabled'}",
        "CHECK_CHECKED": f"Check {t} is {'ticked' if is_true(expected) or not expected else 'not ticked'}",
        "CHECK_SELECTED": f"Check {e} is chosen in {t}", "CHECK_DATE_FORMAT": f"Check {t} is a date like {e}",
        "ASK_USER": f"Ask a person for {s or v or 'a value'}", "PROMPT": f"Ask a person for {s or v or 'a value'}",
        "GET_GOOGLE_TOKEN": "Get the one-time sign-in code", "ALERT_OK": "Accept the alert", "ALERT_CANCEL": "Dismiss the alert",
        "ALERT_TEXT_OUT": f"Save the alert text as {s}" if s else "Read the alert text", "GET_CURRENT_URL": f"Save the page address as {s}" if s else "Read the page address",
        "EXECUTESCRIPT": "Run a script on the page",
    }
    if m == "WAIT_UNTIL":
        kind = output_property.upper()
        return f"Wait until {t} is gone" if kind == "GONE" else (f"Wait until {t} text {how} {e}" if kind == "TEXT" else f"Wait until {t} shows")
    if m == "OUTPUT":
        if save_as:
            return f"Save {t} {prop} as {s}"
        if match:
            return f"Check {t} {prop} {how} {e}"
        return f"Read {t} {prop}"
    if m in names:
        return re.sub(r"\s+", " ", names[m]).strip()
    return f"{m.replace('_', ' ').capitalize()} {target}".strip() if m else ""


# ---------------------------------------------------------------------------------------------
# Reading the workbook
# ---------------------------------------------------------------------------------------------
class _Grid:
    """One sheet's rows as written, with UPPER header → column (last one wins, like the reader)."""

    def __init__(self, editor: WorkbookEditor, sheet: str):
        self.name = sheet
        self.rows = editor.rows(sheet)
        self.header_names = [text_of(v).strip() for v in (self.rows[0] if self.rows else [])]
        self.cols: dict[str, int] = {}
        for i, name in enumerate(self.header_names, start=1):
            if name:
                self.cols[name.upper()] = i

    def get(self, row: int, header: str) -> Any:
        col = self.cols.get(header)
        if not col or row > len(self.rows):
            return None
        values = self.rows[row - 1]
        return values[col - 1] if col <= len(values) else None

    def text(self, row: int, header: str) -> str:
        v = self.get(row, header)
        return "" if is_formula(v) else text_of(v).strip()

    def raw(self, row: int, header: str) -> str:
        return text_of(self.get(row, header)).strip()

    @property
    def max_row(self) -> int:
        return len(self.rows)

    def row_has_content(self, row: int) -> bool:
        return any(v not in (None, "") for v in self.rows[row - 1]) if row <= len(self.rows) else False


def _sheet_key(editor: WorkbookEditor, name: str) -> str | None:
    for s in editor.sheet_names():
        if s.upper() == str(name).upper().strip():
            return s
    return None


def _is_ui_sheet(cols: dict[str, int]) -> bool:
    return all(c in cols for c in UI_COLUMNS)


def _is_api_sheet(cols: dict[str, int]) -> bool:
    return "BLNEXECUTE" in cols and ("WEBSERVICE_URL" in cols or "ENVIRONMENT_PARAMETER" in cols)


def _data_rows(grid: _Grid | None, label_headers: Iterable[str] = ("NOTES", "SCENARIO", "TC_NAME", "TESTCONDITION")) -> list[dict]:
    if grid is None:
        return []
    out = []
    labels = [h for h in label_headers if h in grid.cols]
    for r in range(2, grid.max_row + 1):
        if not grid.row_has_content(r):
            continue
        flag = grid.get(r, "BLNEXECUTE")
        enabled = True if "BLNEXECUTE" not in grid.cols else (None if is_formula(flag) else is_true(flag))
        label = next((grid.text(r, h) for h in labels if grid.text(r, h)), "")
        out.append({"row": r, "enabled": enabled, "label": label})
    return out


def _last_results(runs_dir: Path | None, workbook_name: str) -> tuple[dict[tuple[str, int], dict], dict[str, dict]]:
    """Latest result of every step of this workbook, ``{(SHEET, row): {...}}``, plus the latest run of each test (Q32/33).
    Reads the newest 20 ``runs/*/results.json`` of this workbook; a later run that passes a step clears its failure."""
    steps: dict[tuple[str, int], dict] = {}
    tests: dict[str, dict] = {}
    if runs_dir is None or not Path(runs_dir).is_dir():
        return steps, tests
    files = []
    for folder in Path(runs_dir).iterdir():
        f = folder / "results.json"
        if f.is_file():
            files.append((f.stat().st_mtime_ns, f))
    files.sort(reverse=True)
    seen = 0
    for _, f in files:
        data = _read_results(f)
        if not data or Path(str(data.get("workbook", ""))).name.lower() != workbook_name.lower():
            continue
        seen += 1
        run_id, when = str(data.get("run_id", f.parent.name)), str(data.get("started_at", ""))
        for t in data.get("tests") or []:
            sheet = str(t.get("sheet") or t.get("id") or "").split("#")[0].upper()
            tests.setdefault(sheet, {"runId": run_id, "status": t.get("status", ""), "when": when})
            for s in t.get("steps") or []:
                try:
                    row = int(s.get("row"))
                except (TypeError, ValueError):
                    continue
                if (sheet, row) in steps:
                    continue
                error = str(s.get("error") or "")
                steps[(sheet, row)] = {"status": s.get("status", ""), "error": error[:300], "runId": run_id, "when": when,
                                       "locatorMiss": s.get("status") == "FAILED" and bool(re.search(r"not (be )?found|no element|waiting for locator", error, re.I))}
        if seen >= 20:
            break
    return steps, tests


_RESULTS_CACHE: dict[str, tuple[int, dict | None]] = {}


def _read_results(path: Path) -> dict | None:
    key, stamp = str(path), path.stat().st_mtime_ns
    hit = _RESULTS_CACHE.get(key)
    if hit and hit[0] == stamp:
        return hit[1]
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        data = None
    _RESULTS_CACHE[key] = (stamp, data)
    return data


def _rr_table(editor: WorkbookEditor, sheet: str) -> tuple[dict[str, int], list[list[Any]]]:
    """A hidden runner table: (UPPER header → 0-based index, data rows)."""
    name = _sheet_key(editor, sheet)
    if name is None:
        return {}, []
    rows = editor.rows(name)
    if not rows:
        return {}, []
    header = {text_of(v).strip().upper(): i for i, v in enumerate(rows[0]) if text_of(v).strip()}
    return header, [r for r in rows[1:] if any(v not in (None, "") for v in r)]


def _cell(row: list[Any], header: dict[str, int], name: str) -> str:
    i = header.get(name.upper())
    return text_of(row[i]).strip() if i is not None and i < len(row) else ""


def read_variable_labels(editor: WorkbookEditor) -> dict[str, dict]:
    header, rows = _rr_table(editor, RR_VARIABLES)
    out = {}
    for r in rows:
        token = _cell(r, header, "Token")
        if token:
            out[token.upper()] = {"token": token, "label": _cell(r, header, "Label"), "secret": is_true(_cell(r, header, "Secret")),
                                  "envSpecific": is_true(_cell(r, header, "EnvSpecific")), "notes": _cell(r, header, "Notes")}
    return out


def read_environments(editor: WorkbookEditor) -> dict:
    header, rows = _rr_table(editor, RR_ENVIRONMENTS)
    if header:
        names = [text_of(v).strip() for v in editor.rows(_sheet_key(editor, RR_ENVIRONMENTS))[0]
                 if text_of(v).strip() and text_of(v).strip().upper() not in {h.upper() for h in ENVIRONMENT_FIXED_HEADERS}]
        production: list[str] = []
        table = []
        for r in rows:
            var = _cell(r, header, "Variable")
            if not var:
                continue
            if var.upper() == PRODUCTION_ROW:
                production = [n for n in names if is_true(_cell(r, header, n))]
                continue
            table.append({"variable": var, "required": is_true(_cell(r, header, "Required")), "secret": is_true(_cell(r, header, "Secret")),
                          "values": {n: _cell(r, header, n) for n in names}})
        if not any(_cell(r, header, "Variable").upper() == PRODUCTION_ROW for r in rows):
            production = [n for n in names if n.upper() in ("PROD", "PRODUCTION")]
        return {"source": "rr", "names": names, "production": production, "rows": table}
    legacy = _sheet_key(editor, "Environments")
    if legacy is not None:
        grid = _Grid(editor, legacy)
        names: list[str] = []
        table_by: dict[str, dict] = {}
        for r in range(2, grid.max_row + 1):
            env, param = grid.raw(r, "ENVIRONMENT"), grid.raw(r, "PARAMETER")
            if not env or not param:
                continue
            if env not in names:
                names.append(env)
            table_by.setdefault(param.upper(), {"variable": param, "required": False, "secret": False, "values": {}})["values"][env] = grid.raw(r, "VALUE")
        for row in table_by.values():
            row["values"] = {n: row["values"].get(n, "") for n in names}
        return {"source": "legacy", "names": names, "production": [n for n in names if n.upper() in ("PROD", "PRODUCTION")],
                "rows": list(table_by.values())}
    return {"source": "none", "names": list(DEFAULT_ENVIRONMENTS), "production": ["PROD"], "rows": []}


def read_fingerprints(editor: WorkbookEditor) -> list[dict]:
    header, rows = _rr_table(editor, RR_FINGERPRINTS)
    out = []
    for r in rows:
        name = _cell(r, header, "Name")
        if name:
            out.append({"name": name, "urlContains": _cell(r, header, "UrlContains"), "landmark": _cell(r, header, "Landmark"),
                        "landmarkText": _cell(r, header, "LandmarkText"), "notes": _cell(r, header, "Notes"), "usedBy": []})
    return out


@dataclass
class _DataSheetRow:
    row: int
    sheet: str
    enabled: bool
    param_sheet: str
    comment: str
    tags: list[str]


def _datasheets(editor: WorkbookEditor) -> tuple[str | None, list[_DataSheetRow]]:
    name = _sheet_key(editor, "DataSheets")
    if name is None:
        return None, []
    grid = _Grid(editor, name)
    out = []
    for r in range(2, grid.max_row + 1):
        sheet = grid.text(r, "SHEETSTOEXECUTE") if "SHEETSTOEXECUTE" in grid.cols else text_of(grid.rows[r - 1][0] if grid.rows[r - 1] else "").strip()
        if not sheet:
            continue
        tags = [t.strip() for t in re.split(r"[;,]", grid.text(r, "TAGS")) if t.strip()]
        out.append(_DataSheetRow(r, sheet, is_true(grid.get(r, "BLNEXECUTE")), grid.text(r, "PARAMETERSHEET"), grid.text(r, "COMMENTS"), tags))
    return name, out


# ---------------------------------------------------------------------------------------------
# One keyword sheet → steps and blocks
# ---------------------------------------------------------------------------------------------
def _is_flag_text(text: str) -> bool:
    return text.upper() in _TRUE | _FALSE


def _page_name(page: str, variables: set[str]) -> str:
    p = page.strip()
    if p.upper() in NOT_PAGE_NAMES or p.upper() in variables or p.upper().startswith("DT_") or is_formula(p):
        return ""
    return p


def _step_kind(method: str, legacy: str, save_as: str, match: str) -> str:
    if legacy:
        return "legacy"
    if not method:
        return "empty"
    if method == "OUTPUT":
        return "save" if save_as and not match else "check"
    if method == "CALL_TEST":
        return "call"
    for kind, group in (("nav", _NAV), ("act", _ACT), ("input", _INPUT), ("check", _CHECK), ("save", _SAVE), ("wait", _WAIT), ("flow", _FLOW)):
        if method in group:
            return kind
    return "other"


def parse_test_sheet(editor: WorkbookEditor, sheet: str, variables: set[str], *, secrets: set[str] | None = None,
                     last: dict[tuple[str, int], dict] | None = None) -> dict:
    """Steps, blocks and section rows of a keyword sheet (CONTRACT.md 2.3, 2.4, 1.5).  ``variables``: UPPER tokens that are Params columns of
    this test (whole-cell tokens)."""
    grid = _Grid(editor, sheet)
    secrets = secrets or set()
    last = last or {}
    stored = BLOCK_COLUMN in grid.cols and any(grid.text(r, BLOCK_COLUMN) for r in range(2, grid.max_row + 1))
    steps: list[dict] = []
    sections: list[dict] = []
    frame, window = "", 0
    flow: list[dict] = []
    for r in range(2, grid.max_row + 1):
        if not grid.row_has_content(r):
            continue
        method = grid.text(r, "METHOD").replace("\xa0", " ").strip().upper()
        flag = grid.get(r, "BLNEXECUTE")
        flag_text = text_of(flag).strip()
        if not method:
            ignore = {grid.cols.get(h) for h in ("BLNEXECUTE", "STEP_NUMBER", "INDEX", *RESULT_COLUMNS)} | {None}
            others = [v for i, v in enumerate(grid.rows[r - 1], start=1) if i not in ignore and text_of(v).replace("\xa0", " ").strip()]
            if flag_text and not is_formula(flag) and not _is_flag_text(flag_text) and flag_text.upper() not in variables and not others:
                sections.append({"row": r, "title": flag_text})
                continue
            if not others:
                continue
        locator = grid.raw(r, "FINDBY_VALUE")
        value, expected = grid.raw(r, "VALUE"), grid.raw(r, "EXPECTED_VALUE")
        output = grid.raw(r, "OUTPUT_VALUE")
        exact, contains = is_true(grid.get(r, "EXACT_MATCH")), is_true(grid.get(r, "CONTAINS"))
        match = "exact" if exact else ("contains" if contains else "")
        save_as = ""
        if output:
            m = INLINE_RE.fullmatch(output)
            save_as = m.group(1) if m else (output if output.upper() in variables or method in ("SET_VARIABLE", "ASK_USER", "PROMPT", "ASK",
                                                                                                "ALERT_TEXT_OUT", "GET_CURRENT_URL") and TOKEN_RE.match(output) else "")
        condition = None
        if is_formula(flag):
            enabled = None
            condition = {"kind": "formula", "text": flag_text}
        elif flag_text.upper() in variables:
            enabled = None
            condition = {"kind": "flag", "text": flag_text}
        else:
            enabled = is_true(flag)
        legacy = legacy_reason(method, flag)
        timeout_raw = grid.raw(r, "TIMEOUT")
        try:
            timeout = int(float(timeout_raw)) if timeout_raw and not is_formula(timeout_raw) else None
        except ValueError:
            timeout = None
        index_raw = grid.raw(r, "INDEX")
        try:
            index = int(float(index_raw)) if index_raw and not is_formula(index_raw) else 0
        except ValueError:
            index = 0
        backups: list[str] = []
        raw_backups = grid.raw(r, BACKUP_LOCATORS_COLUMN)
        if raw_backups:
            try:
                parsed = json.loads(raw_backups)
                backups = [str(x) for x in parsed] if isinstance(parsed, list) else [raw_backups]
            except ValueError:
                backups = [raw_backups]
        locator_name = grid.raw(r, "LOCATOR")
        target = describe_target(locator, locator_name, variables) if method in ELEMENT_METHODS or locator else ""
        if method == "SWITCHTOFRAME":
            target = target or _show(value, variables)
        output_property = grid.raw(r, "OUTPUT_PROPERTY")
        generated = auto_name(method, target=target, value=value, expected=expected, save_as=save_as, match=match,
                              output_property=output_property, variables=variables)
        name = grid.raw(r, "STEP_NAME")
        name_auto = is_true(grid.get(r, STEP_NAME_AUTO_COLUMN))
        # variables the step reads / sets
        uses: list[dict] = []
        for column in USE_COLUMNS:
            cell = grid.get(r, column)
            if cell in (None, ""):
                continue
            t = text_of(cell).strip()
            if not is_formula(cell) and t.upper() in variables:
                uses.append({"token": t.upper(), "column": column, "form": "cell"})
            for tok in INLINE_RE.findall(t):
                if column == "VALUE" and method in KEY_METHODS and tok.upper() in KEY_NAMES:
                    continue                                       # {TAB}, {ENTER}: SendKeys key names, not variables
                uses.append({"token": tok.upper(), "column": column, "form": "inline"})
            for tok in SECRET_RE.findall(t):
                uses.append({"token": tok.upper(), "column": column, "form": "secret"})
        if condition and condition["kind"] == "flag":
            uses.append({"token": flag_text.upper(), "column": "BLNEXECUTE", "form": "flag"})
        sets = [save_as.upper()] if save_as else []
        # context (frames / windows) and flow (IF / loops)
        if method == "SWITCHTOFRAME":
            frame = target or "frame"
        elif method == "SWITCHTODEFAULT":
            frame = ""
        elif method == "SWITCHTOWINDOW":
            window += 1
        elif method in ("SWITCHTOMAINWINDOW",):
            window, frame = 0, ""
        context = " · ".join(x for x in ((f"window {window + 1}" if window else ""), (f"frame: {frame}" if frame else "")) if x)
        here_flow = [dict(f) for f in flow]
        if method == "IF":
            flow.append({"kind": "if", "row": r})
        elif method == "ELSE" and flow and flow[-1]["kind"] in ("if", "else"):
            flow[-1] = {"kind": "else", "row": r}
            here_flow = [dict(f) for f in flow[:-1]]
        elif method == "END_IF" and flow and flow[-1]["kind"] in ("if", "else"):
            flow.pop()
            here_flow = [dict(f) for f in flow]
        elif method == "ITERATION_START":
            flow.append({"kind": "loop", "row": r})
        elif method == "ITERATION_END" and flow and flow[-1]["kind"] == "loop":
            flow.pop()
            here_flow = [dict(f) for f in flow]
        steps.append({
            "row": r, "n": len(steps) + 1, "method": method, "kind": _step_kind(method, legacy, save_as, match),
            "name": name, "nameAuto": name_auto, "autoName": generated, "block": grid.raw(r, BLOCK_COLUMN) if stored else "",
            "page": grid.raw(r, "PAGE"),
            "locator": {"findBy": grid.raw(r, "FINDBY"), "value": locator, "index": index, "name": locator_name, "backups": backups, "plainWords": []},
            "target": target, "value": value, "expected": expected, "match": match, "saveAs": save_as, "output": output,
            "outputProperty": output_property, "onFail": "continue" if is_true(grid.get(r, "IGNORE_NOT_EXISTING_OBJECT")) else "stop",
            "timeout": timeout, "enabled": enabled, "condition": condition, "sideEffects": is_true(grid.get(r, SIDE_EFFECTS_COLUMN)),
            "context": context, "flow": here_flow, "call": value if method == "CALL_TEST" else "", "legacy": legacy,
            "uses": uses, "sets": sets, "notes": grid.raw(r, "NOTES"), "lastResult": last.get((sheet.upper(), r)), "problems": [],
        })
        if method in ("CLOSE",):
            window, frame = max(window - 1, 0), ""
    blocks = derive_blocks(steps, sections, stored, variables)
    return {"steps": steps, "blocks": blocks, "sections": sections, "stored": stored, "columns": grid.header_names}


def derive_blocks(steps: list[dict], sections: list[dict], stored: bool, variables: set[str] | None = None) -> list[dict]:
    """Blocks (map nodes) of a test; fills ``step["block"]`` with its block's title (CONTRACT.md 1.5)."""
    variables = variables or set()
    section_at = {s["row"]: s["title"] for s in sections}
    section_rows = sorted(section_at)
    blocks: list[dict] = []
    current: dict | None = None
    title, header_row, page, force_new, si = "Start", None, "", False, 0

    def start(step: dict, kind: str, block_title: str) -> dict:
        block = {"id": f"b{len(blocks) + 1}", "title": block_title, "kind": kind, "firstRow": step["row"], "lastRow": step["row"],
                 "start": step["n"], "end": step["n"], "count": 0, "stored": stored, "headerRow": header_row if kind == "page" else None,
                 "page": page, "gate": step["value"] if step["method"] == "ASSERT_PAGE" else "", "returnsTo": "", "call": step["call"],
                 "dot": "", "open": kind in ("window", "loop")}
        blocks.append(block)
        return block

    for step in steps:
        method = step["method"]
        header = None
        while si < len(section_rows) and section_rows[si] < step["row"]:
            header = section_rows[si]
            si += 1
        if stored:
            if step["block"]:
                title = step["block"]
        else:
            if header is not None:
                title, header_row, force_new = section_at[header], header, True
            p = _page_name(step["page"], variables)
            if p and p != page:
                page = p
                if header is None and title != p:
                    title, header_row, force_new = p, None, True
        inside = current is not None and current["open"] and (not stored or current["title"] == title)
        if inside:
            pass
        elif method == "CALL_TEST":
            current = start(step, "call", title if stored else f"Calls {step['value'] or '?'}")
        elif method == "SWITCHTOWINDOW":
            current = start(step, "window", title if stored else "New window")
            current["returnsTo"] = "" if stored else title
        elif method == "ITERATION_START":
            current = start(step, "loop", title if stored else f"For each row of {step['value'] or 'the data'}")
        elif force_new or current is None or current["kind"] != "page" or current["title"] != title:
            current = start(step, "page", title)
        force_new, header_row = False, None
        current["lastRow"], current["end"] = step["row"], step["n"]
        current["count"] += 1
        step["block"] = current["title"]
        if current["kind"] == "window" and method in ("SWITCHTOMAINWINDOW", "CLOSE") or current["kind"] == "loop" and method == "ITERATION_END":
            current["open"] = False
        if current["kind"] != "page" and not current["open"] or current["kind"] == "call":
            force_new = True
    for b in blocks:
        b.pop("open", None)
    return blocks


# ---------------------------------------------------------------------------------------------
# The whole workbook
# ---------------------------------------------------------------------------------------------
def build_model(editor: WorkbookEditor, *, runs_dir: Path | None = None, environment: str | None = None, version: int = 0,
                status: dict | None = None, with_problems: bool = True) -> dict:
    """The ``Workbook`` JSON of CONTRACT.md 2.1."""
    names = editor.sheet_names()
    ds_name, ds_rows = _datasheets(editor)
    labels = read_variable_labels(editor)
    environments = read_environments(editor)
    fingerprints = read_fingerprints(editor)
    last_steps, last_tests = _last_results(runs_dir, editor.path.name)
    global_name = _sheet_key(editor, "Global")
    globals_: dict[str, Any] = {}
    if global_name:
        for row in editor.rows(global_name)[1:]:
            if row and text_of(row[0]).strip():
                globals_[text_of(row[0]).strip()] = row[1] if len(row) > 1 else None
    env_value = next((v for k, v in globals_.items() if k.upper() == "ENVIRONMENT"), None)
    chosen_env = (environment or ("" if is_formula(env_value) else text_of(env_value))).strip().upper()

    grids: dict[str, _Grid] = {}

    def grid(name: str) -> _Grid | None:
        key = _sheet_key(editor, name)
        if key is None:
            return None
        if key not in grids:
            grids[key] = _Grid(editor, key)
        return grids[key]

    env_keys = {r["variable"].upper() for r in environments["rows"]}
    secret_keys = {k for k, v in labels.items() if v["secret"]} | {r["variable"].upper() for r in environments["rows"] if r["secret"]}
    listed = {d.sheet.upper(): d for d in ds_rows}
    param_sheets: set[str] = set()
    tests: list[dict] = []
    order = [d.sheet for d in ds_rows] + [n for n in names if n.upper() not in listed]
    seen: set[str] = set()
    for sheet_name in order:
        key = _sheet_key(editor, sheet_name)
        if key is None or key.upper() in seen or key.upper().startswith(RR_PREFIX.upper()):
            continue
        cols = {h.upper(): c for h, c in editor.headers(key).items()}
        ui, api = _is_ui_sheet(cols), _is_api_sheet(cols)
        d = listed.get(key.upper())
        if not ui and not api:
            continue
        g = grid(key)
        seen.add(key.upper())
        params = grid(d.param_sheet) if d and d.param_sheet else None
        if params is not None:
            param_sheets.add(params.name)
        variables = set(params.cols) if params is not None else set()
        if ui:
            parsed = parse_test_sheet(editor, key, variables, secrets=secret_keys, last=last_steps)
            data_rows = _data_rows(params)
            kind = "web"
        else:
            parsed = {"steps": [], "blocks": [], "sections": [], "stored": False, "columns": g.header_names}
            data_rows = _data_rows(g)
            json_format = [g.text(r["row"], "JSON_FORMAT") for r in data_rows if r["enabled"]]
            kind = "api" if "JSON_FORMAT" in g.cols and any(is_true(x) for x in json_format) else "xml"
        tests.append({
            "id": key, "sheet": key, "kind": kind, "listed": d is not None, "dataSheetsRow": d.row if d else None,
            "enabled": bool(d.enabled) if d else False, "paramSheet": params.name if params is not None else (d.param_sheet or None if d else None),
            "tags": d.tags if d else [], "comment": d.comment if d else "", "dataRows": data_rows,
            "buildingWith": next((r["row"] for r in data_rows if r["enabled"] is not False), None),
            "blocks": parsed["blocks"], "steps": parsed["steps"], "sections": parsed["sections"], "needs": [], "provides": [],
            "calls": sorted({s["call"].split("#")[0] for s in parsed["steps"] if s["call"]}), "columns": parsed["columns"],
            "lastRun": last_tests.get(key.upper()), "_params": params, "_variables": variables,
        })
    # loops over other Params sheets make them variable sources too
    for t in tests:
        for s in t["steps"]:
            if s["method"] == "ITERATION_START" and s["value"] and grid(s["value"]) is not None:
                param_sheets.add(grid(s["value"]).name)

    # -- variables index -----------------------------------------------------------------------------
    var_index: dict[str, dict] = {}

    def var(token: str) -> dict:
        k = token.upper()
        if k not in var_index:
            info = labels.get(k, {})
            var_index[k] = {"token": info.get("token") or token, "key": k, "label": info.get("label") or humanize(token),
                            "labelSet": bool(info.get("label")), "kind": "data", "secret": k in secret_keys,
                            "envSpecific": bool(info.get("envSpecific")) or k in env_keys, "sources": [], "setBy": [], "usedBy": [],
                            "neededBy": [], "providedBy": []}
        return var_index[k]

    for name in sorted(param_sheets):
        g = grid(name)
        for header in g.header_names:
            if header and header.upper() != "BLNEXECUTE":
                v = var(header)
                v["sources"].append({"sheet": g.name, "column": header})
    for row in environments["rows"]:
        v = var(row["variable"])
        v["kind"] = "env"
        v["sources"].append({"sheet": RR_ENVIRONMENTS if environments["source"] == "rr" else "Environments", "column": row["variable"]})
    for k, info in labels.items():
        var(info["token"])
    for t in tests:
        for s in t["steps"]:
            for u in s["uses"]:
                v = var(u["token"])
                v["usedBy"].append({"test": t["id"], "row": s["row"], "n": s["n"], "column": u["column"]})
                if u["form"] == "flag" and v["kind"] == "data":
                    v["kind"] = "flag"
                if u["form"] == "secret":
                    v["secret"] = True
            for tok in s["sets"]:
                var(tok)["setBy"].append({"test": t["id"], "row": s["row"], "n": s["n"]})
    for v in var_index.values():
        if not v["sources"] and v["setBy"] and v["kind"] == "data":
            v["kind"] = "set"
    for f in fingerprints:
        for t in tests:
            for s in t["steps"]:
                if s["method"] == "ASSERT_PAGE" and s["value"].upper() == f["name"].upper():
                    f["usedBy"].append({"test": t["id"], "row": s["row"], "n": s["n"]})

    # -- needs / provides (Q21) --------------------------------------------------------------------------
    for t in tests:
        params: _Grid | None = t.pop("_params")
        variables = t.pop("_variables")
        enabled_rows = [r["row"] for r in t["dataRows"] if r["enabled"] is not False]
        own_sets: set[str] = set()
        needs: dict[str, dict] = {}
        for s in t["steps"]:
            if s["enabled"] is False:
                continue
            cond = s["condition"] or {}
            if cond.get("kind") == "flag" and params is not None:
                rows_here = [r for r in enabled_rows if is_true(params.get(r, cond["text"].upper()))]    # runs only where the flag is Y
            else:
                rows_here = enabled_rows
            maybe = cond.get("kind") == "formula"                                                   # a formula decides: cannot tell here
            for u in s["uses"]:
                k = u["token"]
                if k in own_sets or k in needs or u["form"] == "secret" or k in secret_keys:
                    continue
                if k in env_keys and k not in variables:
                    continue
                if u["form"] == "flag":
                    continue
                if k in variables:
                    empty = [r for r in rows_here if params.get(r, k) in (None, "") or text_of(params.get(r, k)).strip().upper() == "NONE"]
                    if not empty:
                        continue
                    needs[k] = {"rows": empty, "row": s["row"], "n": s["n"], "maybe": maybe}
                elif u["form"] in ("inline",):
                    needs[k] = {"rows": [], "row": s["row"], "n": s["n"], "maybe": maybe}
            own_sets.update(s["sets"])
        t["needs"] = sorted(needs)
        t["provides"] = sorted({tok for s in t["steps"] for tok in s["sets"]})
        t["_needs"] = needs
        for k in t["needs"]:
            var(k)["neededBy"].append(t["id"])
        for k in t["provides"]:
            var(k)["providedBy"].append(t["id"])

    sheets = []
    params_upper = {p.upper() for p in param_sheets}
    tests_upper = {t["sheet"].upper() for t in tests}
    for n in names:
        u = n.upper()
        role = ("runner" if u.startswith(RR_PREFIX.upper()) else "test" if u in tests_upper else "params" if u in params_upper
                else "datasheets" if u == "DATASHEETS" else "global" if u == "GLOBAL" else "environments" if u == "ENVIRONMENTS"
                else "inputoutput" if u == "INPUTOUTPUT" else "other")
        sheets.append({"name": n, "hidden": editor.is_hidden(n), "role": role})

    model = {
        "name": editor.path.name, "version": version, "environment": chosen_env, "globals": {k: _jsonable(v) for k, v in globals_.items()},
        "sheets": sheets, "tests": tests, "variables": sorted(var_index.values(), key=lambda v: v["label"].lower()),
        "environments": environments, "fingerprints": fingerprints, "scenarios": [], "problems": [],
        "problemCounts": {"error": 0, "warning": 0, "info": 0}, "status": status or {},
    }
    if with_problems:
        from ..lint import builder_problems
        attach_problems(model, builder_problems(model))
    for t in model["tests"]:
        t.pop("_needs", None)
    return model


def attach_problems(model: dict, problems: list[dict]) -> None:
    """Put the problems on the model: the list, the counts, a severity per step and the worst one per block (the red / amber dots)."""
    rank = {"error": 3, "warning": 2, "info": 1, "": 0}
    model["problems"] = problems
    counts = {"error": 0, "warning": 0, "info": 0}
    by_step: dict[tuple[str, int], list[str]] = {}
    for p in problems:
        counts[p["severity"]] = counts.get(p["severity"], 0) + 1
        if p.get("test") and p.get("row"):
            by_step.setdefault((p["test"], p["row"]), []).append(p["severity"])
    model["problemCounts"] = counts
    for t in model["tests"]:
        for s in t["steps"]:
            s["problems"] = by_step.get((t["id"], s["row"]), [])
        for b in t["blocks"]:
            worst = ""
            for s in t["steps"][b["start"] - 1:b["end"]]:
                for sev in s["problems"]:
                    if rank[sev] > rank[worst]:
                        worst = sev
            b["dot"] = worst


def _jsonable(value: Any) -> Any:
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return text_of(value)


def _camel(name: str) -> str:
    head, *rest = name.split("_")
    return head + "".join(p.title() for p in rest)


# ---------------------------------------------------------------------------------------------
# Ops (CONTRACT.md 3.1)
# ---------------------------------------------------------------------------------------------
class _Ops:
    def __init__(self, editor: WorkbookEditor):
        self.e = editor

    # -- lookups -------------------------------------------------------------------------------
    def sheet(self, name: Any) -> str:
        key = _sheet_key(self.e, str(name or ""))
        if key is None:
            raise BuildError(f"There is no sheet called {name!r}.", "not_found", 404)
        return key

    def params_of(self, sheet: str) -> _Grid | None:
        _, rows = _datasheets(self.e)
        for d in rows:
            if d.sheet.upper() == sheet.upper() and d.param_sheet and _sheet_key(self.e, d.param_sheet):
                return _Grid(self.e, _sheet_key(self.e, d.param_sheet))
        return None

    def parsed(self, sheet: str) -> dict:
        params = self.params_of(sheet)
        return parse_test_sheet(self.e, sheet, set(params.cols) if params else set())

    def test_sheet(self, op: dict) -> str:
        sheet = self.sheet(op.get("test"))
        if not _is_ui_sheet({k.upper(): v for k, v in self.e.headers(sheet).items()}):
            raise BuildError(f"{sheet} is not a keyword (web) test sheet.")
        return sheet

    def step_at(self, parsed: dict, row: Any) -> dict:
        try:
            r = int(row)
        except (TypeError, ValueError):
            raise BuildError(f"Not a row number: {row!r}") from None
        for s in parsed["steps"]:
            if s["row"] == r:
                return s
        raise BuildError(f"Row {r} is not a step.", "not_found", 404)

    def column(self, sheet: str, header: str, create: bool = True) -> int | None:
        col = self.e.column(sheet, header)
        if col is None and create:
            col = self.e.add_column(sheet, header)
        return col

    def put(self, sheet: str, row: int, header: str, value: Any) -> None:
        col = self.column(sheet, header, create=value not in (None, ""))
        if col is not None:
            self.e.set(sheet, row, col, value)

    # -- step fields -----------------------------------------------------------------------------
    def write_fields(self, sheet: str, row: int, fields: dict) -> None:
        for key, value in fields.items():
            if key not in FIELD_COLUMNS:
                raise BuildError(f"Unknown step field {key!r}.")
            if key == "match":
                m = str(value or "").lower()
                if m not in ("", "exact", "contains"):
                    raise BuildError("match is exact, contains or empty.")
                self.put(sheet, row, "Exact_Match", "Y" if m == "exact" else None)
                self.put(sheet, row, "Contains", "Y" if m == "contains" else None)
                continue
            if key == "enabled":
                value = "Y" if value else "N"
            elif key == "onFail":
                if value not in ("stop", "continue"):
                    raise BuildError("onFail is stop or continue.")
                value = "Y" if value == "continue" else None
            elif key in ("sideEffects", "nameAuto"):
                value = "Y" if value else None
            elif key == "timeout":
                value = None if value in (None, "") else _int(value, "timeout")
            elif key == "index":
                value = None if value in (None, "") else _int(value, "index")
            elif key == "backups":
                if value in (None, "", []):
                    value = None
                elif not isinstance(value, list):
                    raise BuildError("backups is a list of locators.")
                else:
                    value = json.dumps([str(x) for x in value])
            elif key == "method" and value is not None:
                value = str(value).strip()
            self.put(sheet, row, FIELD_COLUMNS[key], value)

    def refresh_name(self, sheet: str, row: int) -> None:
        step = self.step_at(self.parsed(sheet), row)
        if step["nameAuto"] and step["name"] != step["autoName"]:
            self.put(sheet, row, "Step_Name", step["autoName"])

    def materialize(self, sheet: str) -> dict:
        """Write the BLOCK column for every step (the first block op on an old sheet), so the column is the truth from now on."""
        parsed = self.parsed(sheet)
        if not parsed["stored"]:
            col = self.column(sheet, BLOCK_COLUMN)
            for s in parsed["steps"]:
                self.e.set(sheet, s["row"], col, s["block"])
            parsed = self.parsed(sheet)
        return parsed

    def block_of(self, parsed: dict, row: int) -> dict:
        step = self.step_at(parsed, row)
        for b in parsed["blocks"]:
            if b["start"] <= step["n"] <= b["end"]:
                return b
        raise BuildError(f"Row {row} is in no block.")

    def set_block(self, sheet: str, rows: Iterable[int], title: str) -> None:
        title = str(title or "").strip()
        if not title:
            raise BuildError("A block needs a title.")
        col = self.column(sheet, BLOCK_COLUMN)
        for r in rows:
            self.e.set(sheet, r, col, title)

    # -- ops -------------------------------------------------------------------------------------------
    def insert_step(self, op: dict) -> dict:
        sheet = self.test_sheet(op)
        parsed = self.parsed(sheet)
        fields = dict(op.get("step") or {})
        if not str(fields.get("method") or "").strip():
            raise BuildError("A new step needs a method.")
        if op.get("before") is not None:
            at = _int(op["before"], "before")
        elif op.get("after") is not None:
            at = _int(op["after"], "after") + 1
        else:
            at = (parsed["steps"][-1]["row"] + 1) if parsed["steps"] else max(self.e.max_row(sheet), 1) + 1
        at = max(at, 2)
        neighbour = next((s for s in reversed(parsed["steps"]) if s["row"] < at), None) or next(iter(parsed["steps"]), None)
        block = fields.pop("block", None)
        if block and neighbour and block != neighbour["block"] and not parsed["stored"]:
            self.materialize(sheet)
        self.e.insert_rows(sheet, at, 1)
        step_col = self.e.column(sheet, "Step_Number")
        if step_col:
            source = next((r for r in range(at - 1, 1, -1) if is_formula(self.e.get(sheet, r, step_col))), None)
            if source is not None:
                formula = self.e.get(sheet, source, step_col)
                self.e.set(sheet, at, step_col, Translator(formula, origin=f"{get_column_letter(step_col)}{source}")
                           .translate_formula(f"{get_column_letter(step_col)}{at}"))
        fields.setdefault("enabled", True)
        if "name" in fields and fields["name"]:
            fields.setdefault("nameAuto", False)
        else:
            fields.pop("name", None)
            fields["nameAuto"] = True
        self.write_fields(sheet, at, fields)
        if block or self.parsed(sheet)["stored"]:
            self.set_block(sheet, [at], block or (neighbour["block"] if neighbour else "Start"))
        self.refresh_name(sheet, at)
        return {"op": "insert_step", "test": sheet, "row": at}

    def update_step(self, op: dict) -> dict:
        sheet = self.test_sheet(op)
        parsed = self.parsed(sheet)
        step = self.step_at(parsed, op.get("row"))
        fields = dict(op.get("set") or {})
        if step["legacy"]:
            raise BuildError(f"Row {step['row']} is a legacy row: edit it in the Excel grid. {step['legacy']}", "legacy")
        if "block" in fields:
            block = fields.pop("block")
            if not parsed["stored"]:
                self.materialize(sheet)
            self.set_block(sheet, [step["row"]], block)
        if fields.get("name"):
            fields.setdefault("nameAuto", False)
        elif "name" in fields:
            fields.pop("name")
            fields["nameAuto"] = True
        self.write_fields(sheet, step["row"], fields)
        self.refresh_name(sheet, step["row"])
        return {"op": "update_step", "test": sheet, "row": step["row"]}

    def bulk_edit(self, op: dict) -> dict:
        sheet = self.test_sheet(op)
        fields = dict(op.get("set") or {})
        bad = set(fields) - BULK_FIELDS
        if bad:
            raise BuildError(f"Bulk edit changes only {', '.join(sorted(BULK_FIELDS))} (not {', '.join(sorted(bad))}).")
        parsed = self.parsed(sheet)
        steps = [self.step_at(parsed, r) for r in op.get("rows") or []]
        block = fields.pop("block", None)
        if block is not None:
            if not parsed["stored"]:
                self.materialize(sheet)
            self.set_block(sheet, [s["row"] for s in steps], block)
        skipped = []
        for s in steps:
            if s["legacy"]:
                skipped.append(s["row"])
                continue
            if fields:
                self.write_fields(sheet, s["row"], fields)
                self.refresh_name(sheet, s["row"])
        return {"op": "bulk_edit", "test": sheet, "rows": [s["row"] for s in steps], "skipped": skipped}

    def delete_steps(self, op: dict) -> dict:
        sheet = self.test_sheet(op)
        parsed = self.parsed(sheet)
        rows = sorted({self.step_at(parsed, r)["row"] for r in op.get("rows") or []}, reverse=True)
        for r in rows:
            self.e.delete_rows(sheet, r, 1)
        return {"op": "delete_steps", "test": sheet, "rows": sorted(rows)}

    def move_steps(self, op: dict) -> dict:
        sheet = self.test_sheet(op)
        parsed = self.parsed(sheet)
        rows = sorted({self.step_at(parsed, r)["row"] for r in op.get("rows") or []})
        if not rows:
            raise BuildError("Say which steps to move.")
        block = op.get("block")
        if block and not parsed["stored"]:
            self.materialize(sheet)
        if op.get("before") is not None:
            anchor = _int(op["before"], "before")
        else:
            anchor = parsed["steps"][-1]["row"] + 1
        pending = list(rows)
        landed: list[int] = []
        while pending:
            r = pending.pop(0)
            rm = RowMap.move(r, 1, anchor)
            if not rm.is_identity():
                self.e.move_rows(sheet, r, 1, anchor)
                pending = [rm.row(x) for x in pending]
                landed = [rm.row(x) for x in landed]
                anchor = rm.row(anchor)
                landed.append(rm.row(r))
            else:
                landed.append(r)
        if block:
            self.set_block(sheet, landed, block)
        return {"op": "move_steps", "test": sheet, "rows": sorted(landed)}

    def rename_block(self, op: dict) -> dict:
        sheet = self.test_sheet(op)
        parsed = self.materialize(sheet)
        b = self.block_of(parsed, _int(op.get("row"), "row"))
        rows = [s["row"] for s in parsed["steps"][b["start"] - 1:b["end"]]]
        self.set_block(sheet, rows, op.get("title"))
        return {"op": "rename_block", "test": sheet, "rows": rows}

    def split_block(self, op: dict) -> dict:
        sheet = self.test_sheet(op)
        parsed = self.materialize(sheet)
        row = _int(op.get("row"), "row")
        b = self.block_of(parsed, row)
        step = self.step_at(parsed, row)
        if step["n"] == b["start"]:
            raise BuildError("Split above a step that is not the first of its block.")
        rows = [s["row"] for s in parsed["steps"][step["n"] - 1:b["end"]]]
        self.set_block(sheet, rows, op.get("title"))
        return {"op": "split_block", "test": sheet, "rows": rows}

    def merge_blocks(self, op: dict) -> dict:
        sheet = self.test_sheet(op)
        parsed = self.materialize(sheet)
        b = self.block_of(parsed, _int(op.get("row"), "row"))
        i = parsed["blocks"].index(b)
        if i == 0:
            raise BuildError("The first block has no block before it to merge into.")
        rows = [s["row"] for s in parsed["steps"][b["start"] - 1:b["end"]]]
        self.set_block(sheet, rows, parsed["blocks"][i - 1]["title"])
        return {"op": "merge_blocks", "test": sheet, "rows": rows}

    def rename_variable(self, op: dict) -> dict:
        old, new = str(op.get("from") or "").strip(), str(op.get("to") or "").strip()
        if not TOKEN_RE.match(old) or not TOKEN_RE.match(new):
            raise BuildError("Variable names are letters, digits and _ (starting with a letter).")
        if old.upper() == new.upper() and old == new:
            return {"op": "rename_variable", "changed": 0}
        ou = old.upper()
        inline = re.compile(r"\{(\??)" + re.escape(old) + r"\}", re.I)
        secret = re.compile(r"\{SECRET:" + re.escape(old) + r"\}", re.I)
        changed = 0
        _, ds_rows = _datasheets(self.e)
        param_names = {d.param_sheet.upper() for d in ds_rows if d.param_sheet}
        for sheet in self.e.sheet_names():
            su = sheet.upper()
            if su.startswith(RR_PREFIX.upper()):
                continue
            headers = {k.upper(): v for k, v in self.e.headers(sheet).items()}
            if su in param_names:
                if new.upper() in headers and new.upper() != ou:
                    raise BuildError(f"{sheet} already has a column {new}.", "exists", 409)
                if ou in headers:
                    self.e.set(sheet, 1, headers[ou], new)
                    changed += 1
                continue
            if not _is_ui_sheet(headers):
                continue
            skip = {headers.get(h) for h in ("STEP_NAME", "METHOD", "NOTES", *RESULT_COLUMNS)}
            for r, values in enumerate(self.e.rows(sheet)[1:], start=2):
                for c, v in enumerate(values, start=1):
                    if v in (None, "") or c in skip or not isinstance(v, str):
                        continue
                    nv = v
                    if not is_formula(v) and v.strip().upper() == ou:
                        nv = new
                    else:
                        nv = secret.sub("{SECRET:" + new + "}", inline.sub(lambda m: "{" + m.group(1) + new + "}", nv))
                    if nv != v:
                        self.e.set(sheet, r, c, nv)
                        changed += 1
        for table, key in ((RR_VARIABLES, "TOKEN"), (RR_ENVIRONMENTS, "VARIABLE")):
            name = _sheet_key(self.e, table)
            if name is None:
                continue
            rows = self.e.rows(name)
            header = {text_of(v).strip().upper(): i for i, v in enumerate(rows[0] if rows else [])}
            col = header.get(key)
            for r, values in enumerate(rows[1:], start=2):
                if col is not None and col < len(values) and text_of(values[col]).strip().upper() == ou:
                    self.e.set(name, r, col + 1, new)
                    changed += 1
        return {"op": "rename_variable", "from": old, "to": new, "changed": changed}

    def _upsert(self, table: str, headers: tuple[str, ...], key: str, key_value: str, values: dict[str, Any],
                old_key: str | None = None) -> None:
        self.e.ensure_rr_sheet(table)
        name = _sheet_key(self.e, table)
        rows = self.e.rows(name)
        if not rows:
            rows = [list(headers)]
        header = [text_of(v).strip() for v in rows[0]]
        for h in headers:
            if h.upper() not in {x.upper() for x in header}:
                header.append(h)
        idx = {h.upper(): i for i, h in enumerate(header)}
        body = [list(r) + [None] * (len(header) - len(r)) for r in rows[1:]]
        find = (old_key or key_value).upper()
        target = next((r for r in body if text_of(r[idx[key.upper()]]).strip().upper() == find), None)
        if target is None:
            target = [None] * len(header)
            body.append(target)
        target[idx[key.upper()]] = key_value
        for h, v in values.items():
            if v is not None:
                target[idx[h.upper()]] = v
        self.e.write_table(name, [header] + body)

    def set_variable(self, op: dict) -> dict:
        token = str(op.get("token") or "").strip()
        if not TOKEN_RE.match(token):
            raise BuildError("Variable names are letters, digits and _ (starting with a letter).")
        yn = lambda k: None if k not in op else ("Y" if op[k] else "")
        self._upsert(RR_VARIABLES, VARIABLE_HEADERS, "Token", token,
                     {"Label": op.get("label"), "Secret": yn("secret"), "EnvSpecific": yn("envSpecific"), "Notes": op.get("notes")})
        return {"op": "set_variable", "token": token}

    def add_variable(self, op: dict) -> dict:
        token = str(op.get("token") or "").strip() or make_token(op.get("label") or "")
        if not TOKEN_RE.match(token):
            raise BuildError("Variable names are letters, digits and _ (starting with a letter).")
        sheet = self.sheet(op.get("sheet"))
        if self.e.column(sheet, token) is not None:
            raise BuildError(f"{sheet} already has a column {token}.", "exists", 409)
        col = self.e.add_column(sheet, token)
        if op.get("value") not in (None, ""):
            for r in range(2, self.e.max_row(sheet) + 1):
                self.e.set(sheet, r, col, op["value"])
        if op.get("label"):
            self.set_variable({"token": token, "label": op["label"]})
        return {"op": "add_variable", "token": token, "sheet": sheet}

    def set_cell(self, op: dict) -> dict:
        sheet = self.sheet(op.get("sheet"))
        row = _int(op.get("row"), "row")
        column = op.get("column")
        col = column if isinstance(column, int) else self.e.column(sheet, str(column or ""))
        if col is None:
            raise BuildError(f"{sheet} has no column {column!r}.", "not_found", 404)
        try:
            self.e.set(sheet, row, int(col), op.get("value"))
        except WriterError as err:
            raise BuildError(str(err)) from err
        return {"op": "set_cell", "sheet": sheet, "row": row, "column": int(col)}

    def set_environments(self, op: dict) -> dict:
        names = [str(n).strip() for n in op.get("names") or [] if str(n).strip()]
        if not names:
            raise BuildError("Keep at least one environment.")
        if len({n.upper() for n in names}) != len(names):
            raise BuildError("Two environments have the same name.")
        production = {str(n).upper() for n in op.get("production") or []}
        rows = [list(ENVIRONMENT_FIXED_HEADERS) + names, [PRODUCTION_ROW, "", ""] + ["Y" if n.upper() in production else "" for n in names]]
        seen = set()
        for r in op.get("rows") or []:
            var_name = str(r.get("variable") or "").strip()
            if not TOKEN_RE.match(var_name) or var_name.upper() in seen:
                raise BuildError(f"Environment variable {var_name!r} is not a valid, unique name.")
            seen.add(var_name.upper())
            values = r.get("values") or {}
            rows.append([var_name, "Y" if r.get("required") else "", "Y" if r.get("secret") else ""] +
                        [values.get(n, "") for n in names])
        self.e.ensure_rr_sheet(RR_ENVIRONMENTS)
        self.e.write_table(_sheet_key(self.e, RR_ENVIRONMENTS), rows)
        return {"op": "set_environments", "names": names}

    def set_fingerprint(self, op: dict) -> dict:
        name = str(op.get("name") or "").strip()
        if not name:
            raise BuildError("A page fingerprint needs a name.")
        existing = {f["name"].upper() for f in read_fingerprints(self.e)}
        old = str(op.get("rename") or "").strip() or None
        if old and name.upper() in existing and name.upper() != old.upper():
            raise BuildError(f"A page called {name!r} already exists.", "exists", 409)
        self._upsert(RR_FINGERPRINTS, FINGERPRINT_HEADERS, "Name", name,
                     {"UrlContains": op.get("urlContains"), "Landmark": op.get("landmark"), "LandmarkText": op.get("landmarkText"),
                      "Notes": op.get("notes")}, old_key=old)
        return {"op": "set_fingerprint", "name": name}

    def delete_fingerprint(self, op: dict) -> dict:
        name = str(op.get("name") or "").strip().upper()
        table = _sheet_key(self.e, RR_FINGERPRINTS)
        if table is None:
            raise BuildError("There are no page fingerprints.", "not_found", 404)
        rows = self.e.rows(table)
        kept = [rows[0]] + [r for r in rows[1:] if not (r and text_of(r[0]).strip().upper() == name)]
        if len(kept) == len(rows):
            raise BuildError(f"No page fingerprint called {op.get('name')!r}.", "not_found", 404)
        self.e.write_table(table, kept)
        return {"op": "delete_fingerprint", "name": op.get("name")}

    def _datasheets_sheet(self) -> str:
        name = _sheet_key(self.e, "DataSheets")
        if name is None:
            self.e.add_sheet("DataSheets")
            name = "DataSheets"
            for c, h in enumerate(("SheetsToExecute", "blnExecute", "ParameterSheet", "Comments"), start=1):
                self.e.set(name, 1, c, h)
        return name

    def set_test(self, op: dict) -> dict:
        sheet = self.sheet(op.get("test"))
        ds = self._datasheets_sheet()
        _, rows = _datasheets(self.e)
        d = next((x for x in rows if x.sheet.upper() == sheet.upper()), None)
        row = d.row if d else self.e.max_row(ds) + 1
        if d is None:
            self.put(ds, row, "SheetsToExecute", sheet)
            self.put(ds, row, "blnExecute", "N")
        if "enabled" in op:
            self.put(ds, row, "blnExecute", "Y" if op["enabled"] else "N")
        if "paramSheet" in op:
            if op["paramSheet"] and _sheet_key(self.e, op["paramSheet"]) is None:
                raise BuildError(f"There is no sheet called {op['paramSheet']!r}.", "not_found", 404)
            self.put(ds, row, "ParameterSheet", op["paramSheet"] or None)
        if "comment" in op:
            self.put(ds, row, "Comments", op["comment"] or None)
        if "tags" in op:
            tags = op["tags"] if isinstance(op["tags"], list) else re.split(r"[;,]", str(op["tags"] or ""))
            self.put(ds, row, "Tags", ", ".join(t.strip() for t in tags if str(t).strip()) or None)
        return {"op": "set_test", "test": sheet, "row": row}

    def add_test(self, op: dict) -> dict:
        name = str(op.get("name") or "").strip()
        if str(op.get("kind") or "web") != "web":
            raise BuildError("Only web tests can be created here for now (API and XML tests come with the API editor).")
        if _sheet_key(self.e, name) is not None:
            raise BuildError(f"A sheet called {name!r} already exists.", "exists", 409)
        try:
            self.e.add_sheet(name)
        except WriterError as err:
            raise BuildError(str(err)) from err
        for c, h in enumerate(STANDARD_COLUMNS, start=1):
            self.e.set(name, 1, c, h)
        param = str(op.get("paramSheet") or "").strip()
        if param and _sheet_key(self.e, param) is None:
            try:
                self.e.add_sheet(param)
            except WriterError as err:
                raise BuildError(str(err)) from err
            self.e.set(param, 1, 1, "blnExecute")
            self.e.set(param, 2, 1, "Y")
        self.set_test({"test": name, "enabled": True, **({"paramSheet": param} if param else {})})
        return {"op": "add_test", "test": name}


def _int(value: Any, what: str) -> int:
    try:
        return int(value)
    except (TypeError, ValueError):
        raise BuildError(f"{what} must be a whole number, not {value!r}.") from None


OPS = ("insert_step", "update_step", "delete_steps", "move_steps", "bulk_edit", "rename_block", "split_block", "merge_blocks",
       "rename_variable", "set_variable", "add_variable", "set_cell", "set_environments", "set_fingerprint", "delete_fingerprint",
       "set_test", "add_test")


def apply_ops(editor: WorkbookEditor, ops: list[dict]) -> list[dict]:
    """Apply ``ops`` in order.  On an error the editor is left half-changed: ``BuildDocument.apply`` restores its snapshot."""
    runner = _Ops(editor)
    applied = []
    for i, op in enumerate(ops):
        kind = str((op or {}).get("op") or "")
        if kind not in OPS:
            raise BuildError(f"Unknown op {kind!r}.", "op", 422, index=i)
        try:
            applied.append(getattr(runner, kind)(op))
        except BuildError as err:
            err.extra.setdefault("index", i)
            raise
        except WriterError as err:
            raise BuildError(str(err), "op", 422, index=i) from err
    return applied


# ---------------------------------------------------------------------------------------------
# Diff between two versions (History, Q30; file changed on disk, Q29)
# ---------------------------------------------------------------------------------------------
_SIG = ("method", "name", "value", "expected", "match", "saveAs", "enabled", "onFail", "timeout", "sideEffects", "block")


def _summary(step: dict) -> dict:
    return {"row": step["row"], "n": step["n"], "method": step["method"], "name": step["name"] or step["autoName"],
            "locator": step["locator"]["value"], "value": step["value"], "expected": step["expected"], "enabled": step["enabled"],
            "block": step["block"]}


def _signature(step: dict) -> tuple:
    return tuple(json.dumps(step.get(k), sort_keys=True) for k in _SIG) + (step["locator"]["value"],)


def diff_models(before: dict, after: dict) -> list[dict]:
    """Per-step changes between two models (steps aligned by content, so an insert shows as one added step, not a cascade)."""
    changes: list[dict] = []
    old = {t["id"].upper(): t for t in before["tests"]}
    new = {t["id"].upper(): t for t in after["tests"]}
    for key, t in new.items():
        if key not in old:
            changes.append({"test": t["id"], "kind": "test_added", "row": None, "n": None, "before": None, "after": None})
            continue
        a, b = old[key]["steps"], t["steps"]
        sm = difflib.SequenceMatcher(a=[_signature(s) for s in a], b=[_signature(s) for s in b], autojunk=False)
        for tag, i1, i2, j1, j2 in sm.get_opcodes():
            if tag == "equal":
                continue
            pairs = min(i2 - i1, j2 - j1) if tag == "replace" else 0
            for k in range(pairs):
                changes.append({"test": t["id"], "kind": "changed", "row": b[j1 + k]["row"], "n": b[j1 + k]["n"],
                                "before": _summary(a[i1 + k]), "after": _summary(b[j1 + k])})
            for s in a[i1 + pairs:i2]:
                changes.append({"test": t["id"], "kind": "removed", "row": s["row"], "n": s["n"], "before": _summary(s), "after": None})
            for s in b[j1 + pairs:j2]:
                changes.append({"test": t["id"], "kind": "added", "row": s["row"], "n": s["n"], "before": None, "after": _summary(s)})
    for key, t in old.items():
        if key not in new:
            changes.append({"test": t["id"], "kind": "test_removed", "row": None, "n": None, "before": None, "after": None})
    return changes


# ---------------------------------------------------------------------------------------------
# A new workbook (Q45): a minimal package written by hand, then filled through the editor
# ---------------------------------------------------------------------------------------------
_NEW_PARTS = {
    "[Content_Types].xml": '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
        '</Types>',
    "_rels/.rels": '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
        '</Relationships>',
    "xl/workbook.xml": '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<bookViews><workbookView/></bookViews><sheets><sheet name="Global" sheetId="1" r:id="rId1"/></sheets></workbook>',
    "xl/_rels/workbook.xml.rels": '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
        '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
        '</Relationships>',
    "xl/styles.xml": '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<fonts count="1"><font><sz val="11"/><name val="Calibri"/><family val="2"/></font></fonts>'
        '<fills count="2"><fill><patternFill patternType="none"/></fill><fill><patternFill patternType="gray125"/></fill></fills>'
        '<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>'
        '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
        '<cellXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/></cellXfs>'
        '<cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles></styleSheet>',
    "xl/worksheets/sheet1.xml": '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><dimension ref="A1"/>'
        '<sheetViews><sheetView workbookViewId="0"/></sheetViews><sheetFormatPr defaultRowHeight="15"/><sheetData/>'
        '<pageMargins left="0.7" right="0.7" top="0.75" bottom="0.75" header="0.3" footer="0.3"/></worksheet>',
}


def new_workbook(path: str | Path, environments: list[dict]) -> SaveResult:
    """Create ``path`` (Q45): Global (Environment = the first environment), DataSheets, and the environment table with DOMAIN per
    environment.  Refuses to overwrite a file."""
    path = Path(path)
    if path.exists():
        raise BuildError(f"{path.name} already exists.", "exists", 409)
    envs = [e for e in environments if str(e.get("name") or "").strip()] or [{"name": n, "production": n == "PROD"} for n in DEFAULT_ENVIRONMENTS]
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for part, text in _NEW_PARTS.items():
            z.writestr(part, text)
    editor = WorkbookEditor(path, buf.getvalue(), None)
    for c, h in enumerate(("Parameter", "Value", "Comments"), start=1):
        editor.set("Global", 1, c, h)
    editor.set("Global", 2, 1, "Environment")
    editor.set("Global", 2, 2, str(envs[0]["name"]).strip())
    ops = _Ops(editor)
    ops._datasheets_sheet()
    names = [str(e["name"]).strip() for e in envs]
    ops.set_environments({"names": names, "production": [str(e["name"]).strip() for e in envs if e.get("production")],
                          "rows": [{"variable": "DOMAIN", "required": True, "values": {str(e["name"]).strip(): str(e.get("domain") or "") for e in envs}}]})
    return editor.save_as(path)


# ---------------------------------------------------------------------------------------------
# One open workbook on the server
# ---------------------------------------------------------------------------------------------
class BuildDocument:
    """The workbook the Build tab edits: undo/redo, autosaved draft, save, history, diff, restore (Q29, Q30).

    Undo keeps one snapshot of the package per batch of ops (writing it takes ~0.1 s on the real workbooks). Every change is autosaved to the
    draft (``.drafts/<name>``); when undo takes the workbook all the way back, the draft goes and the file on disk is used again, so "no edits"
    always means "nothing written"."""

    UNDO_DEPTH = 50

    def __init__(self, path: str | Path, runs_dir: str | Path | None = None):
        self.path = Path(path)
        self.runs_dir = Path(runs_dir) if runs_dir else None
        self.lock = threading.RLock()
        self.version = 0
        self._undo: list[bytes] = []
        self._redo: list[bytes] = []
        self._model_cache: tuple[tuple, dict] | None = None
        self._open(prefer_draft=True)

    # -- opening -------------------------------------------------------------------------------
    def _open(self, *, prefer_draft: bool) -> None:
        if prefer_draft and has_draft(self.path):
            self.editor = WorkbookEditor.open_draft(self.path)
            self.from_draft = True
        else:
            self.editor = WorkbookEditor.open(self.path)
            self.from_draft = False
        self._undo, self._redo = [], []
        self.version += 1

    def _load_bytes(self, data: bytes) -> None:
        stamp = self.editor.stamp
        self.editor = WorkbookEditor(self.path, data, stamp)
        self.editor._modified = True                                # content differs from the file on disk (same as open_draft)

    @property
    def modified(self) -> bool:
        return self.editor.modified

    def refresh_if_changed_outside(self) -> bool:
        """An unedited workbook changed on disk (saved in Excel): reopen it silently.  Edited ones keep their draft (status says so)."""
        with self.lock:
            if not self.editor.modified and self.editor.external_change():
                self._open(prefer_draft=False)
                return True
        return False

    # -- model ---------------------------------------------------------------------------------
    def status(self) -> dict:
        draft = draft_path(self.path)
        saved = datetime.fromtimestamp(draft.stat().st_mtime).isoformat(timespec="seconds") if draft.is_file() else None
        return {"version": self.version, "modified": self.editor.modified, "hasDraft": draft.is_file(), "draftSaved": saved,
                "externalChange": self.editor.external_change(), "locked": is_locked(self.path), "canUndo": bool(self._undo),
                "canRedo": bool(self._redo), "file": self.path.name}

    def model(self, environment: str | None = None) -> dict:
        with self.lock:
            key = (self.version, (environment or "").upper(), _runs_stamp(self.runs_dir))
            if self._model_cache is None or self._model_cache[0] != key:
                self._model_cache = (key, build_model(self.editor, runs_dir=self.runs_dir, environment=environment, version=self.version))
            model = dict(self._model_cache[1])
            model["status"] = self.status()
            return model

    # -- editing -------------------------------------------------------------------------------
    def apply(self, ops: list[dict], version: int | None = None) -> list[dict]:
        if not isinstance(ops, list) or not ops:
            raise BuildError("Send a list of ops.")
        with self.lock:
            if version is not None and int(version) != self.version:
                raise BuildError("The workbook changed since you loaded it (another window?). Reload to see the latest.", "stale", 409,
                                 version=self.version)
            before = self.editor.to_bytes()
            was_modified = self.editor.modified
            try:
                applied = apply_ops(self.editor, ops)
            except Exception:
                self.editor = WorkbookEditor(self.path, before, self.editor.stamp)
                self.editor._modified = was_modified
                raise
            self._undo.append(before)
            del self._undo[:-self.UNDO_DEPTH]
            self._redo.clear()
            self.version += 1
            self.editor.save_draft()
            return applied

    def undo(self) -> None:
        with self.lock:
            if not self._undo:
                raise BuildError("Nothing to undo.", "empty", 409)
            self._redo.append(self.editor.to_bytes())
            data = self._undo.pop()
            if not self._undo and not self.from_draft:
                redo = self._redo
                discard_draft(self.path)
                self._open(prefer_draft=False)                         # all the way back: the file on disk, untouched
                self._redo = redo
                return
            self._load_bytes(data)
            self.version += 1
            self.editor.save_draft()

    def redo(self) -> None:
        with self.lock:
            if not self._redo:
                raise BuildError("Nothing to redo.", "empty", 409)
            self._undo.append(self.editor.to_bytes())
            self._load_bytes(self._redo.pop())
            self.version += 1
            self.editor.save_draft()

    def save(self, *, force: bool = False) -> SaveResult:
        """Write the real ``.xlsx`` (backup first).  Nothing edited = nothing written."""
        from .writer import WorkbookChangedOutside, WorkbookLocked
        with self.lock:
            if not self.editor.modified:
                return SaveResult(self.path, None)
            try:
                result = self.editor.save(force=force)
            except WorkbookLocked as err:
                raise BuildError(str(err), "locked", 423) from err
            except WorkbookChangedOutside as err:
                raise BuildError(str(err), "changed_outside", 409) from err
            self._open(prefer_draft=False)
            return result

    def reload(self) -> None:
        with self.lock:
            discard_draft(self.path)
            self._open(prefer_draft=False)

    # -- history -------------------------------------------------------------------------------
    def history(self) -> list[dict]:
        folder = backups_dir(self.path)
        if not folder.is_dir():
            return []
        out = []
        for f in folder.iterdir():
            if f.is_file() and f.name.startswith(self.path.stem + ".") and f.suffix.lower() == self.path.suffix.lower():
                st = f.stat()
                out.append({"id": f.name, "when": datetime.fromtimestamp(st.st_mtime).isoformat(timespec="seconds"), "size": st.st_size,
                            "_m": st.st_mtime_ns})
        out.sort(key=lambda x: x.pop("_m"), reverse=True)
        return out

    def _backup(self, backup_id: str) -> Path:
        name = str(backup_id or "")
        if not name or "/" in name or "\\" in name or name.startswith(".") or not name.startswith(self.path.stem + "."):
            raise BuildError(f"No saved version called {backup_id!r}.", "not_found", 404)
        f = backups_dir(self.path) / name
        if not f.is_file():
            raise BuildError(f"No saved version called {backup_id!r}.", "not_found", 404)
        return f

    def diff(self, against: str = "disk") -> list[dict]:
        with self.lock:
            other = WorkbookEditor.open(self.path if against in ("", "disk") else self._backup(against))
            before = build_model(other, with_problems=False)
            after = build_model(self.editor, with_problems=False)
            return diff_models(before, after)

    def restore(self, backup_id: str) -> None:
        with self.lock:
            data = self._backup(backup_id).read_bytes()
            self._undo.append(self.editor.to_bytes())
            del self._undo[:-self.UNDO_DEPTH]
            self._redo.clear()
            self._load_bytes(data)
            self.version += 1
            self.editor.save_draft()


def _runs_stamp(runs_dir: Path | None) -> tuple:
    if runs_dir is None or not runs_dir.is_dir():
        return ()
    try:
        return tuple(sorted((f.parent.name, f.stat().st_mtime_ns) for f in runs_dir.glob("*/results.json")))
    except OSError:
        return ()


class BuildStore:
    """One ``BuildDocument`` per workbook for the whole server (one builder per workbook at a time)."""

    def __init__(self):
        self._docs: dict[str, BuildDocument] = {}
        self._lock = threading.Lock()

    def get(self, path: Path, runs_dir: Path | None = None) -> BuildDocument:
        key = str(Path(path).resolve())
        with self._lock:
            doc = self._docs.get(key)
            if doc is None or not doc.path.exists():
                doc = self._docs[key] = BuildDocument(path, runs_dir)
            doc.runs_dir = Path(runs_dir) if runs_dir else doc.runs_dir
        doc.refresh_if_changed_outside()
        return doc

    def forget(self, path: Path) -> None:
        with self._lock:
            self._docs.pop(str(Path(path).resolve()), None)


def keyword_catalogue() -> dict:
    """Every method grouped the way the "All actions" menu shows them (Q9)."""
    groups = [("Do", _ACT | {"CLICK"}), ("Type", _INPUT), ("Check", _CHECK | {"OUTPUT"}), ("Save", _SAVE | {"OUTPUT"}), ("Wait", _WAIT),
              ("Window & frames", _NAV), ("Flow", _FLOW | {"CALL_TEST"})]
    known = KNOWN_METHODS - set(LEGACY_METHODS)
    out, placed = [], set()
    for label, members in groups:
        methods = sorted(m for m in members if m in known)
        placed.update(methods)
        out.append({"group": label, "methods": [{"method": m, "label": auto_name(m, target="…", value="…", expected="…", save_as="X"),
                                                 "element": m in ELEMENT_METHODS, "kind": _step_kind(m, "", "", "")} for m in methods]})
    rest = sorted(known - placed)
    out.append({"group": "Advanced", "methods": [{"method": m, "label": m, "element": m in ELEMENT_METHODS, "kind": _step_kind(m, "", "", "")}
                                                 for m in rest]})
    return {"groups": out, "newKeywords": list(NEW_KEYWORDS), "legacy": sorted(LEGACY_METHODS)}
