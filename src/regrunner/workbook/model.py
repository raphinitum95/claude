"""Workbook -> tests, and the row-by-row *resolution* loop.

This module is a faithful re-implementation of the legacy runner's data handling
(``GUI_Test_Main.run`` + ``GUI_Functions.replacetemplatevalues`` / ``writeresults``), minus Excel:

* Each ``DataSheets`` row is a test; each enabled row of its parameter sheet is one iteration.
* Rows are processed in order.  For every row, any cell (except the result columns) whose text
  equals a parameter-sheet header is replaced by that parameter's value *in the sheet state*, so
  later formulas that reference the cell see the substituted value.
* ``blnExecute`` may be ``Y``/``N``, a literal, a formula, or a parameter *flag column name*.
* After a step runs, ``Status``/``Error_Check``/``Output_Value`` are written back to the sheet
  (so formulas such as ``=S49`` or ``=IF($D$13="PASSED",...)`` work) and, when the result cell
  names a parameter column, into the parameter row (append semantics, like the legacy runner).
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Mapping

from .api import is_api_sheet
from .sheet import BookState, ErrorText, SheetState, WorkbookData, cell_text
from .variables import (INLINE_COLUMNS, INLINE_RE, MARKERS, SAVE_METHODS, TOKEN_RE, EnvironmentTable, FlowMap, VariablePool, declared_variables,
                        flow_map, read_environment_table, secret_value, substitute)

UI_COLUMNS = ("METHOD", "PAGE", "FINDBY", "FINDBY_VALUE")
_KEY_METHODS = {"SENDKEYS", "SENDKEY", "SEND_KEY", "SEND_KEYS", "SPECIALKEY", "CLICK_SENDKEYS", "CUSTOMINPUTDATA"}   # {TAB}, {ENTER} in their Value are keys
RESULT_COLUMNS = ("STATUS", "ERROR_CHECK", "OUTPUT_VALUE", "START_TIME", "END_TIME", "DURATION")
BLANK_OK_METHODS = {"GET_GOOGLE_TOKEN", "ASK_USER", "PROMPT", "ASK"}     # steps that deal with an empty Value themselves (they ask a person for the code / the answer)
DATA_COLUMNS = ("VALUE", "FINDBY_VALUE", "EXPECTED_VALUE")       # where a row's data comes from: a parameter named here that is empty is a problem
_TRUE = {"1", "Y", "YES", "TRUE"}


def is_true(value: Any) -> bool:
    return cell_text(value).upper().strip() in _TRUE


def is_blank(value: Any) -> bool:
    return value is None or cell_text(value).upper().strip() in ("", "NONE")


def to_int(value: Any, default: int = 0) -> int:
    try:
        return int(float(cell_text(value)))
    except (TypeError, ValueError):
        return default


def slug(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", text).strip("_") or "test"


@dataclass
class TestCase:
    """One runnable unit: an action sheet paired with one enabled parameter row."""

    id: str
    sheet: str
    param_sheet: str | None
    param_row: int
    title: str
    description: str = ""
    scenario: str = ""
    tags: list[str] = field(default_factory=list)
    enabled: bool = True
    is_ui: bool = True
    kind: str = "ui"                 # ui | api  (an API test is one data row of a WEBSERVICE_URL sheet: see workbook/api.py)
    iteration: int = 1
    iterations: int = 1
    param_enabled: bool = True       # False: the parameter sheet has no row with blnExecute=Y

    __test__ = False  # not a pytest class

    @property
    def slug(self) -> str:
        return slug(self.id)

    @property
    def runnable(self) -> bool:
        """Something regrunner can execute: a UI keyword sheet row set, or an API data row."""
        return self.is_ui or self.kind == "api"


@dataclass
class PreparedStep:
    """A fully resolved row, ready to execute."""

    row: int
    values: dict[str, Any]                       # UPPER column -> value after substitution/formulas
    raw_output_name: str = ""                    # the Output_Value cell *as authored* (a variable name)
    secret_columns: set[str] = field(default_factory=set)
    blank_params: list[tuple[str, str]] = field(default_factory=list)   # (column, parameter as authored) whose Params cell is empty: the token stays in the cell
    param_reads: list[tuple[str, str]] = field(default_factory=list)    # (column, parameter as authored) for every parameter the row's data comes from
    missing_vars: list[str] = field(default_factory=list)                # {NAME}s nobody has a value for: the step fails, the name is never typed
    unmapped: list[str] = field(default_factory=list)                    # {?NAME} template placeholders nobody mapped: the step fails
    inline_names: list[str] = field(default_factory=list)                # every {NAME} the row's cells use (for Needs / Provides)
    output_name: str = ""                                                # the variable this step saves into ("" = none)

    def text(self, column: str) -> str:
        return cell_text(self.values.get(column))

    @property
    def method(self) -> str:
        return self.text("METHOD").upper().strip()

    @property
    def name(self) -> str:
        return self.text("STEP_NAME").strip()

    @property
    def page(self) -> str:
        return self.text("PAGE").upper().strip()

    @property
    def findby(self) -> str:
        return self.text("FINDBY").upper().strip()

    @property
    def findby_value(self) -> str:
        return self.text("FINDBY_VALUE")

    @property
    def index(self) -> int:
        return to_int(self.values.get("INDEX"))

    @property
    def value(self) -> Any:
        return self.values.get("VALUE")

    @property
    def output_property(self) -> str:
        return self.text("OUTPUT_PROPERTY").upper().strip()

    @property
    def ignore_missing(self) -> bool:
        return is_true(self.values.get("IGNORE_NOT_EXISTING_OBJECT"))

    @property
    def timeout(self) -> int | None:
        raw = self.values.get("TIMEOUT")
        return None if is_blank(raw) else max(to_int(raw, 0), 0) or None

    @property
    def expected(self) -> Any:
        return self.values.get("EXPECTED_VALUE")

    @property
    def exact_match(self) -> bool:
        return is_true(self.values.get("EXACT_MATCH"))

    @property
    def contains(self) -> bool:
        return is_true(self.values.get("CONTAINS"))

    @property
    def locator_column(self) -> str:
        """Optional generic-selector column (``Locator``); empty when the sheet has none."""
        return self.text("LOCATOR").strip()

    @property
    def test_case(self) -> str:
        return self.text("TEST_CASE")


# ---------------------------------------------------------------------------------------------
class Workbook:
    """A loaded workbook plus run-level settings; creates one ``TestRuntime`` per test."""

    def __init__(self, path: str | Path, *, environment: str | None = None, headless: bool = True,
                 seed: int | None = None, now: datetime | None = None,
                 secrets: Mapping[str, str] | None = None):
        self.path = Path(path)
        self.data = WorkbookData.load(self.path)
        self.environment = environment
        self.headless = headless
        self.seed = seed
        self.now = now
        self.secrets = {str(k).upper(): str(v) for k, v in (secrets or {}).items() if v not in (None, "")}
        self.warnings: list[str] = []
        self._discovered: list[TestCase] = []           # what discover() found last (param_writers reads it)
        self._env_table: EnvironmentTable | None = None
        self._variable_names: set[str] | None = None

    # -- variables (CONTRACT.md 1.2 / 1.4) ----------------------------------------------------------
    def environment_table(self) -> EnvironmentTable:
        if self._env_table is None:
            self._env_table = read_environment_table(self.data)
        return self._env_table

    def missing_required(self, environment: str) -> list[str]:
        """Required environment variables that have no value for ``environment``: a run there must not start."""
        return self.environment_table().missing_required(environment)

    def variable_names(self) -> set[str]:
        """Every name a step can save into by writing it bare in Output_Value: the headers of the parameter sheets, and ``_rr_variables``."""
        if self._variable_names is None:
            names = set(declared_variables(self.data))
            ds = self.data.sheet("DataSheets")
            col = ds.headers.get("PARAMETERSHEET") if ds is not None else None
            for row in range(2, (ds.max_row if ds is not None else 1) + 1):
                params = self.data.sheet(_raw(ds, row, col)) if col else None
                if params is not None:
                    names |= {h for h in params.headers if TOKEN_RE.match(h) and h != "BLNEXECUTE"}
            self._variable_names = names
        return self._variable_names

    def find_case(self, test_id: str) -> TestCase:
        """The test ``CALL_TEST`` names: its id (``Login``, ``Login#2``), or a sheet with a single test.  Raises ``ValueError``."""
        cases = self._discovered or self.discover()
        want = test_id.strip().upper()
        exact = [c for c in cases if c.id.upper() == want]
        if exact:
            return exact[0]
        by_sheet = [c for c in cases if c.sheet.upper() == want]
        if len(by_sheet) == 1:
            return by_sheet[0]
        if by_sheet:
            raise ValueError(f"{test_id} has {len(by_sheet)} tests (one per data row): name one of {', '.join(c.id for c in by_sheet)}")
        raise ValueError(f"there is no test {test_id!r} in this workbook")

    # -- discovery ---------------------------------------------------------------------------
    def global_settings(self) -> dict[str, Any]:
        """The evaluated ``Global`` sheet with run overrides applied."""
        book = BookState(self.data, now=self.now, seed=self.seed)
        return self._build_global(book)

    def _build_global(self, book: BookState) -> dict[str, Any]:
        if not book.has_sheet("Global"):
            return {}
        sheet = book.sheet("Global")
        overrides = {"ENVIRONMENT": self.environment, "HEADLESSMODE": "Y" if self.headless else "N"}
        for row in range(2, sheet.data.max_row + 1):
            key = cell_text(sheet.get(row, 1)).strip().upper()
            if key in overrides and overrides[key] is not None:
                sheet.set(row, 2, overrides[key])
        sheet.freeze()
        settings: dict[str, Any] = {}
        for row in range(2, sheet.data.max_row + 1):
            name = cell_text(sheet.get(row, 1)).strip()
            if name:
                settings[name] = sheet.read(row, 2)
        return settings

    def discover(self) -> list[TestCase]:
        self.warnings = []                                  # rebuilt on every call: asking twice must not say everything twice
        if "EXPORT SUMMARY" in self.data.sheets:
            self.warnings.append('This workbook was exported from Apple Numbers (it has an "Export Summary" sheet). Numbers replaces formulas with the values they had on the day '
                                 'of the export: dates such as TODAY()+1 stay frozen in the past and calculated cells stop calculating. Edit the .xlsx in Excel (or LibreOffice) '
                                 'instead, or restore the copy from before it was saved in Numbers.')
        ds = self.data.sheet("DataSheets")
        if ds is None:
            raise ValueError("Workbook has no 'DataSheets' sheet listing the tests to run.")
        headers = ds.headers
        col_sheet = headers.get("SHEETSTOEXECUTE", 1)
        col_flag = headers.get("BLNEXECUTE", 2)
        col_param = headers.get("PARAMETERSHEET")
        col_comment = headers.get("COMMENTS")
        col_tags = headers.get("TAGS")
        cases: list[TestCase] = []
        for row in range(2, ds.max_row + 1):
            sheet_name = _raw(ds, row, col_sheet)
            if not sheet_name:
                continue
            action = self.data.sheet(sheet_name)
            if action is None:
                self.warnings.append(f"DataSheets row {row}: sheet {sheet_name!r} does not exist")
                continue
            is_ui = all(c in action.headers for c in UI_COLUMNS)
            param_name = _raw(ds, row, col_param) if col_param else ""
            comment = _raw(ds, row, col_comment) if col_comment else ""
            tags = [t.strip() for t in re.split(r"[;,]", _raw(ds, row, col_tags))] if col_tags else []
            if not is_ui and is_api_sheet(action.headers):
                cases.extend(self._api_cases(sheet_name, comment, tags, is_true(_raw(ds, row, col_flag))))
                continue
            if not is_ui:
                self.warnings.append(f"Sheet {sheet_name!r} is not a UI keyword sheet; skipped")
            enabled = is_true(_raw(ds, row, col_flag))
            iterations = self._enabled_param_rows(param_name)
            no_enabled_row = bool(param_name) and not iterations
            if no_enabled_row:
                self.warnings.append(f"Parameter sheet {param_name!r} has no row with blnExecute=Y; "
                                     f"test {sheet_name!r} cannot run")
            for n, prow in enumerate(iterations or [2], start=1):
                scenario = self._scenario_label(param_name, prow)
                tid = sheet_name if len(iterations) <= 1 else f"{sheet_name}#{n}"
                title = sheet_name if not scenario or len(iterations) <= 1 else f"{sheet_name} · {scenario}"
                cases.append(TestCase(
                    id=tid, sheet=sheet_name, param_sheet=param_name or None, param_row=prow,
                    title=title, description=comment, scenario=scenario, tags=[t for t in tags if t],
                    enabled=enabled and not no_enabled_row, is_ui=is_ui, iteration=n,
                    iterations=max(len(iterations), 1), param_enabled=not no_enabled_row))
        self._discovered = cases
        return cases

    def _api_cases(self, sheet_name: str, comment: str, tags: list[str], enabled: bool) -> list[TestCase]:
        """One test per data row of an API sheet whose ``blnExecute`` is Y."""
        state = BookState(self.data, now=self.now, seed=self.seed).sheet(sheet_name)
        bln = state.data.headers["BLNEXECUTE"]
        rows = [r for r in range(2, state.data.max_row + 1) if is_true(state.read(r, bln))]
        if not rows:
            self.warnings.append(f"API sheet {sheet_name!r} has no row with blnExecute=Y; nothing to run")
        out = []
        for n, r in enumerate(rows, start=1):
            scenario = ""
            for header in ("TC_NAME", "TESTCONDITION"):
                col = state.data.headers.get(header)
                if col and (text := cell_text(state.read(r, col)).strip()):
                    scenario = text
                    break
            out.append(TestCase(id=sheet_name if len(rows) == 1 else f"{sheet_name}#{n}", sheet=sheet_name, param_sheet=None, param_row=r,
                                title=sheet_name if not scenario or len(rows) == 1 else f"{sheet_name} · {scenario}", description=comment,
                                scenario=scenario, tags=[t for t in tags if t], enabled=enabled, is_ui=False, kind="api",
                                iteration=n, iterations=len(rows)))
        return out

    def _enabled_param_rows(self, param_name: str) -> list[int]:
        if not param_name:
            return []
        params = self.data.sheet(param_name)
        if params is None:
            self.warnings.append(f"Parameter sheet {param_name!r} does not exist")
            return []
        col = params.headers.get("BLNEXECUTE")
        if col is None:
            return []
        book = BookState(self.data, now=self.now, seed=self.seed)
        state = book.sheet(param_name)
        return [r for r in range(2, params.max_row + 1) if is_true(state.read(r, col))]

    def _scenario_label(self, param_name: str, row: int) -> str:
        if not param_name:
            return ""
        params = self.data.sheet(param_name)
        if params is None:
            return ""
        for header in ("NOTES", "SCENARIO"):
            col = params.headers.get(header)
            if col:
                value = params.cells.get((row, col))
                if value is not None and value.formula is None and value.value:
                    return str(value.value)
        return ""

    # -- runtime -----------------------------------------------------------------------------
    def api_runtime(self, case: TestCase, shared: dict | None = None, *, seed: int | None = None):
        """Runtime of an API row.  ``shared``: cells the UI tests of this run wrote into their parameter sheets, ``{(sheet, row, col): value}``,
        so that ``=AgentStandAlone_Params!E2`` in an API row sees the policy number the UI test produced."""
        from .api import ApiRuntime
        book = BookState(self.data, now=self.now, seed=self.seed if seed is None else seed)
        globals_ = self._build_global(book)
        for (sheet, row, col), value in (shared or {}).items():
            if book.has_sheet(sheet):
                book.sheet(sheet).set(row, col, value)
        return ApiRuntime(book, case, globals_, self.secrets)

    def runtime(self, case: TestCase, *, seed: int | None = None, shared: dict | None = None, pool: VariablePool | None = None) -> "TestRuntime":
        """``shared``: parameter cells other tests of the run wrote (or a person supplied), ``{(sheet, row, col): value}``; this test sees them.
        ``pool``: the run's shared variables (values any test saved under a name); ``None`` = a pool of its own (planning)."""
        if case.kind == "api":
            return self.api_runtime(case, shared, seed=seed)               # type: ignore[return-value]
        book = BookState(self.data, now=self.now, seed=self.seed if seed is None else seed)
        globals_ = self._build_global(book)
        params_state = None
        if case.param_sheet:
            params_state = book.sheet(case.param_sheet)
            for (sheet, row, col), value in (shared or {}).items():
                if sheet == case.param_sheet and row == case.param_row:
                    params_state.set(row, col, value)
            for (sheet, row, col), value in (shared or {}).items():      # a Params cell may be =PurchaseUS!FB3: the cell an API test filled in
                other = self.data.sheet(sheet)
                if other is not None and sheet != case.param_sheet and is_api_sheet(other.headers) and book.has_sheet(sheet):
                    book.sheet(sheet).set(row, col, value)
            params_state.freeze()            # one consistent value per volatile formula for this test
        return TestRuntime(book, case, globals_, params_state, self.secrets, writers=self.param_writers(), pool=pool,
                           env_table=self.environment_table(), variable_names=self.variable_names())

    def param_writers(self) -> dict[str, list[str]]:
        """Which sheets have a step that writes a parameter (its Output_Value cell holds the parameter's name): ``{PARAMETER: [sheet, ...]}``.
        Read from the tests found by ``discover()``; empty before that."""
        writers: dict[str, list[str]] = {}
        seen: set[tuple[str, str]] = set()
        for case in self._discovered:
            if not case.is_ui or not case.param_sheet or (case.sheet, case.param_sheet) in seen:
                continue
            seen.add((case.sheet, case.param_sheet))
            params, sheet = self.data.sheet(case.param_sheet), self.data.sheet(case.sheet)
            col = sheet.headers.get("OUTPUT_VALUE") if sheet is not None else None
            if params is None or not col:
                continue
            for row in range(2, sheet.max_row + 1):
                name = _raw(sheet, row, col).upper()
                if name in params.headers and case.sheet not in writers.setdefault(name, []):
                    writers[name].append(case.sheet)
        return writers


def _raw(sheet, row: int, col: int | None) -> str:
    if not col:
        return ""
    cell = sheet.cells.get((row, col))
    if cell is None:
        return ""
    return cell_text(cell.value).strip() if cell.formula is None else ""


class TestRuntime:
    """Per-test row processor. Not shared between tests."""

    __test__ = False  # not a pytest class

    def __init__(self, book: BookState, case: TestCase, globals_: dict[str, Any],
                 params: SheetState | None, secrets: Mapping[str, str], writers: Mapping[str, list[str]] | None = None,
                 pool: VariablePool | None = None, env_table: EnvironmentTable | None = None, variable_names: set[str] | None = None):
        self.book = book
        self.writers = writers or {}                             # PARAMETER -> sheets that fill it in (for the message about an empty one)
        self.needs: list[dict[str, Any]] = []                    # filled by plan(): parameters this test uses that are empty and nothing in it sets
        self.reads: dict[str, dict[str, Any]] = {}               # filled by plan(): PARAMETER -> {param, rows, blank} for every parameter it uses that no earlier step of it sets
        self.sets: set[str] = set()                              # filled by plan(): parameters its steps set (Output_Value names them)
        self.case = case
        self.globals = globals_
        self.sheet: SheetState = book.sheet(case.sheet)
        self.params_sheet = params
        self.secrets = secrets
        self.columns = self.sheet.data.headers                   # UPPER -> col
        self.column_names = sorted(self.columns.items(), key=lambda kv: kv[1])
        self.params: dict[str, Any] = {}
        if params is not None:
            for header, col in params.data.headers.items():
                self.params[header] = params.read(case.param_row, col)
        self.total_rows = self.sheet.data.max_row
        self.written: dict[tuple[str, int, int], Any] = {}         # parameter cells this test wrote: handed to the API tests of the run
        self.pool = pool if pool is not None else VariablePool()   # the run's shared variables (CONTRACT.md 1.2)
        self.secret_values: set[str] = self.pool.secret_values     # secrets (typed in by a person, {SECRET:X}): masked wherever a step is shown, in every test of the run
        self.env_table = env_table or EnvironmentTable()
        self.variable_names = variable_names or set()
        self.loops: list[dict[str, Any]] = []                      # the data rows of the loops the test is in (innermost last): {HEADER: value}
        self.pool_reads: set[str] = set()                          # filled by plan(): {NAME}s it needs from another test (Needs, Q21)
        self.pool_sets: set[str] = set()                           # filled by plan(): names its steps save (Provides)
        self.calls: list[str] = []                                 # filled by plan(): the tests its CALL_TEST steps run
        self._authored_output: dict[int, str] = {}                 # Output_Value of each row before a result was written over it (a loop runs a row again)
        self._flow: FlowMap | None = None
        self.blocked_reason = "" if case.param_enabled else (
            f"Parameter sheet {case.param_sheet!r} has no row with blnExecute=Y")
        if not self.blocked_reason and self.flow().errors:
            self.blocked_reason = "The IF / loop steps do not match up: " + "; ".join(self.flow().errors)

    # -- global helpers -----------------------------------------------------------------------
    def global_flag(self, name: str, default: bool = False) -> bool:
        for key, value in self.globals.items():
            if key.upper() == name.upper():
                return is_true(value)
        return default

    def global_text(self, name: str, default: str = "") -> str:
        for key, value in self.globals.items():
            if key.upper() == name.upper():
                return cell_text(value)
        return default

    @property
    def environment(self) -> str:
        return self.global_text("Environment", "").strip().upper()

    # -- row processing -------------------------------------------------------------------------
    def prepare_row(self, row: int) -> PreparedStep | None:
        """Substitute tokens for ``row`` and return it if it should execute, else ``None``."""
        blanks, reads = self._replace_template_values(row)
        # STEP_NUMBER is `=COUNTA($B$1:B<n>)` (== row-1). Evaluating it is quadratic and it never
        # influences behaviour, so it is synthesised instead.
        values = {name: (row - 1 if name == "STEP_NUMBER" else self.sheet.read(row, col))
                  for name, col in self.columns.items()}
        if "OUTPUT_VALUE" in values:
            if row in self._authored_output:
                values["OUTPUT_VALUE"] = self._authored_output[row]           # (the result a loop's earlier pass wrote over it is not the cell)
            else:
                self._authored_output[row] = values["OUTPUT_VALUE"]
        if not is_true(values.get("BLNEXECUTE")):
            return None
        raw_output = cell_text(values.get("OUTPUT_VALUE")).strip()
        prepared = PreparedStep(row=row, values=values, raw_output_name=raw_output, blank_params=blanks, param_reads=reads)
        self._assign_secrets(prepared)
        if not prepared.method:
            return None
        prepared.output_name = self.output_name(prepared)
        self._substitute_inline(prepared)
        return prepared

    # -- variables (CONTRACT.md 1.2) ------------------------------------------------------------------------
    def output_name(self, step: PreparedStep) -> str:
        """The variable a step saves into: its Output_Value is ``{NAME}``, or a bare name that is a variable (a Params column, ``_rr_variables``)
        or is written by a step whose job is to save (``SET_VARIABLE``, ``ASK_USER``...).  "" otherwise (a literal the step is compared with)."""
        text = step.raw_output_name
        if not text:
            return ""
        m = INLINE_RE.fullmatch(text)
        if m:
            return m.group(1)
        if TOKEN_RE.match(text) and (text.upper() in self.params or text.upper() in self.variable_names or step.method in SAVE_METHODS):
            return text
        return ""

    def value_of(self, name: str, secrets: list[str] | None = None) -> tuple[str | None, bool]:
        """``(value, is it a Params column of this test)`` of ``{NAME}``: the loop row, the Params row, the run's pool, the environment table.
        A value from the environment table may be ``{SECRET:X}``: the secret is put in, and ``secrets`` gets it (masked wherever the step shows)."""
        key = name.upper().strip()
        for row in reversed(self.loops):
            if key in row and not is_blank(row[key]):
                return cell_text(row[key]), False
        is_param = key in self.params
        if is_param and not is_blank(self.params[key]):
            return cell_text(self.params[key]), True
        pooled = self.pool.get(key)
        if pooled is not None:
            return pooled, is_param
        env = self.env_table.value(key, self.environment)
        if env is not None:
            from .variables import SECRET_RE
            found: list[str] = []

            def put(m) -> str:
                value = secret_value(m.group(1), self.environment)
                if value is None:
                    return m.group(0)
                found.append(value)
                return value
            env = SECRET_RE.sub(put, env)
            if found:
                self.secret_values.update(found)
                if secrets is not None:
                    secrets.extend(found)
            return env, is_param
        return None, is_param

    def _substitute_inline(self, step: PreparedStep) -> None:
        """``{NAME}`` / ``{SECRET:NAME}`` / ``{?NAME}`` in the text columns (CONTRACT.md 1.2).  An IF keeps its condition as written (it reads the
        variables itself); a SendKeys value keeps ``{TAB}``, ``{ENTER}``... (keys, not variables)."""
        from ..engine.keys import SPECIAL
        keys = {k.upper() for k in SPECIAL}
        keep = keys if step.method in _KEY_METHODS else None
        for column in INLINE_COLUMNS:
            if column not in step.values:
                continue
            if column == "VALUE" and step.method == "IF":                   # (read by the IF itself; listed for Needs / Provides)
                step.inline_names.extend(n.upper() for n in INLINE_RE.findall(cell_text(step.values[column])))
                continue
            cell = step.values[column]
            if is_blank(cell) or isinstance(cell, ErrorText):
                continue
            text = cell_text(cell)
            if "{" not in text:
                continue
            found: list[str] = []
            done = substitute(text, lambda n: self.value_of(n, found),
                              lambda n: secret_value(n, self.environment), keep=keep if column == "VALUE" else None, soft=keys)
            step.inline_names.extend(n.upper() for n in INLINE_RE.findall(text) if not (keep and column == "VALUE" and n.upper() in keep))
            step.unmapped.extend(done.unmapped)
            step.missing_vars.extend(done.missing)
            for name in done.blank_params:
                step.blank_params.append((column, name))
            secrets = done.secrets + found
            if secrets:
                self.secret_values.update(secrets)
                step.secret_columns.add(column)
            if done.text != text:
                step.values[column] = done.text

    def iteration_rows(self, sheet_name: str) -> list[dict[str, Any]]:
        """The data rows an ``ITERATION_START`` repeats over: every row of ``sheet_name`` (blank = the test's own Params sheet) whose blnExecute
        is Y (every non-empty row when it has no blnExecute column), as ``{HEADER: value}``.  Raises ``ValueError`` for a sheet that does not exist."""
        name = sheet_name.strip() or (self.case.param_sheet or "")
        if not name:
            raise ValueError("the loop names no data sheet and the test has no Params sheet")
        if not self.book.has_sheet(name):
            raise ValueError(f"the loop's data sheet {name!r} does not exist")
        sheet = self.book.sheet(name)
        headers = sheet.data.headers
        flag = headers.get("BLNEXECUTE")
        rows = []
        for r in range(2, sheet.data.max_row + 1):
            values = {h: sheet.read(r, c) for h, c in headers.items()}
            if flag is not None:
                if not is_true(values.get("BLNEXECUTE")):
                    continue
            elif all(is_blank(v) for v in values.values()):
                continue
            rows.append(values)
        return rows

    def flow(self) -> FlowMap:
        """Where each IF / ELSE / loop of the sheet ends (from the Method column as written)."""
        if self._flow is None:
            col = self.columns.get("METHOD")
            methods = {}
            for row in range(2, self.total_rows + 1) if col else ():
                m = self.method_at(row)
                if m in ("IF", "ELSE", "END_IF", "ITERATION_START", "ITERATION_END"):
                    methods[row] = m
            self._flow = flow_map(methods)
        return self._flow

    def _authored_text(self, row: int, column: str) -> str:
        col = self.columns.get(column)
        return cell_text(self.sheet.read(row, col)).strip() if col else ""

    def _replace_template_values(self, row: int) -> tuple[list[tuple[str, str]], list[tuple[str, str]]]:
        """Replace each parameter named in the row by its value.  Returns ``(blanks, reads)``: the (column, parameter) pairs whose parameter is empty
        (the legacy runner left the token in the cell then and typed its name into the page; regrunner reports it instead, see ``TestRunner``), and
        every parameter the row's data comes from."""
        blanks: list[tuple[str, str]] = []
        reads: list[tuple[str, str]] = []
        params = self.params
        if not params:
            return blanks, reads
        for name, col in self.column_names:
            if name in RESULT_COLUMNS or name == "STEP_NUMBER":
                continue
            authored = cell_text(self.sheet.read(row, col)).strip()
            cell_upper = authored.upper() or "NONE"
            if cell_upper not in params:
                if name == "BLNEXECUTE" and not is_true(cell_upper):
                    return blanks, reads
                continue
            if name == "BLNEXECUTE" and not is_true(params[cell_upper]):
                return blanks, reads
            if name in DATA_COLUMNS:
                reads.append((name, authored))
            replacement = params[cell_upper]
            if replacement is None:
                if name in DATA_COLUMNS:
                    blanks.append((name, authored))
                continue          # keep the token: blank parameters do not overwrite the sheet
            self.sheet.set(row, col, cell_text(replacement))
        return blanks, reads

    def _assign_secrets(self, step: PreparedStep) -> None:
        if not self.secrets:
            return
        for name, value in list(step.values.items()):
            if name in RESULT_COLUMNS or value in (None, ""):
                continue
            key = cell_text(value).upper().strip()
            if key in self.secrets:
                step.values[name] = self.secrets[key]
                step.secret_columns.add(name)

    def record(self, step: PreparedStep, status: str, error: str, output: Any) -> list[dict[str, str]]:
        """Write the step result back into the sheet and (when named) into the parameter row, and put what the step saved into the run's pool.
        Returns the variables this step set: ``[{name, value, stored, cell}]`` (``value`` what the step wrote, ``stored`` what the cell holds now:
        legacy writes to one parameter add up; ``cell`` is ``run variable`` for a name that is not a Params column)."""
        out_text = "" if output is None else cell_text(output)
        for column, value in (("STATUS", status), ("ERROR_CHECK", error), ("OUTPUT_VALUE", out_text)):
            col = self.columns.get(column)
            if col:
                self.sheet.set(step.row, col, value)
        sets: list[dict[str, str] | None] = []
        saved = step.output_name.upper()
        if self.params_sheet is not None:
            for column, cell in step.values.items():
                if is_blank(cell):
                    continue
                token = cell_text(cell).upper().strip()
                if column == "OUTPUT_VALUE" and saved:
                    token = saved                                   # (``{NAME}`` saves into NAME like a bare NAME)
                if token not in self.params or column in step.secret_columns:
                    continue
                authored = cell_text(cell).strip()
                if column == "ERROR_CHECK":
                    if not is_blank(error):
                        sets.append(self._write_param(token, error, authored))
                elif column == "STATUS":
                    sets.append(self._write_param(token, status, authored))
                elif column == "OUTPUT_VALUE":
                    if not is_blank(out_text) and out_text.upper() != token:
                        sets.append(self._write_param(token, out_text, step.output_name or authored, replace=step.method == "SET_VARIABLE"))
        keeps = status == "PASSED" if step.method == "SET_VARIABLE" else not is_blank(out_text)      # SET_VARIABLE may set "" on purpose
        if saved and keeps and out_text.upper() != saved:
            self.pool.set(saved, out_text, self.case.id)
            if not any(x and x["name"].upper() == saved for x in sets):
                sets.append({"name": step.output_name, "value": out_text, "stored": out_text, "cell": "run variable"})
        return [x for x in sets if x]

    def param_cell(self, token: str) -> str:
        """Where a parameter lives, for a message: ``AgentPortal_Params!R3``."""
        col = self.params_sheet.data.headers.get(token.upper().strip()) if self.params_sheet is not None else None
        if not col:
            return token
        from openpyxl.utils import get_column_letter
        return f"{self.params_sheet.name}!{get_column_letter(col)}{self.case.param_row}"

    def provide(self, token: str, value: str) -> None:
        """A value a person supplied for an empty parameter: this test uses it from now on, and it is handed on like any value a step wrote."""
        key = token.upper().strip()
        col = self.params_sheet.data.headers.get(key) if self.params_sheet is not None else None
        if col is None:
            return
        self.params[key] = value
        self.params_sheet.set(self.case.param_row, col, value)
        self.written[(self.params_sheet.name, self.case.param_row, col)] = value

    def _write_param(self, token: str, value: str, authored: str = "", replace: bool = False) -> dict[str, str] | None:
        col = self.params_sheet.data.headers.get(token)
        if col is None:
            return None
        current = self.params.get(token)
        new = str(value) if is_blank(current) or replace else f"{cell_text(current)};{value}"
        self.params[token] = new
        self.params_sheet.set(self.case.param_row, col, new)
        self.written[(self.params_sheet.name, self.case.param_row, col)] = new
        return {"name": authored or token, "value": str(value), "stored": new, "cell": self.param_cell(token)}

    def method_at(self, row: int) -> str:
        """The Method of ``row`` as written (UPPER), whatever its blnExecute says: ELSE / END_IF / ITERATION_END are structure and always count."""
        col = self.columns.get("METHOD")
        cell = self.sheet.data.cells.get((row, col)) if col else None
        return "" if cell is None or cell.formula else cell_text(cell.value).strip().upper()

    def empty_flag_step(self, row: int) -> tuple[str, str] | None:
        """(method, step name) when ``row`` holds a step but its ``blnExecute`` cell is *empty*: an ``N`` is a decision, an empty cell is easily
        an oversight, and either way the row does not run.  A flag that is a formula is not reported (it may legitimately give "")."""
        bln, method = self.columns.get("BLNEXECUTE"), self.columns.get("METHOD")
        if not bln or not method or self.sheet.is_formula(row, bln):
            return None
        if cell_text(self.sheet.read(row, bln)).strip():
            return None
        what = cell_text(self.sheet.read(row, method)).replace("\xa0", " ").strip()
        if not what:
            return None
        name = self.columns.get("STEP_NAME")
        return what, cell_text(self.sheet.read(row, name)).strip() if name else ""

    # -- planning ------------------------------------------------------------------------------
    def plan(self) -> list[tuple[int, str, str]]:
        """Dry-run the row loop (no browser) and return ``(row, step name, method)`` for each step.

        Statuses are assumed PASSED so status-gated rows are planned; used for progress totals.  Both sides of an IF are planned (which one runs
        is only known then); a loop's steps are planned once per data row; ELSE / END_IF / ITERATION_END are structure, not steps.
        """
        planned = []
        if self.blocked_reason:
            return planned
        setters: set[str] = set()                                # parameters an earlier step of this test fills in
        needs: dict[str, dict[str, Any]] = {}
        reads: dict[str, dict[str, Any]] = {}
        saved: set[str] = set()                                  # names an earlier step of this test saves (into the pool)
        pool_reads: set[str] = set()
        loops: dict[int, int] = {}                               # ITERATION_START row -> how many data rows
        loop_columns: list[tuple[int, set[str]]] = []            # (ITERATION_END row, columns of that loop's data sheet)
        env_names = {r.variable.upper() for r in self.env_table.rows}
        fm = self.flow()
        for row in range(2, self.total_rows + 1):
            loop_columns = [(end, cols) for end, cols in loop_columns if row < end]
            step = self.prepare_row(row)
            if step is None or step.method in MARKERS:
                continue
            planned.append((row, step.name, step.method))
            if step.method == "ITERATION_START" and row in fm.end_of:
                try:
                    data = self.iteration_rows(step.text("VALUE"))
                except ValueError:
                    data = []
                loops[row] = len(data)
                loop_columns.append((fm.end_of[row], {h for d in data[:1] for h in d}))
            if step.method == "CALL_TEST" and step.text("VALUE").strip():
                self.calls.append(step.text("VALUE").strip())
            in_loop = set().union(*(cols for _, cols in loop_columns)) if loop_columns else set()
            for name in step.inline_names:
                key = name.upper()
                if key in in_loop or key in saved or key in env_names or (key in self.params and not is_blank(self.params[key])):
                    continue
                pool_reads.add(key)
            for _, token in step.param_reads:
                if token.upper() not in setters:
                    col = self.params_sheet.data.headers.get(token.upper()) if self.params_sheet is not None else None
                    read = reads.setdefault(token.upper(), {"param": token, "rows": [], "blank": False, "source": self.param_cell(token),
                                                            "key": (self.params_sheet.name, self.case.param_row, col) if col else None})
                    read["rows"].append(row)
            for _, token in ([] if step.method in BLANK_OK_METHODS else step.blank_params):
                if token.upper() not in setters and token.upper() not in pool_reads:
                    need = needs.setdefault(token.upper(), {"param": token, "cell": self.param_cell(token), "rows": [], "steps": [],
                                                            "set_by": [w for w in self.writers.get(token.upper(), []) if w != self.case.sheet]})
                    need["rows"].append(row)
                    need["steps"].append(step.name)
                    if token.upper() in reads:
                        reads[token.upper()]["blank"] = True
            if step.raw_output_name and step.raw_output_name.upper().strip() in self.params:
                setters.add(step.raw_output_name.upper().strip())
            if step.output_name:
                saved.add(step.output_name.upper())
            self.record(step, "PASSED", "", "")
        self.needs = list(needs.values())
        self.reads, self.sets = reads, setters
        self.pool_reads, self.pool_sets = pool_reads, saved
        return _unroll(planned, fm, loops)


def _unroll(planned: list[tuple[int, str, str]], fm: FlowMap, loops: dict[int, int]) -> list[tuple[int, str, str]]:
    """A loop's steps once per data row (none when it has no row), loops inside loops included."""
    out: list[tuple[int, str, str]] = []
    i = 0
    while i < len(planned):
        item = planned[i]
        out.append(item)
        i += 1
        if item[2] == "ITERATION_START" and item[0] in loops:
            end = fm.end_of[item[0]]
            body = []
            while i < len(planned) and planned[i][0] < end:
                body.append(planned[i])
                i += 1
            out.extend(_unroll(body, fm, loops) * loops[item[0]])
    return out
